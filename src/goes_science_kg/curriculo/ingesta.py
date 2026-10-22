"""Ingest all configured Excel sheets from provenance-checked byte snapshots."""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import unicodedata
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from goes_science_kg import config
from goes_science_kg.curriculo._configuracion import ConfigCurriculo, ConfigHoja
from goes_science_kg.curriculo.contratos import (
    FASES,
    AgrupacionCurricular,
    CeldaCurricular,
    ColumnaCurricular,
    Curriculo,
    FaseCurricular,
    FilaCurricular,
    FuenteCurricular,
    HojaCurricular,
    RegistroHoja,
)

_CODIGO_LOCAL = re.compile(r"(?<![\w])(?:EP|PR|E|D|I)\d+(?:\.\d+)*(?![\w])")
LOGGER = logging.getLogger(__name__)


def _sha(contenido: bytes) -> str:
    return hashlib.sha256(contenido).hexdigest()


def _json(valor: Any) -> bytes:
    return json.dumps(valor, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def _normalizar(valor: Any) -> str:
    texto = str(valor or "").strip().split("\n", 1)[0]
    return " ".join(unicodedata.normalize("NFKC", texto).casefold().split())


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _dentro(base: Path, relativa: str) -> Path:
    candidato = base / relativa
    resuelta = candidato.resolve()
    if not resuelta.is_relative_to(base):
        raise ValueError(f"Ruta fuera de la raíz de ingesta: {relativa}")
    return resuelta


def _catalogo_alias(cfg: ConfigCurriculo) -> dict[str, str]:
    resultado: dict[str, str] = {}
    for campo, alias in cfg.alias_campos.items():
        for etiqueta in alias:
            clave = _normalizar(etiqueta)
            if not clave or (clave in resultado and resultado[clave] != campo):
                raise ValueError(f"Alias vacío o ambiguo: {etiqueta!r}")
            resultado[clave] = campo
    return resultado


def _columnas(hoja: Worksheet, cfg: ConfigHoja, alias: dict[str, str]) -> tuple[ColumnaCurricular, ...]:
    columnas: list[ColumnaCurricular] = []
    usados: set[str] = set()
    for indice in range(1, hoja.max_column + 1):
        letra = get_column_letter(indice)
        encabezados = tuple(hoja.cell(fila, indice).value for fila in cfg.filas_encabezado)
        campo = cfg.campos_por_columna.get(letra)
        if campo is None and cfg.tratamiento == "curricular":
            candidatos = {alias[n] for v in encabezados if (n := _normalizar(v)) in alias}
            if len(candidatos) > 1:
                raise ValueError(f"Encabezados ambiguos en {hoja.title}!{letra}: {sorted(candidatos)}")
            campo = next(iter(candidatos), None)
        if campo is not None:
            if campo in usados:
                raise ValueError(f"Campo duplicado {campo!r} en hoja {hoja.title!r}; no se sobrescribe.")
            usados.add(campo)
        columnas.append(ColumnaCurricular(columna=letra, indice=indice, encabezados=encabezados, campo=campo))
    if fuera := set(cfg.campos_por_columna) - {c.columna for c in columnas}:
        raise ValueError(f"Columnas configuradas inexistentes en {hoja.title}: {sorted(fuera)}")
    if cfg.tratamiento == "curricular" and "procedimental" not in usados:
        raise ValueError(f"No hay encabezado Procedimental en {hoja.title!r}.")
    return tuple(columnas)


def _combinadas(hoja: Worksheet) -> dict[tuple[int, int], str]:
    """Expand only explicitly merged ranges without forward-filling ordinary blank cells."""
    resultado = {}
    for rango in hoja.merged_cells.ranges:
        ancla = hoja.cell(rango.min_row, rango.min_col).coordinate
        for fila in range(rango.min_row, rango.max_row + 1):
            for columna in range(rango.min_col, rango.max_col + 1):
                resultado[fila, columna] = ancla
    return resultado


def _registros(hoja: Worksheet, columnas: tuple[ColumnaCurricular, ...]) -> tuple[RegistroHoja, ...]:
    combinadas = _combinadas(hoja)
    registros = []
    for fila in hoja.iter_rows():
        if not any(c.value is not None or c.comment or c.hyperlink for c in fila):
            continue
        celdas = []
        for descriptor, celda in zip(columnas, fila, strict=True):
            ancla = combinadas.get((celda.row, celda.column))
            efectiva = hoja[ancla] if ancla else celda
            celdas.append(
                CeldaCurricular(
                    coordenada=f"{descriptor.columna}{fila[0].row}",
                    fila=fila[0].row,
                    columna=descriptor.columna,
                    campo=descriptor.campo,
                    valor_original=celda.value,
                    valor=efectiva.value,
                    origen_combinado=ancla,
                    tipo_excel=efectiva.data_type,
                    comentario=celda.comment.text if celda.comment else None,
                    hipervinculo=celda.hyperlink.target if celda.hyperlink else None,
                )
            )
        registros.append(RegistroHoja(fila=fila[0].row, celdas=tuple(celdas)))
    return tuple(registros)


def _fila_curricular(
    registro: RegistroHoja,
    hoja: Worksheet,
    cfg: ConfigHoja,
    fuente: FuenteCurricular,
    campos_presentes: set[str],
) -> FilaCurricular | None:
    if registro.fila <= max(cfg.filas_encabezado):
        return None
    por_campo = {celda.campo: celda for celda in registro.celdas if celda.campo}
    procedimiento = por_campo["procedimental"]
    # Una celda combinada no crea un segundo tema por arrastre del mismo procedimiento.
    if not _texto(procedimiento.valor_original):
        return None
    if procedimiento.tipo_excel == "f":
        raise ValueError(
            f"Procedimental contiene fórmula sin valor verificado: {hoja.title}!{procedimiento.coordenada}"
        )
    if not isinstance(procedimiento.valor_original, str):
        raise ValueError(f"Procedimental debe contener texto: {hoja.title}!{procedimiento.coordenada}")
    valores = {campo: _texto(celda.valor) for campo, celda in por_campo.items() if not campo.startswith("fase_")}
    fases = {}
    for fase in FASES:
        campo = f"fase_{fase}"
        celda = por_campo.get(campo)
        texto = _texto(celda.valor) if celda else ""
        fases[fase] = FaseCurricular(
            presente_en_hoja=campo in campos_presentes,
            texto=texto,
            codigos_locales=tuple(dict.fromkeys(_CODIGO_LOCAL.findall(texto))),
            coordenadas=(celda.coordenada,) if celda else (),
        )
    unidad = cfg.unidad if cfg.unidad is not None else valores.pop("unidad", "")
    valores.pop("unidad", None)
    if not unidad:
        raise ValueError(f"Tema sin unidad explícita en {hoja.title}!{registro.fila}.")
    identidad = _sha(
        _json(
            {
                "fuente": fuente.id,
                "version": fuente.version,
                "hoja": hoja.title,
                "fila": registro.fila,
                "contenido": registro.model_dump(mode="json"),
            }
        )
    )
    return FilaCurricular(
        id=f"CURRICULO:{fuente.version}:{identidad}",
        fuente_id=fuente.id,
        version=fuente.version,
        hoja=hoja.title,
        fila=registro.fila,
        grado=cfg.grado,
        asignatura_malla=cfg.asignatura_malla,
        unidad=unidad,
        fases=fases,
        celdas=registro.celdas,
        **valores,
    )


def _leer_hoja(
    hoja: Worksheet,
    cfg: ConfigHoja,
    fuente: FuenteCurricular,
    alias: dict[str, str],
) -> tuple[HojaCurricular, list[FilaCurricular]]:
    columnas = _columnas(hoja, cfg, alias)
    registros = _registros(hoja, columnas)
    campos = {c.campo for c in columnas if c.campo is not None}
    filas = []
    avisos = []
    if cfg.tratamiento == "curricular":
        for registro in registros:
            if fila := _fila_curricular(registro, hoja, cfg, fuente, campos):
                filas.append(fila)
        if not filas:
            raise ValueError(f"Hoja declarada curricular sin procedimientos: {hoja.title}")
        for fase in FASES:
            if f"fase_{fase}" not in campos:
                avisos.append(f"Fase {fase} sin columna en el insumo; no se completa por inferencia.")
        extras = [
            c.columna
            for c in columnas
            if c.campo is None
            and any(
                r.fila > max(cfg.filas_encabezado) and r.celdas[c.indice - 1].valor_original is not None
                for r in registros
            )
        ]
        if extras:
            avisos.append(f"Columnas sin rol lógico conservadas por coordenada: {', '.join(extras)}.")
    modelo = HojaCurricular(
        fuente_id=fuente.id,
        nombre=hoja.title,
        tratamiento=cfg.tratamiento,
        motivo=cfg.motivo,
        grado=cfg.grado,
        asignatura_malla=cfg.asignatura_malla,
        unidad=cfg.unidad,
        filas_encabezado=cfg.filas_encabezado,
        columnas=columnas,
        combinaciones=tuple(sorted(str(r) for r in hoja.merged_cells.ranges)),
        registros=registros,
        filas_curriculares=tuple(fila.id for fila in filas),
        avisos=tuple(avisos),
    )
    return modelo, filas


def _procedencia(contenido: bytes, cfg: ConfigCurriculo) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    try:
        documento = json.loads(contenido)
        archivos = documento["archivos"]
        if documento["version"] != cfg.version or not isinstance(archivos, list):
            raise ValueError("La versión de procedencia no coincide con la configuración.")
        por_nombre = {}
        for archivo in archivos:
            nombre = archivo["archivo"]
            if nombre in por_nombre:
                raise ValueError(f"Archivo duplicado en procedencia: {nombre}")
            if not re.fullmatch(r"[a-f0-9]{64}", archivo["sha256"]):
                raise ValueError(f"SHA256 inválido en procedencia: {nombre}")
            if type(archivo["bytes"]) is not int or archivo["bytes"] <= 0:
                raise ValueError(f"Tamaño inválido en procedencia: {nombre}")
            por_nombre[nombre] = archivo
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Documento de procedencia inválido.") from exc
    if set(por_nombre) != {f.archivo for f in cfg.fuentes}:
        raise ValueError("El conjunto de archivos de procedencia y configuración no coincide.")
    return {k: v for k, v in documento.items() if k != "archivos"}, por_nombre


def cargar_curriculo(config_path: str | Path | None = None, *, nombre: str = "mallas_v4") -> Curriculo:
    """Verify and load every configured workbook and sheet without modifying sources."""

    # Implementation notes:
    # Las rutas de fuentes se resuelven bajo ``raiz`` de la configuración, relativa al
    # propio YAML. El SHA y el tamaño se cotejan antes de parsear exactamente esos bytes.
    # Cualquier archivo u hoja sin decisión explícita interrumpe la ingesta.
    ruta_config = Path(config_path) if config_path is not None else config.ruta(f"config/{nombre}.yaml")
    ruta_config = ruta_config.resolve()
    bytes_config = ruta_config.read_bytes()
    cfg = ConfigCurriculo.model_validate(yaml.safe_load(bytes_config))
    base = (ruta_config.parent / cfg.raiz).resolve()
    directorio = _dentro(base, cfg.directorio)
    bytes_procedencia = _dentro(base, cfg.procedencia).read_bytes()
    origen, procedencia = _procedencia(bytes_procedencia, cfg)
    with os.scandir(directorio) as entradas:
        archivos = {e.name for e in entradas if e.name.lower().endswith(".xlsx")}
    if archivos != set(procedencia):
        raise ValueError("El directorio contiene un conjunto de libros distinto al declarado en procedencia.")
    alias = _catalogo_alias(cfg)
    fuentes, hojas, filas = [], [], []
    for configurada in cfg.fuentes:
        LOGGER.info("Ingesta %s: verificando %s", cfg.version, configurada.archivo)
        archivo = _dentro(directorio, configurada.archivo)
        contenido = archivo.read_bytes()
        ficha = procedencia[configurada.archivo]
        digest = _sha(contenido)
        if digest != ficha["sha256"] or len(contenido) != ficha["bytes"]:
            LOGGER.error("Integridad rechazada: %s", configurada.archivo)
            raise ValueError(f"SHA256 o tamaño no coincide con procedencia: {configurada.archivo}")
        fuente = FuenteCurricular(
            id=f"FUENTE:{cfg.version}:{_sha(_json([cfg.version, configurada.archivo, digest]))}",
            version=cfg.version,
            archivo=configurada.archivo,
            archivo_local=str(archivo.relative_to(base)),
            sha256=digest,
            bytes=len(contenido),
            procedencia={**origen, "archivo": ficha},
        )
        try:
            libro = load_workbook(BytesIO(contenido), data_only=False, read_only=False, keep_links=False)
        except Exception as exc:
            raise ValueError(f"Libro Excel ilegible: {configurada.archivo}") from exc
        try:
            declaradas = {h.nombre: h for h in configurada.hojas}
            if set(libro.sheetnames) != set(declaradas):
                raise ValueError(f"Hojas sin decisión explícita o inexistentes en {configurada.archivo}.")
            for hoja in libro:
                modelo, temas = _leer_hoja(hoja, declaradas[hoja.title], fuente, alias)
                for aviso in modelo.avisos:
                    LOGGER.warning("%s / %s: %s", configurada.archivo, hoja.title, aviso)
                hojas.append(modelo)
                filas.extend(temas)
        finally:
            libro.close()
        fuentes.append(fuente)
    agrupados = defaultdict(list)
    for fila in filas:
        agrupados[fila.grado, fila.asignatura_malla, fila.unidad].append(fila.id)
    agrupaciones = tuple(
        AgrupacionCurricular(grado=g, asignatura_malla=a, unidad=u, filas=tuple(ids))
        for (g, a, u), ids in sorted(agrupados.items())
    )
    huella = _sha(
        _json(
            {
                "configuracion_sha256": _sha(bytes_config),
                "procedencia_sha256": _sha(bytes_procedencia),
                "fuentes": [f.id for f in fuentes],
                "filas": [f.id for f in filas],
                "formato": 1,
            }
        )
    )
    resultado = Curriculo(
        version=cfg.version,
        huella=huella,
        fuentes=tuple(fuentes),
        filas=tuple(filas),
        hojas=tuple(hojas),
        agrupaciones=agrupaciones,
        avisos=("Los códigos, alineaciones y afirmaciones del insumo son declarados; requieren verificación externa.",),
    )
    LOGGER.info(
        "Ingesta %s completada: %d fuentes, %d hojas, %d temas; sin activar configuración.",
        cfg.version,
        len(fuentes),
        len(hojas),
        len(filas),
    )
    return resultado
