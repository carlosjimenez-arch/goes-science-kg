"""Copias anotadas de las matrices de la V3 con los refuerzos revisados (spec 15, etapa 15.6).

Nunca se toca el original ni el texto oficial: se escribe una copia en `entregas/mallas_v3_reforzadas/` con
- tres columnas nuevas al final de cada hoja de unidad: «Refuerzo GOES», «Evidencia GOES» y «Solidez GOES»;
- una nota en la fila del tema por cada refuerzo `nota_en_tema`;
- una fila nueva después del tema ancla por cada `fila_nueva`, con «Origen del contenido» = «Propuesta · GOES · …»
  (la misma convención del MINED para sus propuestas) y relleno verde claro;
- una hoja «Refuerzo GOES» con el índice de los refuerzos del libro.
Compatible con Numbers: sin «·» en los nombres de pestaña y sin fórmulas nuevas.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from goes_science_kg.config import cargar, ruta
from goes_science_kg.ingesta.mallas import _columnas, extraer

SALIDA = "entregas/mallas_v3_reforzadas"
COLUMNAS_GOES = ("Refuerzo GOES", "Evidencia GOES", "Solidez GOES")
VERDE_ENCABEZADO = "1E4D3A"   # estilo de los libros 2027 del proyecto
VERDE_PROPUESTA = "E2EFDA"
ORIGEN = {"error_malla": "Corrección de un error", "panel_hallazgo": "Panel de expertos",
          "panel_propuesta": "Propuesta curricular", "contraste": "Contraste internacional",
          "quitado_v3": "Recuperar lo que la V3 quitó"}
# Columna de la malla → campo de la fila propuesta
CAMPO_DE_COLUMNA = {"Contenido": "contenido", "Subcontenidos": "subcontenido", "Procedimental": "procedimental",
                    "Indicador avanzado": "indicador", "Indicadores de logro": "indicador",
                    "Evidencia de aprendizaje": "evidencia_aprendizaje"}


def _es_hoja_de_unidad(titulo: str) -> bool:
    return titulo.startswith("Unidad") or bool(re.search(r"\d+\.°", titulo))


def _evidencia_corta(r: dict) -> str:
    partes = []
    for e in (r.get("evidencia") or [])[:3]:
        cita = e.get("cita_literal") or e.get("cita") or ""
        partes.append(f"{e.get('fuente', '')}: «{cita[:140]}»" if cita else str(e.get("fuente", "")))
    return " | ".join(partes)


def _solidez(r: dict) -> str:
    return {"solido": "sólido", "respaldado": "respaldado", "revisada": "revisada", "calculada": "calculada"}.get(
        r.get("solidez", ""), r.get("solidez", "")) + f" · confianza {r.get('confianza', '')}"


def _insertar_fila(ws, despues_de: int) -> int:
    """Inserta una fila después de `despues_de` y corre (o amplía) las celdas combinadas. Devuelve la fila nueva."""
    nueva = despues_de + 1
    rangos = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in rangos:
        ws.unmerge_cells(m.coord)
    ws.insert_rows(nueva)
    for m in rangos:
        r0, r1 = m.min_row, m.max_row
        if r0 >= nueva:
            r0, r1 = r0 + 1, r1 + 1
        elif r0 < nueva <= r1:   # la fila nueva queda dentro del bloque: se amplía
            r1 += 1
        ws.merge_cells(start_row=r0, start_column=m.min_col, end_row=r1, end_column=m.max_col)
    altura = ws.row_dimensions[despues_de].height
    if altura:
        ws.row_dimensions[nueva].height = altura
    for col in range(1, ws.max_column + 1):
        origen, destino = ws.cell(despues_de, col), ws.cell(nueva, col)
        if origen.has_style:
            destino._style = copy.copy(origen._style)
    return nueva


def _fila_para_insertar(ws, ancla: int, columnas: set[int]) -> int:
    """La fila después de la cual va la propuesta: el ancla o, si el ancla está en un bloque combinado de una columna
    que se escribe (Contenido, Procedimental…), el final de ese bloque, para no partirlo."""
    n, cambio = ancla, True
    while cambio:
        cambio = False
        for m in ws.merged_cells.ranges:
            if m.min_row <= n < m.max_row and m.min_row >= 3 and any(m.min_col <= c <= m.max_col for c in columnas):
                n, cambio = m.max_row, True
    return n


def _encabezados(ws) -> dict[str, int]:
    """Nombre de columna (filas 1 y 2) → número de columna."""
    return {nombre: i for i, nombre in enumerate(_columnas(ws), start=1) if nombre}


def _agregar_columnas(ws) -> dict[str, int]:
    """Las tres columnas GOES al final (encabezado combinado en las filas 1 y 2)."""
    base = ws.max_column
    cols = {}
    for k, nombre in enumerate(COLUMNAS_GOES, start=1):
        c = base + k
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
        celda = ws.cell(1, c, nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE_ENCABEZADO)
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 48 if k == 1 else 36 if k == 2 else 18
        cols[nombre] = c
    return cols


def _escribir_goes(ws, fila: int, cols: dict[str, int], refuerzos: list[dict]) -> None:
    textos = {"Refuerzo GOES": " | ".join(f"{r['id']}: {r['refuerzo']}" for r in refuerzos),
              "Evidencia GOES": " | ".join(_evidencia_corta(r) for r in refuerzos),
              "Solidez GOES": " | ".join(_solidez(r) for r in refuerzos)}
    for nombre, texto in textos.items():
        celda = ws.cell(fila, cols[nombre], texto)
        celda.alignment = Alignment(wrap_text=True, vertical="top")


def _poner(ws, fila: int, col: int, valor: object) -> None:
    """Escribe salvo que la celda sea parte de un bloque combinado (su valor es el del bloque)."""
    celda = ws.cell(fila, col)
    if not isinstance(celda, MergedCell):
        celda.value = valor


def _escribir_fila_propuesta(ws, fila: int, encabezados: dict[str, int], ancla: int, r: dict) -> None:
    f = r.get("fila") or {}
    verde = PatternFill("solid", fgColor=VERDE_PROPUESTA)
    for col in range(1, ws.max_column + 1):
        ws.cell(fila, col).fill = verde
    for nombre in ("Unidad",):
        if nombre in encabezados:
            _poner(ws, fila, encabezados[nombre], ws.cell(ancla, encabezados[nombre]).value)
    for nombre, campo in CAMPO_DE_COLUMNA.items():
        if nombre in encabezados and f.get(campo):
            _poner(ws, fila, encabezados[nombre], f[campo])
    for fase, texto in (f.get("microhabilidades") or {}).items():
        if fase in encabezados:
            _poner(ws, fila, encabezados[fase], texto)
    if "Origen del contenido" in encabezados:
        _poner(ws, fila, encabezados["Origen del contenido"],
               f"Propuesta · GOES · {ORIGEN.get(r['origen'], r['origen'])}")


def escribir(refuerzos: list[dict], destino: str = SALIDA) -> list[Path]:
    """Una copia anotada por libro de la malla vigente con los refuerzos que se inyectan. Devuelve las rutas."""
    cfg = cargar("mallas")
    temas = {(t.ruta, t.hoja, t.fila): t.id for t in extraer()}
    fila_de = {i: k for k, i in temas.items()}
    activos = [r for r in refuerzos if r.get("decision") == "inyectar" and r.get("tema_ancla") in fila_de]
    d = ruta(destino)
    d.mkdir(parents=True, exist_ok=True)
    escritos = []
    for item in cfg["archivos"]:
        archivo = item["archivo"]
        del_libro = [r for r in activos if fila_de[r["tema_ancla"]][0] == archivo]
        wb = openpyxl.load_workbook(ruta(cfg["directorio"]) / archivo)
        for ws in wb.worksheets:
            if not _es_hoja_de_unidad(ws.title):
                continue
            en_hoja = [r for r in del_libro if fila_de[r["tema_ancla"]][1] == ws.title]
            encabezados = _encabezados(ws)
            cols = _agregar_columnas(ws)
            notas: dict[int, list[dict]] = {}
            filas_nuevas: dict[int, list[dict]] = {}
            for r in en_hoja:
                n = fila_de[r["tema_ancla"]][2]
                (filas_nuevas if r.get("forma") == "fila_nueva" else notas).setdefault(n, []).append(r)
            for n, rs in notas.items():
                _escribir_goes(ws, n, cols, rs)
            # Solo los bloques de «Contenido» y «Subcontenidos» deciden dónde va la fila; las columnas que la V3
            # combina en toda la unidad (competencia, notación…) se amplían y no se escriben.
            escritas = {c for nombre, c in encabezados.items() if nombre in ("Contenido", "Subcontenidos")}
            destino_de = {n: _fila_para_insertar(ws, n, escritas) for n in filas_nuevas}
            for n in sorted(filas_nuevas, key=lambda n: destino_de[n], reverse=True):   # de abajo hacia arriba
                for r in reversed(filas_nuevas[n]):
                    nueva = _insertar_fila(ws, destino_de[n])
                    _escribir_fila_propuesta(ws, nueva, encabezados, n, r)
                    _escribir_goes(ws, nueva, cols, [r])
        _hoja_indice(wb, del_libro, fila_de)
        p = d / archivo.replace(".xlsx", " - Refuerzo GOES.xlsx")
        wb.save(p)
        escritos.append(p)
    return escritos


def _hoja_indice(wb, refuerzos: list[dict], fila_de: dict) -> None:
    ws = wb.create_sheet("Refuerzo GOES", 0)
    encabezado = ["Id", "Grado", "Hoja", "Tema ancla", "Forma", "Origen", "Tipo", "Refuerzo", "Evidencia", "Solidez"]
    ws.append(encabezado)
    for c in range(1, len(encabezado) + 1):
        celda = ws.cell(1, c)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE_ENCABEZADO)
    for r in sorted(refuerzos, key=lambda r: (fila_de[r["tema_ancla"]][1], fila_de[r["tema_ancla"]][2], r["id"])):
        ws.append([r["id"], ", ".join(f"{g}.°" for g in r.get("grados_destino") or r.get("grados") or []),
                   fila_de[r["tema_ancla"]][1], r["tema_ancla"],
                   "fila nueva" if r.get("forma") == "fila_nueva" else "nota en el tema",
                   ORIGEN.get(r["origen"], r["origen"]), r["tipo"], r["refuerzo"], _evidencia_corta(r), _solidez(r)])
    for c, ancho in zip("ABCDEFGHIJ", (10, 8, 26, 22, 14, 24, 16, 70, 60, 22), strict=True):
        ws.column_dimensions[c].width = ancho
    nota = ("Propuestas del equipo GOES con revisión del panel de expertos IA (spec 14 y 15). No son "
            "confirmaciones del MINED ni modifican el texto oficial.")
    ws.cell(ws.max_row + 2, 1, nota)
