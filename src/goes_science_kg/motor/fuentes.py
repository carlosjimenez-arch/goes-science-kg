"""Evidencia desde snapshots verificados de PDF, Excel y unidades DOCX/HTML por código.

Las raíces proceden del catálogo configurado, nunca del payload de un Documento. Cada archivo se
lee y verifica una vez por instancia; si cambia su identidad en disco, se rechaza la caché completa.
"""

from __future__ import annotations

import copy
import hashlib
import logging
import os
import re
import threading
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from goes_science_kg.config import ruta
from goes_science_kg.modelos import Nodo, TipoNodo
from goes_science_kg.motor.contratos import Evidencia
from goes_science_kg.recuperacion import Documento

LOGGER = logging.getLogger(__name__)
MAX_BYTES_FUENTE = 64 * 1024 * 1024
MAX_BYTES_XLSX_DESCOMPRIMIDOS = 128 * 1024 * 1024
MAX_CELDAS = 64
COORDENADA = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")


@dataclass(frozen=True)
class _Snapshot:
    archivo: Path
    contenido: bytes
    sha256: str
    firma: tuple[int, ...]


def _firma(estado: os.stat_result) -> tuple[int, ...]:
    return (estado.st_dev, estado.st_ino, estado.st_size, estado.st_mtime_ns, estado.st_ctime_ns)


def _entero_positivo(valor: Any) -> int:
    if isinstance(valor, bool) or not re.fullmatch(r"[0-9]+", str(valor)) or int(valor) < 1:
        raise ValueError("El localizador requiere un entero positivo")
    return int(valor)


class FuentesVerificadas:
    """Lectura acotada con caché de bytes, libros y localizadores; no ejecuta fórmulas."""

    def __init__(
        self,
        catalogo: dict[str, dict],
        *,
        raiz_fuentes: Path,
        raices_adicionales: Mapping[str, Path] | None = None,
        bloqueos_curriculares: Mapping[str, str] | None = None,
    ):
        self.catalogo = copy.deepcopy(catalogo)
        self.raiz = raiz_fuentes.resolve()
        if "principal" in (raices_adicionales or {}):
            raise ValueError("Una raíz adicional no puede sustituir la raíz principal")
        self._raices = {"principal": self.raiz}
        self._raices.update({nombre: Path(p).resolve() for nombre, p in (raices_adicionales or {}).items()})
        self._snapshots: dict[tuple[Path, str], _Snapshot] = {}
        self._libros: dict[tuple[Path, str], Any] = {}
        self._pdfs: dict[tuple[Path, str], Any] = {}
        self._textos: dict[tuple, str] = {}
        self._estructurados: dict[tuple, Any] = {}
        self._lock = threading.RLock()
        self._lecturas = 0
        self._bytes_verificados = 0
        self.bloqueos_curriculares = dict(bloqueos_curriculares or {})
        self.bloqueos_derivados: dict[str, str] = {}

    @property
    def bloqueos_efectivos(self) -> dict[str, str]:
        """Une vetos para generación sin confundir su origen en el armado de contexto."""
        return self.bloqueos_curriculares | self.bloqueos_derivados

    def motivos_bloqueo(self, evidencias) -> list[str]:
        ids = {e.id.partition("#CELDA:")[0] for e in evidencias}
        bloqueos = self.bloqueos_efectivos
        return [bloqueos[i] for i in sorted(ids.intersection(bloqueos))]

    @classmethod
    def del_grafo(cls, nodos: Iterable[Nodo]) -> FuentesVerificadas:
        catalogo = {n.id.removeprefix("DOC:"): n.props for n in nodos if n.tipo == TipoNodo.DOCUMENTO}
        return cls(catalogo, raiz_fuentes=ruta("data/fuentes"))

    def _registro(self, id_documento: str) -> dict:
        return self.catalogo[id_documento.removeprefix("DOC:")]

    def _localizar(self, registro: dict) -> tuple[Path, Path]:
        raiz = self._raices[registro.get("raiz_fuente", "principal")]
        relativo = Path(registro["archivo_local"])
        opciones = (
            [relativo] if relativo.is_absolute()
            else [raiz / relativo, raiz / "externos" / relativo, ruta(relativo)]
        )
        for candidato in opciones:
            archivo = candidato.resolve()
            # No consultar siquiera is_file() en candidatos fuera de la raíz de este registro.
            if archivo.is_relative_to(raiz) and archivo.is_file():
                return archivo, raiz
        raise ValueError("La fuente no existe dentro de su raíz autorizada")

    @staticmethod
    def _abrir_seguro(archivo: Path, raiz: Path) -> int:
        """Abre cada componente por descriptor, sin seguir enlaces sustituidos tras resolve()."""
        descriptor = os.open(raiz, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
        try:
            partes = archivo.relative_to(raiz).parts
            for directorio in partes[:-1]:
                siguiente = os.open(
                    directorio, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=descriptor,
                )
                os.close(descriptor)
                descriptor = siguiente
            return os.open(partes[-1], os.O_RDONLY | os.O_NOFOLLOW, dir_fd=descriptor)
        finally:
            os.close(descriptor)

    def _snapshot(self, id_documento: str) -> _Snapshot:
        registro = self._registro(id_documento)
        archivo, raiz = self._localizar(registro)
        esperada = registro.get("sha256")
        if not isinstance(esperada, str) or not re.fullmatch(r"[0-9a-f]{64}", esperada):
            raise ValueError("El catálogo requiere un SHA-256 explícito y válido")
        clave = (archivo, esperada)
        with self._lock:
            anterior = self._snapshots.get(clave)
            if anterior is not None:
                if _firma(archivo.stat()) != anterior.firma:
                    raise ValueError("La fuente cambió después de verificar el snapshot; se rechaza la caché")
                return anterior
            descriptor = self._abrir_seguro(archivo, raiz)
            with os.fdopen(descriptor, "rb") as entrada:
                inicial = os.fstat(entrada.fileno())
                if inicial.st_size > MAX_BYTES_FUENTE:
                    raise ValueError("La fuente supera el límite de lectura")
                contenido = entrada.read(MAX_BYTES_FUENTE + 1)
                final = os.fstat(entrada.fileno())
            if len(contenido) > MAX_BYTES_FUENTE or _firma(inicial) != _firma(final):
                raise ValueError("La fuente excede el límite o cambió durante la lectura")
            if hashlib.sha256(contenido).hexdigest() != esperada:
                raise ValueError("La fuente no coincide con su SHA-256 registrado")
            if registro.get("bytes") is not None and len(contenido) != registro["bytes"]:
                raise ValueError("El tamaño de la fuente no coincide con el catálogo")
            snapshot = _Snapshot(archivo, contenido, esperada, _firma(final))
            self._snapshots[clave] = snapshot
            self._lecturas += 1
            self._bytes_verificados += len(contenido)
            return snapshot

    def _archivo(self, id_documento: str) -> tuple[Path, bytes]:
        """Compatibilidad: siempre devuelve bytes de un snapshot validado."""
        snapshot = self._snapshot(id_documento)
        return snapshot.archivo, snapshot.contenido

    def _libro(self, snapshot: _Snapshot):
        from openpyxl import load_workbook

        clave = (snapshot.archivo, snapshot.sha256)
        if clave not in self._libros:
            with ZipFile(BytesIO(snapshot.contenido)) as paquete:
                if sum(item.file_size for item in paquete.infolist()) > MAX_BYTES_XLSX_DESCOMPRIMIDOS:
                    raise ValueError("El XLSX descomprimido supera el presupuesto")
            self._libros[clave] = load_workbook(
                BytesIO(snapshot.contenido), read_only=False, data_only=False, keep_links=False,
            )
        return self._libros[clave]

    @staticmethod
    def _seleccionar_celdas(fuente: dict, hoja) -> list[str]:
        from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

        if fuente.get("celdas") is not None and fuente.get("rango") is not None:
            raise ValueError("Seleccione celdas o rango, sin localizadores ambiguos")
        if fuente.get("celdas") is not None:
            solicitadas = fuente["celdas"]
            if not isinstance(solicitadas, (list, tuple)) or not 1 <= len(solicitadas) <= MAX_CELDAS:
                raise ValueError("Seleccione entre una y 64 celdas explícitas")
            coordenadas = [str(c).upper().replace("$", "") for c in solicitadas]
        elif fuente.get("rango") is not None:
            rango = fuente["rango"]
            if not isinstance(rango, str) or "!" in rango or ":" not in rango:
                raise ValueError("El rango debe usar coordenadas explícitas de una sola hoja")
            min_col, min_fila, max_col, max_fila = range_boundaries(rango.upper())
            if None in (min_col, min_fila, max_col, max_fila):
                raise ValueError("El rango debe incluir filas y columnas")
            area = (max_col - min_col + 1) * (max_fila - min_fila + 1)
            if min_col > max_col or min_fila > max_fila or area > MAX_CELDAS:
                raise ValueError("El rango está invertido o supera 64 celdas")
            coordenadas = [f"{get_column_letter(c)}{f}" for f in range(min_fila, max_fila + 1)
                           for c in range(min_col, max_col + 1)]
        else:
            fila = _entero_positivo(fuente.get("fila"))
            if not 1 <= hoja.max_column <= MAX_CELDAS:
                raise ValueError("Una fila completa supera el máximo de 64 celdas; declare celdas explícitas")
            coordenadas = [f"{get_column_letter(c)}{fila}" for c in range(1, hoja.max_column + 1)]
        for coordenada in coordenadas:
            if not COORDENADA.fullmatch(coordenada):
                raise ValueError("Coordenada de celda inválida")
            fila, columna = coordinate_to_tuple(coordenada)
            if fila > 1048576 or columna > 16384 or fila > hoja.max_row or columna > hoja.max_column:
                raise ValueError("Celda fuera del rango existente de la hoja")
        return list(dict.fromkeys(coordenadas))

    def _celdas(self, snapshot: _Snapshot, fuente: dict) -> list[dict]:
        libro = self._libro(snapshot)
        hoja = libro[fuente["hoja"]]
        celdas = []
        for coordenada in self._seleccionar_celdas(fuente, hoja):
            original = hoja[coordenada]
            rangos = [r for r in hoja.merged_cells.ranges
                      if r.min_row <= original.row <= r.max_row and r.min_col <= original.column <= r.max_col]
            if len(rangos) > 1:
                raise ValueError("La celda pertenece a combinaciones superpuestas y ambiguas")
            rango = rangos[0] if rangos else None
            origen = hoja.cell(rango.min_row, rango.min_col) if rango is not None else original
            celdas.append({
                "coordenada_solicitada": coordenada,
                "coordenada_origen": origen.coordinate,
                "valor_original": original.value,
                "valor_real": origen.value,
                "tipo_excel": origen.data_type,
                "rango_combinado": str(rango) if rango is not None else None,
                "valor_es_formula": origen.data_type == "f",
            })
        return celdas

    @staticmethod
    def _texto_valor(valor: Any) -> str:
        return "" if valor is None else str(valor)

    def _fuente_verificada(self, fuente: dict, snapshot: _Snapshot) -> dict:
        # No devolver texto ni valores arbitrarios del payload como si fueran procedencia verificada.
        campos = (
            ("documento", "localizador") if snapshot.archivo.suffix.lower() in {".docx", ".html", ".htm"}
            else ("documento", "pagina", "hoja", "fila", "celdas", "rango")
        )
        registro = self._registro(fuente["documento"])
        return {k: copy.deepcopy(fuente[k]) for k in campos if k in fuente} | {
            "sha256": snapshot.sha256,
            "verificacion": "contenido_fuente",
            "raiz_fuente": registro.get("raiz_fuente", "principal"),
            "tipo": "documental",
        }

    def _validar_version(self, documento: Documento) -> None:
        esperada = self._registro(documento.fuente["documento"]).get("version_curriculo")
        if esperada is not None and documento.version != esperada:
            raise ValueError("La versión solicitada no coincide con el catálogo configurado")

    def hidratar_celdas(self, documento: Documento) -> list[Evidencia]:
        """Una evidencia por celda no vacía; una combinación conserva el ancla original."""
        try:
            with self._lock:
                self._validar_version(documento)
                snapshot = self._snapshot(documento.fuente["documento"])
                if snapshot.archivo.suffix.lower() != ".xlsx":
                    raise ValueError("La hidratación por celda requiere un libro XLSX")
                salida = []
                from openpyxl.utils.cell import coordinate_to_tuple

                for celda in self._celdas(snapshot, documento.fuente):
                    texto = self._texto_valor(celda["valor_real"])
                    if not texto.strip():
                        continue
                    coordenada = celda["coordenada_solicitada"]
                    fuente = self._fuente_verificada(documento.fuente, snapshot)
                    fuente.pop("rango", None)
                    fuente.pop("pagina", None)
                    fuente["celdas"] = [coordenada]
                    fuente["fila"] = coordinate_to_tuple(coordenada)[0]
                    fuente["celdas_verificadas"] = [celda]
                    salida.append(Evidencia.crear(
                        f"{documento.id}#CELDA:{coordenada}", texto, fuente, documento.version,
                    ))
                return salida
        except Exception as error:  # Fallo cerrado también ante errores de parser de un archivo no confiable.
            LOGGER.warning("AMARILLO celdas %s omitidas: %s", documento.id, type(error).__name__)
            return []

    def hidratar(self, documento: Documento) -> Evidencia | None:
        """Hidrata localizadores físicos; nunca emplea Documento.texto como respaldo."""
        try:
            with self._lock:
                fuente = documento.fuente
                self._validar_version(documento)
                snapshot = self._snapshot(fuente["documento"])
                if snapshot.archivo.suffix.lower() in {".docx", ".html", ".htm"}:
                    from goes_science_kg.ingesta.estructurados import huella_implementacion, indexar_documento

                    formato = snapshot.archivo.suffix.lower()[1:]
                    clave = (snapshot.archivo, snapshot.sha256, huella_implementacion())
                    if clave not in self._estructurados:
                        self._estructurados[clave] = indexar_documento(snapshot.contenido, formato)
                    unidad = self._estructurados[clave].extraer(fuente.get("localizador"))
                    texto = unidad.texto
                    verificada = self._fuente_verificada(fuente, snapshot) | {
                        "localizador": unidad.localizador, "ubicacion": unidad.ubicacion,
                        "extractor": unidad.metodo, "huella_implementacion": unidad.huella_implementacion,
                        "limites_visibilidad": unidad.limites,
                    }
                elif snapshot.archivo.suffix.lower() == ".xlsx" and fuente.get("hoja"):
                    from openpyxl.utils.cell import coordinate_to_tuple

                    celdas = self._celdas(snapshot, fuente)
                    texto = "\n".join(self._texto_valor(c["valor_real"]) for c in celdas if c["valor_real"] is not None)
                    verificada = self._fuente_verificada(fuente, snapshot) | {"celdas_verificadas": celdas}
                    verificada["celdas"] = [c["coordenada_solicitada"] for c in celdas]
                    verificada.pop("pagina", None)
                    filas = {coordinate_to_tuple(c)[0] for c in verificada["celdas"]}
                    verificada.pop("fila", None)
                    if len(filas) == 1:
                        verificada["fila"] = filas.pop()
                elif snapshot.archivo.suffix.lower() == ".pdf" and fuente.get("pagina"):
                    from pypdf import PdfReader

                    pagina = _entero_positivo(fuente["pagina"])
                    archivo_clave = (snapshot.archivo, snapshot.sha256)
                    if archivo_clave not in self._pdfs:
                        self._pdfs[archivo_clave] = PdfReader(BytesIO(snapshot.contenido))
                    lector = self._pdfs[archivo_clave]
                    if not 1 <= pagina <= len(lector.pages):
                        raise ValueError("Página fuera de rango")
                    clave = (*archivo_clave, pagina)
                    if clave not in self._textos:
                        self._textos[clave] = lector.pages[pagina - 1].extract_text() or ""
                    texto = self._textos[clave]
                    verificada = self._fuente_verificada(fuente, snapshot)
                    for campo in ("hoja", "fila", "celdas", "rango"):
                        verificada.pop(campo, None)
                    verificada["pagina"] = pagina
                else:
                    raise ValueError("La fuente requiere un adaptador verificable para su formato/localizador")
                return Evidencia.crear(documento.id, texto, verificada, documento.version) if texto.strip() else None
        except Exception as error:  # Un parser o localizador inválido no autoriza texto de metadatos.
            LOGGER.warning("AMARILLO evidencia %s omitida: %s", documento.id, type(error).__name__)
            return None

    def estadisticas(self) -> dict[str, int]:
        """Contadores locales para comprobar que no se relee un XLSX por cada celda."""
        with self._lock:
            return {
                "archivos_leidos": self._lecturas,
                "bytes_verificados": self._bytes_verificados,
                "libros_parseados": len(self._libros),
                "pdfs_parseados": len(self._pdfs),
            }

    def estadisticas_estructuradas(self) -> dict[str, int]:
        """Contador separado para mantener el contrato histórico de estadisticas()."""
        with self._lock:
            return {"indices_parseados": len(self._estructurados)}
