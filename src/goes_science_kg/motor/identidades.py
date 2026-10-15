"""Puente exacto análisis ↔ V3: valida fuente y columnas; nunca infiere equivalencia por similitud."""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

import openpyxl

from goes_science_kg.config import cargar, relativa, ruta
from goes_science_kg.ingesta.mallas import _columnas, extraer
from goes_science_kg.motor.archivos import destino_seguro, escribir, firma, leer, sha

COPIAR = ("clasificaciones_consolidadas.json", "catalogo_timss.json", "catalogo_pisa.json", "catalogo_acara.json",
          "fuentes_marcos.json", "fuentes_mallas.json")


def limpio(valor) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def columnas_vigentes() -> dict[tuple[str, str, int], tuple[dict, int]]:
    """Lee el contenido completo del bloque de cada procedimental, incluidas celdas combinadas y filas asociadas."""
    cfg = cargar("mallas")
    salida = {}
    for item in cfg["archivos"]:
        archivo = item["archivo"]
        wb = openpyxl.load_workbook(ruta(cfg["directorio"]) / archivo, data_only=True)
        try:
            for ws in wb:
                if not (ws.title.startswith("Unidad") or re.search(r"\d+\.°", ws.title)):
                    continue
                columnas = _columnas(ws)
                nombres = [c for c in columnas if c]
                if len(nombres) != len(set(nombres)):
                    raise ValueError(f"encabezados ambiguos en {archivo}/{ws.title}")
                pc = columnas.index("Procedimental") + 1
                combinadas = {(r, c): (m.min_row, m.min_col, m.max_row) for m in ws.merged_cells.ranges
                              for r in range(m.min_row, m.max_row + 1)
                              for c in range(m.min_col, m.max_col + 1)}
                for r in range(cfg["fila_datos"], ws.max_row + 1):
                    if not ws.cell(r, pc).value:
                        continue
                    fin = combinadas.get((r, pc), (r, pc, r))[2]
                    valores = {}
                    for c, nombre in enumerate(columnas, 1):
                        if not nombre:
                            continue
                        partes = []
                        for rr in range(r, fin + 1):
                            x, y, _ = combinadas.get((rr, c), (rr, c, rr))
                            partes.append(limpio(ws.cell(x, y).value))
                        valores[nombre] = " | ".join(dict.fromkeys(v for v in partes if v))
                    salida[(archivo, ws.title, r)] = (valores, fin)
        finally:
            wb.close()
    return salida


def comprobar_fuentes(fuentes: list[dict]) -> None:
    cfg = cargar("mallas")
    esperadas = {x["archivo"] for x in cfg["archivos"]}
    if len(fuentes) != len(esperadas) or {f["archivo"] for f in fuentes} != esperadas:
        raise ValueError("el análisis no describe exactamente los archivos de la V3")
    for f in fuentes:
        if sha(ruta(cfg["directorio"]) / f["archivo"]) != f["sha256"]:
            raise ValueError(f"SHA de la fuente no coincide: {f['archivo']}")


def comprobar_identidades(identidades: list[dict], fuentes: list[dict], columnas: dict | None = None) -> dict:
    comprobar_fuentes(fuentes)
    columnas = columnas if columnas is not None else columnas_vigentes()
    por_fila = {(t.ruta, t.hoja, t.fila): t for t in extraer()}
    if len(identidades) != len(por_fila):
        raise ValueError("el puente no cubre todos los temas vigentes")
    aliases = [t["id"] for t in identidades]
    ubicaciones = [(t["archivo"], t["hoja"], t["fila"]) for t in identidades]
    if len(set(aliases)) != len(aliases) or len(set(ubicaciones)) != len(ubicaciones):
        raise ValueError("alias o ubicación repetida en el puente")
    salida = {}
    for d, lugar in zip(identidades, ubicaciones, strict=True):
        t = por_fila.get(lugar)
        if t is None or lugar not in columnas:
            raise ValueError(f"fila del análisis sin tema vigente: {lugar}")
        c, fin = columnas[lugar]
        if (d["grado"] != t.grado or limpio(d["procedimental"]) != limpio(t.procedimental)
                or d["fila_fin"] != fin or d["sha256_columnas"] != firma(c)):
            raise ValueError(f"texto, grado, rango o columnas distintos en {d['id']}")
        salida[d["id"]] = t
    if set(ubicaciones) != set(por_fila):
        raise ValueError("el puente no es una correspondencia uno a uno")
    return salida


def incorporar(origen: str | Path, destino: str | Path | None = None) -> dict:
    """Promueve solo los insumos necesarios del análisis independiente; el original no se escribe."""
    origen = ruta(origen)
    destino = destino_seguro(ruta(destino or cargar("motor")["analisis"]))
    if origen.resolve() == destino or origen.resolve().is_relative_to(destino):
        raise ValueError("el destino contiene al análisis de origen")
    datos = {n: leer(origen / n) for n in (*COPIAR, "temas.json")}
    identidades = [{**{k: t[k] for k in ("id", "archivo", "hoja", "fila", "fila_fin", "grado", "procedimental")},
                    "sha256_columnas": firma(t["columnas"])} for t in datos["temas.json"]]
    puente = comprobar_identidades(identidades, datos["fuentes_mallas.json"])
    if any(d.get("id") not in puente for ds in datos["clasificaciones_consolidadas.json"].values() for d in ds):
        raise ValueError("clasificación sin identidad verificada")
    destino.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".analisis-", dir=destino.parent) as temporal:
        carpeta = Path(temporal)
        for nombre in COPIAR:
            escribir(carpeta / nombre, datos[nombre])
        escribir(carpeta / "identidades.json", sorted(identidades, key=lambda d: d["id"]))
        escribir(carpeta / "procedencia.json", {
            "origen": relativa(origen), "archivos_origen": {n: sha(origen / n) for n in (*COPIAR, "temas.json")},
            "metodo": "archivo, hoja, fila, grado, procedimental, rango y SHA de todas las columnas",
            "temas": len(puente), "fuentes": len(datos["fuentes_mallas.json"]),
            "clasificaciones": {k: len(v) for k, v in datos['clasificaciones_consolidadas.json'].items()}})
        if destino.exists():
            if {p.name: sha(p) for p in destino.iterdir()} != {p.name: sha(p) for p in carpeta.iterdir()}:
                raise ValueError("el insumo incorporado ya existe con otro contenido; usar otra versión/destino")
        else:
            os.rename(carpeta, destino)
    return {"temas": len(puente), "fuentes": len(datos["fuentes_mallas.json"]), "destino": relativa(destino)}
