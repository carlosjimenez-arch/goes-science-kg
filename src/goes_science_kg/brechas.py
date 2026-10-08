"""Análisis de brechas por asignatura (fase 4; skill analisis-brechas).

Funciones puras sobre el grafo. Por asignatura:
1. Cobertura de los marcos por ciclo: objetivos del pivote del ciclo con ≥1 tema del ciclo (obj1 u obj2) y
   profundidad (0 · 1 débil · 2–3 básica · 4+ sólida).
2. Oportunidad por CONCEPTO: primer grado en El Salvador − mediana del primer grado en los países de referencia
   (etiquetado directo de sus objetivos). Positivo = El Salvador llega tarde.
3. Secuencia: prerrequisitos que la malla enseña después del concepto que los necesita (o nunca).
4. Conceptos que los países trabajan y El Salvador nunca.

Salidas: asignaturas/<x>/brechas/{brechas.json, brechas.md, oportunidad_conceptos.csv}.
"""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from goes_science_kg.config import cargar, relativa, ruta
from goes_science_kg.excel import guardar as guardar_excel
from goes_science_kg.modelos import Arista, Nodo, TipoArista, TipoNodo
from goes_science_kg.prerrequisitos import evidencia_orden, paises_alto_desempeno

CICLOS = {"T4_27": (2, 4), "T8_27": (5, 8), "PISA25": (7, 9), "TA15": (10, 11), "AUSS": (10, 11)}


def _profundidad(n: int) -> str:
    return "sin cobertura" if n == 0 else "débil" if n == 1 else "básica" if n <= 3 else "sólida"


def analizar(asig: str, nodos: list[Nodo], aristas: list[Arista], ev: dict) -> dict:
    """Brechas de `asig`: cobertura de los marcos por ciclo, oportunidad por concepto, errores de secuencia y
    conceptos que los países trabajan y El Salvador nunca. `ev` es la salida de `evidencia_orden`."""
    por_id = {n.id: n for n in nodos}
    temas = {n.id: n for n in nodos if n.tipo == TipoNodo.TEMA}
    # 1. Cobertura por ciclo
    temas_por_obj = defaultdict(list)
    for a in aristas:
        if a.tipo == TipoArista.CUBRE and a.origen in temas:
            temas_por_obj[a.destino].append(temas[a.origen])
    cobertura = {}
    for marco, (g0, g1) in CICLOS.items():
        objs = [n for n in nodos if n.tipo == TipoNodo.OBJETIVO_MARCO and n.props.get("marco") == marco
                and n.asignatura == asig and (marco != "AUSS" or n.props.get("eje") == "contenido")]
        if not objs:
            continue
        filas = []
        for o in sorted(objs, key=lambda x: x.id):
            k = sum(1 for t in temas_por_obj[o.id] if g0 <= t.grado <= g1)
            filas.append({"objetivo": o.id.removeprefix("OBJ:"), "texto": o.etiqueta, "temas_ciclo": k,
                          "profundidad": _profundidad(k)})
        cobertura[marco] = {"ciclo": f"{g0}.°–{g1}.°", "objetivos": len(filas),
                            "cubiertos": sum(f["temas_ciclo"] > 0 for f in filas),
                            "profundidad": dict(Counter(f["profundidad"] for f in filas)),
                            "sin_cobertura": [f for f in filas if f["temas_ciclo"] == 0],
                            "debiles": [f for f in filas if f["temas_ciclo"] == 1]}
    # 2 y 4. Oportunidad por concepto
    conceptos = [n for n in nodos if n.tipo == TipoNodo.CONCEPTO and n.asignatura == asig]
    oportunidad, nunca = [], []
    for c in sorted(conceptos, key=lambda x: x.id):
        e = ev[c.id]
        if len(e["paises"]) >= 2 and e["paises_mediana"] is not None:
            fila = {"concepto": c.id, "nombre": c.etiqueta, "sv": e["sv"], "mediana_paises": e["paises_mediana"],
                    "mediana_alto_desempeno": e.get("mediana_alto_desempeno"), "paises": e["paises"]}
            if e["sv"] is None:
                nunca.append(fila)
            else:
                oportunidad.append(fila | {"oportunidad": e["sv"] - e["paises_mediana"]})
    tarde = sorted((f for f in oportunidad if f["oportunidad"] >= 2), key=lambda f: (-f["oportunidad"], f["nombre"]))
    adelantado = sorted((f for f in oportunidad if f["oportunidad"] <= -2),
                        key=lambda f: (f["oportunidad"], f["nombre"]))
    # 3. Secuencia
    secuencia = []
    for a in aristas:
        if a.tipo != TipoArista.PRERREQUISITO_DE or por_id[a.destino].asignatura != asig:
            continue
        g_dest, g_pre = ev[a.destino]["sv"], ev[a.origen]["sv"]
        if g_dest is not None and (g_pre is None or g_pre > g_dest):
            secuencia.append({"concepto": por_id[a.destino].etiqueta, "grado_concepto": g_dest,
                              "prerrequisito": por_id[a.origen].etiqueta, "grado_prerrequisito": g_pre,
                              "brecha": (g_pre - g_dest) if g_pre else None,
                              "confianza": a.confianza.value if a.confianza else None})
    secuencia.sort(key=lambda s: (s["confianza"] != "alta", -(s["brecha"] or 99), s["concepto"]))
    return {"asignatura": asig, "cobertura": cobertura, "conceptos": len(conceptos),
            "comparables_con_paises": len(oportunidad) + len(nunca),
            "llega_tarde": tarde, "adelantado": adelantado, "nunca_en_sv": nunca, "secuencia": secuencia,
            "oportunidad_todos": oportunidad}


def _md(r: dict) -> str:
    nombre = cargar("asignaturas")["asignaturas"][r["asignatura"]]["nombre"]
    md = [f"# Brechas · {nombre}", "", "> Generado por `gskg brechas`. No editar a mano. Método: "
          "`src/goes_science_kg/brechas.py` y specs 06 y 09.", "", "## 1. Cobertura de los marcos por ciclo", "",
          "| Marco | Ciclo | Objetivos | Cubiertos | Débiles (1 tema) | Sin cobertura |", "|---|---|---|---|---|---|"]
    for m, c in r["cobertura"].items():
        md.append(f"| {m} | {c['ciclo']} | {c['objetivos']} | {c['cubiertos']} | {len(c['debiles'])} | "
                  f"{len(c['sin_cobertura'])} |")
    for m, c in r["cobertura"].items():
        if c["sin_cobertura"]:
            md += ["", f"**{m} sin cobertura en {c['ciclo']}:** "
                   + "; ".join(f"{f['objetivo']} {f['texto']}" for f in c["sin_cobertura"][:12])]
    md += ["", "## 2. ¿Cuándo llega El Salvador? Oportunidad por concepto", "",
           f"Se compara el primer grado de cada concepto en El Salvador con la mediana de los países de referencia del "
           f"grafo (grado equivalente por edad de ingreso; ver `config/referentes.yaml`) "
           f"(conceptos que trabajan al menos dos países: {r['comparables_con_paises']} de {r['conceptos']}). "
           "Oportunidad = grado SV − mediana; positivo = El Salvador llega tarde.", "",
           "### Llega 2 o más grados tarde", "",
           "La columna «Alto desempeño» es la mediana solo de los países entre los 10 primeros de TIMSS 2023 Ciencias "
           "(con datos en el grafo: "
           + ", ".join(sorted(paises_alto_desempeno() & {p for o in r["oportunidad_todos"] for p in o["paises"]}))
           + "; ver `config/referentes.yaml`).", "",
           "| Concepto | SV | Mediana países | Alto desempeño | Países | Oportunidad |", "|---|---|---|---|---|---|",
           *[f"| {f['nombre']} | {f['sv']}.° | {f['mediana_paises']:g} | "
             + (f"{f['mediana_alto_desempeno']:g}" if f.get("mediana_alto_desempeno") is not None else "—") + " | "
             + ", ".join(f"{p} {g}" for p, g in sorted(f["paises"].items())) + f" | +{f['oportunidad']:g} |"
             for f in r["llega_tarde"]], "", "### Llega 2 o más grados antes", "",
           "| Concepto | SV | Mediana países | Oportunidad |", "|---|---|---|---|",
           *[f"| {f['nombre']} | {f['sv']}.° | {f['mediana_paises']:g} | {f['oportunidad']:g} |"
             for f in r["adelantado"]],
           "", "### Lo trabajan los países y El Salvador nunca", "",
           *([f"- {f['nombre']} (" + ", ".join(f"{p} {g}" for p, g in sorted(f["paises"].items())) + ")"
              for f in r["nunca_en_sv"]] or ["Ninguno."]),
           "", "## 3. Secuencia: prerrequisitos que llegan después del concepto que los necesita", "",
           "| Concepto | Grado | Necesita | Grado del prerrequisito | Confianza |", "|---|---|---|---|---|",
           *[f"| {s['concepto']} | {s['grado_concepto']}.° | {s['prerrequisito']} | "
             f"{str(s['grado_prerrequisito']) + '.°' if s['grado_prerrequisito'] else 'nunca'} | {s['confianza']} |"
             for s in r["secuencia"][:40]],
           "" if len(r["secuencia"]) <= 40 else f"\n… y {len(r['secuencia']) - 40} más en `brechas.json`.", "",
           "## Cómo usar esto", "",
           "- **Sección 1** dice qué objetivos internacionales no tienen ningún tema o tienen uno solo en su ciclo.",
           "- **Sección 2** dice cuándo enseña El Salvador cada idea frente a otros países "
           "(por concepto, no por objetivo).",
           "- **Sección 3** son errores de secuencia candidatos a «mover» en la propuesta (spec 06). Las de confianza "
           "alta van primero; antes de mover un tema conviene revisar el caso con el equipo de Ciencias.", ""]
    return "\n".join(md)


def escribir_brechas(nodos: list[Nodo], aristas: list[Arista]) -> list[dict]:
    """Analiza cada asignatura y escribe brechas.json, brechas.md, oportunidad_conceptos.csv y el libro de Excel
    en asignaturas/<x>/brechas/. Devuelve, por asignatura, cuántos hallazgos hay de cada tipo."""
    ev = evidencia_orden(nodos, aristas)
    resumen = []
    for asig, a in cargar("asignaturas")["asignaturas"].items():
        r = analizar(asig, nodos, aristas, ev)
        d = ruta(a["carpeta"]) / "brechas"
        d.mkdir(parents=True, exist_ok=True)
        (d / "brechas.json").write_text(json.dumps(r, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        (d / "brechas.md").write_text(_md(r), encoding="utf-8")
        with open(d / "oportunidad_conceptos.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, lineterminator="\n")
            paises = sorted({p for o in r["oportunidad_todos"] for p in o["paises"]})
            w.writerow(["concepto", "nombre", "primer_grado_sv", "mediana_paises", "mediana_alto_desempeno",
                        *paises, "oportunidad"])
            for o in sorted(r["oportunidad_todos"], key=lambda o: -o["oportunidad"]):
                w.writerow([o["concepto"], o["nombre"], o["sv"], o["mediana_paises"], o.get("mediana_alto_desempeno"),
                            *[o["paises"].get(p, "") for p in paises], o["oportunidad"]])
        escribir_excel(asig, r)
        resumen.append({"asignatura": asig, "llega_tarde": len(r["llega_tarde"]), "adelantado": len(r["adelantado"]),
                        "nunca_en_sv": len(r["nunca_en_sv"]), "secuencia": len(r["secuencia"])})
    return resumen


def escribir_excel(asig: str, r: dict) -> str:
    """Brechas_<Asignatura>.xlsx: hoja de datos «Oportunidad» y «Resumen» con FÓRMULAS sobre ella
    (estilo de los libros 2027: verde bosque 1E4D3A, sin azul; sin XLOOKUP/FILTER; compatible con Numbers)."""

    nombre = cargar("asignaturas")["asignaturas"][asig]["nombre"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidad"
    paises = sorted({p for o in r["oportunidad_todos"] for p in o["paises"]})
    enc = ["Concepto", "Primer grado SV", "Mediana países", "Mediana alto desempeño", *paises, "Oportunidad"]
    ws.append(enc)
    for o in sorted(r["oportunidad_todos"], key=lambda o: (-o["oportunidad"], o["nombre"])):
        ws.append([o["nombre"], o["sv"], o["mediana_paises"], o.get("mediana_alto_desempeno"),
                   *[o["paises"].get(p) for p in paises], o["oportunidad"]])
    n = ws.max_row
    col_op = chr(ord("A") + len(enc) - 1)
    for c in ws[1]:
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E4D3A")
    rs = wb.create_sheet("Resumen")
    filas = [
        ("Conceptos comparables con los países", f"=COUNTA(Oportunidad!A2:A{n})"),
        ("Llegan 2 o más grados tarde", f'=COUNTIF(Oportunidad!{col_op}2:{col_op}{n},">=2")'),
        ("Llegan 2 o más grados antes", f'=COUNTIF(Oportunidad!{col_op}2:{col_op}{n},"<=-2")'),
        ("Oportunidad promedio (grados)", f"=AVERAGE(Oportunidad!{col_op}2:{col_op}{n})"),
        ("Errores de secuencia (todos)", len(r["secuencia"])),
        ("Errores de secuencia (confianza alta)", sum(s["confianza"] == "alta" for s in r["secuencia"])),
    ]
    rs.append([f"Brechas · {nombre}", "Valor"])
    for f in filas:
        rs.append(list(f))
    rs.append([])
    rs.append(["Oportunidad = primer grado en El Salvador − mediana de los países de referencia (grado equivalente "
               "por edad de ingreso). Positivo: El Salvador llega tarde. Fuente: gskg brechas."])
    for c in rs[1]:
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E4D3A")
    rs.column_dimensions["A"].width, ws.column_dimensions["A"].width = 46, 44
    p = ruta(cargar("asignaturas")["asignaturas"][asig]["carpeta"]) / "brechas" / f"Brechas_{asig.capitalize()}.xlsx"
    guardar_excel(wb, p)
    return relativa(p)
