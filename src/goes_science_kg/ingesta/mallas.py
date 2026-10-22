"""Extract curriculum topics from MINED spreadsheet rows with procedural objectives."""

# Historical implementation notes:
# Migrado de reportes/cobertura_curricular/scripts/extract_temas.py con dos cambios:
# - El archivo de cada malla se fija en config/mallas.yaml (no se elige por fecha de modificación).
# - El id del tema es estable: grado + asignatura + unidad + código del procedimental
#   (p. ej. «G07-CIE-U2-2.3»), no la posición de lectura.
#
# Hojas de unidad: las que empiezan con «Unidad» (2.°–9.°) o contienen «N.°» (Bachillerato).
# Nunca se modifican los archivos de origen.
#
# Versiones (spec 15): si la config tiene `anterior` (la V3 frente a la V2), los ids no son de posición. Un tema
# idéntico
# o parecido a uno de la versión anterior conserva su id y su llave de clasificación (archivo, hoja, fila de la
# anterior); uno nuevo toma su id de posición, con el sufijo «-v3» si ese id existió antes para otro tema. La
# correspondencia la calcula `correspondencia()` y se guarda versionada.

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from difflib import SequenceMatcher
from functools import cache
from typing import TYPE_CHECKING

import openpyxl

from goes_science_kg.config import cargar, ruta

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

CODIGO_ASIG = {"ciencias": "CIE", "fisica": "FIS", "quimica": "QUI", "biologia": "BIO"}


@dataclass
class TemaMalla:
    """Represent one procedural-objective row with its source location and original fields."""

    id: str
    archivo: str  # nombre canónico (sin «(1)»): es la llave de las clasificaciones previas
    ruta: str  # archivo realmente leído
    hoja: str
    fila: int
    grado: int
    asignatura: str  # ciencias | fisica | quimica | biologia (asignatura de la MALLA)
    unidad: str
    contenido: str
    subcontenido: str
    procedimental: str
    indicador: str
    evidencia: str
    habilidad_timss: str
    competencia_pisa: str
    micro_pisa: dict[str, list[str]] = field(default_factory=dict)
    llave: tuple[str, str, int] | tuple[()] = ()  # llave de las clasificaciones previas (ver el docstring)
    origen: str = ""  # frente a la versión anterior: identico | parecido | nuevo («» si no hay anterior)
    id_posicional: str = ""  # el id que daría la posición del tema en ESTA versión

    def dict(self) -> dict:
        """Return the topic's fields as a dictionary."""
        return asdict(self)


def _limpio(v: object) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def _num(texto: str, patron: str) -> str | None:
    m = re.match(patron, texto or "")
    return m.group(1) if m else None


def _columnas(ws: ReadOnlyWorksheet) -> list[str]:
    h1 = [c.value or "" for c in ws[1]]
    h2 = [c.value or "" for c in ws[2]]
    cols = [(str(a) or str(b)).split("\n")[0].strip() for a, b in zip(h1, h2, strict=False)]
    for i, s in enumerate(h2):
        if s:
            cols[i] = str(s).split("\n")[0].strip()
    return cols


@cache
def extraer(nombre_config: str = "mallas") -> tuple[TemaMalla, ...]:
    """Read curricula selected by an explicit configuration name."""

    # Implementation notes:
    # «mallas_v1», las anteriores). Con `anterior` en la config, aplica la correspondencia (ids estables).
    #
    # Con caché: leer los Excel es lo más lento del proceso y varios pasos piden la misma malla en una corrida.
    cfg = cargar(nombre_config)
    temas = _leer(cfg)
    if cfg.get("anterior"):
        _aplicar_correspondencia(temas, cfg)
    return tuple(temas)


def _leer(cfg: dict) -> list[TemaMalla]:
    """Read topic rows with positional identities and physical-location keys."""
    base = ruta(cfg["directorio"])
    fases = cfg["fases_pisa"]
    temas: list[TemaMalla] = []
    for item in cfg["archivos"]:
        archivo = item["archivo"]
        canonico = item.get("canonico", archivo)
        asig = item["asignatura"]
        wb = openpyxl.load_workbook(base / archivo, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if not (ws.title.startswith("Unidad") or re.search(r"\d+\.°", ws.title)):
                continue
            grado = item.get("grado") or int(re.search(r"(\d+)\.°", ws.title).group(1))
            cols = _columnas(ws)
            unidad = contenido = ""
            for fila, row in enumerate(
                ws.iter_rows(min_row=cfg["fila_datos"], values_only=True), start=cfg["fila_datos"]
            ):
                d = dict(zip(cols, row, strict=False))
                if d.get("Unidad"):
                    unidad = _limpio(d["Unidad"])
                if d.get("Contenido"):
                    contenido = _limpio(d["Contenido"])
                if not d.get("Procedimental"):
                    continue
                proc = _limpio(d["Procedimental"])
                temas.append(
                    TemaMalla(
                        id="",
                        archivo=canonico,
                        ruta=archivo,
                        hoja=ws.title,
                        fila=fila,
                        grado=grado,
                        asignatura=asig,
                        unidad=unidad,
                        contenido=contenido,
                        subcontenido=_limpio(d.get("Subcontenidos")),
                        procedimental=proc,
                        indicador=_limpio(d.get("Indicadores de logro") or d.get("Indicador avanzado")),
                        evidencia=_limpio(d.get("Evidencia de aprendizaje")),
                        habilidad_timss=_limpio(d.get("Habilidad TIMSS")),
                        competencia_pisa=_limpio(d.get("Competencia PISA 2025")),
                        micro_pisa={f: re.findall(r"\b([EDI]\d)\s*·", str(d.get(f) or "")) for f in fases},
                    )
                )
        wb.close()
    _asignar_ids(temas)
    for t in temas:
        t.id_posicional, t.llave = t.id, (t.archivo, t.hoja, t.fila)
    return temas


# Verbos y palabras de la redacción de la malla que no dicen QUÉ se enseña (raíces de 6 letras sin tilde).
_GENERICAS = frozenset(
    {
        "identi",
        "recono",
        "experi",
        "aplica",
        "analis",
        "explic",
        "descri",
        "indaga",
        "compro",
        "partir",
        "median",
        "traves",
        "funcio",
        "relaci",
        "utiliz",
        "determ",
        "repres",
        "cotidi",
        "situac",
        "concep",
        "caract",
        "fenome",
    }
)


def _raices(texto: str) -> set[str]:
    """Extract bounded word roots for historical topic-content matching."""
    t = "".join(c for c in unicodedata.normalize("NFKD", texto.lower()) if not unicodedata.combining(c))
    return {w[:6] for w in re.findall(r"[a-zñ]{4,}", t)} - _GENERICAS


def _contencion(previo: TemaMalla, raices_nuevo: set[str]) -> float:
    """Measure how much of a historical topic occurs in a new topic's text."""
    a = _raices(f"{previo.procedimental} {previo.subcontenido}")
    return round(len(a & raices_nuevo) / len(a), 3) if a else 0.0


def _norma(texto: str) -> str:
    """Normalize procedural text for historical version comparison."""
    return re.sub(r"\W+", " ", re.sub(r"^\s*\d+(?:\.\d+)+\.?\s*", "", texto.lower())).strip()


def correspondencia(nombre_config: str = "mallas") -> dict:
    """Match historical curriculum topics within the same grade and subject context."""

    # Implementation notes:
    # asignatura de la malla (determinista):
    # 1. idéntico: el mismo procedimental (normalizado);
    # 2. parecido: similitud del procedimental ≥ `similitud_minima`, de mayor a menor;
    # 3. ampliado: un tema todavía sin par que contiene (≥ `contencion_minima` de sus raíces, en procedimental,
    #    subcontenido, contenido e indicador) a un tema de la anterior sin par: la V3 lo fusionó o lo amplió.
    # Los temas de la anterior que quedan sin par llevan `absorbido_por` (el tema que los contiene, si lo hay) y su
    # `contencion`; los demás son candidatos a «quitado», que se juzga por concepto (spec 15).
    cfg = cargar(nombre_config)
    nuevos, previos = _leer(cfg), extraer(cfg["anterior"])
    umbral = cfg.get("similitud_minima", 0.75)
    par: dict[str, tuple[TemaMalla, str, float]] = {}
    usados: set[str] = set()
    por_texto: dict[tuple, list[TemaMalla]] = defaultdict(list)
    for p in previos:
        por_texto[(p.grado, p.asignatura, _norma(p.procedimental))].append(p)
    for t in nuevos:
        for p in por_texto.get((t.grado, t.asignatura, _norma(t.procedimental)), []):
            if p.id not in usados:
                par[t.id_posicional] = (p, "identico", 1.0)
                usados.add(p.id)
                break
    candidatos = []
    for t in nuevos:
        if t.id_posicional in par:
            continue
        a = _norma(t.procedimental)
        for p in previos:
            if p.id in usados or (p.grado, p.asignatura) != (t.grado, t.asignatura):
                continue
            m = SequenceMatcher(None, a, _norma(p.procedimental))
            if m.quick_ratio() >= umbral and (r := m.ratio()) >= umbral:
                candidatos.append((-round(r, 3), t.id_posicional, p.id, t, p))
    for r, _, _, t, p in sorted(candidatos, key=lambda c: c[:3]):
        if t.id_posicional not in par and p.id not in usados:
            par[t.id_posicional] = (p, "parecido", -r)
            usados.add(p.id)
    contencion = cfg.get("contencion_minima", 0.8)
    raices = {
        t.id_posicional: _raices(" ".join([t.procedimental, t.subcontenido, t.contenido, t.indicador])) for t in nuevos
    }
    candidatos = []
    for t in nuevos:
        if t.id_posicional in par:
            continue
        for p in previos:
            if (
                p.id not in usados
                and (p.grado, p.asignatura) == (t.grado, t.asignatura)
                and (c := _contencion(p, raices[t.id_posicional])) >= contencion
            ):
                candidatos.append((-c, t.id_posicional, p.id, t, p))
    for c, _, _, t, p in sorted(candidatos, key=lambda x: x[:3]):
        if t.id_posicional not in par and p.id not in usados:
            par[t.id_posicional] = (p, "ampliado", -c)
            usados.add(p.id)
    ids_previos = {p.id for p in previos}
    filas = []
    for t in nuevos:
        base = {
            "ruta": t.ruta,
            "hoja": t.hoja,
            "fila": t.fila,
            "grado": t.grado,
            "asignatura": t.asignatura,
            "id_posicional": t.id_posicional,
        }
        if t.id_posicional in par:
            p, estado, sim = par[t.id_posicional]
            filas.append({**base, "id": p.id, "estado": estado, "similitud": sim, "llave": [p.archivo, p.hoja, p.fila]})
        else:
            sufijo = "-v3" if t.id_posicional in ids_previos else ""
            filas.append(
                {
                    **base,
                    "id": t.id_posicional + sufijo,
                    "estado": "nuevo",
                    "llave": [f"{cfg.get('version', 'nueva')}/{t.archivo}", t.hoja, t.fila],
                }
            )
    ids = [f["id"] for f in filas]
    if len(set(ids)) != len(ids):
        raise ValueError(f"ids repetidos en la correspondencia: {[i for i, n in Counter(ids).items() if n > 1][:5]}")
    id_de = {f["id_posicional"]: f["id"] for f in filas}
    # Un tema que reaparece idéntico en OTRO grado se movió (no se quitó). No hereda el id: el id dice el grado.
    por_texto_nuevo: dict[str, list[TemaMalla]] = defaultdict(list)
    for t in nuevos:
        por_texto_nuevo[_norma(t.procedimental)].append(t)
    sin_par = []
    for p in previos:
        if p.id in usados:
            continue
        mismos = [t for t in nuevos if (t.grado, t.asignatura) == (p.grado, p.asignatura)]
        mejor = max(mismos, key=lambda t: _contencion(p, raices[t.id_posicional]), default=None)
        c = _contencion(p, raices[mejor.id_posicional]) if mejor else 0.0
        movido = next((t for t in por_texto_nuevo.get(_norma(p.procedimental), []) if t.grado != p.grado), None)
        sin_par.append(
            {
                "id": p.id,
                "grado": p.grado,
                "asignatura": p.asignatura,
                "procedimental": p.procedimental,
                "absorbido_por": id_de[mejor.id_posicional] if mejor and c >= contencion else None,
                "mas_cercano": id_de[mejor.id_posicional] if mejor else None,
                "contencion": c,
                "movido_a": id_de[movido.id_posicional] if movido else None,
            }
        )
    return {
        "version": "correspondencia-v1",
        "anterior": cfg["anterior"],
        "similitud_minima": umbral,
        "resumen": {
            **Counter(f["estado"] for f in filas),
            "sin_par_en_la_anterior": len(sin_par),
            "absorbidos": sum(bool(x["absorbido_por"]) for x in sin_par),
            "movidos_de_grado": sum(bool(x["movido_a"]) for x in sin_par),
        },
        "temas": filas,
        "sin_par": sin_par,
    }


def escribir_correspondencia(nombre_config: str = "mallas") -> dict:
    """Write the configured correspondence artifact and return its summary."""
    cfg = cargar(nombre_config)
    datos = correspondencia(nombre_config)
    p = ruta(cfg["correspondencia"])
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(datos, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    extraer.cache_clear()
    textos_completos.cache_clear()
    absorbidos.cache_clear()
    return datos["resumen"]


@cache
def absorbidos() -> dict[str, str]:
    """Map historical topics absorbed into explicitly identified current topics."""
    cfg = cargar("mallas")
    if not cfg.get("correspondencia") or not (p := ruta(cfg["correspondencia"])).exists():
        return {}
    return {
        x["id"]: x["absorbido_por"] for x in json.loads(p.read_text(encoding="utf-8"))["sin_par"] if x["absorbido_por"]
    }


@cache
def por_id_con_anteriores() -> dict[str, TemaMalla]:
    """Index current and historical topics while prioritizing current records."""

    # Implementation notes:
    # todavía citan temas que la vigente ya no tiene.
    salida: dict[str, TemaMalla] = {}
    cfg, versiones = cargar("mallas"), ["mallas"]
    while (anterior := cfg.get("anterior")) and anterior not in versiones:
        versiones.append(anterior)
        cfg = cargar(anterior)
    for version in reversed(versiones):
        salida.update({t.id: t for t in extraer(version)})
    return salida


def vigente_de(id_tema: str) -> str | None:
    """Resolve an existing or explicitly absorbed topic to its current identity."""

    # Implementation notes:
    # el tema que lo absorbió si la V3 lo fusionó, o None si ya no está. Para leer capas anteriores a la V3.
    pref = "TEMA:" if id_tema.startswith("TEMA:") else ""
    i = id_tema.removeprefix("TEMA:")
    if i in {t.id for t in extraer()}:
        return id_tema
    return f"{pref}{a}" if (a := absorbidos().get(i)) else None


def _aplicar_correspondencia(temas: list[TemaMalla], cfg: dict) -> None:
    p = ruta(cfg["correspondencia"])
    if not p.exists():
        raise FileNotFoundError(f"falta {cfg['correspondencia']}: córrelo con `uv run gskg malla-correspondencia`")
    por_fila = {(f["ruta"], f["hoja"], f["fila"]): f for f in json.loads(p.read_text(encoding="utf-8"))["temas"]}
    for t in temas:
        f = por_fila.get((t.ruta, t.hoja, t.fila))
        if not f or f["id_posicional"] != t.id_posicional:
            raise ValueError(
                f"la correspondencia no corresponde a la malla ({t.ruta}, {t.hoja}, fila {t.fila}): "
                "córrela de nuevo con `uv run gskg malla-correspondencia`"
            )
        t.id, t.llave, t.origen = f["id"], tuple(f["llave"]), f["estado"]


@cache
def textos_completos(nombre_config: str = "mallas") -> dict[str, str]:
    """Index all original row text, including assessment and scientific-practice fields."""

    # Implementation notes:
    # indicador avanzado), más su contenido y subcontenido. Para búsquedas de texto que no dependan de las columnas que
    # lee el etiquetado (spec 12, verificación de 2.°–8.°).
    cfg = cargar(nombre_config)
    base = ruta(cfg["directorio"])
    temas = extraer(nombre_config)
    por_fila = {(t.ruta, t.hoja, t.fila): t for t in temas}
    salida: dict[str, str] = {}
    for archivo in sorted({t.ruta for t in temas}):
        wb = openpyxl.load_workbook(base / archivo, data_only=True, read_only=True)
        for ws in wb.worksheets:
            for fila, row in enumerate(
                ws.iter_rows(min_row=cfg["fila_datos"], values_only=True), start=cfg["fila_datos"]
            ):
                if t := por_fila.get((archivo, ws.title, fila)):
                    celdas = " ".join(_limpio(v) for v in row if v not in (None, ""))
                    salida[t.id] = f"{t.contenido} {t.subcontenido} {celdas}"
        wb.close()
    return salida


def _asignar_ids(temas: list[TemaMalla]) -> None:
    """Build stable grade, subject, unit, and objective identities with row disambiguation."""
    base: dict[str, list[TemaMalla]] = {}
    for t in temas:
        u = _num(t.unidad, r"\s*(\d+)") or _num(t.hoja, r"\D*(\d+)") or "0"
        c = _num(t.procedimental, r"\s*(\d+(?:\.\d+)+)") or f"f{t.fila}"
        base.setdefault(f"G{t.grado:02d}-{CODIGO_ASIG[t.asignatura]}-U{u}-{c}", []).append(t)
    for clave, grupo in base.items():
        for t in grupo:
            t.id = clave if len(grupo) == 1 else f"{clave}-f{t.fila}"
