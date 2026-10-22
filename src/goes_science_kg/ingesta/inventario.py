"""Inventario de nuevas entregas Excel sin activar mallas ni modificar fuentes.

Los campos de cabecera son sugerencias por coincidencia exacta de etiquetas. Una
hoja sin esquema conocido o con columnas ambiguas requiere mapeo antes de ingesta.
``listo_inspeccion`` significa que el manifiesto puede revisarse; no certifica el
contenido curricular ni convierte una versión declarada en la malla vigente.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import stat
import unicodedata
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

LOGGER = logging.getLogger(__name__)
FORMATOS_ADMITIDOS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
FORMATOS_NO_ADMITIDOS = {".xls", ".xlsb", ".xlt"}


def _normalizar(texto: str) -> str:
    plano = unicodedata.normalize("NFKD", texto.casefold())
    plano = "".join(c for c in plano if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", plano).strip()


ETIQUETAS = {
    "unidad": ("unidad", "unidades", "numero de unidad"),
    "contenido": ("contenido", "contenidos"),
    "tema": ("tema", "temas"),
    "subcontenido": ("subcontenido", "subcontenidos", "subtema", "subtemas"),
    "procedimental": ("procedimental", "contenidos procedimentales", "contenido procedimental"),
    "objetivo": ("objetivo", "objetivos", "objetivo de aprendizaje", "objetivos de aprendizaje"),
    "aprendizaje_esperado": ("aprendizaje esperado", "aprendizajes esperados"),
    "indicador": ("indicador", "indicadores", "indicadores de logro", "indicador de logro"),
    "indicador_avanzado": ("indicador avanzado", "indicadores avanzados"),
    "competencia": ("competencia", "competencias"),
    "competencia_pisa": ("competencia pisa", "competencia pisa 2025"),
    "habilidad_timss": ("habilidad timss", "habilidades timss"),
    "actitudinal": ("actitudinal", "contenidos actitudinales"),
    "evidencia": ("evidencia de aprendizaje", "evidencias de aprendizaje", "evidencia de logro"),
    "grado": ("grado", "grados"),
    "asignatura": ("asignatura", "disciplina"),
    "clase": ("n clase", "no clase", "numero de clase", "clase"),
    "duracion": ("duracion", "horas", "minutos"),
    "origen": ("origen del contenido", "fuente"),
    "fase_aproximar": ("1 aproximar", "aproximar"),
    "fase_explorar": ("2 explorar", "explorar"),
    "fase_explicar": ("3 explicar", "explicar"),
    "fase_indagar": ("4 indagar", "indagar"),
    "fase_reforzar": ("5 reforzar", "reforzar"),
}
CAMPOS_POR_ETIQUETA = {_normalizar(etiqueta): campo for campo, etiquetas in ETIQUETAS.items() for etiqueta in etiquetas}


class ContratoInventario(BaseModel):
    """Metadatos inmutables que rechazan campos inesperados."""

    model_config = ConfigDict(extra="forbid", frozen=True)


class ColumnaDetectada(ContratoInventario):
    """Etiqueta observada y sugerencia opcional, sin imponer una posición fija."""

    indice: int = Field(ge=1)
    letra: str
    texto: str
    campo_sugerido: str | None


class CabeceraDetectada(ContratoInventario):
    """Fila con al menos dos etiquetas curriculares reconocidas."""

    fila: int = Field(ge=1)
    columnas: tuple[ColumnaDetectada, ...]


class HojaInventariada(ContratoInventario):
    """Dimensiones reales de celdas con valor; no usa dimensiones declaradas por Excel."""

    nombre: str
    visibilidad: str
    filas_con_valores: int = Field(ge=0)
    celdas_con_valores: int = Field(ge=0)
    primera_fila: int | None
    ultima_fila: int | None
    ultima_columna: int | None
    cabeceras: tuple[CabeceraDetectada, ...]
    ambiguedades: tuple[str, ...]
    estado: Literal["detectada", "requiere_mapeo", "vacia"]


class ArchivoInventariado(ContratoInventario):
    """Un libro identificado por ruta relativa y hash de sus bytes originales."""

    archivo: str
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    bytes: int = Field(ge=1)
    formato: str
    hojas: tuple[HojaInventariada, ...]


class InventarioCurricular(ContratoInventario):
    """Manifiesto declarativo de una entrega; no cambia configuración vigente."""

    esquema_version: Literal["1"] = "1"
    version_declarada: str
    directorio: str
    estado: Literal["incompleto", "listo_inspeccion"]
    huella_contenido: str = Field(pattern=r"^[0-9a-f]{64}$")
    archivos: tuple[ArchivoInventariado, ...]
    archivos_omitidos: tuple[str, ...]
    observaciones: tuple[str, ...]


def _sha256(archivo: Path) -> str:
    huella = hashlib.sha256()
    with archivo.open("rb") as flujo:
        for bloque in iter(lambda: flujo.read(1024 * 1024), b""):
            huella.update(bloque)
    return huella.hexdigest()


def _listar_archivos(base: Path) -> tuple[list[Path], list[Path]]:
    """No oculta errores de permisos ni sigue enlaces simbólicos de las fuentes."""
    archivos, enlaces = [], []

    def recorrer(directorio: Path) -> None:
        # Path.rglob puede silenciar PermissionError en Python recientes. Scandir
        # debe propagarlo: un directorio inaccesible no es una entrega vacía.
        with os.scandir(directorio) as explorador:
            entradas = sorted(explorador, key=lambda entrada: entrada.name)
        for entrada in entradas:
            archivo = Path(entrada.path)
            if entrada.is_symlink():
                enlaces.append(archivo)
            elif entrada.is_dir(follow_symlinks=False):
                recorrer(archivo)
            elif entrada.is_file(follow_symlinks=False):
                archivos.append(archivo)

    recorrer(base)
    return (
        sorted(archivos, key=lambda archivo: archivo.relative_to(base).as_posix()),
        sorted(enlaces, key=lambda archivo: archivo.relative_to(base).as_posix()),
    )


def _cabecera(fila: int, valores: tuple) -> CabeceraDetectada | None:
    etiquetas = [
        (i, str(valor).strip(), CAMPOS_POR_ETIQUETA.get(_normalizar(str(valor))))
        for i, valor in enumerate(valores, start=1)
        if valor is not None and str(valor).strip()
    ]
    if sum(campo is not None for _, _, campo in etiquetas) < 2:
        return None
    columnas = tuple(
        ColumnaDetectada(
            indice=i,
            letra=get_column_letter(i),
            texto=texto,
            campo_sugerido=campo,
        )
        for i, texto, campo in etiquetas
    )
    return CabeceraDetectada(fila=fila, columnas=columnas)


def _inspeccionar_hoja(hoja) -> HojaInventariada:
    """Recorre la hoja una vez y conserva solo metadatos y candidatos de cabecera."""
    # Algunos productores omiten o declaran mal <dimension>; el XML es la autoridad.
    hoja.reset_dimensions()
    filas = celdas = 0
    primera = ultima = ultima_columna = None
    cabeceras = []
    campos: dict[str, set[str]] = {}
    columnas: dict[str, set[str]] = {}
    for numero, valores in enumerate(hoja.iter_rows(values_only=True), start=1):
        presentes = [i for i, valor in enumerate(valores, start=1) if valor is not None and str(valor).strip()]
        if not presentes:
            continue
        filas += 1
        celdas += len(presentes)
        primera = numero if primera is None else primera
        ultima = numero
        ultima_columna = max(ultima_columna or 0, presentes[-1])
        cabecera = _cabecera(numero, valores)
        if cabecera is None:
            continue
        cabeceras.append(cabecera)
        for columna in cabecera.columnas:
            if columna.campo_sugerido:
                campos.setdefault(columna.campo_sugerido, set()).add(columna.letra)
                columnas.setdefault(columna.letra, set()).add(columna.campo_sugerido)
    ambiguedades = [
        f"El campo sugerido {campo} aparece en columnas distintas: {', '.join(sorted(ubicaciones))}"
        for campo, ubicaciones in sorted(campos.items())
        if len(ubicaciones) > 1
    ]
    ambiguedades += [
        f"La columna {columna} tiene campos sugeridos distintos: {', '.join(sorted(nombres))}"
        for columna, nombres in sorted(columnas.items())
        if len(nombres) > 1
    ]
    if filas and not cabeceras:
        ambiguedades.append("No se detectó una cabecera curricular: requiere mapeo o exclusión explícita de esta hoja")
    estado = "vacia" if not filas else "requiere_mapeo" if ambiguedades else "detectada"
    return HojaInventariada(
        nombre=hoja.title,
        visibilidad=hoja.sheet_state,
        filas_con_valores=filas,
        celdas_con_valores=celdas,
        primera_fila=primera,
        ultima_fila=ultima,
        ultima_columna=ultima_columna,
        cabeceras=tuple(cabeceras),
        ambiguedades=tuple(ambiguedades),
        estado=estado,
    )


def _inspeccionar_archivo(archivo: Path, base: Path) -> ArchivoInventariado:
    antes = archivo.stat()
    digest = _sha256(archivo)
    try:
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=False, keep_links=False)
        try:
            hojas = tuple(_inspeccionar_hoja(hoja) for hoja in libro.worksheets)
        finally:
            libro.close()
    except Exception as error:
        raise ValueError(f"Libro Excel ilegible: {archivo.relative_to(base)} ({type(error).__name__})") from error
    despues = archivo.stat()
    if (antes.st_size, antes.st_mtime_ns, antes.st_ino) != (despues.st_size, despues.st_mtime_ns, despues.st_ino):
        raise ValueError(f"El archivo cambió durante la inspección: {archivo.relative_to(base)}")
    return ArchivoInventariado(
        archivo=archivo.relative_to(base).as_posix(),
        sha256=digest,
        bytes=antes.st_size,
        formato=archivo.suffix.lower(),
        hojas=hojas,
    )


def inventariar(directorio: Path, version: str) -> InventarioCurricular:
    """Inspecciona Excel recursivamente y devuelve un manifiesto sin escribir archivos.

    Rechaza directorios sin libros y formatos Excel incompatibles o corruptos. Las
    hojas sin cabecera reconocida o con un mapeo ambiguo dejan el estado incompleto.
    Las hojas vacías se registran; no hacen fallar una entrega con tablas utilizables.
    Los permisos denegados se propagan. Los enlaces simbólicos se registran como
    omisiones pendientes de resolver y nunca se siguen, incluso dentro del árbol.
    """
    base = directorio.expanduser().resolve()
    try:
        es_directorio = stat.S_ISDIR(base.stat().st_mode)
    except (FileNotFoundError, NotADirectoryError) as error:
        raise ValueError(f"No es un directorio de currículos: {base}") from error
    if not es_directorio:
        raise ValueError(f"No es un directorio de currículos: {base}")
    if not isinstance(version, str) or not version.strip() or len(version.strip()) > 50:
        raise ValueError("La versión declarada debe tener entre 1 y 50 caracteres")
    todos, enlaces = _listar_archivos(base)
    elegidos = [p for p in todos if p.suffix.lower() in FORMATOS_ADMITIDOS and not p.name.startswith("~$")]
    incompatibles = [p.relative_to(base).as_posix() for p in todos if p.suffix.lower() in FORMATOS_NO_ADMITIDOS]
    if incompatibles:
        raise ValueError(f"Formato Excel no admitido; convierte una copia a XLSX: {', '.join(incompatibles)}")
    if not elegidos:
        raise ValueError("El directorio no contiene libros Excel compatibles")
    archivos = tuple(_inspeccionar_archivo(p, base) for p in elegidos)
    observaciones = [f"Enlace simbólico sin inspeccionar: {p.relative_to(base).as_posix()}" for p in enlaces]
    for archivo in archivos:
        for hoja in archivo.hojas:
            observaciones.extend(f"{archivo.archivo} / {hoja.nombre}: {aviso}" for aviso in hoja.ambiguedades)
    if not any(h.estado == "detectada" for a in archivos for h in a.hojas):
        observaciones.append("No hay ninguna hoja con esquema curricular reconocido y sin ambigüedades")
    identidad = [(a.archivo, a.sha256) for a in archivos]
    huella = hashlib.sha256(json.dumps(identidad, ensure_ascii=False).encode()).hexdigest()
    estado = "incompleto" if observaciones else "listo_inspeccion"
    LOGGER.info("Inventario %s: %d libros, estado=%s", version.strip(), len(archivos), estado)
    return InventarioCurricular(
        version_declarada=version.strip(),
        directorio=str(base),
        estado=estado,
        huella_contenido=huella,
        archivos=archivos,
        archivos_omitidos=tuple(
            sorted([p.relative_to(base).as_posix() for p in [*todos, *enlaces] if p not in elegidos])
        ),
        observaciones=tuple(observaciones),
    )
