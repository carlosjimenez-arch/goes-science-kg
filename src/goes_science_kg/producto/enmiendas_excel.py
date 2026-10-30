"""Create derived workbook copies with explicitly supported cell corrections."""

from __future__ import annotations

import hashlib
import json
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from goes_science_kg.archivos import destino_generado
from goes_science_kg.motor.contratos import huella

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def enmiendas_respaldadas(archivo):
    """Select only corrections supported by the exact published unit reviews."""

    todas, aceptadas = {}, set()
    for c in archivo.unidades:
        for e in c.preparacion.correcciones_requeridas:
            if e["id"] != "ENMIENDA:" + huella({k: v for k, v in e.items() if k != "id"}):
                raise ValueError("La enmienda cambió")
            if e["id"] in todas and todas[e["id"]] != e:
                raise ValueError("Una misma enmienda tiene dos contenidos")
            todas[e["id"]] = e
        if c.resultado.propuesta:
            aceptadas.update(c.resultado.propuesta.enmiendas_aceptadas)
    return tuple(todas[i] for i in sorted(aceptadas))


def exportar_enmiendas(curriculo, archivo, *, raiz_v4: Path, destino: Path) -> dict:
    """Write derived corrections without changing originals or unrelated OOXML package entries."""
    destino_generado(destino / "enmiendas.json")
    if destino.exists():
        raise ValueError("Las copias enmendadas requieren un directorio nuevo")
    fuentes = {f.id: f for f in curriculo.fuentes}
    enmiendas = enmiendas_respaldadas(archivo)
    por_fuente = {}
    for e in enmiendas:
        por_fuente.setdefault(e["fuente_id"], []).append(e)
    preparados = []
    for id_, cambios in sorted(por_fuente.items()):
        fuente = fuentes[id_]
        origen = raiz_v4 / fuente.archivo
        if origen.is_symlink() or not origen.resolve().is_relative_to(raiz_v4.resolve()):
            raise ValueError("Fuente Excel fuera de su raíz")
        contenido = origen.read_bytes()
        digest = hashlib.sha256(contenido).hexdigest()
        if digest != fuente.sha256 or any(e["sha256"] != digest for e in cambios):
            raise ValueError("La enmienda no corresponde al Excel original")
        original = load_workbook(origen, data_only=False)
        try:
            if any(original[e["hoja"]][e["celda_ancla"]].value != e["antes"] for e in cambios):
                raise ValueError("El contenido anterior de una celda cambió")
        finally:
            original.close()
        from io import BytesIO

        salida = BytesIO()
        # nosec B314: el XML sale del Excel original verificado por SHA-256, no de una petición HTTP.
        with zipfile.ZipFile(BytesIO(contenido)) as z, zipfile.ZipFile(salida, "w") as nuevo:
            libro = ET.fromstring(z.read("xl/workbook.xml"))  # nosec B314
            relaciones = {
                r.attrib["Id"]: r.attrib["Target"]
                for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))  # nosec B314
            }
            hojas = {
                s.attrib["name"]: posixpath.normpath("xl/" + relaciones[s.attrib[f"{{{REL}}}id"]])
                if not relaciones[s.attrib[f"{{{REL}}}id"]].startswith("/")
                else relaciones[s.attrib[f"{{{REL}}}id"]].lstrip("/")
                for s in libro.findall("s:sheets/s:sheet", NS)
            }
            xml = {}
            for e in cambios:
                ruta = hojas[e["hoja"]]
                if ruta not in xml:
                    xml[ruta] = ET.fromstring(z.read(ruta))  # nosec B314
                celda = xml[ruta].find(f".//s:c[@r='{e['celda_ancla']}']", NS)
                if celda is None:
                    raise ValueError("La celda enmendada no existe físicamente")
                for hijo in list(celda):
                    if hijo.tag.rsplit("}", 1)[-1] in {"v", "f", "is"}:
                        celda.remove(hijo)
                celda.set("t", "inlineStr")
                inline = ET.SubElement(celda, f"{{{NS['s']}}}is")
                texto = ET.SubElement(inline, f"{{{NS['s']}}}t")
                texto.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                texto.text = e["despues"]
            for miembro in z.infolist():
                nuevo.writestr(
                    miembro,
                    ET.tostring(xml[miembro.filename], encoding="utf-8", xml_declaration=True)
                    if miembro.filename in xml
                    else z.read(miembro.filename),
                )
        preparados.append((origen, cambios, salida.getvalue()))
    destino.mkdir(parents=True)
    recibos = []
    for origen, cambios, contenido in preparados:
        path = destino / (origen.stem + "_enmienda_derivada.xlsx")
        with path.open("xb") as f:
            f.write(contenido)
        # Verifica todos los valores y estilos; las partes ajenas al XML cambiado son bytes idénticos.
        antes, despues = load_workbook(origen, data_only=False), load_workbook(path, data_only=False)
        esperados = {(e["hoja"], e["celda_ancla"]): e["despues"] for e in cambios}
        try:
            if antes.sheetnames != despues.sheetnames:
                raise ValueError("La copia enmendada perdió hojas")
            for hoja in antes:
                otra = despues[hoja.title]
                if (
                    str(hoja.merged_cells) != str(otra.merged_cells)
                    or hoja.max_row != otra.max_row
                    or hoja.max_column != otra.max_column
                ):
                    raise ValueError("La copia enmendada cambió la estructura")
                for fila in hoja:
                    for c in fila:
                        otra_celda = otra[c.coordinate]
                        if (
                            otra_celda.value != esperados.get((hoja.title, c.coordinate), c.value)
                            or otra_celda._style != c._style
                        ):
                            raise ValueError("La copia cambió una celda fuera de las enmiendas respaldadas")
        finally:
            antes.close()
            despues.close()
        recibos.append(
            {
                "archivo": path.name,
                "sha256": hashlib.sha256(contenido).hexdigest(),
                "fuente_sha256": hashlib.sha256(origen.read_bytes()).hexdigest(),
                "enmiendas": cambios,
            }
        )
    informe = {
        "archivo_curricular_id": archivo.id,
        "copias": recibos,
        "limite": (
            "Copias derivadas con enmiendas respaldadas automáticamente; los originales y sus alertas se conservan"
        ),
    }
    (destino / "enmiendas.json").write_text(json.dumps(informe, ensure_ascii=False, indent=2) + "\n")
    return informe
