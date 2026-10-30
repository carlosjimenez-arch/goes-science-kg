"""Export original and proposed curricula with states and evidence for offline use."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from goes_science_kg.archivos import destino_generado
from goes_science_kg.curriculo.contratos import FASES
from goes_science_kg.excel import guardar
from goes_science_kg.producto.campana_mejoras import escribir_nuevo_o_igual
from goes_science_kg.producto.vista_mejoras import vista_mejoras

CAMPOS = (
    "objetivo_propuesto",
    "indicador_observable",
    "producto_evaluable",
    "criterio_evaluacion",
    "orientacion_didactica",
    "recursos_condicionales",
    "seguridad",
    "andamiaje",
    "contexto_salvador",
    "justificacion",
)


def filas_exportables(curriculo, archivo):
    """Build original-versus-proposal rows with review states and evidence provenance."""

    por_fila = {}
    for caso in archivo.unidades:
        propuestas = {f.fila_id: f for f in caso.resultado.propuesta.filas} if caso.resultado.propuesta else {}
        for original in caso.preparacion.filas_originales:
            f = propuestas.get(original["id"])
            por_fila[original["id"]] = (caso, f)
    for fila in curriculo.filas:
        caso, f = por_fila[fila.id]
        yield {
            "fila_id": fila.id,
            "grado": fila.grado,
            "asignatura_malla": fila.asignatura_malla,
            "unidad": fila.unidad,
            "hoja_original": fila.hoja,
            "fila_original": fila.fila,
            "objetivo_original": fila.procedimental,
            "indicadores_originales": fila.indicadores_logro,
            "evidencia_original": fila.evidencia_aprendizaje,
            **{
                "fase_" + fase: fila.fases[fase].texto if fila.fases[fase].presente_en_hoja else "Ausente en el insumo"
                for fase in FASES
            },
            "estado_revision": caso.resultado.estado,
            "motivos": "; ".join(caso.resultado.motivos),
            "accion": f.accion if f else "sin_propuesta_respaldada",
            "disciplinas_propuestas": "; ".join(f.disciplinas_propuestas) if f else "",
            **{
                c: ("; ".join(getattr(f, c)) if c == "recursos_condicionales" else getattr(f, c)) if f else ""
                for c in CAMPOS
            },
            "fragmentos_ids": "; ".join(f.fragmentos_ids) if f else "",
            "preparacion_id": caso.preparacion.id,
            "archivo_id": archivo.id,
            "limite": archivo.limite,
        }


def _texto_csv(valor):
    texto = str(valor)
    return "'" + texto if texto.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else texto


def csv_curricular(curriculo, archivo) -> str:
    """Serialize the curricular comparison as a formula-safe CSV."""

    filas = list(filas_exportables(curriculo, archivo))
    salida = io.StringIO(newline="")
    writer = csv.DictWriter(salida, fieldnames=list(filas[0]))
    writer.writeheader()
    writer.writerows({k: _texto_csv(v) for k, v in fila.items()} for fila in filas)
    return "\ufeff" + salida.getvalue()


def exportar_excel(curriculo, archivo, destino: Path) -> None:
    """Export the complete curricular comparison to a deterministic, formula-safe workbook."""

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    destino = destino_generado(destino)
    if destino.exists() or destino.is_symlink():
        raise ValueError("La matriz derivada requiere un archivo nuevo")
    filas = list(filas_exportables(curriculo, archivo))
    wb = Workbook()
    ws = wb.active
    ws.title = "Original y propuesta V4"
    columnas = list(filas[0])
    ws.append(columnas)
    for fila in filas:
        ws.append([fila[c] for c in columnas])
    control = wb.create_sheet("Alcance y control")
    control.append(["Archivo", archivo.id])
    control.append(["Límite", archivo.limite])
    control.append(["Cobertura", json.dumps(archivo.cobertura, ensure_ascii=False)])
    control.append(["Fuentes", "Localizadores, textos íntegros y votos en archivo-curricular.json y la API"])
    for hoja in wb:
        for row in hoja:
            for celda in row:
                if isinstance(celda.value, str):
                    celda.data_type = "s"  # No ejecutar fórmulas procedentes de texto curricular o modelos.
                celda.alignment = Alignment(vertical="top", wrap_text=True)
    for celda in ws[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="174B45")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 42 if col[0].value in CAMPOS else 24
    destino.parent.mkdir(parents=True, exist_ok=True)
    guardar(wb, destino)
    wb.close()
    prueba = load_workbook(destino, read_only=True, data_only=False)
    try:
        if prueba.worksheets[0].max_row != len(filas) + 1:
            raise ValueError("La matriz exportada perdió filas")
        for hoja in prueba:
            for row in hoja.iter_rows():
                if any(c.data_type == "f" for c in row):
                    raise ValueError("La exportación contiene una fórmula no autorizada")
    finally:
        prueba.close()


def exportar_entrega(curriculo, archivo, destino: Path, *, raiz_v4: Path | None = None) -> dict:
    """Write the full offline delivery with its publication identity and source evidence."""

    rutas = [
        destino / n
        for n in ("index.html", "ciencias-v4-propuestas.csv", "ciencias-v4-propuestas.xlsx", "archivo-curricular.json")
    ]
    for p in rutas:
        destino_generado(p)
        if p.exists() or p.is_symlink():
            raise ValueError("La entrega debe escribirse en un destino nuevo")
    destino.mkdir(parents=True, exist_ok=True)
    escribir_nuevo_o_igual(rutas[3], archivo.model_dump(mode="json"))
    rutas[0].write_text(vista_mejoras(curriculo, archivo), encoding="utf-8")
    rutas[1].write_text(csv_curricular(curriculo, archivo), encoding="utf-8")
    exportar_excel(curriculo, archivo, rutas[2])
    salida = {"archivos": [str(p) for p in rutas], "cobertura": archivo.cobertura}
    if raiz_v4 is not None:
        from goes_science_kg.producto.enmiendas_excel import exportar_enmiendas

        salida["enmiendas"] = exportar_enmiendas(curriculo, archivo, raiz_v4=raiz_v4, destino=destino / "enmiendas")
    return salida
