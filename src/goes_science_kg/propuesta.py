"""Propuesta curricular por asignatura y ciclo (fase 5; spec 06; skill propuesta-curricular).

Paso 1 (determinista, este módulo): CANDIDATOS de cambio con su evidencia, calculados sobre el grafo.
- mover_antes: conceptos que El Salvador introduce ≥2 grados después que la mediana de los países, y cuyos
  prerrequisitos ya se enseñan antes del grado destino (si no, se listan como bloqueos).
- corregir_secuencia: prerrequisitos (confianza alta) que la malla enseña después del concepto que los necesita.
- reforzar: objetivos del marco del ciclo con un solo tema (profundidad débil) o sin temas.
- equilibrar: temas de la asignatura por grado del ciclo (un grado casi vacío es señal de mala distribución).
Paso 2 (subagente especialista): redacta la propuesta con esos candidatos (prompts/propuesta_curricular.md).
Paso 3 (determinista): `escribir_excel()` arma el libro de la propuesta para el MINED.
"""

from __future__ import annotations

import json
import math
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from goes_science_kg import conceptos as cp
from goes_science_kg.brechas import analizar
from goes_science_kg.config import ruta
from goes_science_kg.excel import guardar as guardar_excel
from goes_science_kg.modelos import Arista, Confianza, Nodo, TipoArista, TipoNodo
from goes_science_kg.prerrequisitos import evidencia_orden


def candidatos(asig: str, g0: int, g1: int, marco: str, nodos: list[Nodo], aristas: list[Arista]) -> dict:
    ev = evidencia_orden(nodos, aristas)
    por_id = {n.id: n for n in nodos}
    r = analizar(asig, nodos, aristas, ev)
    prereqs = defaultdict(list)
    for a in aristas:
        if a.tipo == TipoArista.PRERREQUISITO_DE:
            prereqs[a.destino].append(a)
    temas_de = defaultdict(list)
    for a in aristas:
        if a.tipo == TipoArista.TRABAJA and a.rol == "principal" and a.origen.startswith("TEMA:"):
            temas_de[a.destino].append(por_id[a.origen])

    mover = []
    for f in r["llega_tarde"]:
        destino = max(g0, math.floor(f["mediana_paises"] + 0.5))  # empates hacia arriba (como Math.round)
        if not (g0 <= destino <= g1) or f["sv"] <= destino:
            continue
        bloqueos = [{"prerrequisito": por_id[a.origen].etiqueta, "primer_grado_sv": ev[a.origen]["sv"]}
                    for a in prereqs[f["concepto"]]
                    if ev[a.origen]["sv"] is None or ev[a.origen]["sv"] >= destino]
        temas = [t for t in temas_de[f["concepto"]] if t.grado == f["sv"]]
        mover.append({"concepto": f["nombre"], "concepto_id": f["concepto"], "grado_actual": f["sv"],
                      "grado_destino": destino, "paises": f["paises"], "bloqueos": bloqueos,
                      "temas_actuales": [{"id": t.id.removeprefix("TEMA:"), "tema": t.etiqueta,
                                          "unidad": t.props.get("unidad")} for t in temas]})
    secuencia = [s for s in r["secuencia"] if s["confianza"] == "alta" and g0 <= s["grado_concepto"] <= g1]
    cob = r["cobertura"].get(marco, {})
    reforzar = cob.get("sin_cobertura", []) + cob.get("debiles", [])
    por_grado = Counter(n.grado for n in nodos if n.tipo == TipoNodo.TEMA and n.asignatura == asig
                        and g0 <= n.grado <= g1)
    temas_ciclo = [{"id": n.id.removeprefix("TEMA:"), "grado": n.grado, "unidad": n.props.get("unidad"),
                    "tema": n.etiqueta,
                    "conceptos": [por_id[a.destino].etiqueta for a in aristas if a.origen == n.id
                                  and a.tipo == TipoArista.TRABAJA and a.destino.startswith("CON:")]}
                   for n in sorted(nodos, key=lambda x: x.id)
                   if n.tipo == TipoNodo.TEMA and n.asignatura == asig and g0 <= (n.grado or 0) <= g1]
    return {"asignatura": asig, "ciclo": [g0, g1], "marco": marco,
            "temas_por_grado": {g: por_grado.get(g, 0) for g in range(g0, g1 + 1)},
            "total_temas_ciclo": sum(por_grado.values()),
            "mover_antes": mover, "corregir_secuencia": secuencia, "reforzar": reforzar,
            "temas_del_ciclo": temas_ciclo}


def escribir_candidatos(asig: str, g0: int, g1: int, marco: str, nodos, aristas) -> str:
    c = candidatos(asig, g0, g1, marco, nodos, aristas)
    d = ruta(f"asignaturas/{asig}/propuesta")
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"candidatos_G{g0:02d}-G{g1:02d}.json"
    p.write_text(json.dumps(c, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return str(p.relative_to(ruta(".")))


# -- Paso 3: libro Excel de la propuesta ----------------------------------------------------------
VERDE, VERDE_CLARO, GRIS = "1E4D3A", "DDEBDD", "F2F2F2"


def escribir_excel(asig: str, g0: int, g1: int) -> str:
    """Libro para el MINED: pestañas «Acciones» y «Distribucion». Estilo de los libros 2027 (sin azul);
    los totales y diferencias son fórmulas; compatible con Numbers (sin «·» en pestañas ni CHAR(10))."""

    d = ruta(f"asignaturas/{asig}/propuesta")
    p = json.loads((d / f"propuesta_G{g0:02d}-G{g1:02d}.json").read_text(encoding="utf-8"))
    c = json.loads((d / f"candidatos_G{g0:02d}-G{g1:02d}.json").read_text(encoding="utf-8"))
    wb = Workbook()
    enc = Font(name="Arial", bold=True, color="FFFFFF")
    normal = Font(name="Arial", size=10)
    ajuste = Alignment(wrap_text=True, vertical="top")

    def cabecera(ws, titulos, anchos):
        ws.append(titulos)
        for i, w in enumerate(anchos, 1):
            celda = ws.cell(row=1, column=i)
            celda.font, celda.fill, celda.alignment = enc, PatternFill("solid", fgColor=VERDE), ajuste
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Acciones"
    cabecera(ws, ["N", "Acción", "Temas afectados", "N.° de temas", "Grado actual", "Grado propuesto",
                  "Unidad propuesta", "Tema propuesto", "Indicador de logro", "Conceptos", "Evidencia",
                  "Justificación"], [5, 11, 22, 9, 9, 10, 26, 48, 44, 32, 34, 48])
    for a in p["acciones"]:
        afectados = a.get("temas_afectados", [])
        ws.append([a.get("n"), a.get("accion"), ", ".join(afectados), len(afectados),
                   a.get("grado_actual"), a.get("grado_propuesto"), a.get("unidad_propuesta"), a.get("tema_propuesto"),
                   a.get("indicador_logro"), ", ".join(a.get("conceptos", [])), "; ".join(a.get("evidencia", [])),
                   a.get("justificacion")])
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font, celda.alignment = normal, ajuste
            if celda.row % 2 == 0:
                celda.fill = PatternFill("solid", fgColor=GRIS)

    wd = wb.create_sheet("Distribucion")
    cabecera(wd, ["Grado", "Temas actuales", "Temas propuestos", "Diferencia", "Acciones que salen del grado",
                  "Acciones que llegan al grado"], [10, 15, 17, 12, 22, 22])
    grados = list(range(g0, g1 + 1))
    for i, g in enumerate(grados, 2):
        wd.append([g, c["temas_por_grado"][str(g)], p["temas_por_grado_propuesto"][str(g)], f"=C{i}-B{i}",
                   f'=COUNTIFS(Acciones!E:E,A{i},Acciones!B:B,"mover")',
                   f'=COUNTIFS(Acciones!F:F,A{i},Acciones!B:B,"mover")'])
    t = len(grados) + 2
    wd.append(["Total", f"=SUM(B2:B{t - 1})", f"=SUM(C2:C{t - 1})", f"=C{t}-B{t}", f"=SUM(E2:E{t - 1})",
               f"=SUM(F2:F{t - 1})"])
    for celda in wd[t]:
        celda.fill, celda.font = PatternFill("solid", fgColor=VERDE_CLARO), Font(name="Arial", bold=True)
    wd.append([])
    wd.append(["Regla: el total del ciclo no cambia (Diferencia del total = 0). Fuente: gskg propuesta + revisión "
               "del especialista; estado: borrador, pendiente de revisión del equipo de Ciencias del MINED."])
    for fila in wd.iter_rows(min_row=2, max_row=t - 1):
        for celda in fila:
            celda.font = normal
    nombre = f"Propuesta_{asig.capitalize()}_G{g0:02d}-G{g1:02d}.xlsx"
    guardar_excel(wb, d / nombre)
    return str((d / nombre).relative_to(ruta(".")))


# -- Simulación: impacto de la propuesta sobre el grafo ---------------------------------------------
def _metricas(asig: str, g0: int, g1: int, nodos, aristas) -> dict:
    ev = evidencia_orden(nodos, aristas)
    r = analizar(asig, nodos, aristas, ev)
    en_ciclo = [f for f in r["oportunidad_todos"] if f["sv"] is not None and g0 <= f["sv"] <= g1
                or f["mediana_paises"] is not None and g0 <= f["mediana_paises"] <= g1]
    return {"secuencia_en_ciclo": sum(1 for s in r["secuencia"] if g0 <= s["grado_concepto"] <= g1),
            "secuencia_alta_en_ciclo": sum(1 for s in r["secuencia"]
                                           if g0 <= s["grado_concepto"] <= g1 and s["confianza"] == "alta"),
            "llega_tarde_2mas": sum(1 for f in en_ciclo if f.get("oportunidad", 0) >= 2),
            "desfase_medio": round(sum(abs(f["oportunidad"]) for f in en_ciclo if "oportunidad" in f)
                                   / max(1, sum(1 for f in en_ciclo if "oportunidad" in f)), 2)}


class _Simulacion:
    """Grafo en memoria sobre el que se aplican, una a una, las acciones de una propuesta."""

    def __init__(self, asig: str, nodos: list[Nodo], aristas: list[Arista], version: str):
        voc = cp.vocabulario()
        self.asig, self.version = asig, version
        self.por_nombre = {c["nombre"].lower(): cid for cid, c in sorted(voc.items())}
        # Cinco nombres se repiten entre asignaturas («Cambios de estado» en Física y Química): se prefiere el de la
        # asignatura que propone la acción.
        self.por_asig_nombre = {(c["asignatura"], c["nombre"].lower()): cid for cid, c in voc.items()}
        self.nodos = {n.id: n.model_copy(deep=True) for n in nodos}
        self.aristas = list(aristas)
        self.sin_resolver: list[str] = []
        self.retiradas = 0

    def resolver(self, nombre: str, de: str | None = None) -> str | None:
        return self.por_asig_nombre.get((de or self.asig, nombre.lower())) or self.por_nombre.get(nombre.lower())

    def _trabaja(self, t: str, cid: str, rol: str, justificacion: str) -> Arista:
        return Arista(origen=t, destino=cid, tipo=TipoArista.TRABAJA, metodo="ia", rol=rol, confianza=Confianza.ALTA,
                      justificacion=justificacion, version=self.version)

    def _agregar_conceptos(self, t: str, nombres: list[str]) -> None:
        for nombre in nombres:
            cid = self.resolver(nombre)
            if cid and t in self.nodos:
                self.aristas.append(self._trabaja(t, cid, "principal", "Acción propuesta (simulación)."))
            elif not cid:
                self.sin_resolver.append(nombre)

    def _sin_trabaja(self, temas: list[str], conceptos: set[str]) -> None:
        """Retira las aristas TRABAJA de esos temas hacia esos conceptos."""
        self.aristas = [x for x in self.aristas if not (x.origen in temas and x.destino in conceptos
                                                        and x.tipo == TipoArista.TRABAJA)]

    def aplicar(self, a: dict) -> None:
        afectados = [f"TEMA:{t}" for t in a.get("temas_afectados", [])]
        if a.get("quitar_conceptos"):
            self._quitar(a, afectados)
        accion = {"mover": self._mover, "fusionar": self._fusionar, "revisar": self._revisar,
                  "dividir": self._dividir, "nuevo": self._nuevo}.get(a["accion"])
        if accion:
            accion(a, afectados)

    def _quitar(self, a: dict, afectados: list[str]) -> None:
        """Cualquier acción puede retirar etiquetas que no corresponden al tema reformulado."""
        quitar = {self.resolver(n, a.get("origen")) for n in a["quitar_conceptos"]} - {None}
        self.sin_resolver += [n for n in a["quitar_conceptos"] if self.resolver(n) is None]
        antes = len(self.aristas)
        self._sin_trabaja(afectados, quitar)
        self.retiradas += antes - len(self.aristas)

    def _mover(self, a: dict, afectados: list[str]) -> None:
        if not a.get("grado_propuesto"):
            return
        for t in afectados:
            if t in self.nodos:
                self.nodos[t].grado = a["grado_propuesto"]
                self._agregar_conceptos(t, a.get("conceptos", []))

    def _fusionar(self, a: dict, afectados: list[str]) -> None:
        """El tema que queda hereda los conceptos de los absorbidos (y los que nombra la acción)."""
        if len(afectados) <= 1:
            return
        queda = afectados[0]
        heredadas = [x.model_copy(update={"origen": queda}) for x in self.aristas
                     if x.origen in afectados[1:] and x.tipo == TipoArista.TRABAJA]
        for t in afectados[1:]:
            self.nodos.pop(t, None)
        self.aristas = [x for x in self.aristas if x.origen in self.nodos and x.destino in self.nodos] + heredadas
        self._agregar_conceptos(queda, a.get("conceptos", []))
        if a.get("grado_propuesto") and queda in self.nodos:  # la fusión también puede cambiar de grado
            self.nodos[queda].grado = a["grado_propuesto"]

    def _revisar(self, a: dict, afectados: list[str]) -> None:
        """Reformular un tema: pasa a trabajar también los conceptos de la acción."""
        for t in afectados:
            self._agregar_conceptos(t, a.get("conceptos", []))

    def _dividir(self, a: dict, afectados: list[str]) -> None:
        """Los conceptos de la acción se van al tema nuevo (en su grado) y salen del original."""
        if not afectados:
            return
        self._sin_trabaja(afectados, {self.resolver(n) for n in a.get("conceptos", [])} - {None})
        tid = self._tema_nuevo(a, a.get("grado_propuesto") or self.nodos[afectados[0]].grado)
        self._agregar_conceptos(tid, a.get("conceptos", []))

    def _nuevo(self, a: dict, afectados: list[str]) -> None:
        if not a.get("grado_propuesto"):
            return
        tid = self._tema_nuevo(a, a["grado_propuesto"])
        for i, nombre in enumerate(a.get("conceptos", [])):
            cid = self.resolver(nombre)
            if cid:
                self.aristas.append(self._trabaja(tid, cid, "principal" if i == 0 else "secundario",
                                                  "Tema propuesto (simulación)."))
            else:
                self.sin_resolver.append(nombre)

    def _tema_nuevo(self, a: dict, grado: int) -> str:
        tid = f"TEMA:PROP-{self.asig}-{a['n']}"
        self.nodos[tid] = Nodo(id=tid, tipo=TipoNodo.TEMA, etiqueta=a.get("tema_propuesto", ""), asignatura=self.asig,
                               grado=grado)
        return tid


def _quitas_de_otras_asignaturas(asig: str, g0: int, g1: int, sim: _Simulacion) -> list[dict]:
    """Etiquetas que retiran las propuestas de OTRAS asignaturas del mismo ciclo (p. ej., Biología deja de tratar
    «Química de los carbohidratos» en 10.°): afectan el primer grado de conceptos de esta asignatura.

    Se aplican todas las que resuelven (retirar una etiqueta es local y seguro, y puede afectar a esta asignatura vía
    un concepto equivalente); solo se reportan las que tocan conceptos de esta asignatura o sus equivalentes. Los
    nombres que no resuelven los reporta la otra propuesta."""
    relevantes = {cid for cid, c in cp.vocabulario().items() if c["asignatura"] == asig}
    relevantes |= {x for e in cp.equivalencias() if {e["a"], e["b"]} & relevantes for x in (e["a"], e["b"])}
    cruzadas = []
    for otra in sorted(ruta("asignaturas").glob(f"*/propuesta/propuesta_G{g0:02d}-G{g1:02d}.json")):
        de = otra.parent.parent.name
        if de == asig:
            continue
        for a in json.loads(otra.read_text(encoding="utf-8"))["acciones"]:
            validos = [n for n in a.get("quitar_conceptos", []) if sim.resolver(n, de)]
            if validos:
                ids = {sim.resolver(n, de) for n in validos}
                cruzadas.append({**a, "accion": "_quitar", "origen": de, "quitar_conceptos": validos,
                                 "relevante": bool(ids & relevantes)})
    return cruzadas


def simular(asig: str, g0: int, g1: int, nodos: list[Nodo], aristas: list[Arista]) -> dict:
    """Aplica mover / nuevo / fusionar / dividir / revisar (y quitar_conceptos) al grafo en memoria y compara
    métricas antes y después."""
    p = json.loads((ruta(f"asignaturas/{asig}/propuesta") / f"propuesta_G{g0:02d}-G{g1:02d}.json")
                   .read_text(encoding="utf-8"))
    sim = _Simulacion(asig, nodos, aristas, p["version"])
    cruzadas = _quitas_de_otras_asignaturas(asig, g0, g1, sim)
    for a in cruzadas + p["acciones"]:
        sim.aplicar(a)
    nuevos = list(sim.nodos.values())
    antes = _metricas(asig, g0, g1, nodos, aristas)
    despues = _metricas(asig, g0, g1, nuevos, sim.aristas)
    ev = evidencia_orden(nuevos, sim.aristas)
    pendientes = [s for s in analizar(asig, nuevos, sim.aristas, ev)["secuencia"]
                  if g0 <= s["grado_concepto"] <= g1 and s["confianza"] == "alta"]
    return {"antes": antes, "despues": despues, "secuencia_alta_pendiente": pendientes,
            "conceptos_no_resueltos": sorted(set(sim.sin_resolver)), "etiquetas_retiradas": sim.retiradas,
            "quitas_de_otras_asignaturas": [{"asignatura": a["origen"], "accion": a["n"], "temas": a["temas_afectados"],
                                             "conceptos": a["quitar_conceptos"]} for a in cruzadas if a["relevante"]]}


# -- Resumen de todas las propuestas -----------------------------------------------------------------
NOMBRES = {"biologia": "Biología", "fisica": "Física", "quimica": "Química",
           "ciencias_tierra_espacio": "Tierra y Espacio"}


def escribir_resumen() -> str:
    filas = []
    for p in sorted(ruta("asignaturas").glob("*/propuesta/simulacion_G*.json")):
        asig, ciclo = p.parent.parent.name, p.stem.removeprefix("simulacion_")
        s = json.loads(p.read_text(encoding="utf-8"))
        pr = json.loads((p.parent / f"propuesta_{ciclo}.json").read_text(encoding="utf-8"))
        acc = Counter(a["accion"] for a in pr["acciones"])
        g0, g1 = (int(x) for x in ciclo.removeprefix("G").split("-G"))
        ant, des = s["antes"], s["despues"]
        tarde = "n/a" if g0 >= 10 else f"{ant['llega_tarde_2mas']} → **{des['llega_tarde_2mas']}**"
        desfase = "n/a" if g0 >= 10 else f"{ant['desfase_medio']} → **{des['desfase_medio']}**"
        filas.append(f"| {NOMBRES[asig]} | {g0}.°–{g1}.° | {len(pr['acciones'])} ("
                     + ", ".join(f"{k} {v}" for k, v in sorted(acc.items())) + ") | "
                     f"{ant['secuencia_alta_en_ciclo']} → **{des['secuencia_alta_en_ciclo']}** | {tarde} | {desfase} | "
                     f"[informe]({asig}/propuesta/informe_{ciclo}.md) · "
                     f"[Excel]({asig}/propuesta/Propuesta_{asig.capitalize()}_{ciclo}.xlsx) |")
    md = ["# Propuestas curriculares (borrador) · resumen", "",
          "Borradores por asignatura y ciclo, generados con el grafo de conocimiento (spec 06). Estado: **borrador**, "
          "pendiente de la revisión del equipo de Ciencias del MINED. Ninguno quita temas y todos conservan el "
          "total de su ciclo.", "",
          "**Verificación:** `gskg propuesta simular` aplica las acciones sobre el grafo (mover, fusionar, crear, "
          "reformular, dividir y retirar etiquetas que un tema reformulado deja de trabajar), recalcula el diagnóstico "
          "y compara antes y después. También aplica las etiquetas que retiran las propuestas de otras asignaturas del "
          "mismo ciclo (p. ej., Biología 10.° deja de tratar la química de los carbohidratos):", "",
          "- *Errores de secuencia (alta)*: prerrequisitos de confianza alta que la malla enseña después del concepto "
          "que los necesita. Los que quedan son decisiones abiertas en cada informe.",
          "- *Llegan ≥2 grados tarde*: conceptos del ciclo que El Salvador introduce dos o más grados después que la "
          "mediana de los seis países del grafo (Uruguay, Colombia, Singapur, Inglaterra, Australia y Japón), con "
          "grados equivalentes por edad. No aplica (n/a) en 10.°–11.°, porque los currículos de los países en el grafo "
          "llegan solo a 9.°.",
          "- *Desfase medio*: promedio de |grado SV − mediana de países| de los conceptos comparables del ciclo.", "",
          "| Asignatura | Ciclo | Acciones | Errores de secuencia (alta) | Llegan ≥2 grados tarde | Desfase medio "
          "| Documentos |", "|---|---|---|---|---|---|---|", *filas, "",
          "Notas:",
          *([] if ruta("asignaturas/quimica/propuesta/propuesta_G02-G04.json").exists() else [
              "- **Química 2.°–4.°** no tiene propuesta: la malla no tiene Química en 4.° y no hay conceptos que "
              "lleguen tarde."]),
          "- **Tierra y Espacio 5.°–8.°** casi no reduce los conceptos que llegan tarde: a propósito mueve la unidad "
          "del espacio de 5.° a 6.° para que la gravedad y las órbitas vayan antes.",
          "- Las propuestas de 2.°–8.° usan los 6 países (Singapur, Inglaterra, Australia y Japón son de alto "
          "desempeño); las de 10.°–11.° se basan en los marcos (ACARA y TIMSS Advanced).",
          "- Los errores de secuencia que quedan se explican en cada informe (introducciones cualitativas o etiquetas "
          "por revisar). Cada informe lista sus **decisiones abiertas para el MINED**.",
          "- Las propuestas de un ciclo afectan a los siguientes (por ejemplo, si un tema baja a 5.°, el de 9.° debe "
          "profundizar). Coordinarlas es parte de la revisión (ver `HANDOFF.md`).", ""]
    p = ruta("asignaturas/PROPUESTAS.md")
    p.write_text("\n".join(md), encoding="utf-8")
    return str(p.relative_to(ruta(".")))
