"""Generate international comparison reports and formula-based Excel workbooks."""

# Historical implementation notes:
# CSV de revisión humana.

from __future__ import annotations

import csv
import json
import re
import unicodedata
from collections import defaultdict
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as col

from goes_science_kg.config import cargar, ruta
from goes_science_kg.excel import guardar as guardar_excel
from goes_science_kg.internacional import contraste, etiquetar, profundidad, tramo, visor
from goes_science_kg.internacional.consenso import objetivos, paises

if TYPE_CHECKING:
    from pathlib import Path

NOMBRE = {
    "fisica": "Física",
    "quimica": "Química",
    "biologia": "Biología",
    "tierra_espacio": "Ciencias de la Tierra y del Espacio",
}


def clases(t: tramo.Tramo | None = None) -> list[tuple[str, str, str]]:
    """Describe the classifications valid within the active grade band."""

    # Implementation notes:
    # tramo ni electivas; en 9.°–11.° no hay nada después.
    t = t or tramo.actual()
    e, g0 = t.etiqueta, t.grados[0]
    salida = [
        (
            "faltante",
            "Faltantes del núcleo común",
            "El consenso (≥ 5 países) lo enseña a todos los estudiantes hasta ese grado; la V2 no lo tiene en 2.°–11.°.",
        ),
        (
            "no_retomado",
            f"No retomados en {e}",
            f"La V2 solo lo ve antes de {g0}.°; al menos 5 países lo profundizan en su núcleo de {e}.",
        ),
        (
            "tardio",
            "Llegan tarde",
            f"La V2 lo introduce en {e}, dos o más grados después de que el núcleo alcanza el consenso.",
        ),
        (
            "posterior",
            f"Solo después de {e}",
            f"La V2 lo enseña después de {e} y el núcleo internacional no lo exige antes.",
        ),
        (
            "solo_especializacion",
            "Nivel de especialización en un curso obligatorio",
            f"La V2 lo introduce en {e}, pero los países solo lo enseñan en cursos electivos.",
        ),
        (
            "adelantado",
            "Adelantados",
            f"La V2 lo introduce en {e}, 1,5 grados o más antes que la mediana del núcleo de los países.",
        ),
        (
            "sin_referente",
            "Sin referente",
            f"La V2 lo enseña en {e} y menos de {contraste.MIN_PAISES_MOMENTO} países lo tienen en su núcleo: no hay "
            "referente para fecharlo."
            if t.lectura_doble
            else f"La V2 lo enseña en {e} y ninguno de los 9 países lo tiene.",
        ),
        (
            "alineado",
            "Alineados",
            f"La V2 lo introduce en {e} con un momento coherente con el consenso (± 1,5 grados).",
        ),
        ("retomado", "Retomados", f"La V2 lo introduce antes de {g0}.° y lo vuelve a trabajar en {e} (espiral)."),
        (
            "previo",
            f"Solo antes de {g0}.°",
            f"La V2 lo ve antes de {g0}.° y el núcleo internacional no lo exige en {e}.",
        ),
    ]
    if t.lectura_doble:
        salida.append(
            (
                "incierto",
                "Inciertos (a revisión)",
                "La lectura amplia y la estricta del etiquetado dan clases distintas: no se afirma nada hasta "
                "revisarlo (revision_clases.csv).",
            )
        )
    imposibles = {"posterior"} if t.grados[-1] >= 11 else set()
    if t.grados[0] <= 2:
        imposibles |= {"no_retomado", "retomado", "previo", "solo_especializacion"}
    return [c for c in salida if c[0] not in imposibles]


SIN_TABLA = {"alineado", "retomado", "previo", "incierto"}


def _cita(o: dict) -> str:
    lugar = f"p. {o['pagina']}" if o.get("pagina") else (o.get("localizador") or "")
    nivel = "núcleo" if o["nivel"] == "nucleo" else "esp."
    return f"{o['pais']} · {o['curso']} [{nivel}] ({o['documento']}, {lugar}): «{o['texto']}»"


def _tema(t: dict) -> str:
    return f"{t['id']} ({t['archivo'].split('/')[-1]}, hoja «{t['hoja']}», fila {t['fila']})"


def _grados_sv(f: dict, r: dict) -> str:
    if not f["sv_primer_grado"]:
        return "no está"
    # Con doble lectura, el primer grado puede salir solo de la búsqueda de texto (sv_temas vacío, sv_temas_texto no).
    ts = [*f.get("sv_temas", []), *f.get("sv_temas_texto", [])]
    gs = sorted({r["temas"][t]["grado"] for t in ts if t in r["temas"]}) or [f["sv_primer_grado"]]
    return f"{gs[0]}.°" if len(gs) == 1 else f"{gs[0]}.°–{gs[-1]}.°"


def _tabla(filas: list[dict], r: dict, objs: dict, max_filas: int = 40) -> list[str]:
    lineas = [
        "| Concepto | Grados SV (V2) | Consenso del núcleo | Países (grado de 1.ª aparición; esp. = solo electiva) "
        "| Evidencia |",
        "|---|---|---|---|---|",
    ]
    for f in filas[:max_filas]:
        g = f["grado_consenso"]
        prop = f["proporcion_nucleo"].get(str(g)) if g else None
        med = f["mediana_primer_grado_nucleo"]
        cons = f"{g}.° ({prop:.0%} de los países con núcleo)" if g else (f"mediana {med:g}" if med else "—")
        cons += f"; {f['n_nucleo_en_tramo']} con núcleo en {tramo.actual().etiqueta}"
        ev = []
        if f["sv_temas"]:
            ev.append("SV: " + _tema(r["temas"][f["sv_temas"][0]]))
        citados = 0
        for d in f["paises"].values():
            if citados == 2:
                break
            if d["objetivos"] and d["objetivos"][0] in objs:
                ev.append(_cita(objs[d["objetivos"][0]]))
                citados += 1
        ev += [f"⚠ Revisar: {a}" for a in visor._alertas(f.get("alertas") or {})]
        if f.get("revision"):
            ev.append(f"Revisión experta: {f['revision']['justificacion']}")
        lineas.append(
            f"| {f['nombre']} | {_grados_sv(f, r)} | {cons} | {contraste.evidencia_paises(f)} | {'<br>'.join(ev)} |"
        )
    if len(filas) > max_filas:
        lineas.append(f"\n_… y {len(filas) - max_filas} más en el Excel._")
    return lineas


def _orden(f: dict) -> tuple:
    if f["clase"] in ("no_retomado", "faltante"):
        return (-f["n_nucleo_en_tramo"], f["nombre"])
    if f["clase"] == "solo_especializacion":
        return (-f["n_especializacion"], f["nombre"])
    g = f["grado_consenso"]
    return (-(f["proporcion_nucleo"].get(str(g), 0) if g else 0), f["nombre"])


def _sensibilidad(filas: list[dict]) -> str:
    """Measure classification sensitivity to the country-consensus threshold."""
    n4 = sum(f["clase"] == "previo" and f["n_nucleo_en_tramo"] == contraste.MIN_PAISES - 1 for f in filas)
    return (
        f"Sensibilidad: con 4 países en lugar de {contraste.MIN_PAISES}, {n4} conceptos más pasarían de «solo "
        f"antes de {tramo.actual().grados[0]}.°» a «no retomados»."
    )


def _seccion_profundidad(asig: str, prof: dict) -> list[str]:
    d = prof["asignaturas"].get(asig)
    if not d:
        return []
    md = [
        f"## Profundidad: demanda cognitiva en {tramo.actual().etiqueta}",
        "",
        "Mismo clasificador por verbos (escala TIMSS) para la V2 y para el núcleo de cada país; los enunciados sin "
        "verbo reconocible no entran en el porcentaje. Coincidencia del clasificador con la etiqueta de Gemini: "
        f"{prof['acuerdo_con_ia']['proporcion']:.0%} (la diferencia principal: TIMSS pone «explicar» en aplicar).",
        "",
        "| Currículo | Enunciados | Recordar | Aplicar | Razonar | Sin verbo |",
        "|---|---|---|---|---|---|",
    ]
    for quien in ["SV", *d["paises"]]:
        f = d["filas"].get(quien)
        if not f:
            continue
        con = sum(f["demanda"].get(n, 0) for n in profundidad.NIVELES) or 1
        pct = [f"{f['demanda'].get(n, 0) / con:.0%}" for n in profundidad.NIVELES]
        nombre = "**El Salvador (V2)**" if quien == "SV" else quien
        md.append(f"| {nombre} | {f['n']} | {' | '.join(pct)} | {f['demanda'].get('sin_verbo', 0)} |")
    return [*md, ""]


def _seccion_practicas(asig: str, prof: dict) -> list[str]:
    d = prof["asignaturas"].get(asig)
    if not d:
        return []
    md = [
        f"## Prácticas científicas en {tramo.actual().etiqueta}",
        "",
        "La V2 trae una columna procedimental en cada fila y los países separan prácticas de contenidos, así que la "
        "proporción de enunciados con práctica no es comparable. Se compara la cobertura del catálogo.",
        "",
        "| Currículo | Indagación | Modelación | Argumentación | Naturaleza y contexto | Prácticas distintas |",
        "|---|---|---|---|---|---|",
    ]
    for quien in ["SV", *d["paises"]]:
        f = d["filas"].get(quien)
        if not f:
            continue
        fam = [
            str(f["familias"].get(k, 0)) for k in ("indagacion", "modelacion", "argumentacion", "naturaleza_y_contexto")
        ]
        nombre = "**El Salvador (V2)**" if quien == "SV" else quien
        md.append(f"| {nombre} | {' | '.join(fam)} | {len(f['practicas'])} |")
    aus = d["practicas_ausentes"]
    md += [
        "",
        f"Prácticas que exige el núcleo de ≥ 5 países y la V2 no nombra en {tramo.actual().etiqueta}: "
        + (", ".join(f"`{p}`" for p in aus) if aus else "ninguna."),
        "",
    ]
    return md


def escribir() -> dict:
    """Write subject reports, formula-based workbooks, review CSVs, and the viewer."""

    # Implementation notes:
    # contraste.json. Devuelve el resumen por asignatura, los países, los casos de secuencia y las filas a revisar.
    r = contraste.clasificar()
    prof = profundidad.analizar()
    objs = {o["id"]: o for o in objetivos()}
    t = tramo.actual()
    cfg = cargar(t.config)
    base = ruta(t.informe)
    base.mkdir(parents=True, exist_ok=True)
    res = contraste.resumen(r)
    propuestos = _propuestos(objs)
    for asig in contraste.ASIGNATURAS:
        filas = [f for f in r["filas"] if f["asignatura"] == asig]
        md = [
            f"# {NOMBRE[asig]} · contraste internacional de {t.etiqueta} (malla V2)",
            "",
            f"> Generado por `gskg internacional construir`. Método y límites: `{t.spec}` y `{t.informe}/README.md`.",
            "> Los datos de los países los extrajo y validó Gemini (Vertex AI, proyecto GOES), con un modelo para",
            "> extraer y otro para validar. Las clases dependen de etiquetas de IA y requieren revisión del MINED.",
            "",
            "## Resumen",
            "",
            "| Clase | Conceptos | Qué significa |",
            "|---|---|---|",
        ]
        for clave, titulo, desc in clases(t):
            md.append(f"| {titulo} | {res[asig].get(clave, 0)} | {desc} |")
        md += [
            "",
            "Países con núcleo común por grado SV (denominador del consenso): "
            + "; ".join(f"{g}.°: {', '.join(r['paises_con_nucleo'][str(g)])}" for g in t.grados),
            "",
            _sensibilidad(filas),
            "",
        ]
        rev = [f for f in filas if f.get("revision")]
        if rev:
            md += [
                f"Revisión experta: {len(rev)} conceptos cambiaron de clase frente a la regla automática "
                f"(`{t.revisiones}`): "
                + "; ".join(f"{f['nombre']} ({f['clase_ia']} → {f['clase']})" for f in rev)
                + ".",
                "",
            ]
        inciertos = [f for f in filas if f["clase"] == "incierto"]
        if inciertos:
            md += [
                f"Inciertos: {len(inciertos)} conceptos tienen clases distintas en la lectura amplia y en la "
                "estricta; no se afirma nada de ellos hasta revisarlos (`revision_clases.csv`, con las dos clases).",
                "",
            ]
        for clave, titulo, desc in clases(t):
            if clave in SIN_TABLA:
                continue
            sel = sorted((f for f in filas if f["clase"] == clave), key=_orden)
            if sel:
                md += [f"## {titulo} ({len(sel)})", "", desc, "", *_tabla(sel, r, objs), ""]
        sec = [s for s in r["secuencia"] if s["concepto"] in {f["concepto"] for f in filas}]
        if sec:
            prudente = (
                " Lectura prudente: el concepto con presencia sólida en la V2 y el prerrequisito sin ninguna "
                "etiqueta hasta después; la alerta marca los temas cuyo texto ya nombra el prerrequisito."
                if t.lectura_doble
                else ""
            )
            md += [
                f"## Candidatos a error de secuencia en la V2 ({len(sec)})",
                "",
                "El prerrequisito llega después que el concepto que lo necesita (grafo de prerrequisitos del repo). "
                "Son candidatos: en la revisión manual de 2026-10-08, 2 de 4 resultaron falsos porque el etiquetado "
                f"no vio el prerrequisito en primaria. Los confirmados están en `{t.informe}/README.md`.{prudente}",
                "",
                "| Concepto (grado) | Prerrequisito que llega después (grado) | Confianza | Tema SV |",
                "|---|---|---|---|",
            ]
            for s in sec:
                alerta = f" ⚠ texto en {', '.join(s['alertas']['texto'])}" if s.get("alertas") else ""
                md.append(
                    f"| {s['nombre_concepto']} ({s['grado_concepto']}.°) | {s['nombre_prerrequisito']} "
                    f"({s['grado_prerrequisito']}.°) | {s['confianza']} | {s['tema_concepto']}{alerta} |"
                )
            md.append("")
        md += _seccion_profundidad(asig, prof) + _seccion_practicas(asig, prof)
        top = [p for p in propuestos.get(asig, []) if p["n_paises"] >= 3][:15]
        if top:
            md += [
                "## Conceptos que el vocabulario no tiene (propuestos para triaje)",
                "",
                "El vocabulario nació de la malla salvadoreña, así que un contenido que El Salvador no enseña puede no "
                "existir en él. Estos nombres los propuso el etiquetado en ≥ 3 países.",
                "",
                "| Propuesto | Países | Con núcleo |",
                "|---|---|---|",
            ]
            md += [f"| {p['nombre']} | {', '.join(p['paises'])} | {', '.join(p['nucleo']) or '—'} |" for p in top]
            md.append("")
        (base / f"{asig}.md").write_text("\n".join(md) + "\n", encoding="utf-8")
        _excel(asig, filas, base)
    revision = _revision(objs, base)
    if t.lectura_doble:
        revision["clases"] = _revision_clases(r["filas"], base)
    visor.escribir(
        r, [*clases(t), ("no_aplica", "Sin exigencia", "Ni la V2 ni el núcleo internacional lo exigen.")], NOMBRE, base
    )
    (base / "contraste.json").write_text(
        json.dumps(
            {
                "resumen": res,
                "filas": r["filas"],
                "secuencia": r["secuencia"],
                "profundidad": prof,
                "propuestos": propuestos,
            },
            ensure_ascii=False,
            indent=1,
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "resumen": res,
        "paises": list(cfg["paises"]),
        "secuencia": len(r["secuencia"]),
        "revision_humana": revision,
    }


def _norma(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFKD", s.lower()) if not unicodedata.combining(c))
    return " ".join(re.sub(r"\b(de|del|la|las|los|el|y|en|sus)\b", " ", s).split())


def _propuestos(objs: dict) -> dict[str, list[dict]]:
    """Group proposed concept names by subject without treating them as accepted identities."""
    por: dict[tuple[str, str], dict] = {}
    for i, e in etiquetar.cargar_etiquetas("pais_").items():
        o = objs.get(i)
        if not o or not o["asignatura"]:
            continue
        for p in e["propuestos"]:
            if p.startswith(("CON:", "PRAC:")):
                continue
            d = por.setdefault((o["asignatura"], _norma(p)), {"nombre": p, "paises": set(), "nucleo": set()})
            d["paises"].add(o["pais"])
            if o["nivel"] == "nucleo":
                d["nucleo"].add(o["pais"])
    salida: dict[str, list[dict]] = defaultdict(list)
    for (asig, _), d in sorted(por.items()):
        salida[asig].append(
            {
                "nombre": d["nombre"],
                "paises": sorted(d["paises"]),
                "nucleo": sorted(d["nucleo"]),
                "n_paises": len(d["paises"]),
            }
        )
    for v in salida.values():
        v.sort(key=lambda d: (-len(d["nucleo"]), -d["n_paises"], d["nombre"]))
    return dict(sorted(salida.items()))


def _revision(objs: dict, base: Path) -> dict:
    """Export partial coverage and low-confidence labels for human review."""
    etq = etiquetar.cargar_etiquetas("pais_")
    filas = []
    for i, o in sorted(objs.items()):
        e = etq.get(i, {})
        motivos = []
        if o["validacion"]["veredicto"] == "parcial":
            motivos.append(f"extracción parcial: {o['validacion']['motivo']}")
        if e.get("confianza") == "baja":
            motivos.append(f"etiqueta de confianza baja: {e.get('justificacion', '')}")
        if motivos:
            filas.append(
                {
                    "id": i,
                    "pais": o["pais"],
                    "documento": o["documento"],
                    "pagina": o.get("pagina") or "",
                    "curso": o["curso"],
                    "nivel": o["nivel"],
                    "texto": o["texto"],
                    "cita": o["cita"],
                    "conceptos": " ".join(e.get("principales", []) + e.get("secundarios", [])),
                    "motivo": " | ".join(motivos),
                    "decision": "",
                    "revisado_por": "",
                }
            )
    with (base / "revision_humana.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0]) if filas else ["id"])
        w.writeheader()
        w.writerows(filas)
    return {"filas": len(filas)}


def _revision_clases(filas: list[dict], base: Path) -> int:
    """Export uncertain dual-reading classifications and flagged findings for review."""

    # Implementation notes:
    # su clase a mano. La decisión va a revisiones.json del tramo con su evidencia.
    sel = [f for f in filas if f["clase"] == "incierto" or (f.get("alertas") and not f.get("revision"))]
    campos = [
        "concepto",
        "nombre",
        "asignatura",
        "clase",
        "lectura_amplia",
        "lectura_estricta",
        "sv_grado_amplio",
        "sv_grado_estricto",
        "consenso_amplio",
        "consenso_estricto",
        "mediana_paises",
        "mediana_paises_solida",
        "paises_nucleo",
        "paises_nucleo_solido",
        "temas_sv",
        "alertas",
        "decision",
        "comentario",
        "revisado_por",
    ]
    with (base / "revision_clases.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos)
        w.writeheader()
        for f in sorted(sel, key=lambda f: (f["asignatura"], -f["n_nucleo"], f["concepto"])):
            w.writerow(
                {
                    "concepto": f["concepto"],
                    "nombre": f["nombre"],
                    "asignatura": f["asignatura"],
                    "clase": f["clase"],
                    "lectura_amplia": f["lecturas"]["amplia"],
                    "lectura_estricta": f["lecturas"]["estricta"],
                    "sv_grado_amplio": f["sv_primer_grado"],
                    "sv_grado_estricto": f["sv_primer_grado_solido"],
                    "consenso_amplio": f["grado_consenso"],
                    "consenso_estricto": f["grado_consenso_solido"],
                    "mediana_paises": f["mediana_primer_grado_nucleo"],
                    "mediana_paises_solida": f["mediana_primer_grado_nucleo_solido"],
                    "paises_nucleo": f["n_nucleo"],
                    "paises_nucleo_solido": f["n_nucleo_solido"],
                    "temas_sv": " ".join(f["sv_temas"][:4]),
                    "alertas": "; ".join(visor._alertas(f.get("alertas") or {})),
                    "decision": "",
                    "comentario": "",
                    "revisado_por": "",
                }
            )
    return len(sel)


def _despues(i: int, t: tramo.Tramo) -> str:
    """Build the formula branch for content introduced after the active grade band."""
    if t.grados[-1] >= 11:
        return ""
    return f'IF(B{i}>{t.grados[-1]},IF(C{i}>=5,"SV tarde","despues del tramo"),'


def _excel(asig: str, filas: list[dict], base: Path) -> None:
    """Build source-data, comparison-formula, and classification-summary worksheets."""

    ps = paises()
    n = len(ps)
    wb = Workbook()
    d = wb.active
    d.title = "Datos"
    t = tramo.actual()
    doble = ["Grado SV solido", "Lectura amplia", "Lectura estricta"] if t.lectura_doble else []
    d.append(
        [
            "Concepto",
            "Clase (IA)",
            "Grado SV V2",
            f"Grado SV en {t.rango}",
            *[f"{p} nucleo" for p in ps],
            *[f"{p} nucleo {t.rango}" for p in ps],
            *[f"{p} esp" for p in ps],
            *doble,
        ]
    )
    for f in filas:
        fp = [f["paises"].get(p, {}) for p in ps]
        lect = f.get("lecturas") or {}
        extra = [f.get("sv_primer_grado_solido"), lect.get("amplia"), lect.get("estricta")] if t.lectura_doble else []
        d.append(
            [
                f["nombre"],
                f["clase"],
                f["sv_primer_grado"],
                f.get("sv_en_tramo"),
                *[x.get("primer_grado_nucleo") for x in fp],
                *[1 if x.get("nucleo_en_tramo") else None for x in fp],
                *[1 if x.get("en_especializacion") else None for x in fp],
                *extra,
            ]
        )
    a0, a1 = col(5), col(4 + n)
    b0, b1 = col(5 + n), col(4 + 2 * n)
    c0, c1 = col(5 + 2 * n), col(4 + 3 * n)
    c = wb.create_sheet("Contraste")
    c.append(
        [
            "Concepto",
            "Grado SV V2",
            "Paises con nucleo",
            "Mediana nucleo",
            "Paises con especializacion",
            "Diferencia SV - mediana",
            f"Grado SV en {t.rango}",
            f"Paises con nucleo en {t.rango}",
            "Lectura (regla simplificada)",
            "Clase oficial",
        ]
    )
    for i in range(2, len(filas) + 2):
        c.append(
            [
                f"=Datos!A{i}",
                f'=IF(Datos!C{i}="","",Datos!C{i})',
                f"=COUNT(Datos!{a0}{i}:{a1}{i})",
                f'=IF(C{i}=0,"",MEDIAN(Datos!{a0}{i}:{a1}{i}))',
                f"=COUNT(Datos!{c0}{i}:{c1}{i})",
                f'=IF(OR(B{i}="",D{i}=""),"",B{i}-D{i})',
                f'=IF(Datos!D{i}="","",Datos!D{i})',
                f"=COUNT(Datos!{b0}{i}:{b1}{i})",
                f'=IF(B{i}="",IF(C{i}>=5,"falta en SV",""),IF(G{i}="",{_despues(i, t)}IF(H{i}>=5,'
                f'"no retomado en {t.rango}","previo"){")" if t.grados[-1] < 11 else ""},'
                f'IF(B{i}<{t.grados[0]},"retomado",IF(AND(C{i}<=1,E{i}>=3),"solo especializacion",IF(F{i}="","coherente",'
                f'IF(F{i}<=-1.5,"SV adelantado",IF(F{i}>=2,"SV tarde","coherente")))))))',
                f"=Datos!B{i}",
            ]
        )
    r = wb.create_sheet("Resumen")
    r.append(["Clase", "Conceptos"])
    ultima = len(filas) + 1
    for clave, _, _ in clases(t):
        fila = r.max_row + 1
        r.append([clave, f"=COUNTIF(Datos!B2:B{ultima},A{fila})"])
    r.append(["Total", f"=SUM(B2:B{r.max_row})"])
    verde = PatternFill("solid", fgColor="1E4D3A")
    for hoja in (d, c, r):
        for celda in hoja[1]:
            celda.fill, celda.font = verde, Font(color="FFFFFF", bold=True)
        hoja.column_dimensions["A"].width = 48
        hoja.freeze_panes = "B2"
    guardar_excel(wb, base / f"Contraste_{asig}.xlsx")
