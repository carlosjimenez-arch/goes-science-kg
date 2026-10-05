"""Explorador HTML de la cobertura (filtros por grado, evaluadora, dominio y estado) → outputs/Explorador_Cobertura_Ciencias_SV.html

Uso:  python scripts/build_explorador.py
Lee los mismos datos que el libro (temas, clasificaciones TIMSS 2027 y PISA 2025, catálogos, presupuesto) y, para el índice
de apego, los valores RECALCULADOS del libro (outputs/Cobertura_TIMSS2027_Ciencias_SV_numbers.xlsx). Así el HTML y el Excel
dicen lo mismo. Archivo autocontenido: se abre con doble clic, sin internet salvo las fuentes tipográficas.
Reglas de conteo iguales al libro: un objetivo cuenta en un grado si un tema de ese grado lo trabaja como principal o secundario;
TIMSS 4.° → 2.°–4.°, TIMSS 8.° → 5.°–8.°, PISA → 7.°–9.°, Advanced → Física 10.°–11.°.
"""
import json, os, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from hojas import NOMBRES

OUT = 'outputs/Explorador_Cobertura_Ciencias_SV.html'
EXPORT = 'outputs/Cobertura_TIMSS2027_Ciencias_SV_numbers.xlsx'


def datos():
    T = json.load(open('data/interim/temas.json'))
    R = json.load(open('data/interim/clasificacion_2027.json'))
    cat27 = json.load(open('data/referencia/catalogo_timss2027.json'))
    ta = [c for c in json.load(open('data/referencia/catalogo_timss.json')) if c['marco'] == 'TA']
    catp = json.load(open('data/referencia/catalogo_pisa2025.json'))['elementos']
    ix = {(t['archivo'], t['hoja'], t['fila']): t['id'] for t in T}
    P = {ix[(r['archivo'], r['hoja'], r['fila'])]: r for r in json.load(open('data/referencia/clasificacion_pisa2025.json'))}
    pres = json.load(open('data/referencia/presupuesto_candidatos.json'))
    temas = []
    for t in T:
        if t['asignatura'] not in ('Ciencias', 'Física'):
            continue
        c = R.get(t['id'], {}); p = P.get(t['id'])
        temas.append({'id': t['id'], 'g': t['grado'], 'a': t['asignatura'], 'u': t['unidad'], 'p': t['procedimental'],
                      't1': c.get('obj1', ''), 't2': c.get('obj2', ''), 'tc': c.get('confianza', ''), 'tj': c.get('justificacion', ''),
                      'comp': (t.get('competencia_pisa') or '')[:1], 'micro': [x for f in t.get('micro_pisa', {}).values() for x in f],
                      **({'p1': p['cont1'], 'p2': p['cont2'], 'k': p['conocimiento'], 'pr': p['proc'], 'ep': p['epis'],
                          'cx': p['contexto'], 'ar': p['area'], 'dm': p['demanda'], 'pc': p['confianza']} if p else {})})
    obj = [{'c': c['codigo'], 'm': 'TIMSS 2027 4.°' if c['marco'] == 'T4_27' else 'TIMSS 2027 8.°', 'ev': 'TIMSS 2027',
            'd': c['dominio'], 'o': c['objetivo'], 'amb': c['ambiental'], 'inv': c['es_investigacion']} for c in cat27]
    obj += [{'c': c['codigo'], 'm': 'TIMSS Advanced 2015 Física', 'ev': 'TIMSS Advanced', 'd': c['dominio'], 'o': c['objetivo']} for c in ta]
    tipo = {'contenido': 'Contenido', 'procedimental': 'Procedimental', 'epistemico': 'Epistémico', 'subcompetencia': 'Subcompetencia'}
    obj += [{'c': e['codigo'], 'm': 'PISA 2025 · ' + tipo[e['tipo']], 'ev': 'PISA 2025',
             'd': e['grupo'] if e['tipo'] != 'subcompetencia' else e['competencia'], 'o': e['texto'], 'tp': e['tipo']}
            for e in catp if e['tipo'] in tipo]
    apego, entra = [], {}
    if os.path.exists(EXPORT):
        wb = openpyxl.load_workbook(EXPORT, data_only=True)
        ws = wb[NOMBRES['Apego_escenarios']]
        for r in range(8, 32):
            a, b, f_, g_ = (ws.cell(r, c).value for c in (1, 2, 6, 7))
            if a and isinstance(f_, (int, float)) and 'escenario' not in str(a):
                apego.append({'g': a, 'm': b, 'bal': ws.cell(r, 3).value, 'cog': ws.cell(r, 4).value,
                              'amp': ws.cell(r, 5).value if isinstance(ws.cell(r, 5).value, (int, float)) else None, 'i': f_, 'l': g_})
        wp = wb[NOMBRES['Presupuesto']]
        for r in range(1, wp.max_row + 1):
            if wp.cell(r, 16).value and wp.cell(r, 1).value:
                entra[wp.cell(r, 1).value] = wp.cell(r, 16).value
    pres = [{**c, 'entra': entra.get(c['id'], '')} for c in pres]
    return {'temas': temas, 'obj': obj, 'apego': apego, 'pres': pres}


HTML = r"""<title>Cobertura de Ciencias El Salvador</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Literata:opsz,wght@7..72,500;7..72,650&family=Source+Sans+3:wght@400;600;700&family=JetBrains+Mono:wght@400;600&display=swap">
<style>
/* Tablero de consulta: franja de filtros arriba, resumen (índice de apego) y luego el detalle en pestañas. Verde bosque del libro, sin azul. */
:root{
  --bg:#f5f6f2; --surface:#ffffff; --fg:#1d2621; --muted:#5d6a62; --line:#dfe3dc;
  --accent:#1e4d3a; --accent-soft:#e6efe9; --amber:#fff2cc; --amber-fg:#7a4a00;
  --st-rojo:#f3c1aa; --st-amar:#f7e19a; --st-vcl:#dcebd1; --st-ver:#a9d08e; --na:#ececea;
  --display:"Literata",Georgia,"Times New Roman",serif; --body:"Source Sans 3","Helvetica Neue",Arial,sans-serif; --mono:"JetBrains Mono",ui-monospace,Menlo,monospace;
}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){
  --bg:#121815; --surface:#1a221e; --fg:#e4ebe6; --muted:#9aa89f; --line:#2c3832;
  --accent:#8fcfac; --accent-soft:#213229; --amber:#3a2f14; --amber-fg:#f2c97a;
  --st-rojo:#6b3426; --st-amar:#5e4f1c; --st-vcl:#2b4130; --st-ver:#3f6b3b; --na:#232b27; color-scheme:dark}}
:root[data-theme="dark"]{
  --bg:#121815; --surface:#1a221e; --fg:#e4ebe6; --muted:#9aa89f; --line:#2c3832;
  --accent:#8fcfac; --accent-soft:#213229; --amber:#3a2f14; --amber-fg:#f2c97a;
  --st-rojo:#6b3426; --st-amar:#5e4f1c; --st-vcl:#2b4130; --st-ver:#3f6b3b; --na:#232b27; color-scheme:dark}
*{box-sizing:border-box}
body{background:var(--bg);color:var(--fg);font:15px/1.5 var(--body);margin:0}
.wrap{max-width:1240px;margin:0 auto;padding-inline:clamp(16px,3vw,32px);padding-block:24px 48px;display:grid;gap:22px}
header{display:grid;gap:6px}
.eyebrow{font:600 12px/1 var(--body);letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
h1{font:650 clamp(24px,3.4vw,34px)/1.15 var(--display);margin:0;text-wrap:balance;color:var(--accent)}
h2{font:600 19px/1.25 var(--display);margin:0;text-wrap:balance}
.lead{max-width:72ch;color:var(--muted);margin:0}
.filters{position:sticky;top:env(safe-area-inset-top,0px);z-index:5;background:var(--bg);padding-block:10px;border-bottom:1px solid var(--line);
  display:flex;flex-wrap:wrap;gap:10px 14px;align-items:end}
.f{display:grid;gap:4px;min-width:0}
.f label{font:600 11px/1 var(--body);letter-spacing:.06em;text-transform:uppercase;color:var(--muted)}
select,input[type=search]{font:inherit;color:var(--fg);background:var(--surface);border:1px solid var(--line);border-radius:6px;padding:6px 8px;min-width:0;max-width:100%}
input[type=search]{width:min(260px,100%)}
select:focus-visible,input:focus-visible,button:focus-visible{outline:2px solid var(--accent);outline-offset:2px}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{font:600 13px/1 var(--body);border:1px solid var(--line);background:var(--surface);color:var(--fg);border-radius:999px;padding:6px 11px;cursor:pointer}
.chip[aria-pressed="true"]{background:var(--accent);border-color:var(--accent);color:var(--bg)}
.tabs{display:flex;gap:4px;border-bottom:1px solid var(--line);flex-wrap:wrap}
.tab{font:600 14px/1 var(--body);background:none;border:0;border-bottom:3px solid transparent;color:var(--muted);padding:10px 12px;cursor:pointer}
.tab[aria-selected="true"]{color:var(--accent);border-bottom-color:var(--accent)}
.panel{display:grid;gap:14px;min-width:0}
.summary{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px}
.stat{background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:12px 14px;display:grid;gap:2px}
.stat b{font:650 24px/1.1 var(--display);font-variant-numeric:tabular-nums}
.stat span{color:var(--muted);font-size:13px}
.legend{display:flex;flex-wrap:wrap;gap:8px 14px;font-size:13px;color:var(--muted)}
.sw{display:inline-block;width:12px;height:12px;border-radius:3px;vertical-align:-1px;margin-right:5px;border:1px solid var(--line)}
.tbl{overflow-x:auto;background:var(--surface);border:1px solid var(--line);border-radius:8px}
table{border-collapse:collapse;width:100%;font-size:13.5px}
th,td{padding:7px 9px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}
th{font:600 12px/1.2 var(--body);letter-spacing:.03em;color:var(--muted);background:var(--surface);position:sticky;top:0}
td.n,th.n{text-align:center;font-variant-numeric:tabular-nums;white-space:nowrap}
td.code{font:500 12.5px/1.4 var(--mono);white-space:nowrap}
.na{background:var(--na);color:var(--muted)}
.c0{color:var(--muted)} .c1{background:var(--st-vcl)} .c2{background:var(--st-ver)}
.st{font:600 12px/1 var(--body);border-radius:999px;padding:4px 8px;white-space:nowrap;display:inline-block}
.st-rojo{background:var(--st-rojo)} .st-amar{background:var(--st-amar)} .st-vcl{background:var(--st-vcl)} .st-ver{background:var(--st-ver)}
.muted{color:var(--muted)} .small{font-size:12.5px}
.bar{height:8px;border-radius:4px;background:var(--line);overflow:hidden;min-width:80px}
.bar i{display:block;height:100%;background:var(--accent)}
.note{font-size:13px;color:var(--muted);max-width:90ch;margin:0}
.empty{padding:24px;text-align:center;color:var(--muted)}
.edit{background:var(--amber);color:var(--amber-fg);font-weight:600}
@media (max-width:640px){th,td{padding:6px}.hide-sm{display:none}}
@media (prefers-reduced-motion:no-preference){.chip,.tab{transition:background .15s,color .15s,border-color .15s}}
</style>
<div class="wrap">
  <header>
    <span class="eyebrow">Malla de Ciencia y Tecnología · 2.° a 9.° y Física de Bachillerato</span>
    <h1>¿Qué cubre la malla frente a TIMSS 2027 y PISA 2025?</h1>
    <p class="lead">Explorador del libro <span class="code">Cobertura_TIMSS2027_Ciencias_SV.xlsx</span>. Filtra por grado y por prueba; cada número sale de la misma clasificación del libro (asistida por IA y revisable por el equipo de Ciencias).</p>
  </header>

  <div class="filters" role="group" aria-label="Filtros">
    <div class="f"><label for="f-eval">Evaluación</label>
      <select id="f-eval"><option value="">Todas</option><option>TIMSS 2027</option><option>PISA 2025</option><option>TIMSS Advanced</option></select></div>
    <div class="f"><label>Grado</label><div class="chips" id="f-grados"></div></div>
    <div class="f"><label for="f-dom">Dominio / sistema</label><select id="f-dom"><option value="">Todos</option></select></div>
    <div class="f"><label for="f-est">Estado</label>
      <select id="f-est"><option value="">Todos</option><option value="No cubierto">No cubierto</option><option value="Débil (1 tema)">Débil</option><option value="Básica (2–3)">Básica</option><option value="Sólida (4+)">Sólida</option></select></div>
    <div class="f"><label for="f-q">Buscar</label><input type="search" id="f-q" placeholder="luz, célula, D3, T8_27-E4.1…"></div>
  </div>

  <div class="tabs" role="tablist">
    <button class="tab" role="tab" id="t-apego" aria-selected="true">Apego por grado</button>
    <button class="tab" role="tab" id="t-obj" aria-selected="false">Objetivos del marco</button>
    <button class="tab" role="tab" id="t-temas" aria-selected="false">Temas de la malla</button>
    <button class="tab" role="tab" id="t-pres" aria-selected="false">Presupuesto</button>
  </div>

  <section class="panel" id="p-apego"></section>
  <section class="panel" id="p-obj" hidden></section>
  <section class="panel" id="p-temas" hidden></section>
  <section class="panel" id="p-pres" hidden></section>
</div>
<script>
const D = __DATA__;
const GRADOS = [2,3,4,5,6,7,8,9,10,11];
const GL = g => g>=10 ? `Fís. ${g}.°` : `${g}.°`;
const CICLO = m => m.startsWith('TIMSS 2027 4') ? [2,3,4] : m.startsWith('TIMSS 2027 8') ? [5,6,7,8] : m.startsWith('TIMSS Advanced') ? [10,11] : [7,8,9];
const st = {eval:'', grados:new Set(), dom:'', est:'', q:'', tab:'apego'};
try { const s = JSON.parse(localStorage.getItem('cob-sv')||'{}'); if (s.tab) st.tab = s.tab; } catch(e) {}
const $ = id => document.getElementById(id);
const esc = s => String(s??'').replace(/[&<>"]/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const pct = x => x==null ? '—' : Math.round(x*100)+' %';

// conteo por objetivo y grado (mismas reglas que el libro)
function temasDe(o){
  return D.temas.filter(t => {
    if (o.ev==='TIMSS 2027' || o.ev==='TIMSS Advanced') return t.t1===o.c || t.t2===o.c;
    if (o.tp==='contenido') return t.p1===o.c || t.p2===o.c;
    if (o.tp==='procedimental') return (t.pr||[]).includes(o.c);
    if (o.tp==='epistemico') return (t.ep||[]).includes(o.c);
    if (o.tp==='subcompetencia') return t.a==='Ciencias' && t.micro.includes(o.c);
    return false; });
}
const PRE = D.obj.map(o => { const ts = temasDe(o), cic = CICLO(o.m);
  const por = {}; GRADOS.forEach(g => por[g] = cic.includes(g) ? ts.filter(t => t.g===g).length : null);
  const tot = cic.reduce((s,g)=> s + (por[g]||0), 0);
  const estado = tot===0?'No cubierto':tot===1?'Débil (1 tema)':tot<=3?'Básica (2–3)':'Sólida (4+)';
  return {...o, por, tot, estado, cic}; });
const STC = {'No cubierto':'st-rojo','Débil (1 tema)':'st-amar','Básica (2–3)':'st-vcl','Sólida (4+)':'st-ver'};

function initFilters(){
  const box = $('f-grados');
  GRADOS.forEach(g => { const b = document.createElement('button'); b.className='chip'; b.type='button'; b.textContent=GL(g);
    b.setAttribute('aria-pressed','false'); b.onclick=()=>{ st.grados.has(g)?st.grados.delete(g):st.grados.add(g); b.setAttribute('aria-pressed', st.grados.has(g)); render(); };
    box.appendChild(b); });
  [...new Set(D.obj.map(o=>o.d))].sort().forEach(d => $('f-dom').insertAdjacentHTML('beforeend', `<option>${esc(d)}</option>`));
  $('f-eval').onchange=e=>{st.eval=e.target.value;render()}; $('f-dom').onchange=e=>{st.dom=e.target.value;render()};
  $('f-est').onchange=e=>{st.est=e.target.value;render()}; $('f-q').oninput=e=>{st.q=e.target.value.toLowerCase();render()};
  ['apego','obj','temas','pres'].forEach(k => $('t-'+k).onclick=()=>{st.tab=k; try{localStorage.setItem('cob-sv',JSON.stringify({tab:k}))}catch(e){}; render();});
}
const gsel = () => st.grados.size ? [...st.grados] : null;

function viewApego(){
  const sel = gsel();
  const rows = D.apego.filter(a => {
    if (st.eval==='PISA 2025' && !a.m.startsWith('PISA')) return false;
    if (st.eval==='TIMSS 2027' && !a.m.startsWith('TIMSS 2027')) return false;
    if (st.eval==='TIMSS Advanced' && !a.m.startsWith('TIMSS Advanced')) return false;
    if (sel){ const n = a.g.match(/\d+/g)?.map(Number)||[]; const rango = a.g.startsWith('Ciclo') ? (n.length>1? Array.from({length:n[1]-n[0]+1},(_,i)=>n[0]+i):n) : n.map(x=> a.g.startsWith('Física')? x : x);
      if (!rango.some(x=>sel.includes(x))) return false; }
    return true; });
  const cic = D.apego.filter(a=>a.g.startsWith('Ciclo'));
  return `<div class="summary">${cic.map(a=>`<div class="stat"><span>${esc(a.g.replace(' · PISA',''))} · ${esc(a.m)}</span><b>${pct(a.i)}</b><span>índice de apego · ${esc(a.l)}</span></div>`).join('')}</div>
  <p class="note">Índice de apego = promedio de la coincidencia del reparto de contenidos, la coincidencia cognitiva y (en ciclos) la amplitud. 100 % = el reparto de la malla es igual al de la prueba. Valores recalculados del libro (hoja Apego_escenarios).</p>
  <div class="tbl"><table><thead><tr><th>Grado / ciclo</th><th>Marco</th><th class="n">Balance</th><th class="n">Cognitivo</th><th class="n hide-sm">Amplitud</th><th>Índice de apego</th><th>Lectura</th></tr></thead><tbody>
  ${rows.map(a=>`<tr><td>${esc(a.g)}</td><td class="muted small">${esc(a.m)}</td><td class="n">${pct(a.bal)}</td><td class="n">${pct(a.cog)}</td><td class="n hide-sm">${pct(a.amp)}</td>
     <td><div style="display:flex;gap:8px;align-items:center"><div class="bar" style="flex:1"><i style="width:${Math.round((a.i||0)*100)}%"></i></div><span class="n">${pct(a.i)}</span></div></td>
     <td><span class="st ${a.l==='Alto'?'st-ver':a.l==='Medio'?'st-amar':'st-rojo'}">${esc(a.l)}</span></td></tr>`).join('') || `<tr><td colspan="7" class="empty">Ningún grado coincide con los filtros.</td></tr>`}
  </tbody></table></div>`;
}

function viewObj(){
  const sel = gsel();
  const rows = PRE.filter(o => (!st.eval || o.ev===st.eval) && (!st.dom || o.d===st.dom) && (!st.est || o.estado===st.est)
    && (!sel || sel.some(g=>o.cic.includes(g))) && (!st.q || (o.c+' '+o.o+' '+o.d).toLowerCase().includes(st.q)));
  const gs = sel ? GRADOS.filter(g=>sel.includes(g)) : GRADOS.filter(g => rows.some(o=>o.cic.includes(g)));
  const cnt = k => rows.filter(o=>o.estado===k).length;
  return `<div class="summary">${['No cubierto','Débil (1 tema)','Básica (2–3)','Sólida (4+)'].map(k=>`<div class="stat"><span><i class="sw ${STC[k]}"></i>${k}</span><b>${cnt(k)}</b><span>de ${rows.length} objetivos</span></div>`).join('')}</div>
  <div class="legend"><span><i class="sw" style="background:var(--na)"></i>N/A: el grado no es del ciclo que evalúa ese marco</span><span>Número = temas de ese grado que trabajan el objetivo (principal o secundario)</span></div>
  <div class="tbl"><table><thead><tr><th>Código</th><th>Objetivo</th><th class="hide-sm">Dominio / sistema</th>${gs.map(g=>`<th class="n">${GL(g)}</th>`).join('')}<th class="n">En su ciclo</th><th>Estado</th></tr></thead><tbody>
  ${rows.map(o=>`<tr><td class="code">${esc(o.c)}</td><td>${esc(o.o)}<div class="muted small">${esc(o.m)}${o.amb?' · ambiental':''}${o.inv?' · investigaciones':''}</div></td><td class="hide-sm small">${esc(o.d)}</td>
    ${gs.map(g=>{const v=o.por[g]; return v==null?'<td class="n na small">N/A</td>':`<td class="n ${v===0?'c0':v<4?'c1':'c2'}">${v}</td>`}).join('')}
    <td class="n"><b>${o.tot}</b></td><td><span class="st ${STC[o.estado]}">${esc(o.estado)}</span></td></tr>`).join('') || `<tr><td colspan="9" class="empty">Ningún objetivo coincide con los filtros.</td></tr>`}
  </tbody></table></div>`;
}

function viewTemas(){
  const sel = gsel();
  const pisa = st.eval==='PISA 2025';
  const rows = D.temas.filter(t => (!sel || sel.includes(t.g)) && (!pisa || t.p1)
    && (!st.eval || st.eval==='PISA 2025' || (st.eval==='TIMSS Advanced') === (t.a==='Física'))
    && (!st.q || (t.id+' '+t.p+' '+t.u+' '+t.t1+' '+t.t2+' '+(t.p1||'')+' '+t.micro.join(' ')).toLowerCase().includes(st.q))).slice(0,600);
  return `<p class="note">Una fila = un procedimental de la malla. TIMSS: objetivo principal y secundario. PISA (solo 7.°–9.°): categoría de contenido, tipo de conocimiento, contexto y demanda; la competencia y las microhabilidades vienen de la propia malla. Se muestran hasta 600 temas; usa los filtros.</p>
  <div class="tbl"><table><thead><tr><th>Grado</th><th>Tema (procedimental)</th><th>TIMSS</th><th class="hide-sm">PISA</th><th class="hide-sm">Microhabilidades</th><th class="n">Confianza</th></tr></thead><tbody>
  ${rows.map(t=>`<tr><td class="n">${GL(t.a==='Física'?t.g:t.g)}</td><td>${esc(t.p)}<div class="muted small">${esc(t.u||'')} · ${esc(t.id)}</div></td>
    <td class="code">${esc(t.t1)}${t.t2?'<br>'+esc(t.t2):''}</td>
    <td class="hide-sm small">${t.p1?`<span class="code">${esc(t.p1)}${t.p2?' · '+esc(t.p2):''}</span><br>${esc(t.k)} · ${esc(t.cx)} · demanda ${esc((t.dm||'').toLowerCase())}`:'<span class="muted">—</span>'}</td>
    <td class="hide-sm code">${t.comp?'C'+t.comp+' · ':''}${esc(t.micro.join(' '))}</td>
    <td class="n small">${esc(t.tc)}${t.pc?' / '+esc(t.pc):''}</td></tr>`).join('') || `<tr><td colspan="6" class="empty">Ningún tema coincide con los filtros.</td></tr>`}
  </tbody></table></div>`;
}

function viewPres(){
  const sel = gsel();
  const ciclo = {'2-4':'TIMSS 2027 4.°','5-8':'TIMSS 2027 8.°','7-9':'PISA 2025'};
  const rows = D.pres.filter(c => (!st.eval || (st.eval==='PISA 2025') === (c.ciclo==='7-9')) && (!sel || sel.includes(c.grado))
    && (!st.q || (c.propuesta+' '+c.area+' '+c.objetivos.join(' ')).toLowerCase().includes(st.q)) && (!st.dom || c.area===st.dom));
  const n = a => rows.filter(c=>c.accion===a && c.entra.startsWith('Sí')).length;
  return `<div class="summary"><div class="stat"><span>Inserciones que entran</span><b>${n('insertar')}</b><span>temas nuevos propuestos</span></div>
    <div class="stat"><span>Fusiones que entran</span><b>${n('fusionar')}</b><span>temas repetidos que se unen</span></div>
    <div class="stat"><span>Reubicaciones</span><b>${rows.filter(c=>c.accion==='reubicar').length}</b><span>opcionales, no cambian el total</span></div></div>
  <p class="note">Se mantiene el número de temas de cada ciclo: lo que se inserta en el área que falta se paga fusionando temas del área que sobra. «Entra» = está dentro de lo que pide la brecha del ciclo; «Reserva» = alternativa si el equipo descarta una. Propuestas de IA para discutir con el equipo de Ciencias (columna «Decisión del equipo» en el libro).</p>
  <div class="tbl"><table><thead><tr><th>ID</th><th>Acción</th><th>Grado</th><th>Propuesta</th><th class="hide-sm">Área</th><th>¿Entra?</th></tr></thead><tbody>
  ${rows.map(c=>`<tr><td class="code">${esc(c.id)}</td><td>${esc(c.accion)}${c.motivo==='cobertura'?'<div class="muted small">por cobertura</div>':''}</td><td class="n">${c.grado}.°</td>
    <td>${esc(c.propuesta)}<div class="muted small">${esc(c.justificacion)} · ${esc(c.objetivos.join(', '))}${c.temas.length?' · temas '+esc(c.temas.join(', ')):''}</div></td>
    <td class="hide-sm small">${esc(c.area)}<div class="muted">${ciclo[c.ciclo]}</div></td>
    <td><span class="st ${c.entra.startsWith('Sí')?'st-ver':c.entra.startsWith('Opcional')?'st-vcl':''}">${esc(c.entra||'—')}</span></td></tr>`).join('') || `<tr><td colspan="6" class="empty">Ningún candidato coincide con los filtros.</td></tr>`}
  </tbody></table></div>`;
}

function render(){
  ['apego','obj','temas','pres'].forEach(k => { $('t-'+k).setAttribute('aria-selected', st.tab===k); $('p-'+k).hidden = st.tab!==k; });
  $('p-'+st.tab).innerHTML = {apego:viewApego, obj:viewObj, temas:viewTemas, pres:viewPres}[st.tab]();
}
initFilters(); render();
</script>
"""


if __name__ == '__main__':
    d = datos()
    page = HTML.replace('__DATA__', json.dumps(d, ensure_ascii=False).replace('</', '<\\/'))
    os.makedirs('outputs', exist_ok=True)
    open(OUT, 'w', encoding='utf-8').write('<!doctype html>\n<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
                                           + page.replace('<div class="wrap">', '</head><body><div class="wrap">', 1) + '</body></html>\n')
    print(f'→ {OUT} ({os.path.getsize(OUT) // 1024} KB) · temas {len(d["temas"])} · objetivos {len(d["obj"])} · apego {len(d["apego"])} · presupuesto {len(d["pres"])}')
