"""Paso 4 · Crea COPIAS anotadas de cada malla con lo que cubre y no cubre frente a TIMSS.

Para cada hoja de unidad agrega, a la derecha de la última columna, 7 columnas:
  TIMSS · objetivo principal | TIMSS · objetivo secundario | TIMSS · dominio | TIMSS · área temática |
  TIMSS · ¿cubre? | TIMSS · confianza | TIMSS · justificación
y una hoja nueva «Cobertura TIMSS» con el resumen por unidad.
Nunca toca data/raw/: escribe en outputs/mallas_anotadas/<archivo>_TIMSS.xlsx
Uso:  python scripts/anotar_mallas.py                 # = --marco 2023 → outputs/mallas_anotadas/
      python scripts/anotar_mallas.py --marco 2027    # → outputs/mallas_anotadas_2027/ (catálogo 2027 + TA 2015)
"""
import argparse, collections, copy, json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ap = argparse.ArgumentParser()
ap.add_argument('--marco', choices=['2023', '2027'], default='2023')
ap.add_argument('--out', help='carpeta de salida (por defecto la del marco)')
ARGS = ap.parse_args()
ES27 = ARGS.marco == '2027'
RAW = 'data/raw'
OUT = ARGS.out or ('outputs/mallas_anotadas_2027' if ES27 else 'outputs/mallas_anotadas')
temas = json.load(open('data/interim/temas.json'))
clas = json.load(open('data/interim/clasificacion_2027.json' if ES27 else 'data/interim/clasificacion.json'))
cat = {c['codigo']: c for c in json.load(open('data/referencia/catalogo_timss.json'))
       if not ES27 or c['marco'] == 'TA'}  # en 2027, Física sigue con TIMSS Advanced 2015
if ES27:
    cat.update({c['codigo']: c for c in json.load(open('data/referencia/catalogo_timss2027.json'))})
ETQ = 'TIMSS 2027' if ES27 else 'TIMSS'
NEW = [f'{ETQ} · objetivo principal', f'{ETQ} · objetivo secundario', f'{ETQ} · dominio', f'{ETQ} · área temática',
       f'{ETQ} · ¿cubre?', f'{ETQ} · confianza', f'{ETQ} · justificación']
WID = [14, 14, 22, 30, 22, 11, 45]
F = 'Arial'
FILL_NEW = PatternFill('solid', fgColor='E2EFDA')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def marco_txt(t):
    if t['asignatura'] == 'Ciencias':
        return f'{ETQ} 4.°' if t['grado'] <= 4 else f'{ETQ} 8.°'
    return 'TIMSS Advanced Física' if t['asignatura'] == 'Física' else None


def info(t):
    c = clas.get(t['id'])
    if t['asignatura'] in ('Química', 'Biología'):
        return ['', '', '', '', 'Sin marco TIMSS (Bachillerato)', '', '']
    if c is None:
        return ['PENDIENTE', '', '', '', 'Pendiente de clasificar', '', '']
    if c['obj1'] == 'FUERA':
        return ['FUERA', c['obj2'], 'Fuera del marco', '', f'No: fuera del marco {marco_txt(t)}', c['confianza'], c['justificacion']]
    o = cat[c['obj1']]
    return [c['obj1'], c['obj2'], o['dominio'], o['area'], f'Sí: {o["objetivo"]}', c['confianza'], c['justificacion']]


def main():
    os.makedirs(OUT, exist_ok=True)
    by_file = collections.defaultdict(list)
    for t in temas:
        by_file[(t['archivo'], t.get('ruta', t['archivo']))].append(t)
    for (f, ruta), ts in by_file.items():
        wb = openpyxl.load_workbook(os.path.join(RAW, ruta))
        by_sheet = collections.defaultdict(list)
        for t in ts:
            by_sheet[t['hoja']].append(t)
        resumen = []
        for hoja, lst in by_sheet.items():
            ws = wb[hoja]
            c0 = ws.max_column + 1
            ref = ws.cell(row=1, column=2)
            for j, h in enumerate(NEW):
                col = c0 + j
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = copy.copy(ref.font); cell.fill = copy.copy(ref.fill)
                cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); cell.border = BOX
                ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = WID[j]
            for t in lst:
                for j, v in enumerate(info(t)):
                    cell = ws.cell(row=t['fila'], column=c0 + j, value=v)
                    cell.font = Font(name=F, size=9); cell.fill = FILL_NEW; cell.border = BOX
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            for t in lst:
                i = info(t)
                resumen.append((t['grado'], t['asignatura'], t['unidad'], i[2] or i[4], i[0]))
        # hoja resumen por unidad
        hoja_res = f'Cobertura {ETQ}'
        if hoja_res in wb.sheetnames:
            del wb[hoja_res]
        rs = wb.create_sheet(hoja_res)
        rs.append(['Unidad', 'Temas', 'Cubren un objetivo TIMSS', 'Fuera del marco', 'Pendientes / sin marco', 'Dominios TIMSS que trabaja (n.º de temas)'])
        agg = collections.OrderedDict()
        for g, a, u, dom, o1 in resumen:
            k = (g, a, u)
            agg.setdefault(k, collections.Counter())
            agg[k]['_n'] += 1
            if o1 == 'FUERA':
                agg[k]['_fuera'] += 1
            elif o1 in ('', 'PENDIENTE'):
                agg[k]['_pend'] += 1
            else:
                agg[k]['_si'] += 1; agg[k][dom] += 1
        for (g, a, u), c in agg.items():
            doms = '; '.join(f'{d}: {n}' for d, n in c.most_common() if not d.startswith('_'))
            rs.append([u, c['_n'], c['_si'], c['_fuera'], c['_pend'], doms])
        for row in rs.iter_rows():
            for cell in row:
                cell.font = Font(name=F, size=10, bold=(cell.row == 1)); cell.border = BOX
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for col, w in zip('ABCDEF', [40, 8, 12, 10, 12, 60]):
            rs.column_dimensions[col].width = w
        name = f.replace('.xlsx', '_TIMSS2027.xlsx' if ES27 else '_TIMSS.xlsx')
        wb.save(os.path.join(OUT, name))
        print('→', name)


if __name__ == '__main__':
    main()
