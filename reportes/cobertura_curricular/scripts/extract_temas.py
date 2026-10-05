"""Paso 1 · Extrae cada tema (fila con «Procedimental») de las mallas en data/raw/ → data/interim/temas.json

Uso:  python scripts/extract_temas.py
Reglas:
- Hojas de unidad: las que empiezan con «Unidad» (2.°–9.°) o contienen «N.°» (Bachillerato: «Física 10.°»...).
- Encabezados en filas 1 y 2 (la fila 2 trae las 5 fases de la secuencia). Datos desde la fila 3.
- Unidad y Contenido vienen en celdas combinadas: se arrastran hacia abajo (forward-fill).
- Nunca se modifica data/raw/.
"""
import glob, json, os, re
import openpyxl

RAW, OUT = 'data/raw', 'data/interim/temas.json'
ASIG = {'Fisica': 'Física', 'Quimica': 'Química', 'Biologia': 'Biología'}
FASES = ('1. Aproximar', '2. Explorar', '3. Explicar', '4. Indagar', '5. Reforzar')  # microhabilidades PISA por fase (códigos E1–E6, D1–D4, I1–I5)


def clean(v):
    return re.sub(r'\s+', ' ', str(v or '')).strip()


def canon(name):
    """«Matriz_X(1).xlsx» o «Matriz_X (2).xlsx» → «Matriz_X.xlsx» (copias descargadas dos veces)."""
    return re.sub(r'\s*\(\d+\)(?=\.xlsx$)', '', name)


def elegir_archivos():
    """Si hay duplicados de una misma malla, usa el MÁS RECIENTE y avisa.
    El nombre canónico (sin «(1)») se conserva como 'archivo' para que la clasificación se reutilice."""
    grupos = {}
    for p in glob.glob(os.path.join(RAW, 'Matriz_*.xlsx')):
        grupos.setdefault(canon(os.path.basename(p)), []).append(p)
    elegidos = []
    for c, ps in sorted(grupos.items()):
        ps.sort(key=os.path.getmtime, reverse=True)
        if len(ps) > 1:
            print(f'AVISO duplicados de {c}: uso {os.path.basename(ps[0])}; ignoro {[os.path.basename(x) for x in ps[1:]]}')
        elegidos.append((c, ps[0]))
    return elegidos


def main():
    temas = []
    for f, path in elegir_archivos():
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if not (ws.title.startswith('Unidad') or re.search(r'\d+\.°', ws.title)):
                continue
            h1 = [c.value or '' for c in ws[1]]
            h2 = [c.value or '' for c in ws[2]]
            cols = [(str(a) or str(b)).split('\n')[0].strip() for a, b in zip(h1, h2)]
            for i, s in enumerate(h2):
                if s:
                    cols[i] = str(s).split('\n')[0].strip()
            unidad = contenido = None
            for fila, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                d = dict(zip(cols, row))
                if d.get('Unidad'):
                    unidad = clean(d['Unidad'])
                if d.get('Contenido'):
                    contenido = clean(d['Contenido'])
                if not d.get('Procedimental'):
                    continue
                m = re.search(r'_(\d+)(?:do|er|to|mo|vo|no)_', f)
                grado = int(m.group(1)) if m else int(re.search(r'(\d+)\.°', ws.title).group(1))
                asig = 'Ciencias' if m else ASIG[f.split('_')[2]]
                temas.append(dict(
                    archivo=f, ruta=os.path.basename(path), hoja=ws.title, fila=fila, grado=grado, asignatura=asig,
                    unidad=unidad, contenido=contenido,
                    subcontenido=clean(d.get('Subcontenidos')), procedimental=clean(d['Procedimental']),
                    indicador=clean(d.get('Indicadores de logro')), evidencia=clean(d.get('Evidencia de aprendizaje')),
                    habilidad_timss=clean(d.get('Habilidad TIMSS')), pisa=clean(d.get('Competencia PISA 2025'))[:1],
                    competencia_pisa=clean(d.get('Competencia PISA 2025')),
                    micro_pisa={f: re.findall(r'\b([EDI]\d)\s*·', str(d.get(f) or '')) for f in FASES}))
    for i, t in enumerate(temas, 1):
        t['id'] = f"T{i:04d}-G{t['grado']:02d}-{t['asignatura'][:3].upper()}"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(temas, open(OUT, 'w'), ensure_ascii=False, indent=0)
    print(f'{len(temas)} temas → {OUT}')


if __name__ == '__main__':
    main()
