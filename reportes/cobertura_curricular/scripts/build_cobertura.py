"""Construye el libro de cobertura TIMSS de la malla de Ciencias, parametrizado por marco.

Uso:  python scripts/build_cobertura.py                 # = --marco 2023 (libro v1)
      python scripts/build_cobertura.py --marco 2027    # libro TIMSS 2027
      python scripts/build_cobertura.py --marco 2023 --out /ruta/prueba.xlsx
Entradas por marco (ver MARCOS):
- 2023: catalogo_timss.json (T4/T8 2023 + TA 2015) · data/interim/clasificacion.json → outputs/Cobertura_TIMSS_Ciencias_SV.xlsx
- 2027: catalogo_timss2027.json (T4_27/T8_27) + TA 2015 de catalogo_timss.json · data/interim/clasificacion_2027.json
        → outputs/Cobertura_TIMSS2027_Ciencias_SV.xlsx
- temas.json (1.260 temas extraídos de las 11 mallas)
Todas las cifras de las hojas de análisis son fórmulas sobre la hoja Temas_TIMSS.
Con --marco 2023 el libro debe salir idéntico al v1 (no cambiar textos ni fórmulas de esa rama).
"""
import argparse, collections as C_, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

MARCOS = {
    '2023': {'cat': ['data/referencia/catalogo_timss.json'], 'clas': 'data/interim/clasificacion.json',
             'out': 'outputs/Cobertura_TIMSS_Ciencias_SV.xlsx', 'M4': 'T4', 'M8': 'T8',
             'meta_cog': {'T4': (0.40, 0.40, 0.20), 'T8': (0.35, 0.35, 0.30)},
             'nombre': {'T4': 'TIMSS 2023 4.°', 'T8': 'TIMSS 2023 8.°'}},
    '2027': {'cat': ['data/referencia/catalogo_timss2027.json', 'data/referencia/catalogo_timss.json'],
             'clas': 'data/interim/clasificacion_2027.json',
             'out': 'outputs/Cobertura_TIMSS2027_Ciencias_SV.xlsx', 'M4': 'T4_27', 'M8': 'T8_27',
             'meta_cog': {'T4_27': (0.35, 0.40, 0.25), 'T8_27': (0.35, 0.35, 0.30)},  # TIMSS 2027, Exhibit 2.2, p. 24
             'nombre': {'T4_27': 'TIMSS 2027 4.°', 'T8_27': 'TIMSS 2027 8.°'}},
}
ap = argparse.ArgumentParser()
ap.add_argument('--marco', choices=sorted(MARCOS), default='2023')
ap.add_argument('--out', help='ruta de salida (por defecto la del marco)')
ARGS = ap.parse_args()
MC = MARCOS[ARGS.marco]
M4, M8 = MC['M4'], MC['M8']
ES27 = ARGS.marco == '2027'

T = json.load(open('data/interim/temas.json'))
if ES27:  # catálogo 2027 para 4.°/8.° y TIMSS Advanced 2015 para Física (sin versión 2027)
    CAT = json.load(open(MC['cat'][0])) + [c for c in json.load(open(MC['cat'][1])) if c['marco'] == 'TA']
else:
    CAT = json.load(open(MC['cat'][0]))
R = json.load(open(MC['clas']))
if ES27:  # la clasificación y el catálogo 2023 viajan en el libro 2027 para comparar con fórmulas («Cambio vs 2023»)
    V1 = json.load(open(MARCOS['2023']['clas']))
    CAT23 = json.load(open(MARCOS['2023']['cat'][0]))
    R23 = lambda col: f"Referencia_2023!${col}$5:${col}${4 + len(CAT23)}"
# PISA 2025 (7.°–9.°): catálogo y clasificación propia; 9.° deja TIMSS y se compara solo con PISA (decisión del equipo GOES)
HAY_PISA = ES27 and os.path.exists('data/referencia/clasificacion_pisa2025.json')
if HAY_PISA:
    CATP_J = json.load(open('data/referencia/catalogo_pisa2025.json'))
    CATP = CATP_J['elementos']
    _ix = {(t['archivo'], t['hoja'], t['fila']): t['id'] for t in T}
    RP = {_ix[(x['archivo'], x['hoja'], x['fila'])]: x for x in json.load(open('data/referencia/clasificacion_pisa2025.json'))}
    SISTEMAS = list(CATP_J['metas']['sistema'])
    FASES = ('1. Aproximar', '2. Explorar', '3. Explicar', '4. Indagar', '5. Reforzar')

F = 'Arial'
fN = Font(name=F, size=10)
fB = Font(name=F, size=10, bold=True)
fH = Font(name=F, size=10, bold=True, color='FFFFFF')
fT = Font(name=F, size=14, bold=True)
fS = Font(name=F, size=11, bold=True)
fIn = Font(name=F, size=10, color='0000FF')          # inputs
fLink = Font(name=F, size=10, color='008000')        # cross-sheet links
fNote = Font(name=F, size=9, italic=True, color='555555')
HDR = PatternFill('solid', fgColor='1F4E79')
SUB = PatternFill('solid', fgColor='DDEBF7')
INP = PatternFill('solid', fgColor='FFF2CC')
GRY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
if ES27:  # estilo del libro 2027: sin azules (pedido del equipo GOES); el libro 2023 conserva el estilo v1
    VERDE, VERDE_CLARO, BANDA = '1E4D3A', 'E6EFE9', 'F6F7F5'
    fT = Font(name=F, size=14, bold=True, color=VERDE)
    fS = Font(name=F, size=11, bold=True, color=VERDE)
    fIn = Font(name=F, size=10, bold=True, color='7A4A00')   # valores editables: texto café sobre celda ámbar
    fLink = Font(name=F, size=10, color='2E6B30')            # vínculos a otras hojas
    HDR = PatternFill('solid', fgColor=VERDE)
    SUB = PatternFill('solid', fgColor=VERDE_CLARO)
    thin = Side(style='thin', color='D9D9D9')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)
PCT = '0.0%'
PP = '+0.0" pp";-0.0" pp";0.0" pp"'
NUM = '0;-0;0'
NUM1 = '+0;-0;0'

def grp(t):
    return f"{t['grado']}°" if t['asignatura'] == 'Ciencias' else f"{t['asignatura']} {t['grado']}°"

def ciclo(t):
    g, a = t['grado'], t['asignatura']
    if a == 'Ciencias':
        if g <= 4: return 'Ciclo 2°–4°'
        if g <= 8: return 'Ciclo 5°–8°'
        return '9°'
    if a == 'Física': return 'Física 10°–11°'
    return f'{a} 10°–11°'

def marco(t):
    g, a = t['grado'], t['asignatura']
    if a == 'Ciencias' and g == 9 and HAY_PISA: return 'PISA'
    if a == 'Ciencias': return M4 if g <= 4 else M8
    if a == 'Física': return 'TA'
    return 'Sin marco'

def hdr(ws, row, headers, widths=None, height=30):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = fH; c.fill = HDR; c.alignment = CTR; c.border = BOX
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

def style_range(ws, r1, r2, c1, c2, fmt=None, font=fN, fill=None, align=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font; cell.border = BOX
            if fmt: cell.number_format = fmt
            if fill: cell.fill = fill
            if align: cell.alignment = align

wb = Workbook()

# ------------------------------------------------------------------ Catalogo_TIMSS
wc = wb.active; wc.title = 'Catalogo_TIMSS'
wc['A1'] = 'Catálogo de referencia TIMSS (objetivos de contenido)'; wc['A1'].font = fT
wc['A2'] = ('T4 = TIMSS 2023 4.° grado · T8 = TIMSS 2023 8.° grado · TA = TIMSS Advanced 2015 Física. '
            'Objetivos parafraseados al español (no son traducción oficial). Meta = % de puntos de la prueba asignado al dominio en el marco.')
if ES27:
    wc['A2'] = ('T4_27 = TIMSS 2027 4.° grado · T8_27 = TIMSS 2027 8.° grado · TA = TIMSS Advanced 2015 Física (no hay versión 2027). '
                'Objetivos parafraseados al español (no son traducción oficial). Meta = % de puntos de la prueba asignado al dominio en el marco.')
wc['A2'].font = fNote; wc.merge_cells('A2:H2'); wc.row_dimensions[2].height = 30; wc['A2'].alignment = WR
hdr(wc, 4, ['Código', 'Marco', 'Dominio', 'Meta del dominio', 'Código área', 'Área temática', 'Objetivo (paráfrasis)', 'Fuente'],
    [12, 8, 26, 12, 11, 40, 70, 30])
if ES27:
    hdr(wc, 4, ['Código', 'Marco', 'Dominio', 'Meta del dominio', 'Código área', 'Área temática', 'Objetivo (paráfrasis)', 'Fuente',
                'Ambiental (asterisco en el marco)', 'Área de Investigaciones', 'Sub-ítems (paráfrasis)'],
        [12, 8, 22, 10, 11, 34, 52, 30, 12, 13, 80])
src = {'T4': 'IEA, TIMSS 2023 Frameworks, cap. 2 (Science), 4.° grado',
       'T8': 'IEA, TIMSS 2023 Frameworks, cap. 2 (Science), 8.° grado',
       'TA': 'IEA, TIMSS Advanced 2015 Frameworks, cap. 2 (Physics)'}
if ES27:
    src[M4] = 'IEA, TIMSS 2027 Assessment Frameworks, cap. 2 (Science), 4.° grado'
    src[M8] = 'IEA, TIMSS 2027 Assessment Frameworks, cap. 2 (Science), 8.° grado'
for i, c in enumerate(CAT, 5):
    fuente = src[c['marco']] + (f", p. {c['pagina_fuente']}" if 'pagina_fuente' in c else '')
    vals = [c['codigo'], c['marco'], c['dominio'], c['meta_dominio'], c['area_codigo'], c['area'], c['objetivo'], fuente]
    for j, v in enumerate(vals, 1):
        wc.cell(row=i, column=j, value=v)
    style_range(wc, i, i, 1, 8, align=WR)
    if ES27:
        wc.cell(row=i, column=9, value='Sí' if c.get('ambiental') else 'No')
        wc.cell(row=i, column=10, value='Sí' if c.get('es_investigacion') else 'No')
        wc.cell(row=i, column=11, value='\n'.join(f"{x['letra']}. {x['texto']}{' *' if x['ambiental'] else ''}" for x in c.get('subitems', [])))
        style_range(wc, i, i, 9, 11, align=WR)
        wc.cell(row=i, column=9).alignment = CTR; wc.cell(row=i, column=10).alignment = CTR
    wc.cell(row=i, column=4).number_format = '0%'; wc.cell(row=i, column=4).font = fIn
CAT_LAST = 4 + len(CAT)
wc.freeze_panes = 'A5'; wc.auto_filter.ref = f'A4:{"K" if ES27 else "H"}{CAT_LAST}'
CR = lambda col: f"Catalogo_TIMSS!${col}$5:${col}${CAT_LAST}"

# ------------------------------------------------------------------ Catalogo_PISA (libro 2027 con PISA)
if HAY_PISA:
    wcp = wb.create_sheet('Catalogo_PISA')
    wcp['A1'] = 'Catálogo de referencia PISA 2025 (Ciencias)'; wcp['A1'].font = fT
    wcp['A2'] = ('Fuente: OECD (2026), PISA 2025 Assessment and Analytical Framework, cap. 2 (data/referencia/externos/pisa/PISA2025_Marco_Evaluacion_2026.pdf). '
                 'Paráfrasis en español, con página. Contenido = Box 2.6; procedimental = Box 2.7; epistémico = Box 2.8; subcompetencias = Boxes 2.3–2.5 '
                 '(mismos códigos E/D/I que usa la malla); áreas de aplicación = Tabla 2.2; demanda cognitiva = Figura 2.2. Metas: Tablas 2.3 y 2.4 (p. 53).')
    wcp['A2'].font = fNote; wcp.merge_cells('A2:G2'); wcp['A2'].alignment = WR; wcp.row_dimensions[2].height = 44
    hdr(wcp, 4, ['Código', 'Tipo', 'Grupo (sistema, competencia o bloque)', 'Descripción (paráfrasis)', 'Página', 'Meta del grupo', 'Rango del marco'],
        [11, 15, 34, 80, 8, 10, 14], height=36)
    TIPO_N = {'contenido': 'Contenido', 'procedimental': 'Procedimental', 'epistemico': 'Epistémico', 'subcompetencia': 'Subcompetencia',
              'area_aplicacion': 'Área de aplicación', 'demanda': 'Demanda cognitiva'}
    for i, e in enumerate(CATP, 5):
        meta = rango = None
        if e['tipo'] == 'contenido':
            meta = e['meta_grupo']; rango = 'Tabla 2.3: ' + str(round(e['meta_grupo'] * 100)) + ' %'
        if e['tipo'] == 'subcompetencia':
            lo, hi = e['meta_grupo']; meta = None; rango = f'Tabla 2.4: {round(lo * 100)}–{round(hi * 100)} %'
        grupo = e['grupo'] if e['tipo'] != 'subcompetencia' else f"{e['grupo']} · {e['competencia']}"
        for j, v in enumerate([e['codigo'], TIPO_N[e['tipo']], grupo, e['texto'], e['pagina'], meta, rango], 1):
            if v is not None:
                wcp.cell(row=i, column=j, value=v)
        style_range(wcp, i, i, 1, 7, align=WR)
        if meta is not None:
            wcp.cell(row=i, column=6).number_format = '0%'; wcp.cell(row=i, column=6).font = fIn
    CATP_LAST = 4 + len(CATP)
    wcp.freeze_panes = 'A5'; wcp.auto_filter.ref = f'A4:G{CATP_LAST}'
    CP = lambda col: f"Catalogo_PISA!${col}$5:${col}${CATP_LAST}"

# ------------------------------------------------------------------ Temas_TIMSS (trazabilidad)
wt = wb.create_sheet('Temas_TIMSS')  # en el libro 2027 se llama «Temas» (ver pasada final)
wt['A1'] = 'Trazabilidad: cada tema de las mallas y su clasificación contra TIMSS'; wt['A1'].font = fT
wt['A2'] = ('Una fila = un procedimental de la malla (un «tema»). Columnas K-L, P-Q: clasificación asistida por IA (revisable). '
            'M, N, O y S-U son fórmulas. Para auditar cualquier número del libro, filtra esta hoja por Grupo y Dominio.')
wt['A2'].font = fNote; wt.merge_cells('A2:W2'); wt['A2'].alignment = WR; wt.row_dimensions[2].height = 28
H = ['ID tema', 'Grado', 'Asignatura', 'Archivo de origen', 'Hoja', 'Fila', 'Unidad', 'Contenido', 'Procedimental (tema)',
     'Marco', 'Objetivo TIMSS principal', 'Objetivo TIMSS secundario', 'Dominio (principal)', 'Área temática (principal)',
     'Dominio (secundario)', 'Confianza', 'Justificación de la clasificación', 'Habilidad TIMSS (texto de la malla)',
     'Conocer (fracción)', 'Aplicar (fracción)', 'Razonar (fracción)', 'Grupo', 'Ciclo']
hdr(wt, 4, H, [17, 7, 10, 30, 22, 6, 28, 26, 55, 9, 13, 13, 22, 32, 20, 10, 45, 45, 9, 9, 9, 12, 15], height=42)
NCOL_T = 23
if ES27:
    H27 = ['Ambiental (obj. principal o secundario: 1 = sí)', 'Investigaciones como principal (1 = sí)',
           'Investigaciones como secundario (1 = sí)', 'Objetivo principal v1 (TIMSS 2023)',
           'Objetivo secundario v1 (TIMSS 2023)', 'Dominio principal v1 (TIMSS 2023)']
    hdr(wt, 4, H + H27, [17, 7, 10, 30, 22, 6, 28, 26, 55, 9, 13, 13, 22, 32, 20, 10, 45, 45, 9, 9, 9, 12, 15, 11, 12, 12, 13, 13, 22], height=54)
    NCOL_T = 29
    if HAY_PISA:
        HP = ['Competencia PISA (malla)', 'Competencia (código)'] + [f'Microhabilidad PISA · {f[3:]}' for f in FASES] + \
             ['Contenido PISA principal', 'Contenido PISA secundario', 'Sistema PISA (principal)', 'Tipo de conocimiento predominante',
              'Conocimiento procedimental (códigos)', 'Conocimiento epistémico (códigos)', 'Contexto PISA', 'Área de aplicación',
              'Demanda cognitiva', 'Confianza PISA', 'Justificación PISA', 'Ciclo PISA', 'Microhabilidades PISA (las 5 fases)']
        hdr(wt, 4, H + H27 + HP, [17, 7, 10, 30, 22, 6, 28, 26, 55, 9, 13, 13, 22, 32, 20, 10, 45, 45, 9, 9, 9, 12, 15, 11, 12, 12, 13, 13, 22,
                                  40, 9, 9, 9, 9, 9, 9, 11, 11, 22, 14, 18, 18, 14, 22, 11, 10, 40, 12, 18], height=54)
        NCOL_T = 49
    wt['A2'] = ('Una fila = un procedimental de la malla (un «tema»). Columnas K-L, P-Q: clasificación TIMSS 2027 asistida por IA (revisable). '
                'M, N, O, S-U y X-Z son fórmulas. AA-AB = clasificación v1 (TIMSS 2023), solo para comparar. '
                'Para auditar cualquier número del libro, filtra esta hoja por Grupo y Dominio.')
order = sorted(T, key=lambda t: ({'Ciencias': 0, 'Física': 1, 'Química': 2, 'Biología': 3}[t['asignatura']], t['grado'], t['id']))
r = 5
for t in order:
    c = R.get(t['id'])
    m = marco(t)
    pend = (c is None) and t['asignatura'] in ('Ciencias', 'Física')
    o1 = c['obj1'] if c else ('PENDIENTE' if pend else 'N/A')
    o2 = c['obj2'] if c else ''
    vals = [t['id'], t['grado'], t['asignatura'], t['archivo'], t['hoja'], t['fila'], t['unidad'], t['contenido'],
            t['procedimental'], m, o1, o2, None, None, None, c['confianza'] if c else 'N/A',
            c['justificacion'] if c else ('Tema nuevo o modificado: falta clasificar' if pend else 'Sin marco TIMSS de contenido para esta asignatura en Bachillerato'),
            t['habilidad_timss'], None, None, None, grp(t), ciclo(t)]
    for j, v in enumerate(vals, 1):
        if v is not None:
            wt.cell(row=r, column=j, value=v)
    wt[f'M{r}'] = (f'=IF(K{r}="FUERA","Fuera del marco",IF(K{r}="N/A","Sin marco",'
                   f'IFERROR(INDEX({CR("C")},MATCH(K{r},{CR("A")},0)),"Pendiente de clasificar")))')
    wt[f'N{r}'] = (f'=IF(K{r}="FUERA","Fuera del marco",IF(K{r}="N/A","Sin marco",'
                   f'IFERROR(INDEX({CR("F")},MATCH(K{r},{CR("A")},0)),"Pendiente de clasificar")))')
    wt[f'O{r}'] = f'=IF(L{r}="","",IFERROR(INDEX({CR("C")},MATCH(L{r},{CR("A")},0)),"?"))'
    k = (f'(ISNUMBER(SEARCH("Conocer",$R{r}))+ISNUMBER(SEARCH("Aplicar",$R{r}))+ISNUMBER(SEARCH("Razonar",$R{r})))')
    wt[f'S{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Conocer",$R{r}))/{k},0)'
    wt[f'T{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Aplicar",$R{r}))/{k},0)'
    wt[f'U{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Razonar",$R{r}))/{k},0)'
    if ES27:
        v1 = V1.get(t['id'])
        wt[f'AA{r}'] = v1['obj1'] if v1 else ('PENDIENTE' if pend else 'N/A')
        wt[f'AB{r}'] = v1['obj2'] if v1 else ''
        si = lambda col, ref: f'IFERROR(INDEX({CR(col)},MATCH({ref},{CR("A")},0)),"No")="Sí"'
        wt[f'X{r}'] = f'=IF(OR({si("I", f"K{r}")},{si("I", f"L{r}")}),1,0)'
        wt[f'Y{r}'] = f'=IF({si("J", f"K{r}")},1,0)'
        wt[f'Z{r}'] = f'=IF({si("J", f"L{r}")},1,0)'
        wt[f'AC{r}'] = (f'=IF(AA{r}="FUERA","Fuera del marco",IF(AA{r}="N/A","Sin marco",'
                        f'IFERROR(INDEX({R23("C")},MATCH(AA{r},{R23("A")},0)),"Pendiente de clasificar")))')
    if HAY_PISA:
        if t.get('competencia_pisa'):
            wt[f'AD{r}'] = t['competencia_pisa']
            for j, f in enumerate(FASES):
                wt.cell(row=r, column=32 + j, value=', '.join(t['micro_pisa'][f]))
        wt[f'AW{r}'] = f'=AF{r}&"; "&AG{r}&"; "&AH{r}&"; "&AI{r}&"; "&AJ{r}'
        wt[f'AE{r}'] = f'=IF(AD{r}="","",IF(LEFT(AD{r},1)="1","C1",IF(LEFT(AD{r},1)="2","C2",IF(LEFT(AD{r},1)="3","C3",""))))'
        p_ = RP.get(t['id'])
        if p_:
            for col, v in zip(['AK', 'AL', 'AN', 'AO', 'AP', 'AQ', 'AR', 'AS', 'AT', 'AU'],
                              [p_['cont1'], p_['cont2'], p_['conocimiento'], ''.join(x + '; ' for x in p_['proc']).strip(),
                               ''.join(x + '; ' for x in p_['epis']).strip(), p_['contexto'], p_['area'], p_['demanda'],
                               p_['confianza'], p_['justificacion']]):
                if v:
                    wt[f'{col}{r}'] = v
            wt[f'AM{r}'] = f'=IF(AK{r}="FUERA","Fuera del marco",IFERROR(INDEX({CP("C")},MATCH(AK{r},{CP("A")},0)),"?"))'
            wt[f'AV{r}'] = 'Ciclo 7°–9°'
    style_range(wt, r, r, 1, NCOL_T)
    for col in 'STU':
        wt[f'{col}{r}'].number_format = '0.00'
    for col in 'IQR':
        wt[f'{col}{r}'].alignment = Alignment(wrap_text=False, vertical='top')
    r += 1
TL = r - 1
wt.freeze_panes = 'C5'; wt.auto_filter.ref = f'A4:{get_column_letter(NCOL_T)}{TL}'
wt.conditional_formatting.add(f'P5:P{TL}', CellIsRule(operator='equal', formula=['"baja"'], fill=PatternFill('solid', fgColor='F8CBAD')))
wt.conditional_formatting.add(f'K5:K{TL}', CellIsRule(operator='equal', formula=['"FUERA"'], fill=GRY))
TC = lambda col: f"Temas_TIMSS!${col}$5:${col}${TL}"

# ------------------------------------------------------------------ Balance_grado
GROUPS = [('2°', M4), ('3°', M4), ('4°', M4), ('5°', M8), ('6°', M8), ('7°', M8), ('8°', M8), ('9°', M8),
          ('Física 10°', 'TA'), ('Física 11°', 'TA')]
DOMS = {}
for c in CAT:
    DOMS.setdefault(c['marco'], [])
    if c['dominio'] not in [d for d, _ in DOMS[c['marco']]]:
        DOMS[c['marco']].append((c['dominio'], c['meta_dominio']))
MNAME = {**MC['nombre'], 'TA': 'TIMSS Advanced 2015 Física'}
if HAY_PISA:  # 9.° sale de TIMSS: se compara solo con PISA (hojas PISA_*)
    GROUPS = [x for x in GROUPS if x[0] != '9°']
    DOMS['PISA'] = [(k, v[0]) for k, v in CATP_J['metas']['sistema'].items()]
    MNAME['PISA'] = 'PISA 2025'


def balance_sheet(ws, title, key_col, groups, note, dom_col='M', obj_col='K', lab='TIMSS', dom_lab='Dominio'):
    ws['A1'] = title; ws['A1'].font = fT
    ws['A2'] = 'Tolerancia para considerar «en meta» (± puntos porcentuales):'; ws['A2'].font = fB
    ws['H2'] = 0.05; ws['H2'].font = fIn; ws['H2'].fill = INP; ws['H2'].number_format = '0%'; ws['H2'].border = BOX
    ws['H2'].comment = Comment('Supuesto editable. 5 pp = margen razonable para un conteo por temas.', 'Claude')
    ws['A3'] = note; ws['A3'].font = fNote; ws.merge_cells('A3:N3'); ws['A3'].alignment = WR; ws.row_dimensions[3].height = 54
    heads = ['Grupo', 'Marco de referencia', dom_lab, f'Temas actuales en el {dom_lab.lower()}', 'Temas evaluables del grupo (sin «fuera»)',
             '% actual', f'Meta {lab} (%)', 'Temas esperados según meta', 'Brecha (temas): + faltan / − sobran',
             'Brecha (pp)', 'Estado', 'Escala para cumplir sin quitar temas (total evaluables)',
             'Temas a AGREGAR para llegar a la meta sin quitar', 'Temas «fuera del marco» del grupo']
    if ES27:
        heads = heads + ['Desvío |actual − meta| (fila) · COINCIDENCIA con la meta (fila total)']
    hdr(ws, 5, heads, [12, 22, 26, 11, 13, 9, 9, 11, 13, 10, 22, 14, 14, 12, 16], height=66)
    r = 6
    blocks = []
    for item in groups:
        g, m = item[:2]
        kc = item[2] if len(item) > 2 else key_col
        doms = DOMS[m]
        r0 = r
        for d, meta in doms:
            ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]; ws[f'C{r}'] = d
            ws[f'D{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC(dom_col)},$C{r})'
            ws[f'E{r}'] = (f'=COUNTIFS({TC(kc)},$A{r})-COUNTIFS({TC(kc)},$A{r},{TC(obj_col)},"FUERA")'
                           f'-COUNTIFS({TC(kc)},$A{r},{TC(obj_col)},"PENDIENTE")')
            ws[f'F{r}'] = f'=IF(E{r}=0,0,D{r}/E{r})'
            ws[f'G{r}'] = meta; ws[f'G{r}'].font = fIn
            ws[f'H{r}'] = f'=G{r}*E{r}'
            ws[f'I{r}'] = f'=H{r}-D{r}'
            ws[f'J{r}'] = f'=(G{r}-F{r})*100'
            ws[f'K{r}'] = (f'=IF(ABS(G{r}-F{r})<=$H$2,"En meta",IF(I{r}>0,"Faltan "&ROUND(I{r},0)&" temas",'
                           f'"Sobran "&ROUND(-I{r},0)&" temas"))')
            if ES27:
                ws[f'K{r}'] = (f'=IF(ABS(G{r}-F{r})<=$H$2,"En meta",IF(I{r}>0,"Faltan "&FIXED(I{r},0)&" temas",'
                               f'"Sobran "&FIXED(-I{r},0)&" temas"))')
            r += 1
        r1 = r - 1
        mx = ','.join(f'D{x}/G{x}' for x in range(r0, r1 + 1))
        for x in range(r0, r1 + 1):
            ws[f'L{x}'] = f'=MAX({mx})'
            ws[f'M{x}'] = f'=MAX(0,ROUNDUP(G{x}*L{x}-D{x},0))'
            ws[f'N{x}'] = f'=COUNTIFS({TC(kc)},$A{x},{TC(obj_col)},"FUERA")'
            if ES27:  # desvío por fila; en la fila total, coincidencia = 1 − ½ Σ|actual − meta|
                ws[f'O{x}'] = f'=ABS(F{x}-G{x})'; ws[f'O{x}'].number_format = PCT
        style_range(ws, r0, r1, 1, 14)
        for x in range(r0, r1 + 1):
            ws[f'F{x}'].number_format = PCT; ws[f'G{x}'].number_format = '0%'; ws[f'G{x}'].font = fIn
            ws[f'H{x}'].number_format = '0.0'; ws[f'I{x}'].number_format = '+0.0;-0.0;0.0'
            ws[f'J{x}'].number_format = PP; ws[f'L{x}'].number_format = '0.0'
            ws[f'K{x}'].font = fB
        # fila total del bloque
        ws[f'C{r}'] = 'Total del grupo'; ws[f'C{r}'].font = fB
        ws[f'D{r}'] = f'=SUM(D{r0}:D{r1})'; ws[f'G{r}'] = f'=SUM(G{r0}:G{r1})'; ws[f'G{r}'].number_format = '0%'
        ws[f'M{r}'] = f'=SUM(M{r0}:M{r1})'
        style_range(ws, r, r, 1, 14, fill=SUB, font=fB)
        if ES27:
            ws[f'O{r}'] = f'=1-SUM(O{r0}:O{r1})/2'
            style_range(ws, r0, r1, 15, 15); style_range(ws, r, r, 15, 15, fill=SUB, font=fB)
            for x in range(r0, r1 + 1): ws[f'O{x}'].number_format = PCT
            ws[f'O{r}'].number_format = PCT
        ws[f'G{r}'].number_format = '0%'
        blocks.append((g, m, r0, r1, r))
        r += 2
    last = r
    ws.conditional_formatting.add(f'K6:K{last}', FormulaRule(formula=['LEFT(K6,6)="Faltan"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    ws.conditional_formatting.add(f'K6:K{last}', FormulaRule(formula=['LEFT(K6,6)="Sobran"'], fill=PatternFill('solid', fgColor='FFE699')))
    ws.conditional_formatting.add(f'K6:K{last}', CellIsRule(operator='equal', formula=['"En meta"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.freeze_panes = 'D6'
    return blocks

wbg = wb.create_sheet('Balance_grado')
BG = balance_sheet(wbg, 'Balance de contenidos por grado frente a las metas de TIMSS', 'V', GROUPS,
    'Cómo leerlo: «% actual» = temas del dominio ÷ temas evaluables del grado (se excluyen los «fuera del marco»). '
    '«Brecha (temas)» = meta × evaluables − actuales: positivo = faltan temas en ese dominio si se redistribuye la misma carga; negativo = sobran. '
    '«Temas a AGREGAR» = cuántos temas habría que sumar a cada dominio para cumplir todas las metas SIN quitar ninguno (la escala la fija el dominio más sobrerrepresentado). '
    'Ojo: TIMSS fija pesos para la prueba acumulada (4.° y 8.°); aplicarlos a cada grado es una aproximación. La hoja Balance_ciclo hace la comparación acumulada.')

wcy = wb.create_sheet('Balance_ciclo')
CYCLES = [('Ciclo 2°–4°', M4), ('Ciclo 5°–8°', M8), ('9°', M8), ('Física 10°–11°', 'TA')]
if HAY_PISA:
    CYCLES = [x for x in CYCLES if x[0] != '9°']
BC = balance_sheet(wcy, 'Balance de contenidos por ciclo (comparación acumulada, la más fiel a TIMSS)', 'W', CYCLES,
    'TIMSS 4.° evalúa lo aprendido hasta 4.° grado → se compara con el acumulado de 2.° a 4.° (1.° grado no tiene malla). '
    'TIMSS 8.° evalúa lo aprendido hasta 8.° → acumulado de 5.° a 8.°. 9.° se muestra contra TIMSS 8.° solo como referencia. '
    'Física de Bachillerato se compara con TIMSS Advanced 2015 (Física). Química y Biología de Bachillerato no tienen marco TIMSS de contenido.')

# ------------------------------------------------------------------ Cobertura_objetivos
wo = wb.create_sheet('Cobertura_objetivos')
wo['A1'] = 'Amplitud: ¿qué objetivos TIMSS cubre la malla y en qué grado?'; wo['A1'].font = fT
wo['A2'] = ('Cada celda de grado = n.º de temas de ese grado que trabajan el objetivo como principal o secundario (COUNTIFS sobre Temas_TIMSS). '
            '«Total en su ciclo» suma solo los grados que evalúa cada marco: TIMSS 4.° → 2.°–4.°; TIMSS 8.° → 5.°–8.° (la columna 9.° es referencia); '
            'TIMSS Advanced → Física 10.°–11.°. «Profundidad»: 0 = no cubierto; 1 = débil; 2–3 = básica; 4+ = sólida.')
wo['A2'].font = fNote; wo.merge_cells('A2:T2'); wo['A2'].alignment = WR; wo.row_dimensions[2].height = 28
GCOLS = [g for g, _ in GROUPS]
heads = ['Código', 'Marco', 'Dominio', 'Área temática', 'Objetivo'] + GCOLS + ['Total temas en su ciclo', 'Como principal (todos los grados)', 'Cubierto en su ciclo', 'Profundidad en su ciclo', 'Primer grado donde aparece']
hdr(wo, 4, heads, [12, 7, 22, 30, 52] + [7] * len(GCOLS) + [8, 9, 9, 15, 12], height=45)
CICLO_COLS = {M4: ('2°', '3°', '4°'), M8: ('5°', '6°', '7°', '8°'), 'TA': ('Física 10°', 'Física 11°')}
r = 5
for c in CAT:
    wo[f'A{r}'] = c['codigo']; wo[f'B{r}'] = c['marco']; wo[f'C{r}'] = c['dominio']; wo[f'D{r}'] = c['area']; wo[f'E{r}'] = c['objetivo']
    for j, g in enumerate(GCOLS):
        col = get_column_letter(6 + j)
        wo[f'{col}{r}'] = (f'=COUNTIFS({TC("V")},{col}$4,{TC("K")},$A{r})+COUNTIFS({TC("V")},{col}$4,{TC("L")},$A{r})')
        if ES27 and g not in CICLO_COLS[c['marco']]:
            wo[f'{col}{r}'] = 'N/A'  # el grado no pertenece al ciclo que evalúa este marco
    gc1, gc2 = get_column_letter(6), get_column_letter(5 + len(GCOLS))
    tcol = get_column_letter(6 + len(GCOLS))
    # Total SOLO dentro del ciclo que evalúa cada marco: T4 = 2°–4° (F:H); T8 = 5°–8° (I:L, 9° queda como referencia); TA = Física 10°–11° (N:O)
    wo[f'{tcol}{r}'] = f'=IF($B{r}="{M4}",SUM(F{r}:H{r}),IF($B{r}="{M8}",SUM(I{r}:L{r}),SUM(N{r}:O{r})))'
    if ES27:  # rangos según las columnas de grado vigentes (9.° ya no es columna cuando PISA está activo)
        rg = {m_: (get_column_letter(6 + GCOLS.index(cc_[0])), get_column_letter(6 + GCOLS.index(cc_[-1]))) for m_, cc_ in CICLO_COLS.items()}
        wo[f'{tcol}{r}'] = (f'=IF($B{r}="{M4}",SUM({rg[M4][0]}{r}:{rg[M4][1]}{r}),IF($B{r}="{M8}",SUM({rg[M8][0]}{r}:{rg[M8][1]}{r}),'
                            f'SUM({rg["TA"][0]}{r}:{rg["TA"][1]}{r})))')
    pcol = get_column_letter(7 + len(GCOLS)); wo[f'{pcol}{r}'] = f'=COUNTIFS({TC("K")},$A{r})'
    ccol = get_column_letter(8 + len(GCOLS)); wo[f'{ccol}{r}'] = f'=IF({tcol}{r}>0,"Sí","No")'
    dcol = get_column_letter(9 + len(GCOLS))
    wo[f'{dcol}{r}'] = f'=IF({tcol}{r}=0,"No cubierto",IF({tcol}{r}=1,"Débil (1 tema)",IF({tcol}{r}<=3,"Básica (2–3)","Sólida (4+)")))'
    fcol = get_column_letter(10 + len(GCOLS))
    wo[f'{fcol}{r}'] = f'=IFERROR(INDEX(${gc1}$4:${gc2}$4,MATCH(TRUE,INDEX({gc1}{r}:{gc2}{r}>0,0),0)),"—")'
    if ES27:  # primer grado solo dentro del ciclo (las celdas «N/A» no cuentan)
        cc = [get_column_letter(6 + GCOLS.index(g)) for g in CICLO_COLS[c['marco']]]
        wo[f'{fcol}{r}'] = f'=IFERROR(INDEX(${cc[0]}$4:${cc[-1]}$4,MATCH(TRUE,INDEX({cc[0]}{r}:{cc[-1]}{r}>0,0),0)),"—")'
    style_range(wo, r, r, 1, 10 + len(GCOLS), align=Alignment(vertical='top', wrap_text=True))
    r += 1
OL = r - 1
TCOL, CCOL, DCOL = tcol, ccol, dcol
if ES27:  # híbrido con «Faltantes»: semáforo, N/A fuera del ciclo y qué hacer con cada objetivo débil o no cubierto
    wo['A2'] = ('Cómo leer: cada celda de grado = n.º de temas de ese grado que trabajan el objetivo (principal o secundario). '
                '«N/A» = ese grado no pertenece al ciclo que evalúa el marco (TIMSS 4.° → 2.°–4.°; TIMSS 8.° → 5.°–8.°; Advanced → Física 10.°–11.°), no es un problema. '
                'Semáforo de la columna «Estado»: ROJO = no cubierto · AMARILLO = débil (1 tema) · VERDE CLARO = básica (2–3) · VERDE = sólida (4+). '
                'Filtra la columna «Estado» para ver solo lo rojo y amarillo; la última columna dice dónde aparece hoy y qué se propone (hoja Presupuesto).')
    wo.row_dimensions[2].height = 54
    hdr(wo, 4, heads[:-3] + ['Estado en su ciclo (semáforo)', 'Primer grado donde aparece (en su ciclo)', 'Qué hacer · dónde aparece hoy'], None, height=45)
    wo.column_dimensions[dcol].width = 17
    qcol = get_column_letter(11 + len(GCOLS)); wo.column_dimensions[qcol].width = 70
    Tid_ = {t['id']: t for t in T}
    _pres = json.load(open('data/referencia/presupuesto_candidatos.json')) if os.path.exists('data/referencia/presupuesto_candidatos.json') else []
    for x in range(5, OL + 1):
        cod = wo[f'A{x}'].value
        donde = C_.Counter(f"{grp(Tid_[k])} U{Tid_[k]['unidad'].split(' ')[0].rstrip('.').lstrip('U')}" for k, v in R.items()
                           if cod in (v['obj1'], v['obj2']))
        props = [p_ for p_ in _pres if p_['accion'] == 'insertar' and cod in p_['objetivos']]
        txt = ('Hoy: ' + '; '.join(f'{u} ({n})' for u, n in sorted(donde.items()))) if donde else 'Hoy: ningún tema de 2.°–9.° lo trabaja.'
        if props:
            p0 = sorted(props, key=lambda p_: p_['prioridad'])[0]
            txt += f" · Propuesta {p0['id']} (Presupuesto): {p0['grado']}.° · {p0['propuesta']}"
        wo[f'{qcol}{x}'] = txt
        style_range(wo, x, x, 11 + len(GCOLS), 11 + len(GCOLS), align=WR)
    grey = PatternFill('solid', fgColor='EDEDED')
    wo.conditional_formatting.add(f'F5:{gc2}{OL}', CellIsRule(operator='equal', formula=['"N/A"'], fill=grey, font=Font(name=F, size=9, color='A6A6A6')))
    wo.conditional_formatting.add(f'F5:{gc2}{OL}', CellIsRule(operator='between', formula=['2', '3'], fill=PatternFill('solid', fgColor='C6E0B4')))
    wo.conditional_formatting.add(f'F5:{gc2}{OL}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='E2F0D9')))
    wo.conditional_formatting.add(f'F5:{gc2}{OL}', FormulaRule(formula=['AND(ISNUMBER(F5),F5>=4)'], fill=PatternFill('solid', fgColor='A9D08E')))
    wo.conditional_formatting.add(f'{dcol}5:{dcol}{OL}', CellIsRule(operator='equal', formula=['"Básica (2–3)"'], fill=PatternFill('solid', fgColor='E2F0D9')))
    wo.conditional_formatting.add(f'{qcol}5:{qcol}{OL}', FormulaRule(formula=[f'OR(${dcol}5="No cubierto",${dcol}5="Débil (1 tema)")'], font=Font(name=F, size=10, bold=True)))
wo.freeze_panes = 'F5'; wo.auto_filter.ref = f'A4:{get_column_letter(11 + len(GCOLS)) if ES27 else fcol}{OL}'
rng = f'{dcol}5:{dcol}{OL}'
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"No cubierto"'], fill=PatternFill('solid', fgColor='F8CBAD')))
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Débil (1 tema)"'], fill=PatternFill('solid', fgColor='FFE699')))
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Sólida (4+)"'], fill=PatternFill('solid', fgColor='C6EFCE')))
wo.conditional_formatting.add(f'F5:{gc2}{OL}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))

# resumen de amplitud debajo
r = OL + 3
wo[f'A{r}'] = 'Resumen de amplitud por marco y dominio'; wo[f'A{r}'].font = fS
r += 1
hdr_row = r
for j, h in enumerate(['Marco', 'Dominio', 'Objetivos en el marco', 'Objetivos cubiertos', '% cubierto', 'Objetivos no cubiertos', 'Objetivos débiles (1 tema)'], 1):
    cc = wo.cell(row=r, column=j, value=h); cc.font = fH; cc.fill = HDR; cc.alignment = CTR; cc.border = BOX
r += 1
AMP = []
for m in [M4, M8, 'TA']:
    for d, _ in DOMS[m] + [('(todos)', None)]:
        wo[f'A{r}'] = m; wo[f'B{r}'] = d
        crit_d = '"*"' if d == '(todos)' else f'$B{r}'
        wo[f'C{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d})'
        wo[f'D{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d},${CCOL}$5:${CCOL}${OL},"Sí")'
        wo[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
        wo[f'F{r}'] = f'=C{r}-D{r}'
        wo[f'G{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d},${DCOL}$5:${DCOL}${OL},"Débil (1 tema)")'
        style_range(wo, r, r, 1, 7, fill=SUB if d == '(todos)' else None, font=fB if d == '(todos)' else fN)
        wo[f'E{r}'].number_format = PCT
        AMP.append((m, d, r))
        r += 1

# ------------------------------------------------------------------ Faltantes
wf = wb.create_sheet('Faltantes')
wf['A1'] = 'Qué falta: objetivos TIMSS no cubiertos o débiles, y dónde podrían entrar'; wf['A1'].font = fT
wf['A2'] = ('Lista generada desde Cobertura_objetivos (profundidad «No cubierto» o «Débil»). La columna «Temas hoy» es fórmula y se actualiza si cambia la clasificación. '
            'La sugerencia de ubicación es una propuesta para discutir con el equipo de Ciencias, no una decisión.')
wf['A2'].font = fNote; wf.merge_cells('A2:H2'); wf['A2'].alignment = WR; wf.row_dimensions[2].height = 30
hdr(wf, 4, ['Código', 'Marco', 'Dominio', 'Objetivo', 'Temas hoy', 'Profundidad hoy', 'Dónde aparece hoy', 'Sugerencia de ubicación en la malla'],
    [12, 7, 20, 50, 9, 15, 30, 70], height=32)
SUG = {
 'T4-V2.2': ('—', 'Ciclo 2°–4°: agregar en 3.° U4 «Las plantas» o U5 «Los animales» un tema de semejanza crías-progenitores y rasgos heredados vs. adquiridos (antecede a Genética de 7.°).'),
 'T4-V4.3': ('4.° U3 (1 tema)', 'Reforzar en 4.° U3 «Naturaleza y sus interacciones»: competencia por agua, luz y alimento con una actividad de observación.'),
 'T4-F1.1': ('2.° U1 (1 tema)', 'Reforzar en 2.° U1 «Materia» o 3.° U3 «El agua»: estados sólido, líquido y gaseoso con sus características (hoy solo aparece como cambio de estado del agua).'),
 'T4-F2.3': ('2.° U2 (1 tema)', 'Reforzar en 4.° U1 «Máquinas y energía» o 2.° U2 «Energía»: transferencia de calor entre objetos calientes y fríos (hoy se ve recién en 6.°).'),
 'T8-B5.2': ('7.° U4 (1 tema)', 'Reforzar en 7.° U4 «Ecología» o U5 «Recurso hídrico»: ciclos del carbono y del oxígeno además del agua.'),
 'T8-B6.1': ('—', 'Ciclo 5°–8°: agregar en 8.° U7 «Anatomía y fisiología animal» o 6.° U5 «Célula» un tema de enfermedades infecciosas, transmisión, prevención y sistema inmune (hoy aparece hasta 10.° Biología).'),
 'T8-F3.1': ('—', 'Ciclo 5°–8°: agregar en 8.° U3 «Ondas mecánicas» temas de luz (propagación, reflexión, refracción, diagramas de rayos). Hoy la óptica aparece hasta 10.° Física.'),
 'T8-T4.1': ('Fuera del ciclo: fases de la Luna en 4.° (U4) y mareas en 9.° (U4)', 'Agregar en 5.° U5 «Ciencias del espacio»: traslación y estaciones, mareas, fases lunares y eclipses dentro del ciclo 5°–8°.'),
}
r = 5
gc1, gc2 = get_column_letter(6), get_column_letter(5 + len(GCOLS))
# Determine list in python using same logic to choose rows (values verified later by formulas)
import collections as C
cnt = C.Counter()
Tid = {t['id']: t for t in T}
def in_cycle(code, t):
    if code.startswith(M4): return t['asignatura'] == 'Ciencias' and t['grado'] <= 4
    if code.startswith(M8): return t['asignatura'] == 'Ciencias' and 5 <= t['grado'] <= 8
    return t['asignatura'] == 'Física'
for k, v in R.items():
    for o in (v['obj1'], v['obj2']):
        if o and o != 'FUERA' and in_cycle(o, Tid[k]): cnt[o] += 1
falt = [c for c in CAT if cnt[c['codigo']] <= 1]
for c in falt:
    rr = [x for x in range(5, OL + 1) if wo[f'A{x}'].value == c['codigo']][0]
    wf[f'A{r}'] = c['codigo']; wf[f'B{r}'] = c['marco']; wf[f'C{r}'] = c['dominio']; wf[f'D{r}'] = c['objetivo']
    wf[f'E{r}'] = f"=Cobertura_objetivos!{TCOL}{rr}"; wf[f'E{r}'].font = fLink
    wf[f'F{r}'] = f"=Cobertura_objetivos!{DCOL}{rr}"; wf[f'F{r}'].font = fLink
    wf[f'G{r}'] = SUG.get(c['codigo'], ('', ''))[0]
    wf[f'H{r}'] = SUG.get(c['codigo'], ('', 'Revisar con el equipo de Ciencias.'))[1]
    if ES27:  # sin sugerencias redactadas para 2027: se lista dónde aparece hoy (cualquier grado)
        donde = C.Counter(f"{grp(Tid[k])} {Tid[k]['unidad'].split(' ')[0].rstrip('.')}" for k, v in R.items()
                          if c['codigo'] in (v['obj1'], v['obj2']))
        wf[f'G{r}'] = '; '.join(f'{u} ({n})' for u, n in sorted(donde.items())) or '—'

    style_range(wf, r, r, 1, 8, align=WR)
    wf[f'E{r}'].font = fLink; wf[f'F{r}'].font = fLink
    r += 1
FL = r - 1
r += 1
wf[f'A{r}'] = 'Además (balance): los faltantes/sobrantes por dominio están en Balance_ciclo (acumulado) y Balance_grado (por grado), columnas I, K y M.'
wf[f'A{r}'].font = fNote; wf.merge_cells(f'A{r}:H{r}')

# ------------------------------------------------------------------ Cognitivo_grado
wg = wb.create_sheet('Cognitivo_grado')
wg['A1'] = 'Profundidad cognitiva: Conocer / Aplicar / Razonar frente a TIMSS'; wg['A1'].font = fT
wg['A2'] = ('Fuente: columna «Habilidad TIMSS» de cada malla. Si un tema menciona varios dominios, se reparte en partes iguales '
            '(p. ej. «Conocer… Aplicar…» = 0,5 + 0,5). Columnas S-U de Temas_TIMSS. Química y Biología de Bachillerato se comparan con la meta de TIMSS Advanced Física solo como referencia aproximada.')
wg['A2'].font = fNote; wg.merge_cells('A2:N2'); wg['A2'].alignment = WR; wg.row_dimensions[2].height = 40
hdr(wg, 4, ['Grupo', 'Referencia', 'Temas', 'Conocer (temas-eq.)', 'Aplicar (temas-eq.)', 'Razonar (temas-eq.)',
            '% Conocer', '% Aplicar', '% Razonar', 'Meta Conocer', 'Meta Aplicar', 'Meta Razonar',
            'Brecha Razonar (pp)', 'Lectura'], [13, 26, 8, 11, 11, 11, 9, 9, 9, 9, 9, 9, 11, 46], height=45)
COG = [(g, m) for g, m in GROUPS] + [('Química 10°', 'TA*'), ('Química 11°', 'TA*'), ('Biología 10°', 'TA*'), ('Biología 11°', 'TA*')]
META = {**MC['meta_cog'], 'TA': (0.30, 0.40, 0.30), 'TA*': (0.30, 0.40, 0.30)}
REFN = {**MC['nombre'], 'TA': 'TIMSS Advanced 2015 Física', 'TA*': 'TIMSS Advanced Física (ref. aprox.)'}
r = 5
for g, m in COG:
    wg[f'A{r}'] = g; wg[f'B{r}'] = REFN[m]
    wg[f'C{r}'] = f'=COUNTIFS({TC("V")},$A{r})'
    wg[f'D{r}'] = f'=SUMIFS({TC("S")},{TC("V")},$A{r})'
    wg[f'E{r}'] = f'=SUMIFS({TC("T")},{TC("V")},$A{r})'
    wg[f'F{r}'] = f'=SUMIFS({TC("U")},{TC("V")},$A{r})'
    for a, b in zip('GHI', 'DEF'):
        wg[f'{a}{r}'] = f'=IF($C{r}=0,0,{b}{r}/$C{r})'
    for a, v in zip('JKL', META[m]):
        wg[f'{a}{r}'] = v
    wg[f'M{r}'] = f'=(I{r}-L{r})*100'
    wg[f'N{r}'] = (f'=IF(ABS(M{r})<=5,"Razonar en línea con la meta",IF(M{r}>0,"Más Razonar que la meta (revisar si la etiqueta está inflada)",'
                   f'"Menos Razonar que la meta: faltan tareas de análisis y argumentación"))')
    style_range(wg, r, r, 1, 14)
    for a in 'DEF': wg[f'{a}{r}'].number_format = '0.0'
    for a in 'GHI': wg[f'{a}{r}'].number_format = PCT
    for a in 'JKL': wg[f'{a}{r}'].number_format = '0%'; wg[f'{a}{r}'].font = fIn
    wg[f'M{r}'].number_format = PP
    if ES27:
        wg[f'O{r}'] = f'=1-(ABS(G{r}-J{r})+ABS(H{r}-K{r})+ABS(I{r}-L{r}))/2'
        style_range(wg, r, r, 15, 15); wg[f'O{r}'].number_format = PCT
    r += 1
CGL = r - 1
if ES27:
    c_ = wg.cell(row=4, column=15, value='COINCIDENCIA con la meta (1 − ½ Σ|actual − meta|)')
    c_.font = fH; c_.fill = HDR; c_.alignment = CTR; c_.border = BOX; wg.column_dimensions['O'].width = 15

# ------------------------------------------------------------------ Hojas nuevas del libro 2027
if ES27:
    CIENCIAS = [g for g, m in GROUPS if m in (M4, M8)]                 # 2°…9°
    CICLOS_C = [(g, m) for g, m in CYCLES if m in (M4, M8)]            # Ciclo 2°–4°, Ciclo 5°–8°, 9°
    def evaluables(col, ref):
        return (f'COUNTIFS({TC(col)},{ref})-COUNTIFS({TC(col)},{ref},{TC("K")},"FUERA")'
                f'-COUNTIFS({TC(col)},{ref},{TC("K")},"PENDIENTE")')
    COBR = lambda col: f'Cobertura_objetivos!${col}$5:${col}${OL}'

    # ---------------- Referencia_2023: la v1 recalculada con fórmulas, para «Cambio vs 2023»
    wr = wb.create_sheet('Referencia_2023')
    wr['A1'] = 'Referencia TIMSS 2023 (v1): cobertura y balance recalculados con la clasificación v1'; wr['A1'].font = fT
    wr['A2'] = ('Solo para comparar. Usa las columnas AA-AC de Temas_TIMSS (clasificación v1 contra TIMSS 2023) con las mismas reglas del libro v1 '
                '(Cobertura_TIMSS_Ciencias_SV.xlsx). Los resultados deben coincidir con ese libro.')
    wr['A2'].font = fNote; wr.merge_cells('A2:I2'); wr['A2'].alignment = WR; wr.row_dimensions[2].height = 30
    hdr(wr, 4, ['Código 2023', 'Marco', 'Dominio', 'Meta del dominio', 'Área temática', 'Objetivo (paráfrasis)', 'Ciclo que evalúa',
                'Temas en su ciclo (v1)', 'Cubierto (v1)'], [12, 7, 24, 10, 34, 55, 15, 11, 10], height=45)
    CIC23 = {'T4': 'Ciclo 2°–4°', 'T8': 'Ciclo 5°–8°', 'TA': 'Física 10°–11°'}
    r = 5
    for c in CAT23:
        for j, v in enumerate([c['codigo'], c['marco'], c['dominio'], c['meta_dominio'], c['area'], c['objetivo'], CIC23[c['marco']]], 1):
            wr.cell(row=r, column=j, value=v)
        wr[f'H{r}'] = f'=COUNTIFS({TC("W")},$G{r},{TC("AA")},$A{r})+COUNTIFS({TC("W")},$G{r},{TC("AB")},$A{r})'
        wr[f'I{r}'] = f'=IF(H{r}>0,"Sí","No")'
        style_range(wr, r, r, 1, 9, align=WR); wr[f'D{r}'].number_format = '0%'
        r += 1
    R23L = r - 1
    r += 2
    wr[f'A{r}'] = 'Amplitud 2023 por marco'; wr[f'A{r}'].font = fS; r += 1
    hdr(wr, r, ['Marco', 'Objetivos', 'Cubiertos', '% cubierto'], None); r += 1
    AMP23 = {}
    for m in ('T4', 'T8', 'TA'):
        wr[f'A{r}'] = m
        wr[f'B{r}'] = f'=COUNTIFS($B$5:$B${R23L},$A{r})'
        wr[f'C{r}'] = f'=COUNTIFS($B$5:$B${R23L},$A{r},$I$5:$I${R23L},"Sí")'
        wr[f'D{r}'] = f'=IF(B{r}=0,0,C{r}/B{r})'
        style_range(wr, r, r, 1, 4); wr[f'D{r}'].number_format = PCT
        AMP23[m] = r; r += 1
    r += 2
    wr[f'A{r}'] = 'Balance 2023 por ciclo (mismas fórmulas que Balance_ciclo del libro v1)'; wr[f'A{r}'].font = fS; r += 1
    hdr(wr, r, ['Ciclo', 'Marco 2023', 'Dominio 2023', 'Temas actuales', 'Temas evaluables', 'Meta', 'Brecha (temas): + faltan / − sobran'],
        None, height=40); r += 1
    DOM23 = {}
    for c in CAT23:
        DOM23.setdefault(c['marco'], [])
        if c['dominio'] not in [d for d, _ in DOM23[c['marco']]]:
            DOM23[c['marco']].append((c['dominio'], c['meta_dominio']))
    BAL23 = {}
    for g, m23 in (('Ciclo 2°–4°', 'T4'), ('Ciclo 5°–8°', 'T8'), ('9°', 'T8'), ('Física 10°–11°', 'TA')):
        r0 = r
        for d, meta in DOM23[m23]:
            wr[f'A{r}'] = g; wr[f'B{r}'] = m23; wr[f'C{r}'] = d
            wr[f'D{r}'] = f'=COUNTIFS({TC("W")},$A{r},{TC("AC")},$C{r})'
            wr[f'E{r}'] = (f'=COUNTIFS({TC("W")},$A{r})-COUNTIFS({TC("W")},$A{r},{TC("AA")},"FUERA")'
                           f'-COUNTIFS({TC("W")},$A{r},{TC("AA")},"PENDIENTE")')
            wr[f'F{r}'] = meta; wr[f'F{r}'].font = fIn; wr[f'F{r}'].number_format = '0%'
            wr[f'G{r}'] = f'=F{r}*E{r}-D{r}'; wr[f'G{r}'].number_format = '+0.0;-0.0;0.0'
            style_range(wr, r, r, 1, 7); wr[f'F{r}'].font = fIn; wr[f'F{r}'].number_format = '0%'
            wr[f'G{r}'].number_format = '+0.0;-0.0;0.0'
            r += 1
        BAL23[g] = (r0, r - 1)
        r += 1
    wr.freeze_panes = 'A5'

    # ---------------- Ambiental
    wa = wb.create_sheet('Ambiental')
    wa['A1'] = 'Conocimiento ambiental: ¿cuánto de la malla trabaja los objetivos ambientales de TIMSS 2027?'; wa['A1'].font = fT
    wa['A2'] = ('TIMSS 2027 marca con asterisco los objetivos de Biología y Ciencias de la Tierra que forman la subescala de conocimiento ambiental '
                '(≈25 % de los ítems de la prueba). Un tema cuenta como ambiental si su objetivo principal o secundario es uno de ellos (columna X de Temas_TIMSS). '
                'La referencia de 25 % es la proporción de ítems de la prueba, no una meta curricular: úsala como orientación.')
    wa['A2'].font = fNote; wa.merge_cells('A2:H2'); wa['A2'].alignment = WR; wa.row_dimensions[2].height = 44
    wa['A3'] = 'Referencia TIMSS 2027 (≈ % de ítems ambientales):'; wa['A3'].font = fB
    wa['E3'] = 0.25; wa['E3'].font = fIn; wa['E3'].number_format = '0%'; wa['E3'].border = BOX
    wa['E3'].comment = Comment('TIMSS 2027, cap. 2, p. 23: «aproximadamente 25 por ciento» de los ítems de ciencias. Editable.', 'Claude')
    wa['F3'] = 'Tolerancia (± pp):'; wa['F3'].font = fB
    wa['H3'] = 0.05; wa['H3'].font = fIn; wa['H3'].number_format = '0%'; wa['H3'].border = BOX
    hdr(wa, 5, ['Grado / ciclo', 'Marco', 'Temas evaluables (sin «fuera»)', 'Temas que trabajan un objetivo ambiental',
                '% ambiental', 'Referencia', 'Diferencia (pp)', 'Lectura'], [16, 16, 13, 15, 11, 11, 12, 48], height=48)
    r = 6
    AMB = {}
    for etiqueta, filas, col in (('Por grado', [(g, m) for g, m in GROUPS if m in (M4, M8)], 'V'), ('Por ciclo', CICLOS_C, 'W')):
        wa[f'A{r}'] = etiqueta; wa[f'A{r}'].font = fS; r += 1
        for g, m in filas:
            wa[f'A{r}'] = g; wa[f'B{r}'] = MNAME[m]
            wa[f'C{r}'] = '=' + evaluables(col, f'$A{r}')
            wa[f'D{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("X")},1)'
            wa[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
            wa[f'F{r}'] = '=$E$3'
            wa[f'G{r}'] = f'=(E{r}-F{r})*100'
            wa[f'H{r}'] = (f'=IF(ABS(E{r}-F{r})<=$H$3,"Cerca de la referencia",IF(E{r}<F{r},'
                           f'"Menos contenido ambiental que la referencia","Más contenido ambiental que la referencia"))')
            style_range(wa, r, r, 1, 8)
            wa[f'E{r}'].number_format = PCT; wa[f'F{r}'].number_format = '0%'; wa[f'G{r}'].number_format = PP
            wa[f'F{r}'].font = fLink; wa[f'H{r}'].font = fB
            AMB[g] = r; r += 1
        r += 1
    wa.conditional_formatting.add(f'H6:H{r}', FormulaRule(formula=['LEFT(H6,5)="Menos"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    wa.conditional_formatting.add(f'H6:H{r}', FormulaRule(formula=['LEFT(H6,5)="Cerca"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    wa.conditional_formatting.add(f'H6:H{r}', FormulaRule(formula=['LEFT(H6,4)="Más "'], fill=PatternFill('solid', fgColor='FFE699')))
    r += 1
    wa[f'A{r}'] = 'Objetivos ambientales de TIMSS 2027 y su cobertura en el ciclo que evalúa cada marco'; wa[f'A{r}'].font = fS; r += 1
    hdr(wa, r, ['Código', 'Marco', 'Dominio', 'Objetivo', 'Temas en su ciclo', 'Profundidad', '', ''], None, height=32); r += 1
    a0 = r
    for c in [c for c in CAT if c.get('ambiental')]:
        wa[f'A{r}'] = c['codigo']; wa[f'B{r}'] = c['marco']; wa[f'C{r}'] = c['dominio']; wa[f'D{r}'] = c['objetivo']
        wa[f'E{r}'] = f'=INDEX({COBR(TCOL)},MATCH($A{r},{COBR("A")},0))'
        wa[f'F{r}'] = f'=INDEX({COBR(DCOL)},MATCH($A{r},{COBR("A")},0))'
        style_range(wa, r, r, 1, 6, align=WR); wa[f'E{r}'].font = fLink; wa[f'F{r}'].font = fLink
        r += 1
    a1 = r - 1
    wa.conditional_formatting.add(f'F{a0}:F{a1}', CellIsRule(operator='equal', formula=['"No cubierto"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    wa.conditional_formatting.add(f'F{a0}:F{a1}', CellIsRule(operator='equal', formula=['"Débil (1 tema)"'], fill=PatternFill('solid', fgColor='FFE699')))
    wa.conditional_formatting.add(f'F{a0}:F{a1}', CellIsRule(operator='equal', formula=['"Sólida (4+)"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    r += 1
    hdr(wa, r, ['Marco', 'Objetivos ambientales', 'Cubiertos en su ciclo', '% cubierto', 'No cubiertos', 'Débiles (1 tema)'], None, height=32); r += 1
    AMBOBJ = {}
    for m in (M4, M8):
        wa[f'A{r}'] = m
        wa[f'B{r}'] = f'=COUNTIFS($B${a0}:$B${a1},$A{r})'
        wa[f'C{r}'] = f'=COUNTIFS($B${a0}:$B${a1},$A{r},$E${a0}:$E${a1},">0")'
        wa[f'D{r}'] = f'=IF(B{r}=0,0,C{r}/B{r})'
        wa[f'E{r}'] = f'=B{r}-C{r}'
        wa[f'F{r}'] = f'=COUNTIFS($B${a0}:$B${a1},$A{r},$F${a0}:$F${a1},"Débil (1 tema)")'
        style_range(wa, r, r, 1, 6, fill=SUB, font=fB); wa[f'D{r}'].number_format = PCT
        AMBOBJ[m] = r; r += 1
    wa.freeze_panes = 'A6'

    # ---------------- Investigaciones
    wi = wb.create_sheet('Investigaciones')
    wi['A1'] = 'Investigaciones: temas cuyo foco es la práctica científica (áreas nuevas de TIMSS 2027)'; wi['A1'].font = fT
    wi['A2'] = ('TIMSS 2027 agrega en cada dominio un área de «Investigaciones» (planificar, medir, registrar, analizar datos, concluir, usar equipo). '
                'Un tema cuenta como principal si su foco es la práctica; como secundario si trabaja un contenido mediante indagación. '
                'Columnas Y-Z y AA de Temas_TIMSS. «Eran FUERA en v1» = temas que con TIMSS 2023 no tenían objetivo y ahora sí.')
    wi['A2'].font = fNote; wi.merge_cells('A2:I2'); wi['A2'].alignment = WR; wi.row_dimensions[2].height = 44
    hdr(wi, 4, ['Grado / ciclo', 'Temas evaluables', 'Investigaciones como principal', 'Investigaciones solo como secundario',
                'Total de temas con Investigaciones', '% de los evaluables', 'Principales que eran FUERA en v1',
                'FUERA en v1 que ya no lo son (cualquier objetivo)', 'Siguen FUERA'], [16, 11, 14, 14, 14, 11, 14, 16, 10], height=58)
    r = 5
    INV = {}
    for etiqueta, filas, col in (('Por grado', CIENCIAS, 'V'), ('Por ciclo', [g for g, _ in CICLOS_C], 'W')):
        wi[f'A{r}'] = etiqueta; wi[f'A{r}'].font = fS; r += 1
        for g in filas:
            wi[f'A{r}'] = g
            wi[f'B{r}'] = '=' + evaluables(col, f'$A{r}')
            wi[f'C{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("Y")},1)'
            wi[f'D{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("Y")},0,{TC("Z")},1)'
            wi[f'E{r}'] = f'=C{r}+D{r}'
            wi[f'F{r}'] = f'=IF(B{r}=0,0,E{r}/B{r})'
            wi[f'G{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("Y")},1,{TC("AA")},"FUERA")'
            wi[f'H{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("AA")},"FUERA",{TC("K")},"<>FUERA")'
            wi[f'I{r}'] = f'=COUNTIFS({TC(col)},$A{r},{TC("K")},"FUERA")'
            style_range(wi, r, r, 1, 9); wi[f'F{r}'].number_format = PCT
            INV[g] = r; r += 1
        r += 1
    r += 1
    wi[f'A{r}'] = 'Objetivos de Investigaciones × grado (n.º de temas que lo trabajan como principal o secundario)'; wi[f'A{r}'].font = fS; r += 1
    hdr(wi, r, ['Código', 'Área', 'Objetivo'] + CIENCIAS + ['Total en su ciclo', 'Profundidad'], None, height=32)
    gh = r; r += 1
    for c in [c for c in CAT if c.get('es_investigacion')]:
        wi[f'A{r}'] = c['codigo']; wi[f'B{r}'] = c['area']; wi[f'C{r}'] = c['objetivo']
        for j, g in enumerate(CIENCIAS):
            col = get_column_letter(4 + j)
            wi[f'{col}{r}'] = f'=COUNTIFS({TC("V")},{col}${gh},{TC("K")},$A{r})+COUNTIFS({TC("V")},{col}${gh},{TC("L")},$A{r})'
        tc_, dc_ = get_column_letter(4 + len(CIENCIAS)), get_column_letter(5 + len(CIENCIAS))
        wi[f'{tc_}{r}'] = f'=INDEX({COBR(TCOL)},MATCH($A{r},{COBR("A")},0))'
        wi[f'{dc_}{r}'] = f'=INDEX({COBR(DCOL)},MATCH($A{r},{COBR("A")},0))'
        style_range(wi, r, r, 1, 5 + len(CIENCIAS), align=WR)
        wi[f'{tc_}{r}'].font = fLink; wi[f'{dc_}{r}'].font = fLink
        r += 1
    wi.conditional_formatting.add(f'D{gh + 1}:{get_column_letter(3 + len(CIENCIAS))}{r}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))
    wi.conditional_formatting.add(f'{dc_}{gh + 1}:{dc_}{r}', CellIsRule(operator='equal', formula=['"No cubierto"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    wi.conditional_formatting.add(f'{dc_}{gh + 1}:{dc_}{r}', CellIsRule(operator='equal', formula=['"Débil (1 tema)"'], fill=PatternFill('solid', fgColor='FFE699')))
    wi.freeze_panes = 'B5'

# ------------------------------------------------------------------ Hojas PISA 2025 (7.°–9.°)
if HAY_PISA:
    GP = ['7°', '8°', '9°']
    CICLO_P = 'Ciclo 7°–9°'
    GRUPOS_P = [(g, 'V') for g in GP] + [(CICLO_P, 'AV')]
    ROJO, AMAR, VCLARO, VERDE_S = 'F8CBAD', 'FFE699', 'E2F0D9', 'A9D08E'

    def semaforo(ws, rng):
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"No cubierto"'], fill=PatternFill('solid', fgColor=ROJO)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Débil (1 tema)"'], fill=PatternFill('solid', fgColor=AMAR)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Básica (2–3)"'], fill=PatternFill('solid', fgColor=VCLARO)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Sólida (4+)"'], fill=PatternFill('solid', fgColor=VERDE_S)))

    def estado_rango(ws, rng):
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Por debajo del rango"'], fill=PatternFill('solid', fgColor=ROJO)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Por encima del rango"'], fill=PatternFill('solid', fgColor=AMAR)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Dentro del rango"'], fill=PatternFill('solid', fgColor=VCLARO)))

    def profundidad(ref):
        return f'IF({ref}=0,"No cubierto",IF({ref}=1,"Débil (1 tema)",IF({ref}<=3,"Básica (2–3)","Sólida (4+)")))'

    def tabla_rangos(ws, r, titulo, filas, nota):
        """Bloque de metas editables: nombre | mínimo | máximo | meta (punto medio normalizado). Devuelve {nombre: fila}."""
        ws[f'A{r}'] = titulo; ws[f'A{r}'].font = fS; r += 1
        hdr(ws, r, ['', 'Mínimo del marco', 'Máximo del marco', 'Meta (punto medio, normalizado)'], None, height=32)
        for col in 'BCD':
            ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 12)
        r += 1
        r0 = r
        for nombre, (lo, hi) in filas:
            ws[f'A{r}'] = nombre; ws[f'B{r}'] = lo; ws[f'C{r}'] = hi
            style_range(ws, r, r, 1, 4)
            for col in 'BC':
                ws[f'{col}{r}'].font = fIn; ws[f'{col}{r}'].number_format = '0%'
            r += 1
        r1 = r - 1
        for x in range(r0, r1 + 1):
            ws[f'D{x}'] = f'=AVERAGE(B{x}:C{x})/((SUM($B${r0}:$B${r1})+SUM($C${r0}:$C${r1}))/2)'
            ws[f'D{x}'].number_format = PCT
        ws[f'A{r}'] = nota; ws[f'A{r}'].font = fNote; ws.merge_cells(f'A{r}:L{r}'); ws[f'A{r}'].alignment = WR
        ws.row_dimensions[r].height = 28
        return {nombre: x for (nombre, _), x in zip(filas, range(r0, r1 + 1))}, r + 2

    # ---------------- PISA_contenido: balance por sistema + cobertura de las 21 categorías
    wpc = wb.create_sheet('PISA_contenido')
    BP = balance_sheet(wpc, 'PISA 2025 · Contenido: reparto por sistema (7.°, 8.°, 9.° y acumulado 7.°–9.°)', 'V',
        [(g, 'PISA') for g in GP] + [(CICLO_P, 'PISA', 'AV')],
        'PISA evalúa a los 15 años lo acumulado (≈ 9.° en El Salvador): la fila «Ciclo 7°–9°» es la comparación más fiel. '
        'Meta = peso de cada sistema en la prueba (Tabla 2.3, total por tipo de conocimiento: 37 / 37 / 26 %). '
        '«Sistemas físicos» reúne Física y Química. Los temas sin contenido del Box 2.6 (p. ej. solo medición o electrónica digital) quedan fuera del cálculo. '
        'Columnas AK-AM de la hoja Temas (clasificación asistida por IA, revisable).',
        dom_col='AM', obj_col='AK', lab='PISA', dom_lab='Sistema')
    r = BP[-1][4] + 3
    wpc[f'A{r}'] = 'Cobertura de las 21 categorías de contenido de PISA (Box 2.6): n.º de temas por grado (principal o secundario)'; wpc[f'A{r}'].font = fS
    r += 1
    hdr(wpc, r, ['Código', 'Sistema', 'Categoría de contenido (paráfrasis)'] + GP + ['Total 7°–9°', 'Estado (semáforo)', 'Primer grado'], None, height=32)
    hp = r; r += 1
    c0 = r
    for e in [e for e in CATP if e['tipo'] == 'contenido']:
        wpc[f'A{r}'] = e['codigo']; wpc[f'B{r}'] = e['grupo']; wpc[f'C{r}'] = e['texto']
        for j, g in enumerate(GP):
            col = get_column_letter(4 + j)
            wpc[f'{col}{r}'] = f'=COUNTIFS({TC("V")},{col}${hp},{TC("AK")},$A{r})+COUNTIFS({TC("V")},{col}${hp},{TC("AL")},$A{r})'
        wpc[f'G{r}'] = f'=SUM(D{r}:F{r})'
        wpc[f'H{r}'] = '=' + profundidad(f'G{r}')
        wpc[f'I{r}'] = f'=IFERROR(INDEX($D${hp}:$F${hp},MATCH(TRUE,INDEX(D{r}:F{r}>0,0),0)),"—")'
        style_range(wpc, r, r, 1, 9, align=WR)
        r += 1
    c1 = r - 1
    semaforo(wpc, f'H{c0}:H{c1}')
    wpc.conditional_formatting.add(f'D{c0}:F{c1}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))
    r += 1
    hdr(wpc, r, ['Sistema', 'Categorías', 'Cubiertas (≥1 tema)', '% cubierto', 'No cubiertas', 'Débiles (1 tema)'], None, height=30); r += 1
    AMP_P = {}
    for sis in SISTEMAS + ['(todos)']:
        crit = '"*"' if sis == '(todos)' else f'$A{r}'
        wpc[f'A{r}'] = sis
        wpc[f'B{r}'] = f'=COUNTIFS($B${c0}:$B${c1},{crit})'
        wpc[f'C{r}'] = f'=COUNTIFS($B${c0}:$B${c1},{crit},$G${c0}:$G${c1},">0")'
        wpc[f'D{r}'] = f'=IF(B{r}=0,0,C{r}/B{r})'
        wpc[f'E{r}'] = f'=B{r}-C{r}'
        wpc[f'F{r}'] = f'=COUNTIFS($B${c0}:$B${c1},{crit},$H${c0}:$H${c1},"Débil (1 tema)")'
        style_range(wpc, r, r, 1, 6, fill=SUB if sis == '(todos)' else None, font=fB if sis == '(todos)' else fN)
        wpc[f'D{r}'].number_format = PCT
        AMP_P[sis] = r; r += 1
    wpc.column_dimensions['C'].width = 60

    # ---------------- PISA_competencias: competencia del tema (malla) y subcompetencias por fase
    wpk = wb.create_sheet('PISA_competencias')
    wpk['A1'] = 'PISA 2025 · Competencias científicas y subcompetencias (según la propia malla)'; wpk['A1'].font = fT
    wpk['A2'] = ('Fuente: columnas «Competencia PISA 2025» y «Microhabilidades PISA por fase» de cada malla (equipo de Ciencias, MINED). No hay clasificación de IA aquí. '
                 'Cada tema declara UNA competencia y una microhabilidad por fase (Aproximar, Explorar, Explicar, Indagar, Reforzar) con los códigos de las 15 subcompetencias '
                 'del marco (E1–E6 = competencia 1, D1–D4 = competencia 2, I1–I5 = competencia 3). Ojo: la plantilla de la secuencia asigna códigos por fase '
                 '(p. ej. Explorar casi siempre D1), así que las subcompetencias reflejan en parte la plantilla y no solo el contenido (bloque 4).')
    wpk['A2'].font = fNote; wpk.merge_cells('A2:L2'); wpk['A2'].alignment = WR; wpk.row_dimensions[2].height = 56
    NOMC = {'C1': '1. Explicar fenómenos', 'C2': '2. Diseñar indagación e interpretar datos', 'C3': '3. Investigar, evaluar y usar información'}
    RC, r = tabla_rangos(wpk, 4, 'Metas de PISA por competencia (Tabla 2.4, % de puntos de la prueba)',
                         [(NOMC[k], tuple(v)) for k, v in CATP_J['metas']['competencia'].items()],
                         'Celdas ámbar editables. «Dentro del rango» = el % de temas cae entre el mínimo y el máximo del marco. La coincidencia se mide contra el punto medio normalizado (40 / 30 / 30).')
    wpk[f'A{r}'] = 'Bloque 1 · Competencia declarada en cada tema'; wpk[f'A{r}'].font = fS; r += 1
    hdr(wpk, r, ['Grado / ciclo', 'Temas', 'C1 (temas)', 'C2 (temas)', 'C3 (temas)', '% C1', '% C2', '% C3',
                 'Estado C1', 'Estado C2', 'Estado C3', 'COINCIDENCIA con la meta'], [16, 10, 10, 10, 10, 9, 9, 9, 18, 18, 18, 14], height=40)
    r += 1
    COMP_ROW = {}
    for g, kc in GRUPOS_P + [('2°', 'V'), ('3°', 'V'), ('4°', 'V'), ('5°', 'V'), ('6°', 'V')]:
        if g == '2°':
            wpk[f'A{r}'] = 'Grados anteriores (solo referencia: PISA evalúa a los 15 años)'; wpk[f'A{r}'].font = fNote; r += 1
        wpk[f'A{r}'] = g
        wpk[f'B{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("C")},"Ciencias")'
        for j, k in enumerate(('C1', 'C2', 'C3')):
            n_, p_, e_ = get_column_letter(3 + j), get_column_letter(6 + j), get_column_letter(9 + j)
            rr = RC[NOMC[k]]
            wpk[f'{n_}{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("C")},"Ciencias",{TC("AE")},"{k}")'
            wpk[f'{p_}{r}'] = f'=IF($B{r}=0,0,{n_}{r}/$B{r})'
            wpk[f'{e_}{r}'] = f'=IF({p_}{r}<$B${rr},"Por debajo del rango",IF({p_}{r}>$C${rr},"Por encima del rango","Dentro del rango"))'
        wpk[f'L{r}'] = '=1-(' + '+'.join(f'ABS({get_column_letter(6 + j)}{r}-$D${RC[NOMC[k]]})' for j, k in enumerate(('C1', 'C2', 'C3'))) + ')/2'
        style_range(wpk, r, r, 1, 12)
        for col in 'FGHL': wpk[f'{col}{r}'].number_format = PCT
        COMP_ROW[g] = r; r += 1
    estado_rango(wpk, f'I{COMP_ROW["7°"]}:K{r}')
    r += 1
    wpk[f'A{r}'] = 'Bloque 2 · Subcompetencias: en cuántos temas se trabaja cada una (en al menos una de las 5 fases)'; wpk[f'A{r}'].font = fS; r += 1
    wpk[f'A{r}'] = 'Umbrales del semáforo (% de temas del ciclo):'; wpk[f'A{r}'].font = fB
    wpk[f'E{r}'] = 'débil si menos de'; wpk[f'F{r}'] = 0.05; wpk[f'G{r}'] = 'sólida desde'; wpk[f'H{r}'] = 0.15
    for col in 'FH':
        wpk[f'{col}{r}'].font = fIn; wpk[f'{col}{r}'].number_format = '0%'; wpk[f'{col}{r}'].border = BOX
    thr = r; r += 1
    hdr(wpk, r, ['Código', 'Competencia', 'Subcompetencia (paráfrasis)'] + GP + ['Total 7°–9°', '% de temas del ciclo', 'Estado (semáforo)'], None, height=36)
    hs = r; r += 1
    s0 = r
    fases_cols = [get_column_letter(32 + j) for j in range(5)]  # AF..AJ
    for e in [e for e in CATP if e['tipo'] == 'subcompetencia']:
        wpk[f'A{r}'] = e['codigo']; wpk[f'B{r}'] = NOMC[e['grupo']]; wpk[f'C{r}'] = e['texto']
        for j, g in enumerate(GP):
            col = get_column_letter(4 + j)
            wpk[f'{col}{r}'] = f'=COUNTIFS({TC("V")},{col}${hs},{TC("C")},"Ciencias",{TC("AW")},"*"&$A{r}&"*")'
        wpk[f'G{r}'] = f'=SUM(D{r}:F{r})'
        wpk[f'H{r}'] = f'=IF($B${COMP_ROW[CICLO_P]}=0,0,G{r}/$B${COMP_ROW[CICLO_P]})'
        wpk[f'I{r}'] = f'=IF(G{r}=0,"No trabajada",IF(H{r}<$F${thr},"Débil",IF(H{r}<$H${thr},"Básica","Sólida")))'
        style_range(wpk, r, r, 1, 9, align=WR); wpk[f'H{r}'].number_format = PCT
        r += 1
    s1 = r - 1
    for txt, colr in (('No trabajada', ROJO), ('Débil', AMAR), ('Básica', VCLARO), ('Sólida', VERDE_S)):
        wpk.conditional_formatting.add(f'I{s0}:I{s1}', CellIsRule(operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=colr)))
    r += 1
    wpk[f'A{r}'] = 'Bloque 3 · Peso de cada competencia según las microhabilidades (5 por tema), ciclo 7.°–9.°'; wpk[f'A{r}'].font = fS; r += 1
    hdr(wpk, r, ['Competencia', 'Microhabilidades (veces)', '% del total', 'Meta (punto medio)', 'Diferencia (pp)'], None, height=32); r += 1
    f0 = r
    for k in ('C1', 'C2', 'C3'):
        wpk[f'A{r}'] = NOMC[k]
        wpk[f'B{r}'] = '__FAM_' + k  # se completa con la matriz fase × código (bloque 4)
        wpk[f'D{r}'] = f'=$D${RC[NOMC[k]]}'
        r += 1
    for x in range(f0, r):
        wpk[f'C{x}'] = f'=IF(SUM($B${f0}:$B${r - 1})=0,0,B{x}/SUM($B${f0}:$B${r - 1}))'
        wpk[f'E{x}'] = f'=(C{x}-D{x})*100'
        style_range(wpk, x, x, 1, 5); wpk[f'C{x}'].number_format = PCT; wpk[f'D{x}'].number_format = PCT; wpk[f'E{x}'].number_format = PP
    r += 1
    wpk[f'A{r}'] = 'Bloque 4 · Qué código lleva cada fase (ciclo 7.°–9.°): muestra cuánto fija la plantilla'; wpk[f'A{r}'].font = fS; r += 1
    hdr(wpk, r, ['Código', 'Subcompetencia'] + [f[3:] for f in FASES], None, height=30); r += 1
    m0 = r
    for e in [e for e in CATP if e['tipo'] == 'subcompetencia']:
        wpk[f'A{r}'] = e['codigo']; wpk[f'B{r}'] = e['texto']
        for j, fc in enumerate(fases_cols):
            wpk.cell(row=r, column=3 + j, value=f'=COUNTIFS({TC("AV")},"{CICLO_P}",{TC(fc)},"*"&$A{r}&"*")')
        style_range(wpk, r, r, 1, 7)
        r += 1
    wpk.conditional_formatting.add(f'C{m0}:G{r - 1}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))
    fam = {'C1': (m0, m0 + 5), 'C2': (m0 + 6, m0 + 9), 'C3': (m0 + 10, m0 + 14)}  # E1–E6, D1–D4, I1–I5 (orden del catálogo)
    for x in range(f0, f0 + 3):
        k = wpk[f'B{x}'].value[6:]
        wpk[f'B{x}'] = f'=SUM(C{fam[k][0]}:G{fam[k][1]})'
    wpk.conditional_formatting.add(f'C{m0}:G{r - 1}', CellIsRule(operator='greaterThanOrEqual', formula=['30'], fill=PatternFill('solid', fgColor=VERDE_S)))
    wpk.freeze_panes = 'B4'

    # ---------------- PISA_conocimiento: contenido / procedimental / epistémico
    wpn = wb.create_sheet('PISA_conocimiento')
    wpn['A1'] = 'PISA 2025 · Tipos de conocimiento: contenido, procedimental y epistémico'; wpn['A1'].font = fT
    wpn['A2'] = ('PISA reparte la prueba entre conocimiento de CONTENIDO (qué sabemos), PROCEDIMENTAL (cómo se obtienen datos confiables: variables, medición, '
                 'incertidumbre, representación de datos, diseño) y EPISTÉMICO (por qué creemos lo que la ciencia afirma: modelos, evidencia, razonamiento, consenso). '
                 'La malla no tiene una columna para esto: la clasificación es asistida por IA (hoja Temas, columnas AN-AP). Un tema cuenta en el tipo PREDOMINANTE; '
                 'aparte se cuentan los temas que trabajan algún elemento procedimental o epistémico de forma explícita.')
    wpn['A2'].font = fNote; wpn.merge_cells('A2:L2'); wpn['A2'].alignment = WR; wpn.row_dimensions[2].height = 56
    RK, r = tabla_rangos(wpn, 4, 'Metas de PISA por tipo de conocimiento (Tabla 2.3, total sobre sistemas)',
                         [(k, tuple(v)) for k, v in CATP_J['metas']['conocimiento'].items()],
                         'PISA reparte la PRUEBA por tipo de conocimiento; en un currículo es esperable que domine el contenido. Úsalo como referencia de lo que falta, no como meta estricta.')
    wpn[f'A{r}'] = 'Bloque 1 · Tipo de conocimiento predominante por grado'; wpn[f'A{r}'].font = fS; r += 1
    hdr(wpn, r, ['Grado / ciclo', 'Temas', 'Contenido', 'Procedimental', 'Epistémico', '% Contenido', '% Procedimental', '% Epistémico',
                 'Estado procedimental', 'Estado epistémico', 'Temas con algún procedimental explícito', 'Temas con algún epistémico explícito', 'COINCIDENCIA con la meta'],
        [16, 9, 10, 12, 11, 10, 12, 11, 18, 18, 15, 15, 13], height=48)
    r += 1
    KN_ROW = {}
    for g, kc in GRUPOS_P:
        wpn[f'A{r}'] = g
        wpn[f'B{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AN")},"<>")'
        for j, k in enumerate(('Contenido', 'Procedimental', 'Epistémico')):
            n_, p_ = get_column_letter(3 + j), get_column_letter(6 + j)
            wpn[f'{n_}{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AN")},"{k}")'
            wpn[f'{p_}{r}'] = f'=IF($B{r}=0,0,{n_}{r}/$B{r})'
        for col, k in (('I', 'Procedimental'), ('J', 'Epistémico')):
            pc_ = 'G' if k == 'Procedimental' else 'H'; rr = RK[k]
            wpn[f'{col}{r}'] = f'=IF({pc_}{r}<$B${rr},"Por debajo del rango",IF({pc_}{r}>$C${rr},"Por encima del rango","Dentro del rango"))'
        wpn[f'K{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AO")},"<>")'
        wpn[f'L{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AP")},"<>")'
        wpn[f'M{r}'] = '=1-(' + '+'.join(f'ABS({get_column_letter(6 + j)}{r}-$D${RK[k]})' for j, k in enumerate(('Contenido', 'Procedimental', 'Epistémico'))) + ')/2'
        style_range(wpn, r, r, 1, 13)
        for col in 'FGHM': wpn[f'{col}{r}'].number_format = PCT
        KN_ROW[g] = r; r += 1
    estado_rango(wpn, f'I{KN_ROW["7°"]}:J{r}')
    r += 1
    wpn[f'A{r}'] = 'Bloque 2 · Tabla 2.3 de PISA en el acumulado 7.°–9.°: % de temas por tipo de conocimiento × sistema (vs. rango del marco)'; wpn[f'A{r}'].font = fS; r += 1
    hdr(wpn, r, ['Tipo de conocimiento'] + [f'{s_} · actual' for s_ in SISTEMAS] + [f'{s_} · rango PISA' for s_ in SISTEMAS], None, height=48); r += 1
    den = f'(COUNTIFS({TC("AV")},"{CICLO_P}")-COUNTIFS({TC("AV")},"{CICLO_P}",{TC("AK")},"FUERA"))'
    for k in ('Contenido', 'Procedimental', 'Epistémico'):
        wpn[f'A{r}'] = k
        for j, s_ in enumerate(SISTEMAS):
            col = get_column_letter(2 + j)
            wpn[f'{col}{r}'] = f'=IF({den}=0,0,COUNTIFS({TC("AV")},"{CICLO_P}",{TC("AN")},$A{r},{TC("AM")},"{s_}")/{den})'
            lo, hi = CATP_J['tabla_2_3'][k][s_]
            wpn.cell(row=r, column=5 + j, value=f'{round(lo * 100)}–{round(hi * 100)} %')
            wpn[f'{col}{r}'].number_format = PCT
        style_range(wpn, r, r, 1, 7)
        for col in 'BCD': wpn[f'{col}{r}'].number_format = PCT
        r += 1
    r += 1
    wpn[f'A{r}'] = 'Bloque 3 · Elementos procedimentales (Box 2.7) y epistémicos (Box 2.8): temas que los trabajan de forma explícita'; wpn[f'A{r}'].font = fS; r += 1
    hdr(wpn, r, ['Código', 'Tipo', 'Bloque', 'Elemento (paráfrasis)'] + GP + ['Total 7°–9°', 'Estado (semáforo)'], None, height=32)
    he = r; r += 1
    e0 = r
    for e in [e for e in CATP if e['tipo'] in ('procedimental', 'epistemico')]:
        colc = 'AO' if e['tipo'] == 'procedimental' else 'AP'
        wpn[f'A{r}'] = e['codigo']; wpn[f'B{r}'] = 'Procedimental' if e['tipo'] == 'procedimental' else 'Epistémico'
        wpn[f'C{r}'] = e['grupo']; wpn[f'D{r}'] = e['texto']
        for j, g in enumerate(GP):
            col = get_column_letter(5 + j)
            wpn[f'{col}{r}'] = f'=COUNTIFS({TC("V")},{col}${he},{TC(colc)},"*"&$A{r}&";*")'
        wpn[f'H{r}'] = f'=SUM(E{r}:G{r})'
        wpn[f'I{r}'] = '=' + profundidad(f'H{r}')
        style_range(wpn, r, r, 1, 9, align=WR)
        r += 1
    semaforo(wpn, f'I{e0}:I{r - 1}')
    wpn.conditional_formatting.add(f'E{e0}:G{r - 1}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))
    r += 1
    hdr(wpn, r, ['Tipo', 'Elementos', 'Trabajados (≥1 tema)', '% trabajado', 'No trabajados'], None, height=30); r += 1
    PE_ROW = {}
    for k in ('Procedimental', 'Epistémico'):
        wpn[f'A{r}'] = k
        wpn[f'B{r}'] = f'=COUNTIFS($B${e0}:$B${e0 + 27},$A{r})'
        wpn[f'C{r}'] = f'=COUNTIFS($B${e0}:$B${e0 + 27},$A{r},$H${e0}:$H${e0 + 27},">0")'
        wpn[f'D{r}'] = f'=IF(B{r}=0,0,C{r}/B{r})'; wpn[f'E{r}'] = f'=B{r}-C{r}'
        style_range(wpn, r, r, 1, 5, fill=SUB, font=fB); wpn[f'D{r}'].number_format = PCT
        PE_ROW[k] = r; r += 1
    wpn.column_dimensions['D'].width = 50

    # ---------------- PISA_contextos: personal / local-nacional / global, áreas y demanda cognitiva
    wpx = wb.create_sheet('PISA_contextos')
    wpx['A1'] = 'PISA 2025 · Contextos, áreas de aplicación y demanda cognitiva (capas propias de PISA)'; wpx['A1'].font = fT
    wpx['A2'] = ('PISA sitúa sus preguntas en contextos personales, locales/nacionales y globales en proporción ≈ 1:2:1 (p. 53), dentro de cinco áreas de aplicación '
                 '(Tabla 2.2). No es una meta curricular, pero indica si la malla enseña la ciencia conectada con situaciones reales. «Sin contexto» = el tema es '
                 'conceptual o de laboratorio, sin situación de vida real. La demanda cognitiva (baja/media/alta, Figura 2.2) no tiene meta en el marco. '
                 'Clasificación asistida por IA (hoja Temas, columnas AQ-AS).')
    wpx['A2'].font = fNote; wpx.merge_cells('A2:L2'); wpx['A2'].alignment = WR; wpx.row_dimensions[2].height = 56
    RX, r = tabla_rangos(wpx, 4, 'Proporción de contextos en la prueba (1:2:1)', [(k, tuple(v)) for k, v in CATP_J['metas']['contexto'].items()],
                         'Celdas ámbar editables. Se compara el reparto ENTRE los temas que sí tienen contexto.')
    wpx[f'A{r}'] = 'Bloque 1 · Contexto de los temas'; wpx[f'A{r}'].font = fS; r += 1
    CTX = ('Personal', 'Local/nacional', 'Global')
    hdr(wpx, r, ['Grado / ciclo', 'Temas', 'Personal', 'Local/nacional', 'Global', 'Sin contexto', '% sin contexto',
                 '% personal (de los con contexto)', '% local/nacional', '% global', 'COINCIDENCIA 1:2:1'], [16, 9, 10, 12, 10, 11, 11, 13, 13, 11, 13], height=48)
    r += 1
    CX_ROW = {}
    for g, kc in GRUPOS_P:
        wpx[f'A{r}'] = g
        wpx[f'B{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AQ")},"<>")'
        for j, k in enumerate(CTX + ('Sin contexto',)):
            wpx.cell(row=r, column=3 + j, value=f'=COUNTIFS({TC(kc)},$A{r},{TC("AQ")},"{k}")')
        wpx[f'G{r}'] = f'=IF(B{r}=0,0,F{r}/B{r})'
        for j in range(3):
            wpx.cell(row=r, column=8 + j, value=f'=IF(SUM($C{r}:$E{r})=0,0,{get_column_letter(3 + j)}{r}/SUM($C{r}:$E{r}))')
        wpx[f'K{r}'] = '=1-(' + '+'.join(f'ABS({get_column_letter(8 + j)}{r}-$D${RX[k]})' for j, k in enumerate(CTX)) + ')/2'
        style_range(wpx, r, r, 1, 11)
        for col in 'GHIJK': wpx[f'{col}{r}'].number_format = PCT
        CX_ROW[g] = r; r += 1
    wpx.conditional_formatting.add(f'G{CX_ROW["7°"]}:G{r}', CellIsRule(operator='greaterThan', formula=['0.5'], fill=PatternFill('solid', fgColor=ROJO)))
    r += 1
    wpx[f'A{r}'] = 'Bloque 2 · Áreas de aplicación (Tabla 2.2): n.º de temas'; wpx[f'A{r}'].font = fS; r += 1
    hdr(wpx, r, ['Área de aplicación'] + GP + ['Total 7°–9°'], None, height=30); ha = r; r += 1
    for e in [e for e in CATP if e['tipo'] == 'area_aplicacion']:
        wpx[f'A{r}'] = e['texto']
        for j, g in enumerate(GP):
            col = get_column_letter(2 + j)
            wpx[f'{col}{r}'] = f'=COUNTIFS({TC("V")},{col}${ha},{TC("AR")},$A{r})'
        wpx[f'E{r}'] = f'=SUM(B{r}:D{r})'
        style_range(wpx, r, r, 1, 5); r += 1
    AREA_ROWS = (ha + 1, r - 1)
    wpx.conditional_formatting.add(f'B{ha + 1}:D{r - 1}', CellIsRule(operator='equal', formula=['0'], fill=PatternFill('solid', fgColor=AMAR)))
    r += 1
    wpx[f'A{r}'] = 'Bloque 3 · Demanda cognitiva de lo que pide cada tema (Figura 2.2; sin meta en el marco)'; wpx[f'A{r}'].font = fS; r += 1
    hdr(wpx, r, ['Grado / ciclo', 'Temas', 'Baja', 'Media', 'Alta', '% baja', '% media', '% alta', '% Razonar en la malla (TIMSS, mismo grupo)'], None, height=40); r += 1
    DM_ROW = {}
    for g, kc in GRUPOS_P:
        wpx[f'A{r}'] = g
        wpx[f'B{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AS")},"<>")'
        for j, k in enumerate(('Baja', 'Media', 'Alta')):
            wpx.cell(row=r, column=3 + j, value=f'=COUNTIFS({TC(kc)},$A{r},{TC("AS")},"{k}")')
            wpx.cell(row=r, column=6 + j, value=f'=IF($B{r}=0,0,{get_column_letter(3 + j)}{r}/$B{r})').number_format = PCT
        wpx[f'I{r}'] = (f'=IF(COUNTIFS({TC(kc)},$A{r},{TC("C")},"Ciencias")=0,0,'
                        f'SUMIFS({TC("U")},{TC(kc)},$A{r},{TC("C")},"Ciencias")/COUNTIFS({TC(kc)},$A{r},{TC("C")},"Ciencias"))')
        style_range(wpx, r, r, 1, 9)
        for col in 'FGHI': wpx[f'{col}{r}'].number_format = PCT
        DM_ROW[g] = r; r += 1
    wpx.column_dimensions['A'].width = 40

# ------------------------------------------------------------------ Apego_escenarios (índice de apego y escenarios TIMSS/PISA para 7.°–8.°)
if HAY_PISA:
    wap = wb.create_sheet('Apego_escenarios')
    wap['A1'] = 'Apego a los marcos: índice por grado y escenarios TIMSS / PISA para 7.° y 8.°'; wap['A1'].font = fT
    wap['A2'] = ('COINCIDENCIA = 1 − ½ Σ |% actual − % meta|: es la parte del reparto de la malla que ya coincide con el reparto del marco (100 % = idéntico). '
                 'Se calcula para el balance de contenidos (dominios TIMSS o sistemas PISA) y para el reparto cognitivo (Conocer/Aplicar/Razonar en TIMSS; '
                 'las 3 competencias en PISA). En los ciclos se agrega la AMPLITUD (% de objetivos o categorías cubiertos). ÍNDICE DE APEGO = promedio de lo disponible. '
                 '7.° y 8.° se miden con los dos marcos; qué peso dar a cada uno es una decisión de política (celda ámbar E4).')
    wap['A2'].font = fNote; wap.merge_cells('A2:L2'); wap['A2'].alignment = WR; wap.row_dimensions[2].height = 56
    wap['A4'] = 'Peso de PISA en 7.° y 8.° (0 % = solo TIMSS):'; wap['A4'].font = fB
    wap['E4'] = 0.5; wap['E4'].font = fIn; wap['E4'].number_format = '0%'; wap['E4'].border = BOX
    wap['E4'].comment = Comment('Supuesto editable. El equipo GOES anticipa que PISA será prioridad en 7.°–9.°; 50 % es neutral.', 'Claude')
    wap['G4'] = 'Lectura: alto desde'; wap['H4'] = 0.85; wap['I4'] = 'medio desde'; wap['J4'] = 0.70
    for col in 'HJ':
        wap[f'{col}4'].font = fIn; wap[f'{col}4'].number_format = '0%'; wap[f'{col}4'].border = BOX
    r = 6
    wap[f'A{r}'] = 'A · Índice de apego por grado y ciclo'; wap[f'A{r}'].font = fS; r += 1
    hdr(wap, r, ['Grado / ciclo', 'Marco', 'Coincidencia de balance', 'Coincidencia cognitiva', 'Amplitud (ciclo)', 'ÍNDICE DE APEGO', 'Lectura'],
        [18, 22, 13, 13, 12, 13, 22, 12, 12, 12, 12, 12], height=40)
    r += 1
    AP = {}
    cogrow = {wg[f'A{x}'].value: x for x in range(5, CGL + 1)}
    ampmap_ = {m: rr for (m, d, rr) in AMP if d == '(todos)'}

    def fila_apego(etq, marco_txt, bal, cog, amp):
        global r
        wap[f'A{r}'] = etq; wap[f'B{r}'] = marco_txt
        wap[f'C{r}'] = bal; wap[f'D{r}'] = cog; wap[f'E{r}'] = amp
        wap[f'F{r}'] = f'=AVERAGE(C{r}:E{r})'
        wap[f'G{r}'] = f'=IF(F{r}>=$H$4,"Alto",IF(F{r}>=$J$4,"Medio","Bajo: revisar"))'
        style_range(wap, r, r, 1, 7)
        for col in 'CDEF': wap[f'{col}{r}'].number_format = PCT
        for col in 'CDE': wap[f'{col}{r}'].font = fLink
        wap[f'F{r}'].font = fB
        AP[etq] = r; r += 1

    for g, m, r0, r1, rt in BG:
        fila_apego(g if m != 'TA' else g, MNAME[m], f'=Balance_grado!O{rt}', f'=Cognitivo_grado!O{cogrow[g]}', '—')
    for g, m, r0, r1, rt in BC:
        mc = META[m]
        cog = ('=1-(' + '+'.join(f'ABS(SUMIFS({TC(col)},{TC("W")},$A{{r}})/COUNTIFS({TC("W")},$A{{r}})-{v})'
                                 for col, v in zip('STU', mc)) + ')/2')
        fila_apego(g, MNAME[m], f'=Balance_ciclo!O{rt}', cog.format(r=r), f'=Cobertura_objetivos!E{ampmap_[m]}')
    for (g, m, r0, r1, rt) in BP:
        fila_apego(f'{g} · PISA', 'PISA 2025', f'=PISA_contenido!O{rt}', f'=PISA_competencias!L{COMP_ROW[g]}',
                   f'=PISA_contenido!D{AMP_P["(todos)"]}' if g == CICLO_P else '—')
    for g in ('7°', '8°'):
        wap[f'A{r}'] = f'{g} · escenario elegido'; wap[f'B{r}'] = 'Mezcla según el peso de E4'
        wap[f'F{r}'] = f'=(1-$E$4)*F{AP[g]}+$E$4*F{AP[g + " · PISA"]}'
        wap[f'G{r}'] = f'=IF(F{r}>=$H$4,"Alto",IF(F{r}>=$J$4,"Medio","Bajo: revisar"))'
        style_range(wap, r, r, 1, 7, fill=SUB, font=fB); wap[f'F{r}'].number_format = PCT
        AP[f'{g} · escenario'] = r; r += 1
    a_last = r - 1
    for txt, colr in (('Alto', VERDE_S), ('Medio', AMAR), ('Bajo: revisar', ROJO)):
        wap.conditional_formatting.add(f'G8:G{a_last}', CellIsRule(operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=colr)))

    # ---- B · escenarios por sistema para 7.° y 8.°
    r += 2
    wap[f'A{r}'] = 'B · Escenarios para 7.° y 8.°: ¿cuántos temas faltan (+) o sobran (−) en cada sistema según el peso de cada marco?'; wap[f'A{r}'].font = fS; r += 1
    wap[f'A{r}'] = ('El eje común son los 3 sistemas de PISA. Las metas de TIMSS 8.° se agrupan igual: Sistemas físicos = Química 20 % + Física 25 %; vivos = Biología 35 %; '
                    'Tierra y espacio = Ciencias de la Tierra 20 %. En cada escenario: meta = (1 − p)·TIMSS + p·PISA y % actual = (1 − p)·actual TIMSS + p·actual PISA '
                    '(cada marco con su propia clasificación). Brecha (temas) = (meta − actual) × temas evaluables.')
    wap[f'A{r}'].font = fNote; wap.merge_cells(f'A{r}:N{r}'); wap[f'A{r}'].alignment = WR; wap.row_dimensions[r].height = 42; r += 1
    ESC = [('100 % TIMSS', '0'), ('50 / 50', '0.5'), ('100 % PISA', '1'), ('Personalizado', '$E$4')]
    hdr(wap, r, ['Grado', 'Sistema', 'Meta TIMSS (agrupada)', 'Meta PISA', '% actual TIMSS', '% actual PISA'] +
        [f'{n} · brecha (temas)' for n, _ in ESC] + ['Evaluables TIMSS', 'Evaluables PISA'], None, height=48)
    r += 1
    wrow = r
    wap[f'A{r}'] = 'Peso de PISA →'; wap[f'A{r}'].font = fNote
    for j, (_, w) in enumerate(ESC):
        c_ = wap.cell(row=r, column=7 + j, value=f'={w}' if w.startswith('$') else float(w)); c_.number_format = '0%'; c_.font = fNote
    r += 1
    MAPA_T = {'Sistemas físicos': ('Química', 'Física'), 'Sistemas vivos': ('Biología',), 'Sistemas de la Tierra y el espacio': ('Ciencias de la Tierra',)}
    META_T8 = {d: v for d, v in DOMS[M8]}
    ESC_COINC = {}
    for g in ('7°', '8°'):
        g0 = r
        for s_ in SISTEMAS:
            wap[f'A{r}'] = g; wap[f'B{r}'] = s_
            wap[f'C{r}'] = sum(META_T8[d] for d in MAPA_T[s_]); wap[f'C{r}'].font = fIn
            wap[f'D{r}'] = CATP_J['metas']['sistema'][s_][0]; wap[f'D{r}'].font = fIn
            wap[f'K{r}'] = (f'=COUNTIFS({TC("V")},$A{r},{TC("C")},"Ciencias")-COUNTIFS({TC("V")},$A{r},{TC("K")},"FUERA")'
                            f'-COUNTIFS({TC("V")},$A{r},{TC("K")},"PENDIENTE")')
            wap[f'L{r}'] = f'=COUNTIFS({TC("V")},$A{r},{TC("AK")},"<>")-COUNTIFS({TC("V")},$A{r},{TC("AK")},"FUERA")'
            wap[f'E{r}'] = f'=IF(K{r}=0,0,(' + '+'.join(f'COUNTIFS({TC("V")},$A{r},{TC("M")},"{d}")' for d in MAPA_T[s_]) + f')/K{r})'
            wap[f'F{r}'] = f'=IF(L{r}=0,0,COUNTIFS({TC("V")},$A{r},{TC("AM")},$B{r})/L{r})'
            for j in range(4):
                col = get_column_letter(7 + j); w = f'{col}${wrow}'
                wap[f'{col}{r}'] = f'=((1-{w})*C{r}+{w}*D{r}-((1-{w})*E{r}+{w}*F{r}))*((1-{w})*K{r}+{w}*L{r})'
                wap[f'{col}{r}'].number_format = '+0.0;-0.0;0.0'
            style_range(wap, r, r, 1, 12)
            for col in 'CDEF': wap[f'{col}{r}'].number_format = PCT
            wap[f'C{r}'].font = fIn; wap[f'D{r}'].font = fIn
            for col in 'GHIJ': wap[f'{col}{r}'].number_format = '+0.0;-0.0;0.0'
            r += 1
        wap[f'B{r}'] = 'Coincidencia de balance'
        for j in range(4):
            col = get_column_letter(7 + j); w = f'{col}${wrow}'
            wap[f'{col}{r}'] = ('=1-(' + '+'.join(f'ABS((1-{w})*C{x}+{w}*D{x}-((1-{w})*E{x}+{w}*F{x}))' for x in range(g0, r)) + ')/2')
            wap[f'{col}{r}'].number_format = PCT
        style_range(wap, r, r, 1, 12, fill=SUB, font=fB)
        for j in range(4): wap.cell(row=r, column=7 + j).number_format = PCT
        ESC_COINC[g] = r; r += 1
        wap[f'B{r}'] = 'Índice de apego del escenario'
        for j in range(4):
            col = get_column_letter(7 + j); w = f'{col}${wrow}'
            wap[f'{col}{r}'] = f'=(1-{w})*$F${AP[g]}+{w}*$F${AP[g + " · PISA"]}'
            wap[f'{col}{r}'].number_format = PCT
        style_range(wap, r, r, 1, 12, fill=SUB, font=fB)
        for j in range(4): wap.cell(row=r, column=7 + j).number_format = PCT
        r += 2
    wap.conditional_formatting.add(f'G{wrow + 1}:J{r}', FormulaRule(formula=[f'AND(ISNUMBER(G{wrow + 1}),G{wrow + 1}>=3,$B{wrow + 1}<>"Coincidencia de balance",LEFT($B{wrow + 1},6)<>"Índice")'],
                                                                     fill=PatternFill('solid', fgColor=ROJO)))
    wap.conditional_formatting.add(f'G{wrow + 1}:J{r}', FormulaRule(formula=[f'AND(ISNUMBER(G{wrow + 1}),G{wrow + 1}<=-3,$B{wrow + 1}<>"Coincidencia de balance",LEFT($B{wrow + 1},6)<>"Índice")'],
                                                                     fill=PatternFill('solid', fgColor=AMAR)))
    r += 1
    wap[f'A{r}'] = 'C · Cómo usar los escenarios (recomendación)'; wap[f'A{r}'].font = fS; r += 1
    for txt in ('1. Los tres escenarios pedidos (100 % TIMSS, 50/50, 100 % PISA) están lado a lado; la columna «Personalizado» toma el peso de la celda E4, para probar cualquier mezcla (p. ej. 70 % PISA).',
                '2. Mezclar METAS solo tiene sentido en un eje común (los 3 sistemas). Las metas cognitivas NO se mezclan: TIMSS mide Conocer/Aplicar/Razonar y PISA tres competencias distintas; '
                'por eso el escenario mezcla los ÍNDICES de apego, no las metas cognitivas.',
                '3. 9.° ya no se compara con TIMSS (decisión del equipo): solo PISA. 5.° y 6.° siguen solo con TIMSS 8.° (que evalúa lo acumulado de 5.° a 8.°).',
                '4. Si El Salvador participa en TIMSS 2027 (8.° grado), 8.° sigue siendo el grado evaluado por TIMSS: un 100 % PISA en 8.° dejaría de alinear justo el grado que TIMSS mide. '
                'Una mezcla 50/50 en 8.° y 70–100 % PISA en 7.° y 9.° es una opción razonable mientras se decide.'):
        wap[f'A{r}'] = txt; wap[f'A{r}'].font = fN; wap.merge_cells(f'A{r}:N{r}'); wap[f'A{r}'].alignment = WR; wap.row_dimensions[r].height = 30; r += 1
    wap.freeze_panes = 'A6'

# ------------------------------------------------------------------ Presupuesto (inserciones y fusiones con el mismo total de temas)
PRES_F = 'data/referencia/presupuesto_candidatos.json'
if HAY_PISA and os.path.exists(PRES_F):
    PRES = json.load(open(PRES_F))
    wpr = wb.create_sheet('Presupuesto')
    wpr['A1'] = 'Presupuesto de temas: qué insertar y qué fusionar para acercarse a las metas SIN aumentar el número de temas'; wpr['A1'].font = fT
    wpr['A2'] = ('El «presupuesto» es el número de temas de cada ciclo, que se mantiene. Se FUSIONAN temas que se repiten en el área que sobra (libera espacio) y se '
                 'INSERTAN temas en el área que falta (usa ese espacio). CUÁNTOS: fórmula sobre la brecha del ciclo (Balance_ciclo y PISA_contenido; si el área está '
                 '«En meta» no se toca). CUÁLES: candidatos propuestos por IA leyendo los temas, en orden de prioridad; «Entra» = está dentro de los N necesarios. '
                 'Inserción por COBERTURA = objetivo que nadie trabaja, aunque su área sobre: siempre entra y se paga con una fusión extra de la misma área. '
                 'REUBICAR = mover un tema de grado dentro del ciclo (no cambia el total; equilibra grados y adelanta lo que otros países enseñan antes).')
    wpr['A2'].font = fNote; wpr.merge_cells('A2:Q2'); wpr['A2'].alignment = WR; wpr.row_dimensions[2].height = 70
    wpr['A3'] = ('Se trabaja por CICLO y no por grado: la meta del marco es para lo acumulado, y por grado salen brechas grandes que se compensan entre sí '
                 '(p. ej. 5.°–8.°: Biología +19 en 5.° y −15 en 7.°, que se resuelve reubicando, no insertando y fusionando). '
                 '7.° y 8.° aparecen en el ciclo TIMSS 5.°–8.° y en el PISA 7.°–9.°: según el peso elegido en Apego_escenarios, seguir uno u otro bloque.')
    wpr['A3'].font = fNote; wpr.merge_cells('A3:Q3'); wpr['A3'].alignment = WR; wpr.row_dimensions[3].height = 42
    r = 5
    hdr(wpr, r, ['Clave', 'Ciclo', 'Marco', 'Área', 'Brecha del ciclo (temas)', 'Estado', 'Inserciones por balance', 'Inserciones por cobertura',
                 'Temas a liberar con fusiones', 'Temas que liberan las fusiones elegidas', 'Saldo (insertados − liberados)'],
        [16, 13, 16, 30, 11, 18, 11, 11, 11, 13, 13, 18, 30, 50, 40, 10, 14], height=54)
    r += 1
    CICLO_N = {'2-4': 'Ciclo 2°–4°', '5-8': 'Ciclo 5°–8°', '7-9': CICLO_P}
    bloques = [('2-4', 'Balance_ciclo', [b for b in BC if b[0] == 'Ciclo 2°–4°'][0]),
               ('5-8', 'Balance_ciclo', [b for b in BC if b[0] == 'Ciclo 5°–8°'][0]),
               ('7-9', 'PISA_contenido', BP[-1])]
    p0 = r
    PRES_TOT = {}
    cand_first = None  # se completa abajo (filas de candidatos)
    for clave, hoja, (g, m, b0, b1, bt) in bloques:
        rr0 = r
        for x in range(b0, b1 + 1):
            area_ = (wcy if hoja == 'Balance_ciclo' else wpc)[f'C{x}'].value
            wpr[f'A{r}'] = f'{clave}|{area_}'; wpr[f'B{r}'] = CICLO_N[clave]; wpr[f'C{r}'] = MNAME[m]; wpr[f'D{r}'] = area_
            wpr[f'E{r}'] = f'={hoja}!I{x}'; wpr[f'F{r}'] = f'={hoja}!K{x}'
            wpr[f'G{r}'] = '__G__'  # se completa al cerrar el bloque del ciclo
            style_range(wpr, r, r, 1, 11)
            wpr[f'E{r}'].number_format = '+0.0;-0.0;0.0'; wpr[f'E{r}'].font = fLink; wpr[f'F{r}'].font = fLink
            r += 1
        for x in range(rr0, r):  # si algún área del ciclo sale de la tolerancia, se aplican todas las brechas (el total de temas se mantiene)
            wpr[f'G{x}'] = f'=IF(COUNTIF($F${rr0}:$F${r - 1},"En meta")=ROWS($F${rr0}:$F${r - 1}),0,MAX(0,ROUND(E{x},0)))'
        wpr[f'D{r}'] = f'Total {CICLO_N[clave]}'
        PRES_TOT[clave] = (rr0, r - 1, r)
        style_range(wpr, r, r, 1, 11, fill=SUB, font=fB)
        r += 2
    p1 = r - 2
    # ---- candidatos
    wpr[f'A{r}'] = 'Candidatos propuestos (ordenados por ciclo, acción, área y prioridad). La columna «Decisión del equipo» es para que el equipo de Ciencias anote si acepta.'
    wpr[f'A{r}'].font = fS; r += 1
    hdr(wpr, r, ['ID', 'Clave', 'Ciclo', 'Acción', 'Motivo', 'Área', 'Prioridad', 'Grado', 'Unidad', 'Temas involucrados (ID · procedimental)',
                 'Objetivos del marco', 'Propuesta', 'Justificación', 'Confianza', 'Libera (temas)', '¿Entra en el presupuesto?', 'Decisión del equipo'],
        None, height=40)
    r += 1
    q0 = r
    Tid2 = {t['id']: t for t in T}
    ORD = {'insertar': 0, 'fusionar': 1, 'reubicar': 2}
    for c in sorted(PRES, key=lambda c: (c['ciclo'], ORD[c['accion']], c['area'], c['prioridad'])):
        temas_txt = '\n'.join(f"{t} · {Tid2[t]['grado']}.° · {Tid2[t]['procedimental'][:90]}" for t in c['temas'])
        vals = [c['id'], None, CICLO_N[c['ciclo']], c['accion'], c.get('motivo', ''), c['area'], c['prioridad'], c['grado'], c['unidad'],
                temas_txt, ', '.join(c['objetivos']), c['propuesta'], c['justificacion'], c['confianza'], c['libera'], None, '']
        for j, v in enumerate(vals, 1):
            if v is not None:
                wpr.cell(row=r, column=j, value=v)
        wpr[f'B{r}'] = f'="{c["ciclo"]}|"&F{r}'
        style_range(wpr, r, r, 1, 17, align=WR)
        wpr[f'Q{r}'].fill = INP; wpr[f'Q{r}'].font = fIn
        r += 1
    q1 = r - 1
    KEYS = f'$A${p0}:$A${p1}'
    for x in range(q0, q1 + 1):
        wpr[f'P{x}'] = (f'=IF(D{x}="reubicar","Opcional (no cambia el total)",IF(E{x}="cobertura","Sí (cobertura)",'
                        f'IF(D{x}="insertar",IF(G{x}<=IFERROR(INDEX($G${p0}:$G${p1},MATCH(B{x},{KEYS},0)),0),"Sí","Reserva"),'
                        f'IF(SUMIFS($O${q0}:$O${q1},$B${q0}:$B${q1},B{x},$D${q0}:$D${q1},"fusionar",$G${q0}:$G${q1},"<"&G{x})'
                        f'<IFERROR(INDEX($I${p0}:$I${p1},MATCH(B{x},{KEYS},0)),0),"Sí","Reserva"))))')
        wpr[f'P{x}'].font = fB
    QR = lambda col: f'${col}${q0}:${col}${q1}'
    for clave, (rr0, rr1, rt_) in PRES_TOT.items():
        for x in range(rr0, rr1 + 1):
            wpr[f'H{x}'] = f'=COUNTIFS({QR("B")},A{x},{QR("D")},"insertar",{QR("E")},"cobertura")'
            wpr[f'I{x}'] = f'=IF(COUNTIF($F${rr0}:$F${rr1},"En meta")=ROWS($F${rr0}:$F${rr1}),0,MAX(0,ROUND(-E{x},0)))+H{x}'
            wpr[f'J{x}'] = f'=SUMIFS({QR("O")},{QR("B")},A{x},{QR("D")},"fusionar",{QR("P")},"Sí")'
            wpr[f'K{x}'] = f'=G{x}+H{x}-J{x}'
        for col in 'GHIJK':
            wpr[f'{col}{rt_}'] = f'=SUM({col}{rr0}:{col}{rr1})'
    for txt, colr in (('Sí', VERDE_S), ('Sí (cobertura)', VERDE_S), ('Reserva', 'EDEDED')):
        wpr.conditional_formatting.add(f'P{q0}:P{q1}', CellIsRule(operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=colr)))
    wpr.conditional_formatting.add(f'N{q0}:N{q1}', CellIsRule(operator='equal', formula=['"baja"'], fill=PatternFill('solid', fgColor=ROJO)))
    wpr.freeze_panes = f'B{q0}'; wpr.auto_filter.ref = f'A{q0 - 1}:Q{q1}'
    wpr.column_dimensions['J'].width = 60; wpr.column_dimensions['L'].width = 55; wpr.column_dimensions['M'].width = 40

# ------------------------------------------------------------------ Comparación con otros países (Prompt 02)
PAISES_F = 'data/interim/paises/alineacion.json'
HAY_PAISES = ES27 and os.path.exists(PAISES_F)
if HAY_PAISES:
    PA = json.load(open(PAISES_F))
    NOMBRES_P = [p for p in ('Uruguay', 'Colombia', 'Singapur') if any(o['pais'] == p for o in PA)]
    GRADOS = list(range(1, 10))

    # ---------------- Paises_objetivos: un objetivo de país por fila (trazabilidad, como Temas_TIMSS)
    wp = wb.create_sheet('Paises_objetivos')
    wp['A1'] = 'Objetivos de otros países alineados con TIMSS 2027 (el pivote común)'; wp['A1'].font = fT
    wp['A2'] = ('Una fila = un objetivo/contenido del currículo de otro país, con su documento y página. Columnas J-K: alineación con TIMSS 2027 '
                'asistida por IA con el mismo método que los temas de la malla (revisable). M es fórmula. Nunca se compara país contra país: todo pasa por los códigos TIMSS.')
    wp['A2'].font = fNote; wp.merge_cells('A2:N2'); wp['A2'].alignment = WR; wp.row_dimensions[2].height = 30
    hdr(wp, 4, ['ID', 'País', 'Documento', 'Página', 'Grado o tramo', 'Grado (inicio)', 'Eje / área', 'Texto (paráfrasis o literal corto)',
                'Marco', 'Objetivo TIMSS principal', 'Objetivo TIMSS secundario', 'Confianza', 'Dominio (principal)', 'Justificación'],
        [17, 10, 34, 7, 10, 8, 26, 60, 8, 13, 13, 10, 22, 45], height=42)
    r = 5
    for o in sorted(PA, key=lambda o: (NOMBRES_P.index(o['pais']), o['grado_min'], o['id'])):
        vals = [o['id'], o['pais'], o['documento'], o['pagina'], o['grado_o_tramo'], o['grado_min'], o['eje'], o['texto'],
                o['marco'], o['obj1'], o['obj2'], o['confianza'], None, o['justificacion']]
        for j, v in enumerate(vals, 1):
            if v is not None:
                wp.cell(row=r, column=j, value=v)
        wp[f'M{r}'] = f'=IF(J{r}="FUERA","Fuera del marco",IFERROR(INDEX({CR("C")},MATCH(J{r},{CR("A")},0)),"?"))'
        style_range(wp, r, r, 1, 14)
        wp[f'H{r}'].alignment = Alignment(wrap_text=False, vertical='top')
        r += 1
    PL = r - 1
    wp.freeze_panes = 'C5'; wp.auto_filter.ref = f'A4:N{PL}'
    wp.conditional_formatting.add(f'L5:L{PL}', CellIsRule(operator='equal', formula=['"baja"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    wp.conditional_formatting.add(f'J5:J{PL}', CellIsRule(operator='equal', formula=['"FUERA"'], fill=GRY))
    PC = lambda col: f"Paises_objetivos!${col}$5:${col}${PL}"

    # ---------------- Comparacion_paises: primer grado por objetivo TIMSS 2027 y país
    wx = wb.create_sheet('Comparacion_paises')
    wx['A1'] = '¿En qué grado introduce cada país los objetivos de TIMSS 2027? El Salvador frente a ' + ', '.join(NOMBRES_P); wx['A1'].font = fT
    wx['A2'] = ('Primer grado = el grado más bajo con al menos un tema (El Salvador) u objetivo (otros países) que trabaja el objetivo TIMSS como principal o secundario. '
                'Oportunidad = grado El Salvador − mediana de los demás países que lo tienen: positivo = El Salvador llega más tarde; negativo = va adelantado. '
                'Las columnas de la derecha (n.º por grado) son la base de cada «primer grado».')
    wx['A2'].font = fNote; wx.merge_cells('A2:L2'); wx['A2'].alignment = WR; wx.row_dimensions[2].height = 44
    wx['A3'] = 'Grado mínimo comparable (la malla de El Salvador empieza en 2.°; lo anterior de otros países se cuenta como este grado):'; wx['A3'].font = fB
    wx['L3'] = 2; wx['L3'].font = fIn; wx['L3'].border = BOX
    wx['L3'].comment = Comment('Supuesto editable. 1.° grado no tiene malla en El Salvador: si otro país introduce algo en 1.°, se compara como 2.°.', 'Claude')
    quien = ['El Salvador'] + NOMBRES_P
    heads = ['Código', 'Marco', 'Dominio', 'Objetivo'] + [f'Primer grado · {p}' for p in quien] + \
            ['Mediana de los demás', 'Oportunidad (grados)', 'Lectura', 'Países que lo tienen (sin SV)']
    ncol = len(heads)
    # bloque auxiliar: n.º por grado y país, a la derecha
    aux0 = ncol + 2
    for k, p in enumerate(quien):
        for j, g in enumerate(GRADOS):
            col = get_column_letter(aux0 + k * len(GRADOS) + j)
            wx[f'{col}3'] = g; wx[f'{col}3'].font = fNote
            heads_aux = f'{p[:3].upper()} {g}°'
            wx.column_dimensions[col].width = 6
            c = wx.cell(row=4, column=aux0 + k * len(GRADOS) + j, value=heads_aux)
            c.font = fH; c.fill = HDR; c.alignment = CTR; c.border = BOX
    hdr(wx, 4, heads, [12, 7, 20, 50] + [10] * len(quien) + [10, 11, 30, 10], height=48)
    r = 5
    for c in [c for c in CAT if c['marco'] in (M4, M8)]:
        wx[f'A{r}'] = c['codigo']; wx[f'B{r}'] = c['marco']; wx[f'C{r}'] = c['dominio']; wx[f'D{r}'] = c['objetivo']
        for k, p in enumerate(quien):
            c1 = get_column_letter(aux0 + k * len(GRADOS)); c9 = get_column_letter(aux0 + (k + 1) * len(GRADOS) - 1)
            for j, g in enumerate(GRADOS):
                col = get_column_letter(aux0 + k * len(GRADOS) + j)
                if p == 'El Salvador':
                    wx[f'{col}{r}'] = (f'=COUNTIFS({TC("C")},"Ciencias",{TC("B")},{col}$3,{TC("K")},$A{r})'
                                       f'+COUNTIFS({TC("C")},"Ciencias",{TC("B")},{col}$3,{TC("L")},$A{r})')
                else:
                    wx[f'{col}{r}'] = (f'=COUNTIFS({PC("B")},"{p}",{PC("F")},{col}$3,{PC("J")},$A{r})'
                                       f'+COUNTIFS({PC("B")},"{p}",{PC("F")},{col}$3,{PC("K")},$A{r})')
            pc = get_column_letter(5 + k)
            wx[f'{pc}{r}'] = (f'=IFERROR(MAX($L$3,INDEX(${c1}$3:${c9}$3,MATCH(TRUE,INDEX({c1}{r}:{c9}{r}>0,0),0))),"—")')
        sv, o1, o2 = 'E', 'F', get_column_letter(4 + len(quien))
        cm, co, cl, cn = (get_column_letter(5 + len(quien) + i) for i in range(4))
        wx[f'{cn}{r}'] = f'=COUNT({o1}{r}:{o2}{r})'
        wx[f'{cm}{r}'] = f'=IF({cn}{r}=0,"—",MEDIAN({o1}{r}:{o2}{r}))'
        wx[f'{co}{r}'] = f'=IF(OR({sv}{r}="—",{cm}{r}="—"),"—",{sv}{r}-{cm}{r})'
        wx[f'{cl}{r}'] = (f'=IF({sv}{r}="—",IF({cn}{r}=0,"Nadie lo trabaja","El Salvador no llega"),IF({cn}{r}=0,"Solo El Salvador",'
                          f'IF({co}{r}>0,"El Salvador llega más tarde",IF({co}{r}<0,"El Salvador va adelantado","Mismo grado"))))')
        style_range(wx, r, r, 1, ncol, align=Alignment(vertical='top', wrap_text=True))
        style_range(wx, r, r, aux0, aux0 + len(quien) * len(GRADOS) - 1)
        wx[f'{cm}{r}'].number_format = '0.0'; wx[f'{co}{r}'].number_format = '+0.0;-0.0;0.0'; wx[f'{cl}{r}'].font = fB
        r += 1
    XL = r - 1
    CL_ = cl
    wx.conditional_formatting.add(f'{cl}5:{cl}{XL}', CellIsRule(operator='equal', formula=['"El Salvador no llega"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    wx.conditional_formatting.add(f'{cl}5:{cl}{XL}', CellIsRule(operator='equal', formula=['"El Salvador llega más tarde"'], fill=PatternFill('solid', fgColor='FFE699')))
    wx.conditional_formatting.add(f'{cl}5:{cl}{XL}', CellIsRule(operator='equal', formula=['"El Salvador va adelantado"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    wx.conditional_formatting.add(f'{get_column_letter(aux0)}5:{get_column_letter(aux0 + len(quien) * len(GRADOS) - 1)}{XL}',
                                  CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))
    # resumen por país debajo
    r = XL + 3
    wx[f'A{r}'] = 'Resumen por país'; wx[f'A{r}'].font = fS; r += 1
    hdr(wx, r, ['País', 'Objetivos alineados (1.°–9.°)', 'Fuera de TIMSS', 'Confianza baja', 'Objetivos TIMSS 2027 que trabaja (de 80)'], None, height=40); r += 1
    for k, p in enumerate(quien):
        pc = get_column_letter(5 + k)
        wx[f'A{r}'] = p
        if p == 'El Salvador':
            wx[f'B{r}'] = f'=COUNTIFS({TC("C")},"Ciencias")'
            wx[f'C{r}'] = f'=COUNTIFS({TC("C")},"Ciencias",{TC("K")},"FUERA")'
            wx[f'D{r}'] = f'=COUNTIFS({TC("C")},"Ciencias",{TC("P")},"baja")'
        else:
            wx[f'B{r}'] = f'=COUNTIFS({PC("B")},$A{r})'
            wx[f'C{r}'] = f'=COUNTIFS({PC("B")},$A{r},{PC("J")},"FUERA")'
            wx[f'D{r}'] = f'=COUNTIFS({PC("B")},$A{r},{PC("L")},"baja")'
        wx[f'E{r}'] = f'=COUNT({pc}5:{pc}{XL})'
        style_range(wx, r, r, 1, 5); r += 1
    r += 1
    hdr(wx, r, ['Lectura', 'Objetivos TIMSS 2027'], None, height=30); r += 1
    for lec in ('El Salvador llega más tarde', 'El Salvador va adelantado', 'Mismo grado', 'El Salvador no llega', 'Solo El Salvador', 'Nadie lo trabaja'):
        wx[f'A{r}'] = lec; wx[f'B{r}'] = f'=COUNTIFS(${cl}$5:${cl}${XL},$A{r})'
        style_range(wx, r, r, 1, 2); r += 1
    wx.freeze_panes = 'E5'; wx.auto_filter.ref = f'A4:{get_column_letter(ncol)}{XL}'

# ------------------------------------------------------------------ Resumen
ws = wb.create_sheet('Resumen', 0)
ws['A1'] = 'Cobertura de la malla de Ciencias de El Salvador frente a TIMSS — resumen por grado'; ws['A1'].font = fT
ws['A2'] = ('Todo es fórmula sobre Temas_TIMSS. Balance = reparto de temas por dominio vs. metas TIMSS. Amplitud = objetivos TIMSS que aparecen. '
            'Profundidad cognitiva = Conocer/Aplicar/Razonar. Ver «Leeme» para el proceso completo.')
ws['A2'].font = fNote; ws.merge_cells('A2:M2'); ws['A2'].alignment = WR; ws.row_dimensions[2].height = 30
hdr(ws, 4, ['Grupo', 'Marco de referencia', 'Temas totales', 'Temas fuera del marco', '% fuera del marco',
            'Dominio con mayor FALTANTE', 'Faltan (temas)', 'Dominio más SOBRErrepresentado', 'Sobran (temas)',
            'Temas a agregar para cumplir todas las metas sin quitar', 'Dominios en meta (±tolerancia)', '% Razonar actual', '% Razonar meta'],
    [12, 24, 9, 10, 10, 24, 9, 24, 9, 14, 11, 10, 10], height=60)
if ES27:  # encabezados que se explican solos (retroalimentación del equipo GOES)
    ws['A1'] = 'Resumen por grado: TIMSS 2027 (2.°–8.° y Física) arriba · PISA 2025 (7.°–9.°) abajo'
    hdr(ws, 4, ['Grupo', 'Marco de referencia', 'Temas totales',
                'Temas de la malla SIN objetivo en el marco (la malla los tiene; TIMSS no los evalúa)', '% de temas sin objetivo en el marco',
                'Dominio TIMSS con MÁS temas faltantes (frente a su meta)', 'Temas que faltan en ese dominio',
                'Dominio TIMSS con MÁS temas de sobra', 'Temas que sobran en ese dominio',
                'Temas a AGREGAR, sin quitar ninguno, para que TODOS los dominios lleguen a su meta', 'Dominios en meta (±tolerancia)',
                '% Razonar actual', '% Razonar meta'], [12, 24, 9, 16, 11, 24, 11, 24, 11, 18, 11, 10, 10], height=92)
    ws['A3'] = ('Lo contrario de «sin objetivo en el marco» (objetivos que TIMSS evalúa y la malla NO trabaja) está en el bloque por ciclo, columna «Objetivos no cubiertos», '
                'y en rojo en Cobertura_objetivos. Al final de esta hoja hay una guía de cada columna.')
    ws['A3'].font = fNote; ws.merge_cells('A3:M3'); ws['A3'].alignment = WR; ws.row_dimensions[3].height = 28
r = 5
for g, m, r0, r1, rt in BG:
    gi = [x for x in range(5, CGL + 1) if wg[f'A{x}'].value == g][0]
    ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]
    ws[f'C{r}'] = f'=COUNTIFS({TC("V")},$A{r})'
    ws[f'D{r}'] = f'=Balance_grado!N{r0}'
    ws[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
    I = f'Balance_grado!$I${r0}:$I${r1}'; Cc = f'Balance_grado!$C${r0}:$C${r1}'; K = f'Balance_grado!$K${r0}:$K${r1}'
    ws[f'F{r}'] = f'=INDEX({Cc},MATCH(MAX({I}),{I},0))'
    ws[f'G{r}'] = f'=ROUND(MAX({I}),0)'
    ws[f'H{r}'] = f'=INDEX({Cc},MATCH(MIN({I}),{I},0))'
    ws[f'I{r}'] = f'=ROUND(-MIN({I}),0)'
    ws[f'J{r}'] = f'=Balance_grado!M{rt}'
    ws[f'K{r}'] = f'=COUNTIF({K},"En meta")&" de "&ROWS({K})'
    ws[f'L{r}'] = f'=Cognitivo_grado!I{gi}'
    ws[f'M{r}'] = f'=Cognitivo_grado!L{gi}'
    style_range(ws, r, r, 1, 13)
    for a in 'DJLM': ws[f'{a}{r}'].font = fLink
    ws[f'E{r}'].number_format = PCT; ws[f'L{r}'].number_format = PCT; ws[f'M{r}'].number_format = '0%'
    r += 1
r += 1
ws[f'A{r}'] = 'Por ciclo (comparación acumulada)'; ws[f'A{r}'].font = fS; r += 1
hdr(ws, r, ['Ciclo', 'Marco de referencia', 'Temas evaluables', 'Dominio con mayor FALTANTE', 'Faltan (temas)',
            'Dominio más SOBRErrepresentado', 'Sobran (temas)', 'Temas a agregar sin quitar', 'Dominios en meta',
            'Objetivos TIMSS del marco', 'Objetivos cubiertos', '% amplitud', 'Objetivos no cubiertos'], None, height=45)
if ES27:
    hdr(ws, r, ['Ciclo', 'Marco de referencia', 'Temas evaluables', 'Dominio con mayor FALTANTE', 'Faltan (temas)',
                'Dominio más SOBRErrepresentado', 'Sobran (temas)', 'Temas a agregar sin quitar', 'Dominios en meta',
                'Objetivos TIMSS del marco', 'Objetivos cubiertos', '% amplitud', 'Objetivos del marco que la malla NO trabaja',
                'Objetivos cubiertos (2023)', '% amplitud (2023)', 'Dominio con mayor FALTANTE (2023)', 'Cambio vs 2023',
                '% temas ambientales (ref. ≈25 %)', 'Temas en Investigaciones (principal)'], None, height=58)
    for col, w in zip('NOPQRS', [11, 10, 22, 60, 12, 13]):
        ws.column_dimensions[col].width = w
r += 1
ampmap = {m: rr for (m, d, rr) in AMP if d == '(todos)'}
for g, m, r0, r1, rt in BC:
    I = f'Balance_ciclo!$I${r0}:$I${r1}'; Cc = f'Balance_ciclo!$C${r0}:$C${r1}'; K = f'Balance_ciclo!$K${r0}:$K${r1}'
    ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]
    ws[f'C{r}'] = f'=Balance_ciclo!E{r0}'
    ws[f'D{r}'] = f'=INDEX({Cc},MATCH(MAX({I}),{I},0))'; ws[f'E{r}'] = f'=ROUND(MAX({I}),0)'
    ws[f'F{r}'] = f'=INDEX({Cc},MATCH(MIN({I}),{I},0))'; ws[f'G{r}'] = f'=ROUND(-MIN({I}),0)'
    ws[f'H{r}'] = f'=Balance_ciclo!M{rt}'
    ws[f'I{r}'] = f'=COUNTIF({K},"En meta")&" de "&ROWS({K})'
    if g == '9°':
        for a in 'JKLM': ws[f'{a}{r}'] = '—'
    else:
        ar = ampmap[m]
        ws[f'J{r}'] = f'=Cobertura_objetivos!C{ar}'; ws[f'K{r}'] = f'=Cobertura_objetivos!D{ar}'
        ws[f'L{r}'] = f'=Cobertura_objetivos!E{ar}'; ws[f'M{r}'] = f'=Cobertura_objetivos!F{ar}'
        ws[f'L{r}'].number_format = PCT
    style_range(ws, r, r, 1, 13)
    ws[f'L{r}'].number_format = PCT
    if ES27:
        m23 = {'Ciclo 2°–4°': 'T4', 'Ciclo 5°–8°': 'T8', '9°': 'T8', 'Física 10°–11°': 'TA'}[g]
        b0, b1 = BAL23[g]
        G23 = f'Referencia_2023!$G${b0}:$G${b1}'; C23 = f'Referencia_2023!$C${b0}:$C${b1}'
        ws[f'P{r}'] = f'=INDEX({C23},MATCH(MAX({G23}),{G23},0))'
        if g == '9°':
            ws[f'N{r}'] = '—'; ws[f'O{r}'] = '—'
            ws[f'Q{r}'] = f'="Mayor faltante: "&P{r}&" (2023) → "&D{r}&" (2027)"'
        else:
            ws[f'N{r}'] = f'=Referencia_2023!C{AMP23[m23]}'; ws[f'O{r}'] = f'=Referencia_2023!D{AMP23[m23]}'
            # sin números con formato dentro del texto (Numbers/Excel los formatean distinto); el % está en las columnas L y O
            ws[f'Q{r}'] = (f'="Cubiertos: "&N{r}&" de "&Referencia_2023!B{AMP23[m23]}&" (2023) → "&K{r}&" de "&J{r}&" (2027)'
                           f' · Mayor faltante: "&P{r}&" (2023) → "&D{r}&" (2027)"')
            if m23 == 'TA':
                ws[f'Q{r}'] = '="Sin cambio de marco: Física sigue con TIMSS Advanced 2015"'
        if g in AMB:
            ws[f'R{r}'] = f'=Ambiental!E{AMB[g]}'; ws[f'S{r}'] = f'=Investigaciones!C{INV[g]}'
        else:
            ws[f'R{r}'] = '—'; ws[f'S{r}'] = '—'
        style_range(ws, r, r, 14, 19, align=WR)
        for a in 'NOPRS': ws[f'{a}{r}'].font = fLink
        ws[f'O{r}'].number_format = PCT; ws[f'R{r}'].number_format = PCT; ws[f'Q{r}'].font = fB
    r += 1
ws.freeze_panes = 'B5'
if HAY_PISA:  # 7.°–9.° frente a PISA 2025 (9.° solo PISA)
    r += 1
    ws[f'A{r}'] = 'PISA 2025 (7.°, 8.°, 9.° y acumulado 7.°–9.°): 9.° se compara solo con PISA; 7.° y 8.° con los dos marcos (ver Apego_escenarios)'; ws[f'A{r}'].font = fS; r += 1
    hdr(ws, r, ['Grupo', 'Marco de referencia', 'Temas totales', 'Temas sin contenido PISA (la malla los tiene; PISA no los evalúa)', '% sin contenido PISA',
                'Sistema PISA con MÁS temas faltantes', 'Temas que faltan', 'Sistema PISA con MÁS temas de sobra', 'Temas que sobran',
                'Temas a AGREGAR, sin quitar, para que todos los sistemas lleguen a su meta', 'Sistemas en meta',
                '% competencia 3 «Investigar, evaluar y usar información» (meta 24–36 %)', 'ÍNDICE DE APEGO PISA'], None, height=92)
    r += 1
    for (g, m, r0, r1, rt) in BP:
        I = f'PISA_contenido!$I${r0}:$I${r1}'; Cc = f'PISA_contenido!$C${r0}:$C${r1}'; K = f'PISA_contenido!$K${r0}:$K${r1}'
        kc = 'AV' if g == CICLO_P else 'V'
        ws[f'A{r}'] = g; ws[f'B{r}'] = 'PISA 2025'
        ws[f'C{r}'] = f'=COUNTIFS({TC(kc)},$A{r},{TC("AK")},"<>")'
        ws[f'D{r}'] = f'=PISA_contenido!N{r0}'
        ws[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
        ws[f'F{r}'] = f'=INDEX({Cc},MATCH(MAX({I}),{I},0))'; ws[f'G{r}'] = f'=ROUND(MAX({I}),0)'
        ws[f'H{r}'] = f'=INDEX({Cc},MATCH(MIN({I}),{I},0))'; ws[f'I{r}'] = f'=ROUND(-MIN({I}),0)'
        ws[f'J{r}'] = f'=PISA_contenido!M{rt}'
        ws[f'K{r}'] = f'=COUNTIF({K},"En meta")&" de "&ROWS({K})'
        ws[f'L{r}'] = f'=PISA_competencias!H{COMP_ROW[g]}'
        ws[f'M{r}'] = f'=Apego_escenarios!F{AP[g + " · PISA"]}'
        style_range(ws, r, r, 1, 13)
        for a in 'DJLM': ws[f'{a}{r}'].font = fLink
        for a in 'ELM': ws[f'{a}{r}'].number_format = PCT
        RES_P = r
        RES_PG = globals().setdefault('RES_PG', {}); RES_PG[g] = r
        r += 1
if ES27:  # guía de lectura de cada columna
    r += 1
    ws[f'A{r}'] = 'Cómo leer cada columna'; ws[f'A{r}'].font = fS; r += 1
    for txt in ('• Temas SIN objetivo en el marco: temas que la malla SÍ enseña pero que el marco no evalúa (p. ej. medición general, electrónica digital, gestión de riesgos). '
                'No son un error: indican cuánto de la malla va más allá de la prueba. Lo contrario —objetivos del marco que la malla NO trabaja— está en «Objetivos del marco que la malla NO trabaja».',
                '• Dominio TIMSS con MÁS temas faltantes: de los dominios de TIMSS (Biología, Química, Física, Tierra), el que más se queda por debajo del peso que TIMSS le da en la prueba. '
                '«Temas que faltan» = cuántos temas habría que pasar a ese dominio, con el mismo total de temas, para llegar a la meta.',
                '• Temas a AGREGAR sin quitar: si no se quiere quitar ningún tema, cuántos habría que SUMAR en total para que todos los dominios lleguen a su meta. '
                'Ejemplo: con 100 temas, 60 % Biología y meta 45 %, habría que llegar a 60 ÷ 0,45 ≈ 133 temas, sumando 33 en los otros dominios. Por eso suele ser una cifra alta: '
                'la alternativa realista (fusionar e insertar con el mismo total) está en la hoja Presupuesto.',
                '• Dominios en meta: cuántos dominios están a ±5 puntos porcentuales de su meta (tolerancia editable en Balance_grado, celda H2).',
                '• ÍNDICE DE APEGO: parte del reparto de la malla que ya coincide con el del marco (100 % = igual). Detalle y escenarios TIMSS/PISA en Apego_escenarios.'):
        ws[f'A{r}'] = txt; ws[f'A{r}'].font = fN; ws.merge_cells(f'A{r}:M{r}'); ws[f'A{r}'].alignment = WR; ws.row_dimensions[r].height = 30; r += 1

# ------------------------------------------------------------------ Fuentes
wsf = wb.create_sheet('Fuentes')
wsf['A1'] = 'Fuentes y versiones usadas'; wsf['A1'].font = fT
hdr(wsf, 3, ['Archivo (carpeta «Mallas sugeridas», Drive)', 'ID Drive', 'Última modificación', 'Temas extraídos', 'Nota'], [48, 46, 20, 10, 60])
files = [
 ('Matriz_Ciencias_2do_Grado_U1_U6_3.xlsx', '1ivJe3XBmtbIi_Pgjc6Z2Xc4pn1tL0Rw3', '', ''),
 ('Matriz_Ciencias_3er_Grado_U1_U6 (Hoja de cálculo de Google)', '1de6X-m-0tCrVH5g6FSYx5rYdWtCYM9UurTG5mdAWXS0', '2026-09-28', 'Se usó la versión Google Sheets (más reciente que el .xlsx).'),
 ('Matriz_Ciencias_4to_Grado_U1_U6.xlsx', '1ijnmjjWjW1YwhtiHeqh4hX17UX_OqB1y', '2026-09-28', ''),
 ('Matriz_Ciencias_5to_Grado_U1_U6.xlsx', '1V9G2VP2Vbe-GCCYYQAZdGBYH6MPLYJY6', '', ''),
 ('Matriz_Ciencias_6to_Grado_U1_U6.xlsx', '11jnnAGb8NZvI0jQuFxXVvX73dCVR7DKC', '', ''),
 ('Matriz_Ciencias_7mo_Grado_U1_U6.xlsx', '1OR7NdtVVqI8P8NFaHgzIE2kpHCu-621X', '', ''),
 ('Matriz_Ciencias_8vo_Grado_U1_U7.xlsx', '1trSwU3iYDitze9gH-VuOsXcWm9Hhs095', '2026-09-28', ''),
 ('Matriz_Ciencias_9no_Grado_U1_U7.xlsx', '1or_yQdA1e1iLSx-x8XDs1vNSZC_-HM1D', '', ''),
 ('Matriz_Bachillerato_Fisica_10_11.xlsx', '1SMhkw4GQloLiBSHTS3rvnQoroZmHTh38', '2026-09-30', ''),
 ('Matriz_Bachillerato_Quimica_10_11.xlsx', '1PGQKwpI0KnyUQicq4HQY4ghVZSCF4d6V', '', 'Solo se usa para el análisis cognitivo (sin marco TIMSS de contenido).'),
 ('Matriz_Bachillerato_Biologia_10_11.xlsx', '1f3nWhHBpQS1vajP1PZqNGcKZEYLh7oF2', '', 'Solo se usa para el análisis cognitivo (sin marco TIMSS de contenido).'),
]
fname = {'Matriz_Ciencias_3er_Grado_U1_U6 (Hoja de cálculo de Google)': 'Matriz_Ciencias_3er_Grado_U1_U6.xlsx'}
r = 4
for f, i, d, n in files:
    wsf[f'A{r}'] = f; wsf[f'B{r}'] = i; wsf[f'C{r}'] = d or 'descargado 2026-10-01'
    wsf[f'D{r}'] = f'=COUNTIFS({TC("D")},"{fname.get(f, f)}")'
    wsf[f'E{r}'] = n
    style_range(wsf, r, r, 1, 5, align=WR); r += 1
wsf[f'C{r}'] = 'Total'; wsf[f'D{r}'] = f'=SUM(D4:D{r-1})'; FUENTES_TOT = r; style_range(wsf, r, r, 3, 4, fill=SUB, font=fB)
r += 2
wsf[f'A{r}'] = 'Marcos de referencia'; wsf[f'A{r}'].font = fS; r += 1
REFS = ['IEA (2021). TIMSS 2023 Assessment Frameworks, cap. 2 Science. https://timss.bc.edu/timss2023/frameworks/pdf/T23_Frameworks_Ch2_Science.pdf',
        'IEA (2014). TIMSS Advanced 2015 Assessment Frameworks, cap. 2 Physics. https://elvis.bc.edu/timss2015-advanced/downloads/TA15_FW_Chap2.pdf',
        'Mallas: columna «Habilidad TIMSS» (dominios Conocer/Aplicar/Razonar) elaborada por el equipo de Ciencias (MINED).']
if ES27:
    REFS = ['IEA (2025). TIMSS 2027 Assessment Frameworks, cap. 2 Science (data/referencia/externos/timss/TIMSS2027_Marcos_Evaluacion.pdf). '
            'Catálogo: data/referencia/catalogo_timss2027.json (página impresa por objetivo).',
            'Cambios 2023→2027: data/referencia/cambios_timss_2023_2027.md y equivalencias_timss_2023_2027.csv.'] + REFS
if HAY_PAISES:
    REFS += ['Otros países (Comparacion_paises): ANEP Uruguay, programas EBI 2023 tramos 3–6 · MEN Colombia, DBA Ciencias Naturales (2016) y '
             'Estándares Básicos de Competencias (2006) · MOE Singapur, Primary Science Syllabus 2023. URL y SHA-256 en data/referencia/externos/manifest.csv y estado.csv. '
             'Chile no está incluido (descarga bloqueada por el sitio).']
for s in REFS:
    wsf[f'A{r}'] = s; wsf[f'A{r}'].font = fN; wsf.merge_cells(f'A{r}:E{r}'); r += 1

# ------------------------------------------------------------------ Leeme
wl = wb.create_sheet('Leeme', 0)
wl.column_dimensions['A'].width = 4; wl.column_dimensions['B'].width = 120
lines = [
 ('T', 'Cobertura TIMSS de la malla de Ciencia y Tecnología — cómo se construyó y cómo leerla'),
 ('N', 'Versión 1 · 1 de octubre de 2026 · Elaborado con Claude (Cowork) a partir de las mallas compartidas por el equipo de Ciencias (MINED).'),
 ('S', 'Qué responde este libro'),
 ('B', '1. Balance: ¿el reparto de temas entre Vida / Física / Química / Tierra se parece al peso que TIMSS da a cada dominio? ¿Cuántos temas faltan o sobran para llegar a la meta?'),
 ('B', '2. Amplitud: ¿qué objetivos de contenido de TIMSS aparecen en la malla, en qué grado, con cuántos temas, y cuáles no aparecen?'),
 ('B', '3. Profundidad cognitiva: ¿la proporción Conocer / Aplicar / Razonar se parece a la de TIMSS?'),
 ('S', 'Proceso (de dónde salen los números)'),
 ('B', 'Paso 1 · Extracción. Se leyeron las 11 mallas (2.° a 9.° y Física, Química, Biología de 10.° y 11.°). Cada fila con «Procedimental» es un TEMA: 1.260 en total. '
       'Se conservan archivo, hoja y fila de origen (hoja Temas_TIMSS, columnas D-F) para poder volver a la malla.'),
 ('B', 'Paso 2 · Catálogo de referencia. Se armó la lista de objetivos de contenido de TIMSS 2023 (4.° grado: 32 objetivos; 8.° grado: 46) y de TIMSS Advanced 2015 Física (23), '
       'con su dominio y el % que TIMSS asigna a cada dominio (hoja Catalogo_TIMSS, en azul = valores tomados del marco).'),
 ('B', 'Paso 3 · Clasificación. Cada tema de 2.° a 9.° y de Física 10.°–11.° (889 temas) se asignó a UN objetivo principal y, si aplica, uno secundario, leyendo procedimental, contenido, '
       'indicador y evidencia. Lo hicieron agentes de IA con instrucciones fijas, tema por tema, dejando confianza (alta/media/baja) y una justificación corta. '
       'Los temas sin contenido TIMSS (medición general, método científico, tecnología/diseño, gestión de riesgos, electrónica digital, mecánica de fluidos en Bachillerato, etc.) se marcan «FUERA».'),
 ('B', 'Paso 4 · Reglas de grado. 2.°–4.° se comparan con TIMSS 4.°; 5.°–9.° con TIMSS 8.° (9.° solo como referencia, porque TIMSS 8.° evalúa hasta 8.°); Física 10.°–11.° con TIMSS Advanced. '
       'Química y Biología de Bachillerato no tienen marco TIMSS de contenido: solo entran en el análisis cognitivo.'),
 ('B', 'Paso 5 · Cálculo (todo con fórmulas). % actual = temas del dominio ÷ temas evaluables (sin FUERA). Temas esperados = meta × evaluables. '
       'Brecha = esperados − actuales (+ faltan, − sobran). «Temas a agregar sin quitar»: se busca el menor total que permite cumplir todas las metas sumando temas '
       '(lo fija el dominio más sobrerrepresentado: total = máx(actual ÷ meta)), y a cada dominio se le suma meta × total − actual, redondeado hacia arriba.'),
 ('B', 'Paso 6 · Amplitud. Un objetivo está «cubierto» si al menos un tema lo trabaja como principal o secundario. Profundidad: 0 no cubierto, 1 débil, 2–3 básica, 4+ sólida.'),
 ('B', 'Paso 7 · Cognitivo. Se usa la columna «Habilidad TIMSS» de la malla. Si un tema nombra varios dominios se reparte en partes iguales. Se compara con TIMSS (4.°: 40/40/20; 8.°: 35/35/30; Advanced Física: 30/40/30).'),
 ('S', 'Hojas'),
 ('B', 'Resumen: una fila por grado y por ciclo con el dominio que más falta, el que más sobra, temas a agregar, amplitud y % Razonar.'),
 ('B', 'Balance_grado y Balance_ciclo: detalle por dominio. Celda H2 = tolerancia editable para «En meta» (5 pp).'),
 ('B', 'Cobertura_objetivos: matriz objetivo × grado (n.º de temas). Al final, resumen de amplitud por dominio.'),
 ('B', 'Faltantes: objetivos no cubiertos o débiles, con una sugerencia de dónde podrían entrar en la malla actual.'),
 ('B', 'Cognitivo_grado: Conocer / Aplicar / Razonar por grado y asignatura.'),
 ('B', '__AMBIENTAL__'),
 ('B', 'Temas_TIMSS: trazabilidad completa. Si cambias un objetivo en las columnas K o L, todo el libro se recalcula.'),
 ('B', 'Catalogo_TIMSS y Fuentes: referencias y versiones de archivos.'),
 ('S', 'Límites que hay que decir en la reunión'),
 ('B', '• La clasificación es asistida por IA: 406 temas con confianza alta, 364 media y 119 baja (en rojo en Temas_TIMSS). Debe validarla el equipo de Ciencias, empezando por las de confianza baja.'),
 ('B', '• Se cuentan TEMAS, no horas. Un tema puede durar una o varias sesiones de 55 min; cuando exista la distribución de horas conviene ponderar por horas.'),
 ('B', '• Los % de TIMSS son pesos de la prueba, no un mandato curricular. Sirven como referencia de equilibrio, no como regla obligatoria.'),
 ('B', '• Comparar un solo grado con TIMSS es una aproximación (TIMSS mide lo acumulado). La comparación más fiel está en Balance_ciclo.'),
 ('B', '• Los objetivos TIMSS están parafraseados; para decisiones formales, consultar el marco oficial (hoja Fuentes).'),
 ('B', '• 1.° grado no tiene malla, así que no entra en el ciclo 2.°–4.°.'),
]
if ES27:
    from collections import Counter as _Cn
    _cf = _Cn(v['confianza'] for v in R.values())
    _n4 = sum(c['marco'] == M4 for c in CAT); _n8 = sum(c['marco'] == M8 for c in CAT)
    _rep = {
     'Cobertura TIMSS de la malla': ('T', 'Cobertura TIMSS 2027 de la malla de Ciencia y Tecnología — cómo se construyó y cómo leerla'),
     'Versión 1': ('N', 'Versión 2 (TIMSS 2027) · Elaborado con Claude Code a partir de las mallas del equipo de Ciencias (MINED). '
              'La versión 1 (TIMSS 2023) se conserva en Cobertura_TIMSS_Ciencias_SV.xlsx.'),
     'Paso 2 ': ('B', f'Paso 2 · Catálogo de referencia. Objetivos de contenido de TIMSS 2027 (4.° grado: {_n4}; 8.° grado: {_n8}), construidos desde el PDF oficial con página citada, '
              'y TIMSS Advanced 2015 Física (23), con su dominio y el % que TIMSS asigna a cada dominio (hoja Catalogo_TIMSS, celdas ámbar = valores tomados del marco y editables). '
              'Novedades 2027: áreas de Investigaciones en cada dominio y subescala de conocimiento ambiental.'),
     'Paso 3 ': ('B', 'Paso 3 · Clasificación. Los 622 temas de 2.° a 9.° se reclasificaron contra TIMSS 2027 con agentes de IA (instrucciones fijas, un lote por grado), '
               'usando la clasificación 2023 y la tabla de equivalencias solo como referencia. Física 10.°–11.° conserva la clasificación TIMSS Advanced 2015. '
               'Un tema va a un área de Investigaciones solo si su foco es la práctica científica; si trabaja un contenido mediante indagación, la investigación va como secundaria.'),
     'Paso 4 ': ('B', 'Paso 4 · Reglas de grado. 2.°–4.° se comparan con TIMSS 2027 4.°; 5.°–9.° con TIMSS 2027 8.° (9.° solo como referencia); Física 10.°–11.° con TIMSS Advanced 2015. '
               'Química y Biología de Bachillerato no tienen marco TIMSS de contenido: solo entran en el análisis cognitivo.'),
     'Paso 7 ': ('B', 'Paso 7 · Cognitivo. Se usa la columna «Habilidad TIMSS» de la malla. Si un tema nombra varios dominios se reparte en partes iguales. '
               'Se compara con TIMSS 2027 (4.°: 35/40/25; 8.°: 35/35/30) y Advanced Física (30/40/30).'),
     '• La clasificación es asistida': ('B', f"• La clasificación es asistida por IA: {_cf['alta']} temas con confianza alta, {_cf['media']} media y {_cf['baja']} baja (en rojo en Temas_TIMSS). "
               'Debe validarla el equipo de Ciencias, empezando por outputs/revision_confianza_baja_timss2027.csv.'),
    }
    _rep['__AMBIENTAL__'] = ('B', 'Ambiental: % de temas que trabajan objetivos ambientales de TIMSS 2027 (≈25 % de la prueba), por grado y ciclo, '
                                   'y cobertura de cada objetivo ambiental. Investigaciones: temas cuyo foco es la práctica científica. '
                                   'Referencia_2023: la clasificación v1 recalculada para la columna «Cambio vs 2023» del Resumen.')
    if HAY_PAISES:
        _rep['__AMBIENTAL__'] = ('B', _rep['__AMBIENTAL__'][1] + ' Comparacion_paises: primer grado en que El Salvador y ' + ', '.join(NOMBRES_P) +
                                 ' introducen cada objetivo TIMSS 2027 y la «oportunidad» (grado El Salvador − mediana de los demás). '
                                 'Paises_objetivos: los objetivos de cada país con documento, página y su alineación con TIMSS 2027.')
    for _pref, _v in _rep.items():
        _idx = [i for i, (_, t) in enumerate(lines) if t.startswith(_pref)]
        assert len(_idx) == 1, _pref
        lines[_idx[0]] = _v
lines = [x for x in lines if x[1] != '__AMBIENTAL__']  # en 2023 la línea no existe
if HAY_PISA:  # libro TIMSS 2027 + PISA 2025: resumen ejecutivo en prosa (cifras = fórmulas) y hojas nuevas
    P_ = lambda ref: f'FIXED({ref}*100,0)&" %"'
    N_ = lambda ref: f'FIXED({ref},0)'
    A_ = lambda etq: f'Apego_escenarios!F{AP[etq]}'
    b5 = [b for b in BG if b[0] == '5°'][0]; b4 = [b for b in BG if b[0] == '4°'][0]; b8 = [b for b in BG if b[0] == '8°'][0]
    dom8 = [d for d, _ in DOMS[M8]]; dom4 = [d for d, _ in DOMS[M4]]
    p9 = [b for b in BP if b[0] == '9°'][0]; p8 = [b for b in BP if b[0] == '8°'][0]; pc = BP[-1]
    t24, t58, t79 = PRES_TOT['2-4'][2], PRES_TOT['5-8'][2], PRES_TOT['7-9'][2]
    XK = 'Comparacion_paises!$K$5:$K$84'
    resumen_ej = [
     ('S', 'Resumen ejecutivo (las cifras se recalculan con el libro)'),
     ('F', '="1. Apego general. En lo acumulado, la malla se parece mucho a lo que piden las pruebas: índice de apego de "&' + P_(A_('Ciclo 2°–4°')) +
           '&" frente a TIMSS 4.° (2.°–4.°), "&' + P_(A_('Ciclo 5°–8°')) + '&" frente a TIMSS 8.° (5.°–8.°) y "&' + P_(A_('Ciclo 7°–9° · PISA')) +
           '&" frente a PISA (7.°–9.°). Cubre "&Cobertura_objetivos!D' + str(ampmap_[M4]) + '&" de "&Cobertura_objetivos!C' + str(ampmap_[M4]) +
           '&" objetivos de TIMSS 4.°, "&Cobertura_objetivos!D' + str(ampmap_[M8]) + '&" de "&Cobertura_objetivos!C' + str(ampmap_[M8]) +
           '&" de TIMSS 8.° y "&PISA_contenido!C' + str(AMP_P['(todos)']) + '&" de 21 categorías de contenido de PISA. Los problemas están en el reparto entre grados y en las capas propias de PISA."'),
     ('F', '="2. Grados a mejorar. 5.° es el grado con menor apego ("&' + P_(A_('5°')) + '&"): le faltan unos "&' + N_(f'Balance_grado!I{b5[2] + dom8.index("Biología")}') +
           '&" temas de Biología del marco de 8.°, que Uruguay, Colombia y Singapur enseñan en 5.°–6.° y El Salvador recién en 6.°–7.°. En 4.° faltan unos "&' +
           N_(f'Balance_grado!I{b4[2] + dom4.index("Ciencias físicas")}') + '&" temas de Ciencias físicas (luz, sonido, electricidad, imanes). '
           '8.° tiene solo "&PISA_contenido!D' + str(p8[2] + 2) + '&" temas de Tierra y espacio y 9.° solo "&PISA_contenido!D' + str(p9[2] + 1) +
           '&" de sistemas vivos (ni célula, ni genética, ni salud humana). En 6.° y 8.° Razonar queda en "&' + P_(f'Cognitivo_grado!I{cogrow["6°"]}') + '&" y "&' +
           P_(f'Cognitivo_grado!I{cogrow["8°"]}') + '&" (meta TIMSS 30 %)."'),
     ('F', '="3. Competencias PISA. La malla declara una competencia y cinco subcompetencias por tema (columnas de la propia malla). En 7.°–9.° la competencia 3, '
           '«Investigar, evaluar y usar información para decidir», llega a "&' + P_(f'PISA_competencias!H{COMP_ROW[CICLO_P]}') +
           '&" de los temas (PISA: 24–36 %); distinguir evidencia de opinión (I2), criticar argumentos (I4) y formular hipótesis (E5) casi no aparecen. '
           'La competencia 2 está sobrerrepresentada ("&' + P_(f'PISA_competencias!G{COMP_ROW[CICLO_P]}') + '&")."'),
     ('F', '="4. Conocimiento procedimental y epistémico (lo que más agrega PISA). Solo "&' + P_(f'PISA_conocimiento!G{KN_ROW[CICLO_P]}') +
           '&" de los temas de 7.°–9.° tienen como foco el conocimiento procedimental (PISA: 27–33 %) y ninguno el epistémico (24–30 %). "&PISA_conocimiento!E' +
           str(PE_ROW['Epistémico']) + '&" de los 20 elementos epistémicos (modelos frente a realidad, consenso, revisión por pares, límites de la certeza) no se trabajan en ningún tema."'),
     ('F', '="5. Contextos. "&' + P_(f'PISA_contextos!G{CX_ROW[CICLO_P]}') + '&" de los temas de 7.°–9.° no sitúan la ciencia en una situación real ("&' +
           P_(f'PISA_contextos!G{CX_ROW["8°"]}') + '&" en 8.°). Cuando hay contexto, el reparto personal / local / global se parece al 1:2:1 de PISA."'),
     ('F', '="6. Frente a otros países, El Salvador introduce "&COUNTIF(' + XK + ',"El Salvador llega más tarde")&" objetivos TIMSS más tarde que la mediana, "&COUNTIF(' + XK +
           ',"El Salvador va adelantado")&" antes y no trabaja "&COUNTIF(' + XK + ',"El Salvador no llega")&" que otros sí (herencia y propiedades de la luz)."'),
     ('S', 'Recomendaciones'),
     ('B', 'R1. Decidir el peso de PISA en 7.° y 8.° (hoja Apego_escenarios, celda E4). Sugerencia mientras se decide: PISA como guía en 7.° y 9.°, y 50/50 en 8.° '
           'si El Salvador rinde TIMSS 2027 (8.° es el grado que TIMSS evalúa).'),
     ('F', '="R2. 2.°–4.° (hoja Presupuesto): fusionar temas repetidos de Biología para liberar "&FIXED(Presupuesto!J' + str(t24) + ',0)&" espacios e insertar "&FIXED(Presupuesto!G' + str(t24) + ',0)' +
           '&" temas de Ciencias físicas y de la Tierra (luz, sonido, circuitos, imanes, tiempo atmosférico) y "&FIXED(Presupuesto!H' + str(t24) + ',0)&" de herencia (B2.2), sin cambiar el total."'),
     ('B', 'R3. 5.°–8.°: el total por dominio ya está en meta; no hace falta agregar temas. Conviene REUBICAR (Presupuesto, acción «reubicar»): adelantar Biología de 7.° a 5.°, '
           'insertar luz (con una fusión de Física) y traer al ciclo los movimientos de la Tierra y la Luna (hoy en 9.°), y subir tareas de Razonar en 6.° y 8.°.'),
     ('F', '="R4. 7.°–9.° (PISA): fusionar temas de física y química que se repiten (liberan "&FIXED(Presupuesto!J' + str(t79) + ',0)&") para insertar "&FIXED(Presupuesto!G' + str(t79) + ',0)' +
           '&" temas de sistemas vivos (sobre todo en 9.°: célula, salud, genética, evolución) y de Tierra y espacio (sobre todo en 8.°: universo, historia de la Tierra)."'),
     ('B', 'R5. Capas PISA: en cada unidad de 7.°–9.°, al menos un tema o fase que trabaje lo epistémico (por qué se confía en un modelo o en una conclusión) y la competencia 3 '
           '(evaluar fuentes, distinguir evidencia de opinión, criticar argumentos con datos), situado en un contexto local o personal. No requiere temas nuevos: cambia la fase «Reforzar».'),
     ('B', 'R6. Validar antes de decidir: las clasificaciones y los candidatos del presupuesto son asistidos por IA. El equipo de Ciencias debe revisar primero los de confianza baja '
           '(columna «Decisión del equipo» en Presupuesto y outputs/revision_confianza_baja_timss2027.csv).'),
    ]
    lines[2:2] = [('F', '="TIMSS 2027 y PISA 2025"&"\n"&"BIEN · En lo acumulado la malla se parece a las pruebas: "&' + P_(A_('Ciclo 2°–4°')) + '&" (TIMSS 4.°), "&' +
                         P_(A_('Ciclo 5°–8°')) + '&" (TIMSS 8.°) y "&' + P_(A_('Ciclo 7°–9° · PISA')) + '&" (PISA) de apego."&"\n"&'
                         '"OJO · TIMSS: falta Biología en 5.°, luz-sonido-electricidad en 4.° y más Razonar en 6.° y 8.°."&"\n"&'
                         '"OJO · PISA: falta biología en 9.°, Tierra y espacio en 8.°, y enseñar a evaluar información y el porqué de la ciencia."&"\n"&'
                         '"→ Empieza por «Resumen» y «Presupuesto». Las pestañas que empiezan con TIMSS o PISA tienen el detalle de cada prueba."')] + resumen_ej
    _rep2 = {
     'Cobertura TIMSS 2027 de la malla': ('T', 'Cobertura de la malla de Ciencia y Tecnología frente a TIMSS 2027 y PISA 2025 — hallazgos, cómo se construyó y cómo leerla'),
     'Versión 2': ('N', 'Versión 3 (TIMSS 2027 + PISA 2025) · 2 de octubre de 2026 · Elaborado con Claude Code a partir de las mallas del equipo de Ciencias (MINED). '
                        'Las versiones anteriores (TIMSS 2023) se conservan en Cobertura_TIMSS_Ciencias_SV.xlsx.'),
     '3. Profundidad cognitiva': ('B', '3. Profundidad cognitiva: ¿la proporción Conocer / Aplicar / Razonar se parece a la de TIMSS? 4. PISA (7.°–9.°): ¿cubre los sistemas y categorías '
                                      'de contenido, las tres competencias y sus subcompetencias, el conocimiento procedimental y epistémico, y los contextos que PISA evalúa? '
                                      '5. ¿Qué insertar y qué fusionar para acercarse a las metas sin aumentar la carga?'),
     'Paso 4 ': ('B', 'Paso 4 · Reglas de grado. 2.°–4.° → TIMSS 2027 4.°; 5.°–8.° → TIMSS 2027 8.°; 7.°–9.° → PISA 2025 (7.° y 8.° con los dos marcos, 9.° solo PISA); '
                     'Física 10.°–11.° → TIMSS Advanced 2015. Química y Biología de Bachillerato no tienen marco de contenido.'),
     'Cobertura_objetivos:': ('B', 'Cobertura_objetivos: matriz objetivo × grado con semáforo (rojo = no cubierto, amarillo = débil, verde = básica o sólida), «N/A» en grados que no '
                                   'pertenecen al ciclo del marco y una columna «Qué hacer» con dónde aparece hoy y la propuesta del Presupuesto (reemplaza a la hoja Faltantes).'),
     'Faltantes:': ('B', 'Apego_escenarios: índice de apego por grado y ciclo, y escenarios 100 % TIMSS / 50-50 / 100 % PISA / personalizado para 7.° y 8.°. '
                         'Presupuesto: cuántos temas insertar y fusionar por ciclo con el mismo total, y los candidatos concretos propuestos por IA.'),
     'Temas_TIMSS:': ('B', 'Temas: trazabilidad completa (TIMSS en K-AC; PISA en AD-AW). Si cambias un objetivo o una categoría, todo el libro se recalcula.'),
     'Catalogo_TIMSS y Fuentes': ('B', 'PISA_contenido, PISA_competencias, PISA_conocimiento y PISA_contextos: las capas de PISA 2025 para 7.°–9.° '
                                       '(sistemas y categorías de contenido; competencias y subcompetencias de la malla; conocimiento procedimental y epistémico; contextos, áreas y demanda cognitiva). '
                                       'Catalogo_TIMSS, Catalogo_PISA y Fuentes: referencias con página.'),
    }
    for _pref, _v in _rep2.items():
        _idx = [i for i, (_, t) in enumerate(lines) if t.startswith(_pref)]
        assert len(_idx) == 1, _pref
        lines[_idx[0]] = _v
    lines.insert([i for i, (_, t) in enumerate(lines) if t.startswith('Paso 4 ')][0] + 1,
                 ('B', 'Paso 4-bis · PISA 2025. Catálogo desde el marco final (OECD 2026, cap. 2, con página). Los temas de 7.°–9.° se clasificaron con IA por categoría de contenido, '
                       'tipo de conocimiento, elementos procedimentales y epistémicos, contexto, área de aplicación y demanda cognitiva. La competencia y las subcompetencias NO se '
                       'clasificaron: salen de la propia malla (columnas «Competencia PISA 2025» y microhabilidades E/D/I por fase).'))
r = 1
for kind, text in lines:
    if kind == 'S':
        r += 1  # línea en blanco antes de cada sección
    c = wl.cell(row=r, column=2, value=text)
    c.font = {'T': fT, 'N': fNote, 'S': fS, 'B': fN, 'F': fN}[kind]
    if kind == 'F':
        wl.row_dimensions[r].height = 58
    c.alignment = WR
    r += 1
for rr in range(1, r):
    v = wl.cell(row=rr, column=2).value
    if v and len(v) > 120:
        wl.row_dimensions[rr].height = 14 * (len(v) // 115 + 1)

if not ES27:
    wb._sheets = [wb['Leeme'], wb['Resumen'], wb['Balance_ciclo'], wb['Balance_grado'], wb['Cobertura_objetivos'], wb['Faltantes'],
                  wb['Cognitivo_grado'], wb['Temas_TIMSS'], wb['Catalogo_TIMSS'], wb['Fuentes']]
else:
    wb._sheets = [wb[n] for n in ('Leeme', 'Resumen', 'Balance_ciclo', 'Balance_grado', 'Cobertura_objetivos', 'Ambiental',
                                  'Investigaciones', 'Faltantes', 'Cognitivo_grado', 'Temas_TIMSS', 'Catalogo_TIMSS',
                                  'Referencia_2023', 'Fuentes')]
    if HAY_PAISES:  # comparación con otros países (Prompt 02) después de Investigaciones
        i = wb._sheets.index(wb['Investigaciones']) + 1
        wb._sheets[i:i] = [wx]
        wb._sheets.insert(wb._sheets.index(wb['Referencia_2023']), wp)
    if HAY_PISA:  # libro TIMSS 2027 + PISA 2025: lo que se lee primero, luego el detalle y al final los datos
        orden = ['Leeme', 'Resumen', 'Apego_escenarios', 'Presupuesto', 'Balance_ciclo', 'Balance_grado', 'Cobertura_objetivos',
                 'Cognitivo_grado', 'Ambiental', 'Investigaciones', 'Comparacion_paises',
                 'PISA_contenido', 'PISA_competencias', 'PISA_conocimiento', 'PISA_contextos', 'Temas_TIMSS', 'Catalogo_TIMSS', 'Catalogo_PISA', 'Paises_objetivos',
                 'Referencia_2023', 'Fuentes']
        todas = {x.title: x for x in [wl, ws, wap, wbg, wcy, wo, wpc, wpk, wpn, wpx, wg, wt, wc, wcp, wr, wsf, wf] +
                 [x for x in wb._sheets] + ([wpr] if 'wpr' in globals() else []) + ([wa, wi] if ES27 else []) + ([wx, wp] if HAY_PAISES else [])}
        wb._sheets = [todas[n] for n in orden if n in todas]  # «Faltantes» queda integrada en Cobertura_objetivos
for s in wb.worksheets:
    s.sheet_view.showGridLines = False
if ES27:
    for s in wb.worksheets:  # toda celda con texto café (editable) lleva fondo ámbar
        for row in s.iter_rows():
            for c in row:
                if c.font is not None and c.font.color is not None and c.font.color.rgb == '007A4A00':
                    c.fill = INP
    banda = PatternFill('solid', fgColor=BANDA)
    for s, rng in ((wt, f'A5:W{TL}'), (wo, f'A5:{fcol}{OL}'), (wbg, 'A6:N200'), (wcy, 'A6:N60'), (wc, f'A5:K{CAT_LAST}')):
        s.conditional_formatting.add(rng, FormulaRule(formula=['AND(MOD(ROW(),2)=0,$A5<>"")'] if s is not wbg and s is not wcy
                                                      else ['AND(MOD(ROW(),2)=0,$C6<>"",$C6<>"Total del grupo")'], fill=banda))
    TAB = {'Leeme': '7F7F7F', 'Resumen': VERDE, 'Balance_ciclo': '2E6B30', 'Balance_grado': '2E6B30',
           'Cobertura_objetivos': '8C6D1F', 'Faltantes': 'B85C38', 'Cognitivo_grado': '6B4E71',
           'Temas_TIMSS': '595959', 'Catalogo_TIMSS': '595959', 'Fuentes': '7F7F7F',
           'Ambiental': '4F7942', 'Investigaciones': '8C6D1F', 'Referencia_2023': 'A6A6A6',
           'Comparacion_paises': '2E6B30', 'Paises_objetivos': '595959', 'Apego_escenarios': VERDE, 'Presupuesto': 'B85C38',
           'PISA_contenido': '1F6F5C', 'PISA_competencias': '1F6F5C', 'PISA_conocimiento': '1F6F5C', 'PISA_contextos': '1F6F5C',
           'Catalogo_PISA': '595959', 'Temas': '595959'}
    for s in wb.worksheets:
        s.sheet_properties.tabColor = TAB.get(s.title, '7F7F7F')
if HAY_PISA:  # la hoja de temas lleva TIMSS y PISA: se llama «Temas» (fórmulas y textos se actualizan)
    for s_ in wb.worksheets:
        for row in s_.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Temas_TIMSS' in c.value:
                    c.value = c.value.replace('Temas_TIMSS', 'Temas')
    wt.title = 'Temas'
    wt.sheet_properties.tabColor = '595959'

# ------------------------------------------------------------------ Recuadro de hallazgos: uno arriba de cada hoja + nombres de pestaña claros
if HAY_PISA:
    import re as _re, sys as _sys
    _sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from hojas import NOMBRES
    from openpyxl.utils import column_index_from_string
    K24, K58, K79P, K9P, K8 = 'Ciclo 2°–4°', 'Ciclo 5°–8°', CICLO_P + ' · PISA', '9° · PISA', '8°'
    PRO_, EPI_, C3_ = 'Procedimental', 'Epistémico', NOMC['C3']
    EM_, AM_, AN_ = 'En meta', 'A', 'E'
    FX = lambda ref: f'FIXED({ref},0)'
    PC = lambda ref: f'FIXED(({ref})*100,0)&" %"'
    CAJA = PatternFill('solid', fgColor='E6EFE9')
    TAG = lambda cond: f'IF({cond},"BIEN · ","OJO · ")'   # veredicto al inicio de la viñeta
    G2 = lambda expr: f'SUBSTITUTE({expr},"°",".°")'       # «5°» → «5.°» como en el resto del texto

    def caja(ws_, fila, marco, partes):
        """Recuadro de hallazgos (encabezado = marco que mide la hoja) en la fila indicada; la nota técnica que había ahí pasa al final («Cómo se calcula»)."""
        ultima = 12
        for mr in list(ws_.merged_cells.ranges):
            if mr.min_row == fila and mr.min_col == 1:
                ultima = max(ultima, mr.max_col); ws_.unmerge_cells(str(mr))
        ultima = min(ultima, 14)
        nota = ws_.cell(row=fila, column=1).value
        if isinstance(nota, str) and nota and not nota.startswith('='):
            fb = ws_.max_row + 2
            c_ = ws_.cell(row=fb, column=1, value='Cómo se calcula: ' + nota)
            c_.font = fNote; c_.alignment = WR
            ws_.merge_cells(start_row=fb, start_column=1, end_row=fb, end_column=ultima)
            ws_.row_dimensions[fb].height = 14 * (len(nota) // 140 + 2)
        f_ = f'="{marco}"&"\n"&' + '&"\n"&'.join(partes)
        c_ = ws_.cell(row=fila, column=1, value=f_)
        c_.font = Font(name=F, size=11, color=VERDE); c_.fill = CAJA; c_.border = BOX
        c_.alignment = Alignment(wrap_text=True, vertical='top')
        ws_.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
        ws_.row_dimensions[fila].height = 15 * (len(partes) + 1) + 8

    def rg(col, r0, r1, hoja=''):
        return f'{hoja}${col}${r0}:${col}${r1}'

    # --- Resumen (TIMSS 2.°–8.° arriba, PISA 7.°–9.° abajo)
    gr0, gr1 = 5, 4 + len([b for b in BG if b[1] in (M4, M8)])  # solo 2.°–8.° (Física usa otro marco)
    G_, A_r, F_r = rg('G', gr0, gr1), rg('A', gr0, gr1), rg('F', gr0, gr1)
    r9p, rcp = RES_PG['9°'], RES_PG[CICLO_P]
    caja(ws, 2, 'TIMSS 2027 y PISA 2025', [
        f'"OJO · TIMSS (2.°–8.°): faltan "&{FX(f"MAX({G_})")}&" temas de "&INDEX({F_r},MATCH(MAX({G_}),{G_},0))&" en "&{G2(f"INDEX({A_r},MATCH(MAX({G_}),{G_},0))")}&", el grado más desbalanceado."',
        f'"OJO · PISA (7.°–9.°): faltan "&{FX(f"G{r9p}")}&" temas de "&LOWER(F{r9p})&" en 9.°; en 7.°–9.° sobran "&{FX(f"I{rcp}")}&" de "&LOWER(H{rcp})&"."',
        '"● Arriba: tabla TIMSS. Abajo: tabla PISA. Al final: qué significa cada columna."'])

    # --- Apego (mezcla los dos marcos; lo dice cada fila)
    g0_, g1_ = AP['2°'], AP['8°']
    Fg, Ag = rg('F', g0_, g1_), rg('A', g0_, g1_)
    caja(wap, 2, 'TIMSS 2027 y PISA 2025', [
        f'{TAG(f"MIN(F{AP[K24]},F{AP[K58]},F{AP[K79P]})>=$H$4")}&"Acumulado: TIMSS 4.° "&{PC(f"F{AP[K24]}")}&", TIMSS 8.° "&{PC(f"F{AP[K58]}")}&", PISA "&{PC(f"F{AP[K79P]}")}&" de apego (alto = 85 % o más)."',
        f'"OJO · Grado más flojo en TIMSS: "&{G2(f"INDEX({Ag},MATCH(MIN({Fg}),{Fg},0))")}&" ("&{PC(f"MIN({Fg})")}&"). En PISA: 9.° ("&{PC(f"F{AP[K9P]}")}&")."',
        f'"● 7.° y 8.° se miden con los dos marcos: el peso de PISA se elige en la celda E4 (hoy "&{PC("E4")}&")."'])

    # --- Presupuesto (TIMSS en 2.°–4.° y 5.°–8.°; PISA en 7.°–9.°)
    t24, t58, t79 = PRES_TOT['2-4'][2], PRES_TOT['5-8'][2], PRES_TOT['7-9'][2]
    caja(wpr, 2, 'TIMSS 2027 (2.°–8.°) y PISA 2025 (7.°–9.°)', [
        '"● El total de temas no cambia: cada tema nuevo se paga uniendo dos temas que se repiten."',
        f'"● TIMSS 2.°–4.°: "&{FX(f"G{t24}+H{t24}")}&" temas nuevos a cambio de "&{FX(f"J{t24}")}&" uniones de temas repetidos. TIMSS 5.°–8.°: ya equilibrado; solo "&{FX(f"G{t58}+H{t58}")}&" tema nuevo (un objetivo que nadie enseña)."',
        f'"● PISA 7.°–9.°: "&{FX(f"G{t79}+H{t79}")}&" temas nuevos de biología y Tierra a cambio de "&{FX(f"J{t79}")}&" uniones en física y química."',
        '"→ Filtra «¿Entra en el presupuesto?» = Sí y anota tu decisión en la última columna."'])

    # --- Balances TIMSS
    def frase_bal(hoja_ws, bloque, etq):
        g, m, r0, r1, rt = bloque
        I = rg('I', r0, r1); C = rg('C', r0, r1); K = rg('K', r0, r1)
        return (f'{TAG('COUNTIF(%s,"En meta")=ROWS(%s)' % (K, K))}&"{etq}: "&IF(COUNTIF({K},"En meta")=ROWS({K}),"los "&ROWS({K})&" dominios están en meta.",'
                f'"faltan "&{FX(f"MAX({I})")}&" temas de "&INDEX({C},MATCH(MAX({I}),{I},0))&" y sobran "&{FX(f"-MIN({I})")}&" de "&INDEX({C},MATCH(MIN({I}),{I},0))&".")')
    caja(wcy, 3, 'TIMSS 2027', [frase_bal(wcy, BC[0], 'TIMSS 4.° (2.°–4.°)'), frase_bal(wcy, BC[1], 'TIMSS 8.° (5.°–8.°)'),
                                frase_bal(wcy, BC[2], 'TIMSS Advanced (Física 10.°–11.°)'),
                                '"→ Esta es la comparación más justa: TIMSS evalúa lo acumulado, no cada grado. PISA: hoja «PISA - Contenidos»."'])
    b0_, b1_ = 6, [b for b in BG if b[1] in (M4, M8)][-1][4]  # 2.°–8.°; Física (Advanced) en otra frase
    I_, Ab, Cb = rg('I', b0_, b1_), rg('A', b0_, b1_), rg('C', b0_, b1_)
    caja(wbg, 3, 'TIMSS 2027', [
        f'"● 2.°–8.°: la mayor falta es "&{FX(f"MAX({I_})")}&" temas de "&INDEX({Cb},MATCH(MAX({I_}),{I_},0))&" en "&{G2(f"INDEX({Ab},MATCH(MAX({I_}),{I_},0))")}&"; la mayor sobra, "&{FX(f"-MIN({I_})")}&" de "&INDEX({Cb},MATCH(MIN({I_}),{I_},0))&" en "&{G2(f"INDEX({Ab},MATCH(MIN({I_}),{I_},0))")}&"."',
        '"● Por grado las brechas se compensan entre grados: decide con «TIMSS - Balance por ciclo»."'])

    # --- Semáforo TIMSS
    D_ = rg(DCOL, 5, OL); Cc_ = rg(CCOL, 5, OL)
    caja(wo, 2, 'TIMSS 2027 (2.°–8.°) y TIMSS Advanced (Física)', [
        f'{TAG('COUNTIF(%s,"No cubierto")=0' % D_)}&"Rojo: "&COUNTIF({D_},"No cubierto")&" objetivos de TIMSS sin ningún tema. Amarillo: "&COUNTIF({D_},"Débil (1 tema)")&" con un solo tema."',
        f'"● Cubiertos: "&COUNTIF({Cc_},"Sí")&" de "&ROWS({Cc_})&" objetivos. «N/A» = ese grado no lo evalúa ese TIMSS (no es un problema)."',
        '"→ Filtra «Estado» por rojo o amarillo; la última columna dice qué hacer. Lo de PISA está en «PISA - Contenidos»."'])

    # --- PISA
    pc_ = BP[-1]; p9_ = [b for b in BP if b[0] == '9°'][0]; p8_ = [b for b in BP if b[0] == '8°'][0]
    caja(wpc, 3, 'PISA 2025 (7.°–9.°)', [
        f'{TAG('COUNTIF($K$%d:$K$%d,"En meta")=3' % (pc_[2], pc_[2] + 2))}&"7.°–9.°: sobran "&{FX(f"-I{pc_[2]}")}&" temas de sistemas físicos; faltan "&{FX(f"I{pc_[2] + 1}")}&" de sistemas vivos y "&{FX(f"I{pc_[2] + 2}")}&" de Tierra y espacio."',
        f'"OJO · 9.° tiene solo "&D{p9_[2] + 1}&" temas de biología; 8.° solo "&D{p8_[2] + 2}&" de Tierra y espacio."',
        f'"● Contenidos de PISA presentes: "&C{AMP_P["(todos)"]}&" de "&B{AMP_P["(todos)"]}&". La cobertura TIMSS está en «TIMSS - Semáforo de objetivos»."'])
    rc_ = COMP_ROW[CICLO_P]
    caja(wpk, 2, 'PISA 2025 (7.°–9.°)', [
        f'{TAG(f"H{rc_}>=$B${RC[C3_]}")}&"Competencia 3 (evaluar información y decidir): "&{PC(f"H{rc_}")}&" de los temas. PISA pide 24–36 %."',
        f'{TAG("J%d=%sDentro del rango%s" % (rc_, chr(34), chr(34)))}&"Competencia 2 (diseñar y leer datos): "&{PC(f"G{rc_}")}&", "&LOWER(J{rc_})&" (24–36 %)."',
        f'"● Subcompetencias débiles o ausentes: "&(COUNTIF({rg("I", s0, s1)},"Débil")+COUNTIF({rg("I", s0, s1)},"No trabajada"))&" de 15: distinguir evidencia de opinión, criticar argumentos y plantear hipótesis."',
        '"→ Dato de la propia malla (columnas Competencia y Microhabilidades), sin IA."'])
    rk_ = KN_ROW[CICLO_P]
    caja(wpn, 2, 'PISA 2025 (7.°–9.°)', [
        f'{TAG(f"G{rk_}>=$B${RK[PRO_]}")}&"Temas centrados en cómo se obtienen datos (procedimental): "&{PC(f"G{rk_}")}&". PISA: 27–33 %."',
        f'{TAG(f"H{rk_}>=$B${RK[EPI_]}")}&"Temas centrados en por qué confiamos en la ciencia (epistémico): "&{PC(f"H{rk_}")}&". PISA: 24–30 %."',
        f'"● "&E{PE_ROW["Epistémico"]}&" de "&B{PE_ROW["Epistémico"]}&" ideas epistémicas no aparecen (modelo vs. realidad, consenso, revisión por pares)."',
        '"→ No hacen falta temas nuevos: basta una fase «Reforzar» por unidad que trabaje el porqué."'])
    rx_ = CX_ROW[CICLO_P]; a0_, a1_ = AREA_ROWS
    caja(wpx, 2, 'PISA 2025 (7.°–9.°)', [
        f'{TAG(f"G{rx_}<=0.5")}&{PC(f"G{rx_}")}&" de los temas no se conectan con una situación real ("&{PC(f"G{CX_ROW[K8]}")}&" en 8.°)."',
        f'"● Con contexto, el reparto personal / local / global coincide "&{PC(f"K{rx_}")}&" con el 1:2:1 de PISA."',
        f'"● Área menos trabajada: "&INDEX({rg("A", a0_, a1_)},MATCH(MIN({rg("E", a0_, a1_)}),{rg("E", a0_, a1_)},0))&" ("&MIN({rg("E", a0_, a1_)})&" temas)."'])

    # --- TIMSS: medio ambiente, investigación, cognitivo, otros países
    ga0, ga1 = AMB['2°'], AMB['8°']
    caja(wa, 2, 'TIMSS 2027 (2.°–8.°)', [
        f'{TAG("MIN(E%d,E%d)>=$E$3-$H$3" % (AMB[K24], AMB[K58]))}&"Temas sobre medio ambiente: "&{PC(f"E{AMB[K24]}")}&" en 2.°–4.° y "&{PC(f"E{AMB[K58]}")}&" en 5.°–8.° (la prueba: 25 %)."',
        f'{TAG(f"MIN({rg(AN_, ga0, ga1)})>=$E$3-$H$3")}&"Grado con menos: "&{G2(f"INDEX({rg(AM_, ga0, ga1)},MATCH(MIN({rg(AN_, ga0, ga1)}),{rg(AN_, ga0, ga1)},0))")}&" ("&{PC(f"MIN({rg(AN_, ga0, ga1)})")}&")."',
        f'"● Objetivos ambientales de TIMSS presentes: "&(C{AMBOBJ[M4]}+C{AMBOBJ[M8]})&" de "&(B{AMBOBJ[M4]}+B{AMBOBJ[M8]})&"."'])
    caja(wi, 2, 'TIMSS 2027 (2.°–8.°)', [
        f'"● Temas centrados en investigar: "&C{INV[K24]}&" en 2.°–4.° y "&C{INV[K58]}&" en 5.°–8.°."',
        f'"● "&(H{INV[K24]}+H{INV[K58]})&" temas que TIMSS 2023 dejaba fuera ahora cuentan (área nueva de 2027)."',
        '"→ Investigar = el foco es medir, registrar, analizar datos o concluir."'])
    n_c = len([g for g, m in GROUPS if m in (M4, M8)])
    Ir, Ar = rg('I', 5, 4 + n_c), rg('A', 5, 4 + n_c)
    caja(wg, 2, 'TIMSS 2027 (2.°–8.°) y Advanced', [
        f'"OJO · Menos Razonar: "&{G2(f"INDEX({Ar},MATCH(MIN({Ir}),{Ir},0))")}&" ("&{PC(f"MIN({Ir})")}&"). TIMSS pide 25 % en 4.° y 30 % en 8.°."',
        f'"● Grados de 2.° a 8.° con Razonar en línea con TIMSS: "&COUNTIF({rg("N", 5, 4 + n_c)},"Razonar en línea*")&" de "&ROWS({rg("N", 5, 4 + n_c)})&"."',
        '"● Sale de la columna «Habilidad TIMSS» que escribió el equipo de Ciencias. Las competencias PISA están en «PISA - Competencias»."'])
    if HAY_PAISES:
        K_ = rg(CL_, 5, XL); J_ = rg(get_column_letter(column_index_from_string(CL_) - 1), 5, XL); Ax = rg('A', 5, XL)
        caja(wx, 2, 'TIMSS 2027 (objetivos de 4.° y 8.°)', [
            f'"● El Salvador enseña "&COUNTIF({K_},"El Salvador llega más tarde")&" objetivos de TIMSS más tarde que Uruguay, Colombia y Singapur, y "&COUNTIF({K_},"El Salvador va adelantado")&" antes."',
            f'{TAG("COUNTIF(%s,%sEl Salvador no llega%s)=0" % (K_, chr(34), chr(34)))}&"No enseña "&COUNTIF({K_},"El Salvador no llega")&" que otros sí. Mayor atraso ("&{FX(f"MAX({J_})")}&" grados): "&INDEX({rg("D", 5, XL)},MATCH(MAX({J_}),{J_},0))&"."',
            '"● Chile no está: su sitio bloqueó la descarga."'])

    # --- Datos y fuentes
    caja(wt, 2, 'TIMSS 2027 (columnas K–AC) y PISA 2025 (columnas AD–AW)', [
        f'"● Un tema por fila: "&COUNTA({rg("A", 5, TL)})&" en total; "&(COUNTA({rg("K", 5, TL)})-COUNTIF({rg("K", 5, TL)},"N/A"))&" con clasificación TIMSS y "&COUNTA({rg("AK", 5, TL)})&" con PISA (7.°–9.°)."',
        f'"● Confianza baja (para revisar primero): "&COUNTIF({rg("P", 5, TL)},"baja")&" en TIMSS (incluye Física de Bachillerato) y "&COUNTIF({rg("AT", 5, TL)},"baja")&" en PISA."',
        '"→ Cambia un código en K–L (TIMSS) o AK–AL (PISA) y todo el libro se recalcula."'])
    caja(wc, 2, 'TIMSS', [
        f'"● TIMSS 2027: "&COUNTIF({rg("B", 5, CAT_LAST)},"{M4}")&" objetivos de 4.° y "&COUNTIF({rg("B", 5, CAT_LAST)},"{M8}")&" de 8.°. TIMSS Advanced Física: "&COUNTIF({rg("B", 5, CAT_LAST)},"TA")&"."',
        f'"● "&COUNTIF({rg("I", 5, CAT_LAST)},"Sí")&" son de medio ambiente y "&COUNTIF({rg("J", 5, CAT_LAST)},"Sí")&" de investigación. Metas en ámbar (editables)."'])
    caja(wcp, 2, 'PISA 2025', [
        f'"● "&COUNTIF({rg("B", 5, CATP_LAST)},"Contenido")&" contenidos, "&COUNTIF({rg("B", 5, CATP_LAST)},"Procedimental")&" procedimentales, "&COUNTIF({rg("B", 5, CATP_LAST)},"Epistémico")&" epistémicos y "&COUNTIF({rg("B", 5, CATP_LAST)},"Subcompetencia")&" subcompetencias."',
        '"● Todo con la página del marco oficial (OCDE 2026). Metas en ámbar (editables)."'])
    if HAY_PAISES:
        caja(wp, 2, 'TIMSS 2027 como referencia común', [
            f'"● "&COUNTA({rg("A", 5, PL)})&" objetivos: Uruguay "&COUNTIF({rg("B", 5, PL)},"Uruguay")&", Colombia "&COUNTIF({rg("B", 5, PL)},"Colombia")&", Singapur "&COUNTIF({rg("B", 5, PL)},"Singapur")&"."',
            f'"● Sin objetivo TIMSS: "&COUNTIF({rg("J", 5, PL)},"FUERA")&". Confianza baja: "&COUNTIF({rg("L", 5, PL)},"baja")&"."'])
    caja(wr, 2, 'TIMSS 2023 (versión anterior)', [
        f'"● Solo para comparar. Con TIMSS 2023 se cubrían "&C{AMP23["T4"]}&" de "&B{AMP23["T4"]}&" objetivos de 4.° y "&C{AMP23["T8"]}&" de "&B{AMP23["T8"]}&" de 8.°."'])
    caja(wsf, 2, 'TIMSS y PISA', [
        f'"● "&D{FUENTES_TOT}&" temas leídos de 11 mallas del MINED."',
        '"● Marcos oficiales: TIMSS 2027, TIMSS Advanced 2015 y PISA 2025, citados con página."'])
    for c_ in wl['B']:  # recuadro de hallazgos del Inicio
        if isinstance(c_.value, str) and c_.value.startswith('="TIMSS 2027 y PISA 2025"&'):
            c_.font = Font(name=F, size=11, bold=True, color=VERDE); c_.fill = CAJA; c_.border = BOX
            c_.alignment = Alignment(wrap_text=True, vertical='top'); wl.row_dimensions[c_.row].height = 110

    # --- nombres de pestaña: fórmulas («Hoja!» → «'Nombre'!») y textos que nombran hojas
    viejos = sorted(NOMBRES, key=len, reverse=True)
    pat_f = _re.compile(r"(?<![\w'])(" + '|'.join(_re.escape(v) for v in viejos) + r")!")
    pat_t = _re.compile(r'(?<![\w·])(' + '|'.join(_re.escape(v) for v in viejos if '_' in v) + r')(?![\w])')
    for s_ in wb.worksheets:
        for row in s_.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith('='):
                        v = pat_f.sub(lambda m: f"'{NOMBRES[m.group(1)]}'!", v)
                        v = _re.sub(r'(«|hoja |hojas )(Leeme|Resumen|Presupuesto|Ambiental|Investigaciones|Temas|Fuentes)(?=[»,.)])',
                                    lambda m: m.group(1) + NOMBRES[m.group(2)], v)
                    else:
                        v = pat_t.sub(lambda m: NOMBRES[m.group(1)], v)
                        v = _re.sub(r'^(Resumen|Ambiental|Investigaciones|Temas|Presupuesto)(?=:)', lambda m: NOMBRES[m.group(1)], v)
                        v = _re.sub(r'(«|hoja |hojas |, |y )(Leeme|Resumen|Presupuesto|Ambiental|Investigaciones|Temas|Fuentes)(?=[»,.)]|$)',
                                    lambda m: m.group(1) + NOMBRES[m.group(2)], v)
                    c.value = v
    for s_ in wb.worksheets:
        if s_.title in NOMBRES:
            s_.title = NOMBRES[s_.title]
if HAY_PISA:
    _fmt = {'+0.0;-0.0;0.0': '0.0', PP: '0.0" pp"', '+0;-0;0': '0', '+0.0" pp";-0.0" pp";0.0" pp"': '0.0" pp"'}
    for s_ in wb.worksheets:  # un solo formato que Excel y Numbers muestran igual
        for row in s_.iter_rows():
            for c in row:
                if c.number_format in _fmt:
                    c.number_format = _fmt[c.number_format]
    wl.row_dimensions[1].height = 40
    for c_ in wl['B']:  # alto según el texto visible, no según el largo de la fórmula
        if isinstance(c_.value, str) and c_.value.startswith('=') and not c_.value.startswith('="TIMSS 2027 y PISA 2025"&'):
            vis = len(_re.sub(r'"&[^&]*&"', '00 %', c_.value))
            wl.row_dimensions[c_.row].height = 15 * (vis // 150 + 1) + 4
        elif isinstance(c_.value, str) and c_.value.startswith('="TIMSS 2027 y PISA 2025"&'):
            wl.row_dimensions[c_.row].height = 92
wb.active = 0
if ES27:  # Excel recalcula todo al abrir (el archivo de openpyxl no trae valores guardados)
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
out = ARGS.out or MC['out']
os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
wb.save(out)
print('ok', out, 'temas', TL - 4, 'objetivos', OL - 4, 'faltantes', FL - 4)
