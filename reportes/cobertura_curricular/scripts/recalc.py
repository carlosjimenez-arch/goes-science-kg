"""Recalcula TODAS las fórmulas de un .xlsx con LibreOffice (headless) y reporta errores.

Uso: python scripts/recalc.py outputs/Cobertura_TIMSS_Ciencias_SV.xlsx
Busca LibreOffice en: $SOFFICE, PATH (soffice/libreoffice) y /Applications/LibreOffice.app (macOS).
Si no está instalado, NO instalarlo en la Mac del equipo: recalcular con Numbers y exportar a
outputs/<libro>_numbers.xlsx (ver «4-bis» en CLAUDE.md); pytest revisa ese archivo.
"""
import os, shutil, subprocess, sys, tempfile
from pathlib import Path
import openpyxl

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRORES = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!', 'Err:502', 'Err:504', 'Err:508')


def find_soffice():
    for c in (os.environ.get('SOFFICE'), shutil.which('soffice'), shutil.which('libreoffice'),
              '/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        if c and os.path.exists(c):
            return c
    return None


def main(path, timeout=300):
    so = find_soffice()
    if not so:
        print('LibreOffice no encontrado (no instalarlo). Recalcula con Numbers y exporta a outputs/<libro>_numbers.xlsx: ver «4-bis» en CLAUDE.md.')
        sys.exit(2)
    path = os.path.abspath(path)
    with tempfile.TemporaryDirectory() as tmp:
        prof = Path(tmp) / 'profile'
        env_arg = f'-env:UserInstallation={prof.as_uri()}'
        subprocess.run([so, '--headless', '--terminate_after_init', env_arg], capture_output=True, timeout=120)
        mdir = prof / 'user' / 'basic' / 'Standard'
        mdir.mkdir(parents=True, exist_ok=True)
        (mdir / 'Module1.xba').write_text(MACRO)
        subprocess.run([so, '--headless', '--norestore', env_arg,
                        'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application', path],
                       capture_output=True, timeout=timeout)
    wb = openpyxl.load_workbook(path, data_only=True)
    errs = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value in ERRORES]
    resumen = wb['Resumen']['C5'].value if 'Resumen' in wb.sheetnames else 'n/a'
    status = 'success' if not errs and resumen is not None else 'errors_found'
    print({'status': status, 'total_errors': len(errs), 'muestra': errs[:10], 'Resumen!C5': resumen})
    sys.exit(0 if status == 'success' else 1)


if __name__ == '__main__':
    main(sys.argv[1])
