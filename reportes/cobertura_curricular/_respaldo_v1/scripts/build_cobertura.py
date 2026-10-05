"""Construye Cobertura_TIMSS_Ciencias_SV.xlsx a partir de:
- temas.json (1.260 temas extraídos de las 11 mallas)
- catalogo_timss.json (101 objetivos TIMSS 2023 4°/8° y TIMSS Advanced 2015 Física)
- clasificacion.json (tema -> objetivo principal/secundario, confianza, justificación)
Todas las cifras de las hojas de análisis son fórmulas sobre la hoja Temas_TIMSS.
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

T = json.load(open('data/interim/temas.json'))
CAT = json.load(open('data/referencia/catalogo_timss.json'))
R = json.load(open('data/interim/clasificacion.json'))

F = 'Arial'
fN = Font(name=F, size=10)
fB = Font(name=F, size=10, bold=True)
fH = Font(name=F, size=10, bold=True, color='FFFFFF')
fT = Font(name=F, size=14, bold=True)
fS = Font(name=F, size=11, bold=True)
fIn = Font(name=F, size=10, color='0000FF')          # inputs
fLink = Font(name=F, size=10, color='008000')        # cross-sheet links
fNote = Font(name=F, size=9, italic=True, color='555555')
HDR = PatternFill('solid', fgColor='1F4E79')
SUB = PatternFill('solid', fgColor='DDEBF7')
INP = PatternFill('solid', fgColor='FFF2CC')
GRY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)
PCT = '0.0%'
PP = '+0.0" pp";-0.0" pp";0.0" pp"'
NUM = '0;-0;0'
NUM1 = '+0;-0;0'

def grp(t):
    return f"{t['grado']}°" if t['asignatura'] == 'Ciencias' else f"{t['asignatura']} {t['grado']}°"

def ciclo(t):
    g, a = t['grado'], t['asignatura']
    if a == 'Ciencias':
        if g <= 4: return 'Ciclo 2°–4°'
        if g <= 8: return 'Ciclo 5°–8°'
        return '9°'
    if a == 'Física': return 'Física 10°–11°'
    return f'{a} 10°–11°'

def marco(t):
    g, a = t['grado'], t['asignatura']
    if a == 'Ciencias': return 'T4' if g <= 4 else 'T8'
    if a == 'Física': return 'TA'
    return 'Sin marco'

def hdr(ws, row, headers, widths=None, height=30):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = fH; c.fill = HDR; c.alignment = CTR; c.border = BOX
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

def style_range(ws, r1, r2, c1, c2, fmt=None, font=fN, fill=None, align=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font; cell.border = BOX
            if fmt: cell.number_format = fmt
            if fill: cell.fill = fill
            if align: cell.alignment = align

wb = Workbook()

# ------------------------------------------------------------------ Catalogo_TIMSS
wc = wb.active; wc.title = 'Catalogo_TIMSS'
wc['A1'] = 'Catálogo de referencia TIMSS (objetivos de contenido)'; wc['A1'].font = fT
wc['A2'] = ('T4 = TIMSS 2023 4.° grado · T8 = TIMSS 2023 8.° grado · TA = TIMSS Advanced 2015 Física. '
            'Objetivos parafraseados al español (no son traducción oficial). Meta = % de puntos de la prueba asignado al dominio en el marco.')
wc['A2'].font = fNote; wc.merge_cells('A2:H2'); wc.row_dimensions[2].height = 30; wc['A2'].alignment = WR
hdr(wc, 4, ['Código', 'Marco', 'Dominio', 'Meta del dominio', 'Código área', 'Área temática', 'Objetivo (paráfrasis)', 'Fuente'],
    [12, 8, 26, 12, 11, 40, 70, 30])
src = {'T4': 'IEA, TIMSS 2023 Frameworks, cap. 2 (Science), 4.° grado',
       'T8': 'IEA, TIMSS 2023 Frameworks, cap. 2 (Science), 8.° grado',
       'TA': 'IEA, TIMSS Advanced 2015 Frameworks, cap. 2 (Physics)'}
for i, c in enumerate(CAT, 5):
    vals = [c['codigo'], c['marco'], c['dominio'], c['meta_dominio'], c['area_codigo'], c['area'], c['objetivo'], src[c['marco']]]
    for j, v in enumerate(vals, 1):
        wc.cell(row=i, column=j, value=v)
    style_range(wc, i, i, 1, 8, align=WR)
    wc.cell(row=i, column=4).number_format = '0%'; wc.cell(row=i, column=4).font = fIn
CAT_LAST = 4 + len(CAT)
wc.freeze_panes = 'A5'; wc.auto_filter.ref = f'A4:H{CAT_LAST}'
CR = lambda col: f"Catalogo_TIMSS!${col}$5:${col}${CAT_LAST}"

# ------------------------------------------------------------------ Temas_TIMSS (trazabilidad)
wt = wb.create_sheet('Temas_TIMSS')
wt['A1'] = 'Trazabilidad: cada tema de las mallas y su clasificación contra TIMSS'; wt['A1'].font = fT
wt['A2'] = ('Una fila = un procedimental de la malla (un «tema»). Columnas K-L, P-Q: clasificación asistida por IA (revisable). '
            'M, N, O y S-U son fórmulas. Para auditar cualquier número del libro, filtra esta hoja por Grupo y Dominio.')
wt['A2'].font = fNote; wt.merge_cells('A2:W2'); wt['A2'].alignment = WR; wt.row_dimensions[2].height = 28
H = ['ID tema', 'Grado', 'Asignatura', 'Archivo de origen', 'Hoja', 'Fila', 'Unidad', 'Contenido', 'Procedimental (tema)',
     'Marco', 'Objetivo TIMSS principal', 'Objetivo TIMSS secundario', 'Dominio (principal)', 'Área temática (principal)',
     'Dominio (secundario)', 'Confianza', 'Justificación de la clasificación', 'Habilidad TIMSS (texto de la malla)',
     'Conocer (fracción)', 'Aplicar (fracción)', 'Razonar (fracción)', 'Grupo', 'Ciclo']
hdr(wt, 4, H, [17, 7, 10, 30, 22, 6, 28, 26, 55, 9, 13, 13, 22, 32, 20, 10, 45, 45, 9, 9, 9, 12, 15], height=42)
order = sorted(T, key=lambda t: ({'Ciencias': 0, 'Física': 1, 'Química': 2, 'Biología': 3}[t['asignatura']], t['grado'], t['id']))
r = 5
for t in order:
    c = R.get(t['id'])
    m = marco(t)
    pend = (c is None) and t['asignatura'] in ('Ciencias', 'Física')
    o1 = c['obj1'] if c else ('PENDIENTE' if pend else 'N/A')
    o2 = c['obj2'] if c else ''
    vals = [t['id'], t['grado'], t['asignatura'], t['archivo'], t['hoja'], t['fila'], t['unidad'], t['contenido'],
            t['procedimental'], m, o1, o2, None, None, None, c['confianza'] if c else 'N/A',
            c['justificacion'] if c else ('Tema nuevo o modificado: falta clasificar' if pend else 'Sin marco TIMSS de contenido para esta asignatura en Bachillerato'),
            t['habilidad_timss'], None, None, None, grp(t), ciclo(t)]
    for j, v in enumerate(vals, 1):
        if v is not None:
            wt.cell(row=r, column=j, value=v)
    wt[f'M{r}'] = (f'=IF(K{r}="FUERA","Fuera del marco",IF(K{r}="N/A","Sin marco",'
                   f'IFERROR(INDEX({CR("C")},MATCH(K{r},{CR("A")},0)),"Pendiente de clasificar")))')
    wt[f'N{r}'] = (f'=IF(K{r}="FUERA","Fuera del marco",IF(K{r}="N/A","Sin marco",'
                   f'IFERROR(INDEX({CR("F")},MATCH(K{r},{CR("A")},0)),"Pendiente de clasificar")))')
    wt[f'O{r}'] = f'=IF(L{r}="","",IFERROR(INDEX({CR("C")},MATCH(L{r},{CR("A")},0)),"?"))'
    k = (f'(ISNUMBER(SEARCH("Conocer",$R{r}))+ISNUMBER(SEARCH("Aplicar",$R{r}))+ISNUMBER(SEARCH("Razonar",$R{r})))')
    wt[f'S{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Conocer",$R{r}))/{k},0)'
    wt[f'T{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Aplicar",$R{r}))/{k},0)'
    wt[f'U{r}'] = f'=IFERROR(ISNUMBER(SEARCH("Razonar",$R{r}))/{k},0)'
    style_range(wt, r, r, 1, 23)
    for col in 'STU':
        wt[f'{col}{r}'].number_format = '0.00'
    for col in 'IQR':
        wt[f'{col}{r}'].alignment = Alignment(wrap_text=False, vertical='top')
    r += 1
TL = r - 1
wt.freeze_panes = 'C5'; wt.auto_filter.ref = f'A4:W{TL}'
wt.conditional_formatting.add(f'P5:P{TL}', CellIsRule(operator='equal', formula=['"baja"'], fill=PatternFill('solid', fgColor='F8CBAD')))
wt.conditional_formatting.add(f'K5:K{TL}', CellIsRule(operator='equal', formula=['"FUERA"'], fill=GRY))
TC = lambda col: f"Temas_TIMSS!${col}$5:${col}${TL}"

# ------------------------------------------------------------------ Balance_grado
GROUPS = [('2°', 'T4'), ('3°', 'T4'), ('4°', 'T4'), ('5°', 'T8'), ('6°', 'T8'), ('7°', 'T8'), ('8°', 'T8'), ('9°', 'T8'),
          ('Física 10°', 'TA'), ('Física 11°', 'TA')]
DOMS = {}
for c in CAT:
    DOMS.setdefault(c['marco'], [])
    if c['dominio'] not in [d for d, _ in DOMS[c['marco']]]:
        DOMS[c['marco']].append((c['dominio'], c['meta_dominio']))
MNAME = {'T4': 'TIMSS 2023 4.°', 'T8': 'TIMSS 2023 8.°', 'TA': 'TIMSS Advanced 2015 Física'}


def balance_sheet(ws, title, key_col, groups, note):
    ws['A1'] = title; ws['A1'].font = fT
    ws['A2'] = 'Tolerancia para considerar «en meta» (± puntos porcentuales):'; ws['A2'].font = fB
    ws['H2'] = 0.05; ws['H2'].font = fIn; ws['H2'].fill = INP; ws['H2'].number_format = '0%'; ws['H2'].border = BOX
    ws['H2'].comment = Comment('Supuesto editable. 5 pp = margen razonable para un conteo por temas.', 'Claude')
    ws['A3'] = note; ws['A3'].font = fNote; ws.merge_cells('A3:N3'); ws['A3'].alignment = WR; ws.row_dimensions[3].height = 54
    heads = ['Grupo', 'Marco de referencia', 'Dominio', 'Temas actuales en el dominio', 'Temas evaluables del grupo (sin «fuera»)',
             '% actual', 'Meta TIMSS (%)', 'Temas esperados según meta', 'Brecha (temas): + faltan / − sobran',
             'Brecha (pp)', 'Estado', 'Escala para cumplir sin quitar temas (total evaluables)',
             'Temas a AGREGAR para llegar a la meta sin quitar', 'Temas «fuera del marco» del grupo']
    hdr(ws, 5, heads, [12, 22, 26, 11, 13, 9, 9, 11, 13, 10, 22, 14, 14, 12], height=66)
    r = 6
    blocks = []
    for g, m in groups:
        doms = DOMS[m]
        r0 = r
        for d, meta in doms:
            ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]; ws[f'C{r}'] = d
            ws[f'D{r}'] = f'=COUNTIFS({TC(key_col)},$A{r},{TC("M")},$C{r})'
            ws[f'E{r}'] = (f'=COUNTIFS({TC(key_col)},$A{r})-COUNTIFS({TC(key_col)},$A{r},{TC("K")},"FUERA")'
                           f'-COUNTIFS({TC(key_col)},$A{r},{TC("K")},"PENDIENTE")')
            ws[f'F{r}'] = f'=IF(E{r}=0,0,D{r}/E{r})'
            ws[f'G{r}'] = meta; ws[f'G{r}'].font = fIn
            ws[f'H{r}'] = f'=G{r}*E{r}'
            ws[f'I{r}'] = f'=H{r}-D{r}'
            ws[f'J{r}'] = f'=(G{r}-F{r})*100'
            ws[f'K{r}'] = (f'=IF(ABS(G{r}-F{r})<=$H$2,"En meta",IF(I{r}>0,"Faltan "&ROUND(I{r},0)&" temas",'
                           f'"Sobran "&ROUND(-I{r},0)&" temas"))')
            r += 1
        r1 = r - 1
        mx = ','.join(f'D{x}/G{x}' for x in range(r0, r1 + 1))
        for x in range(r0, r1 + 1):
            ws[f'L{x}'] = f'=MAX({mx})'
            ws[f'M{x}'] = f'=MAX(0,ROUNDUP(G{x}*L{x}-D{x},0))'
            ws[f'N{x}'] = f'=COUNTIFS({TC(key_col)},$A{x},{TC("K")},"FUERA")'
        style_range(ws, r0, r1, 1, 14)
        for x in range(r0, r1 + 1):
            ws[f'F{x}'].number_format = PCT; ws[f'G{x}'].number_format = '0%'; ws[f'G{x}'].font = fIn
            ws[f'H{x}'].number_format = '0.0'; ws[f'I{x}'].number_format = '+0.0;-0.0;0.0'
            ws[f'J{x}'].number_format = PP; ws[f'L{x}'].number_format = '0.0'
            ws[f'K{x}'].font = fB
        # fila total del bloque
        ws[f'C{r}'] = 'Total del grupo'; ws[f'C{r}'].font = fB
        ws[f'D{r}'] = f'=SUM(D{r0}:D{r1})'; ws[f'G{r}'] = f'=SUM(G{r0}:G{r1})'; ws[f'G{r}'].number_format = '0%'
        ws[f'M{r}'] = f'=SUM(M{r0}:M{r1})'
        style_range(ws, r, r, 1, 14, fill=SUB, font=fB)
        ws[f'G{r}'].number_format = '0%'
        blocks.append((g, m, r0, r1, r))
        r += 2
    last = r
    ws.conditional_formatting.add(f'K6:K{last}', FormulaRule(formula=['LEFT(K6,6)="Faltan"'], fill=PatternFill('solid', fgColor='F8CBAD')))
    ws.conditional_formatting.add(f'K6:K{last}', FormulaRule(formula=['LEFT(K6,6)="Sobran"'], fill=PatternFill('solid', fgColor='FFE699')))
    ws.conditional_formatting.add(f'K6:K{last}', CellIsRule(operator='equal', formula=['"En meta"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.freeze_panes = 'D6'
    return blocks

wbg = wb.create_sheet('Balance_grado')
BG = balance_sheet(wbg, 'Balance de contenidos por grado frente a las metas de TIMSS', 'V', GROUPS,
    'Cómo leerlo: «% actual» = temas del dominio ÷ temas evaluables del grado (se excluyen los «fuera del marco»). '
    '«Brecha (temas)» = meta × evaluables − actuales: positivo = faltan temas en ese dominio si se redistribuye la misma carga; negativo = sobran. '
    '«Temas a AGREGAR» = cuántos temas habría que sumar a cada dominio para cumplir todas las metas SIN quitar ninguno (la escala la fija el dominio más sobrerrepresentado). '
    'Ojo: TIMSS fija pesos para la prueba acumulada (4.° y 8.°); aplicarlos a cada grado es una aproximación. La hoja Balance_ciclo hace la comparación acumulada.')

wcy = wb.create_sheet('Balance_ciclo')
CYCLES = [('Ciclo 2°–4°', 'T4'), ('Ciclo 5°–8°', 'T8'), ('9°', 'T8'), ('Física 10°–11°', 'TA')]
BC = balance_sheet(wcy, 'Balance de contenidos por ciclo (comparación acumulada, la más fiel a TIMSS)', 'W', CYCLES,
    'TIMSS 4.° evalúa lo aprendido hasta 4.° grado → se compara con el acumulado de 2.° a 4.° (1.° grado no tiene malla). '
    'TIMSS 8.° evalúa lo aprendido hasta 8.° → acumulado de 5.° a 8.°. 9.° se muestra contra TIMSS 8.° solo como referencia. '
    'Física de Bachillerato se compara con TIMSS Advanced 2015 (Física). Química y Biología de Bachillerato no tienen marco TIMSS de contenido.')

# ------------------------------------------------------------------ Cobertura_objetivos
wo = wb.create_sheet('Cobertura_objetivos')
wo['A1'] = 'Amplitud: ¿qué objetivos TIMSS cubre la malla y en qué grado?'; wo['A1'].font = fT
wo['A2'] = ('Cada celda de grado = n.º de temas de ese grado que trabajan el objetivo como principal o secundario (COUNTIFS sobre Temas_TIMSS). '
            '«Total en su ciclo» suma solo los grados que evalúa cada marco: TIMSS 4.° → 2.°–4.°; TIMSS 8.° → 5.°–8.° (la columna 9.° es referencia); '
            'TIMSS Advanced → Física 10.°–11.°. «Profundidad»: 0 = no cubierto; 1 = débil; 2–3 = básica; 4+ = sólida.')
wo['A2'].font = fNote; wo.merge_cells('A2:T2'); wo['A2'].alignment = WR; wo.row_dimensions[2].height = 28
GCOLS = [g for g, _ in GROUPS]
heads = ['Código', 'Marco', 'Dominio', 'Área temática', 'Objetivo'] + GCOLS + ['Total temas en su ciclo', 'Como principal (todos los grados)', 'Cubierto en su ciclo', 'Profundidad en su ciclo', 'Primer grado donde aparece']
hdr(wo, 4, heads, [12, 7, 22, 30, 52] + [7] * len(GCOLS) + [8, 9, 9, 15, 12], height=45)
r = 5
for c in CAT:
    wo[f'A{r}'] = c['codigo']; wo[f'B{r}'] = c['marco']; wo[f'C{r}'] = c['dominio']; wo[f'D{r}'] = c['area']; wo[f'E{r}'] = c['objetivo']
    for j, g in enumerate(GCOLS):
        col = get_column_letter(6 + j)
        wo[f'{col}{r}'] = (f'=COUNTIFS({TC("V")},{col}$4,{TC("K")},$A{r})+COUNTIFS({TC("V")},{col}$4,{TC("L")},$A{r})')
    gc1, gc2 = get_column_letter(6), get_column_letter(5 + len(GCOLS))
    tcol = get_column_letter(6 + len(GCOLS))
    # Total SOLO dentro del ciclo que evalúa cada marco: T4 = 2°–4° (F:H); T8 = 5°–8° (I:L, 9° queda como referencia); TA = Física 10°–11° (N:O)
    wo[f'{tcol}{r}'] = f'=IF($B{r}="T4",SUM(F{r}:H{r}),IF($B{r}="T8",SUM(I{r}:L{r}),SUM(N{r}:O{r})))'
    pcol = get_column_letter(7 + len(GCOLS)); wo[f'{pcol}{r}'] = f'=COUNTIFS({TC("K")},$A{r})'
    ccol = get_column_letter(8 + len(GCOLS)); wo[f'{ccol}{r}'] = f'=IF({tcol}{r}>0,"Sí","No")'
    dcol = get_column_letter(9 + len(GCOLS))
    wo[f'{dcol}{r}'] = f'=IF({tcol}{r}=0,"No cubierto",IF({tcol}{r}=1,"Débil (1 tema)",IF({tcol}{r}<=3,"Básica (2–3)","Sólida (4+)")))'
    fcol = get_column_letter(10 + len(GCOLS))
    wo[f'{fcol}{r}'] = f'=IFERROR(INDEX(${gc1}$4:${gc2}$4,MATCH(TRUE,INDEX({gc1}{r}:{gc2}{r}>0,0),0)),"—")'
    style_range(wo, r, r, 1, 10 + len(GCOLS), align=Alignment(vertical='top', wrap_text=True))
    r += 1
OL = r - 1
TCOL, CCOL, DCOL = tcol, ccol, dcol
wo.freeze_panes = 'F5'; wo.auto_filter.ref = f'A4:{fcol}{OL}'
rng = f'{dcol}5:{dcol}{OL}'
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"No cubierto"'], fill=PatternFill('solid', fgColor='F8CBAD')))
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Débil (1 tema)"'], fill=PatternFill('solid', fgColor='FFE699')))
wo.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Sólida (4+)"'], fill=PatternFill('solid', fgColor='C6EFCE')))
wo.conditional_formatting.add(f'F5:{gc2}{OL}', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=10, color='BFBFBF')))

# resumen de amplitud debajo
r = OL + 3
wo[f'A{r}'] = 'Resumen de amplitud por marco y dominio'; wo[f'A{r}'].font = fS
r += 1
hdr_row = r
for j, h in enumerate(['Marco', 'Dominio', 'Objetivos en el marco', 'Objetivos cubiertos', '% cubierto', 'Objetivos no cubiertos', 'Objetivos débiles (1 tema)'], 1):
    cc = wo.cell(row=r, column=j, value=h); cc.font = fH; cc.fill = HDR; cc.alignment = CTR; cc.border = BOX
r += 1
AMP = []
for m in ['T4', 'T8', 'TA']:
    for d, _ in DOMS[m] + [('(todos)', None)]:
        wo[f'A{r}'] = m; wo[f'B{r}'] = d
        crit_d = '"*"' if d == '(todos)' else f'$B{r}'
        wo[f'C{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d})'
        wo[f'D{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d},${CCOL}$5:${CCOL}${OL},"Sí")'
        wo[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
        wo[f'F{r}'] = f'=C{r}-D{r}'
        wo[f'G{r}'] = f'=COUNTIFS($B$5:$B${OL},$A{r},$C$5:$C${OL},{crit_d},${DCOL}$5:${DCOL}${OL},"Débil (1 tema)")'
        style_range(wo, r, r, 1, 7, fill=SUB if d == '(todos)' else None, font=fB if d == '(todos)' else fN)
        wo[f'E{r}'].number_format = PCT
        AMP.append((m, d, r))
        r += 1

# ------------------------------------------------------------------ Faltantes
wf = wb.create_sheet('Faltantes')
wf['A1'] = 'Qué falta: objetivos TIMSS no cubiertos o débiles, y dónde podrían entrar'; wf['A1'].font = fT
wf['A2'] = ('Lista generada desde Cobertura_objetivos (profundidad «No cubierto» o «Débil»). La columna «Temas hoy» es fórmula y se actualiza si cambia la clasificación. '
            'La sugerencia de ubicación es una propuesta para discutir con el equipo de Ciencias, no una decisión.')
wf['A2'].font = fNote; wf.merge_cells('A2:H2'); wf['A2'].alignment = WR; wf.row_dimensions[2].height = 30
hdr(wf, 4, ['Código', 'Marco', 'Dominio', 'Objetivo', 'Temas hoy', 'Profundidad hoy', 'Dónde aparece hoy', 'Sugerencia de ubicación en la malla'],
    [12, 7, 20, 50, 9, 15, 30, 70], height=32)
SUG = {
 'T4-V2.2': ('—', 'Ciclo 2°–4°: agregar en 3.° U4 «Las plantas» o U5 «Los animales» un tema de semejanza crías-progenitores y rasgos heredados vs. adquiridos (antecede a Genética de 7.°).'),
 'T4-V4.3': ('4.° U3 (1 tema)', 'Reforzar en 4.° U3 «Naturaleza y sus interacciones»: competencia por agua, luz y alimento con una actividad de observación.'),
 'T4-F1.1': ('2.° U1 (1 tema)', 'Reforzar en 2.° U1 «Materia» o 3.° U3 «El agua»: estados sólido, líquido y gaseoso con sus características (hoy solo aparece como cambio de estado del agua).'),
 'T4-F2.3': ('2.° U2 (1 tema)', 'Reforzar en 4.° U1 «Máquinas y energía» o 2.° U2 «Energía»: transferencia de calor entre objetos calientes y fríos (hoy se ve recién en 6.°).'),
 'T8-B5.2': ('7.° U4 (1 tema)', 'Reforzar en 7.° U4 «Ecología» o U5 «Recurso hídrico»: ciclos del carbono y del oxígeno además del agua.'),
 'T8-B6.1': ('—', 'Ciclo 5°–8°: agregar en 8.° U7 «Anatomía y fisiología animal» o 6.° U5 «Célula» un tema de enfermedades infecciosas, transmisión, prevención y sistema inmune (hoy aparece hasta 10.° Biología).'),
 'T8-F3.1': ('—', 'Ciclo 5°–8°: agregar en 8.° U3 «Ondas mecánicas» temas de luz (propagación, reflexión, refracción, diagramas de rayos). Hoy la óptica aparece hasta 10.° Física.'),
 'T8-T4.1': ('Fuera del ciclo: fases de la Luna en 4.° (U4) y mareas en 9.° (U4)', 'Agregar en 5.° U5 «Ciencias del espacio»: traslación y estaciones, mareas, fases lunares y eclipses dentro del ciclo 5°–8°.'),
}
r = 5
gc1, gc2 = get_column_letter(6), get_column_letter(5 + len(GCOLS))
# Determine list in python using same logic to choose rows (values verified later by formulas)
import collections as C
cnt = C.Counter()
Tid = {t['id']: t for t in T}
def in_cycle(code, t):
    if code.startswith('T4'): return t['asignatura'] == 'Ciencias' and t['grado'] <= 4
    if code.startswith('T8'): return t['asignatura'] == 'Ciencias' and 5 <= t['grado'] <= 8
    return t['asignatura'] == 'Física'
for k, v in R.items():
    for o in (v['obj1'], v['obj2']):
        if o and o != 'FUERA' and in_cycle(o, Tid[k]): cnt[o] += 1
falt = [c for c in CAT if cnt[c['codigo']] <= 1]
for c in falt:
    rr = [x for x in range(5, OL + 1) if wo[f'A{x}'].value == c['codigo']][0]
    wf[f'A{r}'] = c['codigo']; wf[f'B{r}'] = c['marco']; wf[f'C{r}'] = c['dominio']; wf[f'D{r}'] = c['objetivo']
    wf[f'E{r}'] = f"=Cobertura_objetivos!{TCOL}{rr}"; wf[f'E{r}'].font = fLink
    wf[f'F{r}'] = f"=Cobertura_objetivos!{DCOL}{rr}"; wf[f'F{r}'].font = fLink
    wf[f'G{r}'] = SUG.get(c['codigo'], ('', ''))[0]
    wf[f'H{r}'] = SUG.get(c['codigo'], ('', 'Revisar con el equipo de Ciencias.'))[1]
    style_range(wf, r, r, 1, 8, align=WR)
    wf[f'E{r}'].font = fLink; wf[f'F{r}'].font = fLink
    r += 1
FL = r - 1
r += 1
wf[f'A{r}'] = 'Además (balance): los faltantes/sobrantes por dominio están en Balance_ciclo (acumulado) y Balance_grado (por grado), columnas I, K y M.'
wf[f'A{r}'].font = fNote; wf.merge_cells(f'A{r}:H{r}')

# ------------------------------------------------------------------ Cognitivo_grado
wg = wb.create_sheet('Cognitivo_grado')
wg['A1'] = 'Profundidad cognitiva: Conocer / Aplicar / Razonar frente a TIMSS'; wg['A1'].font = fT
wg['A2'] = ('Fuente: columna «Habilidad TIMSS» de cada malla. Si un tema menciona varios dominios, se reparte en partes iguales '
            '(p. ej. «Conocer… Aplicar…» = 0,5 + 0,5). Columnas S-U de Temas_TIMSS. Química y Biología de Bachillerato se comparan con la meta de TIMSS Advanced Física solo como referencia aproximada.')
wg['A2'].font = fNote; wg.merge_cells('A2:N2'); wg['A2'].alignment = WR; wg.row_dimensions[2].height = 40
hdr(wg, 4, ['Grupo', 'Referencia', 'Temas', 'Conocer (temas-eq.)', 'Aplicar (temas-eq.)', 'Razonar (temas-eq.)',
            '% Conocer', '% Aplicar', '% Razonar', 'Meta Conocer', 'Meta Aplicar', 'Meta Razonar',
            'Brecha Razonar (pp)', 'Lectura'], [13, 26, 8, 11, 11, 11, 9, 9, 9, 9, 9, 9, 11, 46], height=45)
COG = [(g, m) for g, m in GROUPS] + [('Química 10°', 'TA*'), ('Química 11°', 'TA*'), ('Biología 10°', 'TA*'), ('Biología 11°', 'TA*')]
META = {'T4': (0.40, 0.40, 0.20), 'T8': (0.35, 0.35, 0.30), 'TA': (0.30, 0.40, 0.30), 'TA*': (0.30, 0.40, 0.30)}
REFN = {'T4': 'TIMSS 2023 4.°', 'T8': 'TIMSS 2023 8.°', 'TA': 'TIMSS Advanced 2015 Física', 'TA*': 'TIMSS Advanced Física (ref. aprox.)'}
r = 5
for g, m in COG:
    wg[f'A{r}'] = g; wg[f'B{r}'] = REFN[m]
    wg[f'C{r}'] = f'=COUNTIFS({TC("V")},$A{r})'
    wg[f'D{r}'] = f'=SUMIFS({TC("S")},{TC("V")},$A{r})'
    wg[f'E{r}'] = f'=SUMIFS({TC("T")},{TC("V")},$A{r})'
    wg[f'F{r}'] = f'=SUMIFS({TC("U")},{TC("V")},$A{r})'
    for a, b in zip('GHI', 'DEF'):
        wg[f'{a}{r}'] = f'=IF($C{r}=0,0,{b}{r}/$C{r})'
    for a, v in zip('JKL', META[m]):
        wg[f'{a}{r}'] = v
    wg[f'M{r}'] = f'=(I{r}-L{r})*100'
    wg[f'N{r}'] = (f'=IF(ABS(M{r})<=5,"Razonar en línea con la meta",IF(M{r}>0,"Más Razonar que la meta (revisar si la etiqueta está inflada)",'
                   f'"Menos Razonar que la meta: faltan tareas de análisis y argumentación"))')
    style_range(wg, r, r, 1, 14)
    for a in 'DEF': wg[f'{a}{r}'].number_format = '0.0'
    for a in 'GHI': wg[f'{a}{r}'].number_format = PCT
    for a in 'JKL': wg[f'{a}{r}'].number_format = '0%'; wg[f'{a}{r}'].font = fIn
    wg[f'M{r}'].number_format = PP
    r += 1
CGL = r - 1

# ------------------------------------------------------------------ Resumen
ws = wb.create_sheet('Resumen', 0)
ws['A1'] = 'Cobertura de la malla de Ciencias de El Salvador frente a TIMSS — resumen por grado'; ws['A1'].font = fT
ws['A2'] = ('Todo es fórmula sobre Temas_TIMSS. Balance = reparto de temas por dominio vs. metas TIMSS. Amplitud = objetivos TIMSS que aparecen. '
            'Profundidad cognitiva = Conocer/Aplicar/Razonar. Ver «Leeme» para el proceso completo.')
ws['A2'].font = fNote; ws.merge_cells('A2:M2'); ws['A2'].alignment = WR; ws.row_dimensions[2].height = 30
hdr(ws, 4, ['Grupo', 'Marco de referencia', 'Temas totales', 'Temas fuera del marco', '% fuera del marco',
            'Dominio con mayor FALTANTE', 'Faltan (temas)', 'Dominio más SOBRErrepresentado', 'Sobran (temas)',
            'Temas a agregar para cumplir todas las metas sin quitar', 'Dominios en meta (±tolerancia)', '% Razonar actual', '% Razonar meta'],
    [12, 24, 9, 10, 10, 24, 9, 24, 9, 14, 11, 10, 10], height=60)
r = 5
for g, m, r0, r1, rt in BG:
    gi = [x for x in range(5, CGL + 1) if wg[f'A{x}'].value == g][0]
    ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]
    ws[f'C{r}'] = f'=COUNTIFS({TC("V")},$A{r})'
    ws[f'D{r}'] = f'=Balance_grado!N{r0}'
    ws[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
    I = f'Balance_grado!$I${r0}:$I${r1}'; Cc = f'Balance_grado!$C${r0}:$C${r1}'; K = f'Balance_grado!$K${r0}:$K${r1}'
    ws[f'F{r}'] = f'=INDEX({Cc},MATCH(MAX({I}),{I},0))'
    ws[f'G{r}'] = f'=ROUND(MAX({I}),0)'
    ws[f'H{r}'] = f'=INDEX({Cc},MATCH(MIN({I}),{I},0))'
    ws[f'I{r}'] = f'=ROUND(-MIN({I}),0)'
    ws[f'J{r}'] = f'=Balance_grado!M{rt}'
    ws[f'K{r}'] = f'=COUNTIF({K},"En meta")&" de "&ROWS({K})'
    ws[f'L{r}'] = f'=Cognitivo_grado!I{gi}'
    ws[f'M{r}'] = f'=Cognitivo_grado!L{gi}'
    style_range(ws, r, r, 1, 13)
    for a in 'DJLM': ws[f'{a}{r}'].font = fLink
    ws[f'E{r}'].number_format = PCT; ws[f'L{r}'].number_format = PCT; ws[f'M{r}'].number_format = '0%'
    r += 1
r += 1
ws[f'A{r}'] = 'Por ciclo (comparación acumulada)'; ws[f'A{r}'].font = fS; r += 1
hdr(ws, r, ['Ciclo', 'Marco de referencia', 'Temas evaluables', 'Dominio con mayor FALTANTE', 'Faltan (temas)',
            'Dominio más SOBRErrepresentado', 'Sobran (temas)', 'Temas a agregar sin quitar', 'Dominios en meta',
            'Objetivos TIMSS del marco', 'Objetivos cubiertos', '% amplitud', 'Objetivos no cubiertos'], None, height=45)
r += 1
ampmap = {m: rr for (m, d, rr) in AMP if d == '(todos)'}
for g, m, r0, r1, rt in BC:
    I = f'Balance_ciclo!$I${r0}:$I${r1}'; Cc = f'Balance_ciclo!$C${r0}:$C${r1}'; K = f'Balance_ciclo!$K${r0}:$K${r1}'
    ws[f'A{r}'] = g; ws[f'B{r}'] = MNAME[m]
    ws[f'C{r}'] = f'=Balance_ciclo!E{r0}'
    ws[f'D{r}'] = f'=INDEX({Cc},MATCH(MAX({I}),{I},0))'; ws[f'E{r}'] = f'=ROUND(MAX({I}),0)'
    ws[f'F{r}'] = f'=INDEX({Cc},MATCH(MIN({I}),{I},0))'; ws[f'G{r}'] = f'=ROUND(-MIN({I}),0)'
    ws[f'H{r}'] = f'=Balance_ciclo!M{rt}'
    ws[f'I{r}'] = f'=COUNTIF({K},"En meta")&" de "&ROWS({K})'
    if g == '9°':
        for a in 'JKLM': ws[f'{a}{r}'] = '—'
    else:
        ar = ampmap[m]
        ws[f'J{r}'] = f'=Cobertura_objetivos!C{ar}'; ws[f'K{r}'] = f'=Cobertura_objetivos!D{ar}'
        ws[f'L{r}'] = f'=Cobertura_objetivos!E{ar}'; ws[f'M{r}'] = f'=Cobertura_objetivos!F{ar}'
        ws[f'L{r}'].number_format = PCT
    style_range(ws, r, r, 1, 13)
    ws[f'L{r}'].number_format = PCT
    r += 1
ws.freeze_panes = 'B5'

# ------------------------------------------------------------------ Fuentes
wsf = wb.create_sheet('Fuentes')
wsf['A1'] = 'Fuentes y versiones usadas'; wsf['A1'].font = fT
hdr(wsf, 3, ['Archivo (carpeta «Mallas sugeridas», Drive)', 'ID Drive', 'Última modificación', 'Temas extraídos', 'Nota'], [48, 46, 20, 10, 60])
files = [
 ('Matriz_Ciencias_2do_Grado_U1_U6_3.xlsx', '1ivJe3XBmtbIi_Pgjc6Z2Xc4pn1tL0Rw3', '', ''),
 ('Matriz_Ciencias_3er_Grado_U1_U6 (Hoja de cálculo de Google)', '1de6X-m-0tCrVH5g6FSYx5rYdWtCYM9UurTG5mdAWXS0', '2026-09-28', 'Se usó la versión Google Sheets (más reciente que el .xlsx).'),
 ('Matriz_Ciencias_4to_Grado_U1_U6.xlsx', '1ijnmjjWjW1YwhtiHeqh4hX17UX_OqB1y', '2026-09-28', ''),
 ('Matriz_Ciencias_5to_Grado_U1_U6.xlsx', '1V9G2VP2Vbe-GCCYYQAZdGBYH6MPLYJY6', '', ''),
 ('Matriz_Ciencias_6to_Grado_U1_U6.xlsx', '11jnnAGb8NZvI0jQuFxXVvX73dCVR7DKC', '', ''),
 ('Matriz_Ciencias_7mo_Grado_U1_U6.xlsx', '1OR7NdtVVqI8P8NFaHgzIE2kpHCu-621X', '', ''),
 ('Matriz_Ciencias_8vo_Grado_U1_U7.xlsx', '1trSwU3iYDitze9gH-VuOsXcWm9Hhs095', '2026-09-28', ''),
 ('Matriz_Ciencias_9no_Grado_U1_U7.xlsx', '1or_yQdA1e1iLSx-x8XDs1vNSZC_-HM1D', '', ''),
 ('Matriz_Bachillerato_Fisica_10_11.xlsx', '1SMhkw4GQloLiBSHTS3rvnQoroZmHTh38', '2026-09-30', ''),
 ('Matriz_Bachillerato_Quimica_10_11.xlsx', '1PGQKwpI0KnyUQicq4HQY4ghVZSCF4d6V', '', 'Solo se usa para el análisis cognitivo (sin marco TIMSS de contenido).'),
 ('Matriz_Bachillerato_Biologia_10_11.xlsx', '1f3nWhHBpQS1vajP1PZqNGcKZEYLh7oF2', '', 'Solo se usa para el análisis cognitivo (sin marco TIMSS de contenido).'),
]
fname = {'Matriz_Ciencias_3er_Grado_U1_U6 (Hoja de cálculo de Google)': 'Matriz_Ciencias_3er_Grado_U1_U6.xlsx'}
r = 4
for f, i, d, n in files:
    wsf[f'A{r}'] = f; wsf[f'B{r}'] = i; wsf[f'C{r}'] = d or 'descargado 2026-10-01'
    wsf[f'D{r}'] = f'=COUNTIFS({TC("D")},"{fname.get(f, f)}")'
    wsf[f'E{r}'] = n
    style_range(wsf, r, r, 1, 5, align=WR); r += 1
wsf[f'C{r}'] = 'Total'; wsf[f'D{r}'] = f'=SUM(D4:D{r-1})'; style_range(wsf, r, r, 3, 4, fill=SUB, font=fB)
r += 2
wsf[f'A{r}'] = 'Marcos de referencia'; wsf[f'A{r}'].font = fS; r += 1
for s in ['IEA (2021). TIMSS 2023 Assessment Frameworks, cap. 2 Science. https://timss.bc.edu/timss2023/frameworks/pdf/T23_Frameworks_Ch2_Science.pdf',
          'IEA (2014). TIMSS Advanced 2015 Assessment Frameworks, cap. 2 Physics. https://elvis.bc.edu/timss2015-advanced/downloads/TA15_FW_Chap2.pdf',
          'Mallas: columna «Habilidad TIMSS» (dominios Conocer/Aplicar/Razonar) elaborada por el equipo de Ciencias (MINED).']:
    wsf[f'A{r}'] = s; wsf[f'A{r}'].font = fN; wsf.merge_cells(f'A{r}:E{r}'); r += 1

# ------------------------------------------------------------------ Leeme
wl = wb.create_sheet('Leeme', 0)
wl.column_dimensions['A'].width = 4; wl.column_dimensions['B'].width = 120
lines = [
 ('T', 'Cobertura TIMSS de la malla de Ciencia y Tecnología — cómo se construyó y cómo leerla'),
 ('N', 'Versión 1 · 1 de octubre de 2026 · Elaborado con Claude (Cowork) a partir de las mallas compartidas por el equipo de Ciencias (MINED).'),
 ('S', 'Qué responde este libro'),
 ('B', '1. Balance: ¿el reparto de temas entre Vida / Física / Química / Tierra se parece al peso que TIMSS da a cada dominio? ¿Cuántos temas faltan o sobran para llegar a la meta?'),
 ('B', '2. Amplitud: ¿qué objetivos de contenido de TIMSS aparecen en la malla, en qué grado, con cuántos temas, y cuáles no aparecen?'),
 ('B', '3. Profundidad cognitiva: ¿la proporción Conocer / Aplicar / Razonar se parece a la de TIMSS?'),
 ('S', 'Proceso (de dónde salen los números)'),
 ('B', 'Paso 1 · Extracción. Se leyeron las 11 mallas (2.° a 9.° y Física, Química, Biología de 10.° y 11.°). Cada fila con «Procedimental» es un TEMA: 1.260 en total. '
       'Se conservan archivo, hoja y fila de origen (hoja Temas_TIMSS, columnas D-F) para poder volver a la malla.'),
 ('B', 'Paso 2 · Catálogo de referencia. Se armó la lista de objetivos de contenido de TIMSS 2023 (4.° grado: 32 objetivos; 8.° grado: 46) y de TIMSS Advanced 2015 Física (23), '
       'con su dominio y el % que TIMSS asigna a cada dominio (hoja Catalogo_TIMSS, en azul = valores tomados del marco).'),
 ('B', 'Paso 3 · Clasificación. Cada tema de 2.° a 9.° y de Física 10.°–11.° (889 temas) se asignó a UN objetivo principal y, si aplica, uno secundario, leyendo procedimental, contenido, '
       'indicador y evidencia. Lo hicieron agentes de IA con instrucciones fijas, tema por tema, dejando confianza (alta/media/baja) y una justificación corta. '
       'Los temas sin contenido TIMSS (medición general, método científico, tecnología/diseño, gestión de riesgos, electrónica digital, mecánica de fluidos en Bachillerato, etc.) se marcan «FUERA».'),
 ('B', 'Paso 4 · Reglas de grado. 2.°–4.° se comparan con TIMSS 4.°; 5.°–9.° con TIMSS 8.° (9.° solo como referencia, porque TIMSS 8.° evalúa hasta 8.°); Física 10.°–11.° con TIMSS Advanced. '
       'Química y Biología de Bachillerato no tienen marco TIMSS de contenido: solo entran en el análisis cognitivo.'),
 ('B', 'Paso 5 · Cálculo (todo con fórmulas). % actual = temas del dominio ÷ temas evaluables (sin FUERA). Temas esperados = meta × evaluables. '
       'Brecha = esperados − actuales (+ faltan, − sobran). «Temas a agregar sin quitar»: se busca el menor total que permite cumplir todas las metas sumando temas '
       '(lo fija el dominio más sobrerrepresentado: total = máx(actual ÷ meta)), y a cada dominio se le suma meta × total − actual, redondeado hacia arriba.'),
 ('B', 'Paso 6 · Amplitud. Un objetivo está «cubierto» si al menos un tema lo trabaja como principal o secundario. Profundidad: 0 no cubierto, 1 débil, 2–3 básica, 4+ sólida.'),
 ('B', 'Paso 7 · Cognitivo. Se usa la columna «Habilidad TIMSS» de la malla. Si un tema nombra varios dominios se reparte en partes iguales. Se compara con TIMSS (4.°: 40/40/20; 8.°: 35/35/30; Advanced Física: 30/40/30).'),
 ('S', 'Hojas'),
 ('B', 'Resumen: una fila por grado y por ciclo con el dominio que más falta, el que más sobra, temas a agregar, amplitud y % Razonar.'),
 ('B', 'Balance_grado y Balance_ciclo: detalle por dominio. Celda H2 = tolerancia editable para «En meta» (5 pp).'),
 ('B', 'Cobertura_objetivos: matriz objetivo × grado (n.º de temas). Al final, resumen de amplitud por dominio.'),
 ('B', 'Faltantes: objetivos no cubiertos o débiles, con una sugerencia de dónde podrían entrar en la malla actual.'),
 ('B', 'Cognitivo_grado: Conocer / Aplicar / Razonar por grado y asignatura.'),
 ('B', 'Temas_TIMSS: trazabilidad completa. Si cambias un objetivo en las columnas K o L, todo el libro se recalcula.'),
 ('B', 'Catalogo_TIMSS y Fuentes: referencias y versiones de archivos.'),
 ('S', 'Límites que hay que decir en la reunión'),
 ('B', '• La clasificación es asistida por IA: 406 temas con confianza alta, 364 media y 119 baja (en rojo en Temas_TIMSS). Debe validarla el equipo de Ciencias, empezando por las de confianza baja.'),
 ('B', '• Se cuentan TEMAS, no horas. Un tema puede durar una o varias sesiones de 55 min; cuando exista la distribución de horas conviene ponderar por horas.'),
 ('B', '• Los % de TIMSS son pesos de la prueba, no un mandato curricular. Sirven como referencia de equilibrio, no como regla obligatoria.'),
 ('B', '• Comparar un solo grado con TIMSS es una aproximación (TIMSS mide lo acumulado). La comparación más fiel está en Balance_ciclo.'),
 ('B', '• Los objetivos TIMSS están parafraseados; para decisiones formales, consultar el marco oficial (hoja Fuentes).'),
 ('B', '• 1.° grado no tiene malla, así que no entra en el ciclo 2.°–4.°.'),
]
r = 1
for kind, text in lines:
    if kind == 'S':
        r += 1  # línea en blanco antes de cada sección
    c = wl.cell(row=r, column=2, value=text)
    c.font = {'T': fT, 'N': fNote, 'S': fS, 'B': fN}[kind]
    c.alignment = WR
    r += 1
for rr in range(1, r):
    v = wl.cell(row=rr, column=2).value
    if v and len(v) > 120:
        wl.row_dimensions[rr].height = 14 * (len(v) // 115 + 1)

wb._sheets = [wb['Leeme'], wb['Resumen'], wb['Balance_ciclo'], wb['Balance_grado'], wb['Cobertura_objetivos'], wb['Faltantes'],
              wb['Cognitivo_grado'], wb['Temas_TIMSS'], wb['Catalogo_TIMSS'], wb['Fuentes']]
for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.active = 0
out = 'outputs/Cobertura_TIMSS_Ciencias_SV.xlsx'
import os; os.makedirs('outputs', exist_ok=True)
wb.save(out)
print('ok', out, 'temas', TL - 4, 'objetivos', OL - 4, 'faltantes', FL - 4)
