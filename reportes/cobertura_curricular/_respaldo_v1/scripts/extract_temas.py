"""Paso 1 · Extrae cada tema (fila con «Procedimental») de las mallas en data/raw/ → data/interim/temas.json

Uso:  python scripts/extract_temas.py
Reglas:
- Hojas de unidad: las que empiezan con «Unidad» (2.°–9.°) o contienen «N.°» (Bachillerato: «Física 10.°»...).
- Encabezados en filas 1 y 2 (la fila 2 trae las 5 fases de la secuencia). Datos desde la fila 3.
- Unidad y Contenido vienen en celdas combinadas: se arrastran hacia abajo (forward-fill).
- Nunca se modifica data/raw/.
"""
import glob, json, os, re
import openpyxl

RAW, OUT = 'data/raw', 'data/interim/temas.json'
ASIG = {'Fisica': 'Física', 'Quimica': 'Química', 'Biologia': 'Biología'}


def clean(v):
    return re.sub(r'\s+', ' ', str(v or '')).strip()


def main():
    temas = []
    for path in sorted(glob.glob(os.path.join(RAW, 'Matriz_*.xlsx'))):
        f = os.path.basename(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if not (ws.title.startswith('Unidad') or re.search(r'\d+\.°', ws.title)):
                continue
            h1 = [c.value or '' for c in ws[1]]
            h2 = [c.value or '' for c in ws[2]]
            cols = [(str(a) or str(b)).split('\n')[0].strip() for a, b in zip(h1, h2)]
            for i, s in enumerate(h2):
                if s:
                    cols[i] = str(s).split('\n')[0].strip()
            unidad = contenido = None
            for fila, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                d = dict(zip(cols, row))
                if d.get('Unidad'):
                    unidad = clean(d['Unidad'])
                if d.get('Contenido'):
                    contenido = clean(d['Contenido'])
                if not d.get('Procedimental'):
                    continue
                m = re.search(r'_(\d+)(?:do|er|to|mo|vo|no)_', f)
                grado = int(m.group(1)) if m else int(re.search(r'(\d+)\.°', ws.title).group(1))
                asig = 'Ciencias' if m else ASIG[f.split('_')[2]]
                temas.append(dict(
                    archivo=f, hoja=ws.title, fila=fila, grado=grado, asignatura=asig,
                    unidad=unidad, contenido=contenido,
                    subcontenido=clean(d.get('Subcontenidos')), procedimental=clean(d['Procedimental']),
                    indicador=clean(d.get('Indicadores de logro')), evidencia=clean(d.get('Evidencia de aprendizaje')),
                    habilidad_timss=clean(d.get('Habilidad TIMSS')), pisa=clean(d.get('Competencia PISA 2025'))[:1]))
    for i, t in enumerate(temas, 1):
        t['id'] = f"T{i:04d}-G{t['grado']:02d}-{t['asignatura'][:3].upper()}"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(temas, open(OUT, 'w'), ensure_ascii=False, indent=0)
    print(f'{len(temas)} temas → {OUT}')


if __name__ == '__main__':
    main()
