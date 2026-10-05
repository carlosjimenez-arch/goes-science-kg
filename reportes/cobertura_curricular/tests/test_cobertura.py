import json, os
import openpyxl
import pytest

CAT = {c['codigo'] for c in json.load(open('data/referencia/catalogo_timss.json'))}


def test_catalogo_completo():
    cat = json.load(open('data/referencia/catalogo_timss.json'))
    assert sum(c['marco'] == 'T4' for c in cat) == 32
    assert sum(c['marco'] == 'T8' for c in cat) == 46
    assert sum(c['marco'] == 'TA' for c in cat) == 23


def test_clasificacion_valida():
    temas = json.load(open('data/interim/temas.json'))
    clas = json.load(open('data/interim/clasificacion.json'))
    for t in temas:
        if t['asignatura'] in ('Ciencias', 'Física'):
            assert t['id'] in clas, f"sin clasificar: {t['id']}"
    for k, c in clas.items():
        assert c['obj1'] == 'FUERA' or c['obj1'] in CAT, k
        assert c['obj2'] in ('', None) or c['obj2'] in CAT, k
        assert c['obj1'] != c['obj2'], k
        assert c['confianza'] in ('alta', 'media', 'baja'), k


def test_marco_corresponde_al_grado():
    temas = {t['id']: t for t in json.load(open('data/interim/temas.json'))}
    for k, c in json.load(open('data/interim/clasificacion.json')).items():
        t = temas[k]
        if c['obj1'] == 'FUERA':
            continue
        pref = 'T4' if (t['asignatura'] == 'Ciencias' and t['grado'] <= 4) else 'T8' if t['asignatura'] == 'Ciencias' else 'TA'
        assert c['obj1'].startswith(pref), (k, c['obj1'])


@pytest.mark.skipif(not os.path.exists('outputs/Cobertura_TIMSS_Ciencias_SV.xlsx'), reason='libro no generado')
def test_libro_sin_errores_y_recalculado():
    wb = openpyxl.load_workbook('outputs/Cobertura_TIMSS_Ciencias_SV.xlsx', data_only=True)
    ws = wb['Resumen']
    assert ws['C5'].value is not None, 'el libro no fue recalculado (abrir con LibreOffice/recalc)'
    for s in wb.worksheets:
        for row in s.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, str) and v.startswith('#')), f'error de fórmula en {s.title}: {v}'
