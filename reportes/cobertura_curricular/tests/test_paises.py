"""Pruebas del Prompt 02: objetivos de otros países, su alineación con TIMSS 2027 y la hoja Comparacion_paises."""
import json, os, re
import openpyxl
from conftest import hoja, tiene
import pytest

DIR = 'data/interim/paises'
ALIN = f'{DIR}/alineacion.json'
LIBRO27 = 'outputs/Cobertura_TIMSS2027_Ciencias_SV.xlsx'
EXPORT27 = 'outputs/Cobertura_TIMSS2027_Ciencias_SV_numbers.xlsx'
CAT27 = {c['codigo']: c['marco'] for c in json.load(open('data/referencia/catalogo_timss2027.json'))}
PAISES = [p for p in ('uruguay', 'colombia', 'singapur') if os.path.exists(f'{DIR}/{p}.json')]
pytestmark = pytest.mark.skipif(not PAISES, reason='objetivos de países aún no extraídos')


@pytest.mark.parametrize('pais', PAISES)
def test_objetivos_extraidos(pais):
    objs = json.load(open(f'{DIR}/{pais}.json'))
    assert len({o['id'] for o in objs}) == len(objs)
    for o in objs:
        for k in ('id', 'pais', 'documento', 'pagina', 'grado_min', 'grado_max', 'grado_o_tramo', 'texto'):
            assert o.get(k) not in (None, ''), (o['id'], k)
        assert os.path.exists(f"data/referencia/externos/paises/{pais}/{o['documento']}"), o['id']
        assert 1 <= o['grado_min'] <= o['grado_max'] <= 12, o['id']
        assert len(o['texto'].split()) <= 30, o['id']


@pytest.mark.skipif(not os.path.exists(ALIN), reason='alineación aún no unida')
def test_alineacion_valida():
    al = json.load(open(ALIN))
    assert len({o['id'] for o in al}) == len(al)
    for o in al:
        m = 'T4_27' if o['grado_min'] <= 4 else 'T8_27'
        assert o['marco'] == m and o['grado_min'] <= 9, o['id']
        assert o['obj1'] == 'FUERA' or CAT27.get(o['obj1']) == m, o['id']
        assert o['obj2'] == '' or (CAT27.get(o['obj2']) == m and o['obj2'] != o['obj1']), o['id']
        assert o['confianza'] in ('alta', 'media', 'baja'), o['id']
        assert len(o['justificacion'].split()) <= 15, o['id']
    todos = [o for p in PAISES for o in json.load(open(f'{DIR}/{p}.json')) if o['grado_min'] <= 9]
    assert {o['id'] for o in todos} == {o['id'] for o in al}


@pytest.mark.skipif(not (os.path.exists(ALIN) and os.path.exists(LIBRO27)), reason='libro 2027 sin comparación de países')
def test_hoja_comparacion_es_formula():
    wb = openpyxl.load_workbook(LIBRO27)
    assert tiene(wb, 'Comparacion_paises') and tiene(wb, 'Paises_objetivos')
    ws = hoja(wb, 'Comparacion_paises')
    assert ws['A5'].value in CAT27
    for cc in ('E5', 'F5', 'I5', 'J5', 'K5'):
        assert str(ws[cc].value).startswith('='), cc
    assert not any(re.search(r'\b(MINIFS|MAXIFS|XLOOKUP|FILTER)\(', str(c.value))
                   for row in ws.iter_rows() for c in row if isinstance(c.value, str))


@pytest.mark.skipif(not (os.path.exists(ALIN) and os.path.exists(EXPORT27)), reason='libro 2027 aún no recalculado')
def test_comparacion_recalculada_coincide_con_python():
    """El primer grado de El Salvador que calcula la hoja coincide con un conteo directo en Python."""
    wb = openpyxl.load_workbook(EXPORT27, data_only=True)
    if not tiene(wb, 'Comparacion_paises'):
        pytest.skip('export anterior a la comparación de países')
    temas = {t['id']: t for t in json.load(open('data/interim/temas.json'))}
    clas = json.load(open('data/interim/clasificacion_2027.json'))
    primero = {}
    for i, c in clas.items():
        t = temas[i]
        if t['asignatura'] != 'Ciencias':
            continue
        for o in (c['obj1'], c['obj2']):
            if o in CAT27:
                primero[o] = min(primero.get(o, 99), max(2, t['grado']))
    ws = hoja(wb, 'Comparacion_paises')
    for r in range(5, 5 + len(CAT27)):
        cod, sv = ws[f'A{r}'].value, ws[f'E{r}'].value
        assert (sv if sv != '—' else None) == primero.get(cod), (cod, sv, primero.get(cod))
