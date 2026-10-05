"""Pruebas del marco TIMSS 2027: catálogo, equivalencias, clasificación v2 y libro 2027."""
import csv, json, os, re, subprocess, sys
from collections import defaultdict
import openpyxl
from conftest import hoja, tiene
import pytest

CAT27 = json.load(open('data/referencia/catalogo_timss2027.json'))
COD27 = {c['codigo'] for c in CAT27}
COD23 = {c['codigo'] for c in json.load(open('data/referencia/catalogo_timss.json'))}
TEMAS = json.load(open('data/interim/temas.json'))
V2 = json.load(open('data/referencia/clasificacion_v2_timss2027.json'))
LIBRO27 = 'outputs/Cobertura_TIMSS2027_Ciencias_SV.xlsx'
LIBRO_V1 = 'outputs/Cobertura_TIMSS_Ciencias_SV.xlsx'


# ------------------------------------------------------------------ catálogo 2027
def test_catalogo_conteos():
    assert sum(c['marco'] == 'T4_27' for c in CAT27) == 32
    assert sum(c['marco'] == 'T8_27' for c in CAT27) == 48


def test_codigos_unicos_y_formato():
    assert len(COD27) == len(CAT27)
    for c in CAT27:
        assert re.fullmatch(r'T[48]_27-[BPCE]\d\.\d', c['codigo']), c['codigo']
        assert c['codigo'].startswith(c['area_codigo'] + '.'), c['codigo']


def test_metas_de_dominio():
    metas = {(c['marco'], c['dominio']): c['meta_dominio'] for c in CAT27}
    assert metas == {('T4_27', 'Biología'): 0.45, ('T4_27', 'Ciencias físicas'): 0.35, ('T4_27', 'Ciencias de la Tierra'): 0.20,
                     ('T8_27', 'Biología'): 0.35, ('T8_27', 'Química'): 0.20, ('T8_27', 'Física'): 0.25,
                     ('T8_27', 'Ciencias de la Tierra'): 0.20}
    for m in ('T4_27', 'T8_27'):
        assert abs(sum(v for (mm, _), v in metas.items() if mm == m) - 1) < 1e-9


def test_areas_por_dominio_e_investigaciones():
    areas = defaultdict(set)
    for c in CAT27:
        areas[(c['marco'], c['dominio'])].add(c['area_codigo'])
    assert {k: len(v) for k, v in areas.items()} == {
        ('T4_27', 'Biología'): 6, ('T4_27', 'Ciencias físicas'): 6, ('T4_27', 'Ciencias de la Tierra'): 5,
        ('T8_27', 'Biología'): 7, ('T8_27', 'Química'): 4, ('T8_27', 'Física'): 6, ('T8_27', 'Ciencias de la Tierra'): 5}
    inv = defaultdict(int)
    for c in CAT27:
        inv[(c['marco'], c['dominio'])] += c['es_investigacion']
    assert all(n == 1 for n in inv.values()), inv  # un objetivo de Investigaciones por dominio


def test_paginas_y_subitems_presentes():
    for c in CAT27:
        assert isinstance(c['pagina_fuente'], int) and 26 <= c['pagina_fuente'] <= 45, c['codigo']
        assert c['pagina_pdf'] == c['pagina_fuente'] + 2
        assert c['subitems'] and all(s['texto'] for s in c['subitems']), c['codigo']
        assert c['ambiental'] == any(s['ambiental'] for s in c['subitems']), c['codigo']


def test_ambiental_solo_biologia_y_tierra():
    for c in CAT27:
        if c['ambiental']:
            assert c['dominio'] in ('Biología', 'Ciencias de la Tierra'), c['codigo']
    assert sum(c['ambiental'] for c in CAT27 if c['marco'] == 'T4_27') == 9
    assert sum(c['ambiental'] for c in CAT27 if c['marco'] == 'T8_27') == 15


def test_catalogo_valida_contra_pdf():
    """El script compara títulos, páginas, asteriscos y metas con el PDF oficial."""
    if not os.path.exists('data/referencia/externos/timss/TIMSS2027_Marcos_Evaluacion.pdf'):
        pytest.skip('PDF no disponible')
    r = subprocess.run([sys.executable, 'scripts/build_catalogo_timss2027.py', '--check'], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr


# ------------------------------------------------------------------ equivalencias
def test_equivalencias_codigos_validos():
    filas = list(csv.DictReader(open('data/referencia/equivalencias_timss_2023_2027.csv', encoding='utf-8')))
    tipos = {'igual', 'dividido', 'fusionado', 'sin_equivalente', 'nuevo_2027'}
    for f in filas:
        assert f['tipo'] in tipos
        assert f['codigo_2023'] == '' or f['codigo_2023'] in COD23
        assert f['codigo_2027'] == '' or f['codigo_2027'] in COD27
    assert {f['codigo_2023'] for f in filas if f['codigo_2023']} == {c for c in COD23 if not c.startswith('TA')}


# ------------------------------------------------------------------ clasificación v2
def test_cada_tema_de_ciencias_y_fisica_tiene_clasificacion_2027():
    idx = {(r['archivo'], r['hoja'], r['fila']): r for r in V2}
    for t in TEMAS:
        if t['asignatura'] in ('Ciencias', 'Física'):
            r = idx.get((t['archivo'], t['hoja'], t['fila']))
            assert r is not None, f"sin clasificación 2027: {t['id']}"
            assert r['procedimental'] == t['procedimental'], t['id']
    assert len(V2) == sum(t['asignatura'] in ('Ciencias', 'Física') for t in TEMAS)


def test_codigos_2027_validos_y_del_marco_del_grado():
    temas = {(t['archivo'], t['hoja'], t['fila']): t for t in TEMAS}
    for r in V2:
        t = temas[(r['archivo'], r['hoja'], r['fila'])]
        if t['asignatura'] == 'Física':
            assert r['obj1'] == 'FUERA' or r['obj1'].startswith('TA-'), r  # TIMSS Advanced 2015 sin cambios
            continue
        pref = 'T4_27-' if t['grado'] <= 4 else 'T8_27-'
        assert r['obj1'] == 'FUERA' or (r['obj1'] in COD27 and r['obj1'].startswith(pref)), r
        assert r['obj2'] in ('', None) or (r['obj2'] in COD27 and r['obj2'].startswith(pref)), r
        assert r['obj1'] != r['obj2'], r
        assert r['confianza'] in ('alta', 'media', 'baja'), r
        assert len(r['justificacion'].split()) <= 15, r
        assert r['version'] == 'v2-2027-reclasificado', r


def test_fisica_identica_a_v1():
    v1 = {(r['archivo'], r['hoja'], r['fila']): r for r in json.load(open('data/referencia/clasificacion_v1.json'))}
    temas = {(t['archivo'], t['hoja'], t['fila']): t for t in TEMAS}
    for r in V2:
        k = (r['archivo'], r['hoja'], r['fila'])
        if temas[k]['asignatura'] == 'Física':
            assert r == v1[k], k


def test_revision_baja_completa():
    filas = list(csv.DictReader(open('outputs/revision_confianza_baja_timss2027.csv', encoding='utf-8')))
    bajas = [r for r in V2 if r.get('version') == 'v2-2027-reclasificado' and r['confianza'] == 'baja']
    assert len(filas) == len(bajas)
    assert 'decision_equipo' in filas[0]


# ------------------------------------------------------------------ libros
FORMULAS_PROHIBIDAS = re.compile(r'\b(XLOOKUP|FILTER|UNIQUE|SEQUENCE|LET|LAMBDA|TEXTJOIN|IFS)\(')


@pytest.mark.skipif(not os.path.exists(LIBRO27), reason='libro 2027 no generado')
def test_libro_2027_estructura_y_formulas():
    wb = openpyxl.load_workbook(LIBRO27)
    temas = 'Temas' if tiene(wb, 'Temas') else 'Temas_TIMSS'  # con PISA la hoja se llama «Temas»
    for h in ('Resumen', 'Ambiental', 'Investigaciones', 'Referencia_2023', 'Cobertura_objetivos', temas):
        assert tiene(wb, h), h
    nombres = set(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    assert not FORMULAS_PROHIBIDAS.search(v), (ws.title, c.coordinate)
                    assert '{' not in v and '}' not in v, (ws.title, c.coordinate)  # f-string sin reemplazar
                    for q, n in re.findall(r"'([^']+)'!|([A-Za-z_0-9]+)!", v):
                        assert (q or n) in nombres, (ws.title, c.coordinate, q or n)
    # las cifras de análisis son fórmulas, no números escritos
    for h, celdas in (('Resumen', ['C5', 'F5', 'G5']), ('Ambiental', ['C7', 'D7', 'E7']), ('Investigaciones', ['B6', 'C6'])):
        for cc in celdas:
            assert str(hoja(wb, h)[cc].value).startswith('='), (h, cc)


@pytest.mark.skipif(not os.path.exists(LIBRO27), reason='libro 2027 no generado')
def test_libro_2027_sin_azul():
    import colorsys
    def azul(rgb):
        if not isinstance(rgb, str) or len(rgb) < 6:
            return False
        r, g, b = (int(rgb[-6:][i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        return s > 0.25 and 0.5 <= h <= 0.72 and l < 0.95
    wb = openpyxl.load_workbook(LIBRO27)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert not azul(c.font.color.rgb if c.font and c.font.color else None), (ws.title, c.coordinate)
                assert not azul(c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None), (ws.title, c.coordinate)


@pytest.mark.skipif(not os.path.exists(LIBRO_V1), reason='libro v1 no disponible')
def test_marco_2023_reproduce_libro_v1(tmp_path):
    """--marco 2023 debe dar el mismo contenido (fórmula/constante) que el entregable v1 en todas las hojas."""
    out = tmp_path / 'v1.xlsx'
    r = subprocess.run([sys.executable, 'scripts/build_cobertura.py', '--marco', '2023', '--out', str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    norm = lambda v: v.replace('TRUE()', 'TRUE') if isinstance(v, str) else v  # LibreOffice reescribe TRUE
    a, b = openpyxl.load_workbook(LIBRO_V1), openpyxl.load_workbook(out)
    assert a.sheetnames == b.sheetnames
    for h in a.sheetnames:
        A = {c.coordinate: norm(c.value) for row in a[h].iter_rows() for c in row if c.value is not None}
        B = {c.coordinate: norm(c.value) for row in b[h].iter_rows() for c in row if c.value is not None}
        assert A == B, h


EXPORT27 = 'outputs/Cobertura_TIMSS2027_Ciencias_SV_numbers.xlsx'  # el libro 2027 exportado desde Numbers/Excel (con valores)


@pytest.mark.skipif(not os.path.exists(EXPORT27), reason='libro 2027 aún no recalculado (exportar desde Numbers o Excel)')
def test_libro_2027_recalculado_sin_errores():
    wb = openpyxl.load_workbook(EXPORT27, data_only=True)
    errores = [(s.title, c.coordinate, c.value) for s in wb.worksheets for row in s.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith('#')]
    assert not errores, errores[:10]
