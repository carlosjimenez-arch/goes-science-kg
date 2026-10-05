"""Pruebas de PISA 2025 (7.°–9.°): catálogo, clasificación, microhabilidades, hojas, escenarios y presupuesto."""
import collections as C, json, os, re, subprocess, sys
import openpyxl
from conftest import hoja, tiene
import pytest

CATP = json.load(open('data/referencia/catalogo_pisa2025.json'))
EL = {e['codigo']: e for e in CATP['elementos']}
TEMAS = json.load(open('data/interim/temas.json'))
CLASP = 'data/referencia/clasificacion_pisa2025.json'
PRES = 'data/referencia/presupuesto_candidatos.json'
LIBRO = 'outputs/Cobertura_TIMSS2027_Ciencias_SV.xlsx'
EXPORT = 'outputs/Cobertura_TIMSS2027_Ciencias_SV_numbers.xlsx'


def test_catalogo_conteos_y_paginas():
    n = C.Counter(e['tipo'] for e in CATP['elementos'])
    assert n == {'contenido': 21, 'procedimental': 8, 'epistemico': 20, 'subcompetencia': 15, 'area_aplicacion': 5, 'demanda': 3}
    assert len(EL) == len(CATP['elementos'])
    assert all(32 <= e['pagina'] <= 53 for e in CATP['elementos'])
    assert abs(sum(v[0] for v in CATP['metas']['sistema'].values()) - 1) < 1e-9
    assert sorted(c for c, e in EL.items() if e['tipo'] == 'subcompetencia') == sorted(
        [f'E{i}' for i in range(1, 7)] + [f'D{i}' for i in range(1, 5)] + [f'I{i}' for i in range(1, 6)])


def test_catalogo_valida_contra_pdf():
    if not os.path.exists('data/referencia/externos/pisa/PISA2025_Marco_Evaluacion_2026.pdf'):
        pytest.skip('PDF no disponible')
    r = subprocess.run([sys.executable, 'scripts/build_catalogo_pisa2025.py', '--check'], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr


def test_microhabilidades_de_la_malla():
    for t in TEMAS:
        if t['asignatura'] == 'Ciencias':
            codigos = [c for f in t['micro_pisa'].values() for c in f]
            assert len(t['micro_pisa']) == 5 and codigos, t['id']
            assert all(EL.get(c, {}).get('tipo') == 'subcompetencia' for c in codigos), t['id']
            assert t['competencia_pisa'][:1] in '123', t['id']


@pytest.mark.skipif(not os.path.exists(CLASP), reason='clasificación PISA no generada')
def test_clasificacion_pisa_valida():
    cl = json.load(open(CLASP))
    temas = {(t['archivo'], t['hoja'], t['fila']): t for t in TEMAS}
    esperados = {k for k, t in temas.items() if t['asignatura'] == 'Ciencias' and t['grado'] in (7, 8, 9)}
    assert {(r['archivo'], r['hoja'], r['fila']) for r in cl} == esperados and len(cl) == len(esperados)
    for r in cl:
        assert r['procedimental'] == temas[(r['archivo'], r['hoja'], r['fila'])]['procedimental']
        assert r['cont1'] == 'FUERA' or EL[r['cont1']]['tipo'] == 'contenido', r
        assert r['cont2'] == '' or (EL[r['cont2']]['tipo'] == 'contenido' and r['cont2'] != r['cont1']), r
        assert r['conocimiento'] in ('Contenido', 'Procedimental', 'Epistémico')
        assert all(EL[x]['tipo'] == 'procedimental' for x in r['proc']) and all(EL[x]['tipo'] == 'epistemico' for x in r['epis'])
        assert r['contexto'] in ('Personal', 'Local/nacional', 'Global', 'Sin contexto')
        assert r['demanda'] in ('Baja', 'Media', 'Alta') and r['confianza'] in ('alta', 'media', 'baja')
        assert len(r['justificacion'].split()) <= 15


@pytest.mark.skipif(not os.path.exists(PRES), reason='presupuesto no generado')
def test_presupuesto_candidatos():
    cand = json.load(open(PRES))
    ids = {t['id']: t for t in TEMAS}
    assert len({c['id'] for c in cand}) == len(cand)
    usados = [t for c in cand if c['accion'] == 'fusionar' for t in c['temas']]
    for c in cand:
        assert c['accion'] in ('insertar', 'fusionar', 'reubicar')
        if c['accion'] == 'fusionar':
            assert len(c['temas']) >= 2 and len({ids[t]['grado'] for t in c['temas']}) == 1, c['id']
        if c['accion'] == 'insertar':
            assert c['motivo'] in ('balance', 'cobertura') and not c['temas'], c['id']


@pytest.mark.skipif(not os.path.exists(LIBRO), reason='libro no generado')
def test_libro_pisa_estructura():
    wb = openpyxl.load_workbook(LIBRO)
    for h in ('Apego_escenarios', 'Presupuesto', 'PISA_contenido', 'PISA_competencias', 'PISA_conocimiento', 'PISA_contextos', 'Catalogo_PISA', 'Temas'):
        assert tiene(wb, h), h
    assert not tiene(wb, 'Faltantes')  # integrada en Cobertura_objetivos
    assert '9°' not in {c.value for c in hoja(wb, 'Balance_grado')['A']}  # 9.° solo con PISA
    wo = hoja(wb, 'Cobertura_objetivos')
    heads = [c.value for c in wo[4]]
    fila = {wo.cell(r, 1).value: r for r in range(5, wo.max_row + 1)}
    r4 = fila['T4_27-B1.1']
    assert wo.cell(r4, heads.index('5°') + 1).value == 'N/A' and str(wo.cell(r4, heads.index('2°') + 1).value).startswith('=')
    caja = lambda c: str(c.value or '').startswith('="') and c.fill.fgColor.rgb == '00E6EFE9'  # recuadro verde con cifras en fórmula
    assert any(caja(c) for c in hoja(wb, 'Leeme')['B'])
    for s_ in wb.worksheets:  # cada pestaña abre con su recuadro y dice qué marco mide
        assert any(caja(s_[c]) for c in ('A2', 'A3')) or s_ is hoja(wb, 'Leeme'), s_.title


@pytest.mark.skipif(not os.path.exists(EXPORT), reason='libro no recalculado')
def test_recalculado_escenarios_y_presupuesto():
    wb = openpyxl.load_workbook(EXPORT, data_only=True)
    if not tiene(wb, 'Apego_escenarios'):
        pytest.skip('export anterior a PISA')
    wa = hoja(wb, 'Apego_escenarios')
    fila = {}
    for r in range(1, wa.max_row + 1):
        fila.setdefault(wa.cell(r, 1).value, r)  # primera aparición = bloque A (índice de apego)
    idx = [r for r in range(1, wa.max_row + 1) if wa.cell(r, 2).value == 'Índice de apego del escenario']
    for g, r in zip(('7°', '8°'), idx):  # 0 % PISA = índice TIMSS; 100 % PISA = índice PISA
        assert abs(wa.cell(r, 7).value - wa.cell(fila[g], 6).value) < 1e-9
        assert abs(wa.cell(r, 9).value - wa.cell(fila[f'{g} · PISA'], 6).value) < 1e-9
    wp = hoja(wb, 'Presupuesto')
    totales = [r for r in range(1, 40) if str(wp.cell(r, 4).value or '').startswith('Total Ciclo')]
    assert len(totales) == 3
    for r in totales:
        assert abs(wp.cell(r, 11).value) <= 1, wp.cell(r, 4).value  # saldo ≈ 0: se mantiene el total de temas


@pytest.mark.skipif(not (os.path.exists(EXPORT) and os.path.exists(CLASP)), reason='libro no recalculado')
def test_sistemas_pisa_coinciden_con_python():
    wb = openpyxl.load_workbook(EXPORT, data_only=True)
    if not tiene(wb, 'PISA_contenido'):
        pytest.skip('export anterior a PISA')
    cl = json.load(open(CLASP))
    n = C.Counter(EL[r['cont1']]['grupo'] for r in cl if r['cont1'] != 'FUERA')
    ws = hoja(wb, 'PISA_contenido')
    en_hoja = {ws.cell(r, 3).value: ws.cell(r, 4).value for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == 'Ciclo 7°–9°'}
    assert en_hoja == dict(n)
