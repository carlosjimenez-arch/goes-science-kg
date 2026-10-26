"""Verificación independiente de los 1.128 paquetes y de todas las celdas finales."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from goes_science_kg.api_materiales import crear_router_materiales
from goes_science_kg.curriculo.contratos import Curriculo
from goes_science_kg.materiales.almacen import Biblioteca
from goes_science_kg.materiales.compilacion import sha

BASE = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def entrega_final():
    carpeta = BASE / "data/materiales_produccion_v4"
    if not (carpeta / "biblioteca.sqlite").is_file():
        pytest.skip("Restaurar la biblioteca final desde su recibo para ejecutar la prueba real")
    b = Biblioteca(carpeta)
    b.validar()
    registros, archivos = b.exportar()
    return b, registros, archivos


def test_todas_las_celdas_y_fases_corresponden_a_bytes_originales(entrega_final):
    b, registros, archivos = entrega_final
    c = Curriculo.model_validate_json(archivos["curriculo.json"])
    libros = {s.id: load_workbook(BytesIO(archivos[s.archivo_local]), data_only=False) for s in c.fuentes}
    try:
        assert len(c.fuentes) == 11 and len(c.hojas) == 126
        assert len(c.filas) == 1020 and len(c.agrupaciones) == 108
        por_id = {r["id"]: r for r in registros if r["tipo"] == "paquete_objetivo"}
        assert set(por_id) == {f.id for f in c.filas}
        for fila in c.filas:
            hoja = libros[fila.fuente_id][fila.hoja]
            d = por_id[fila.id]["documento"]
            assert d["original"] == fila.model_dump(mode="json")
            assert not d["revisiones"] and d["propuesta"] is None
            assert fila.fases["aproximar"].presente_en_hoja == (
                fila.grado == 9 and fila.hoja == "Unidad 11 - Biosfera y sostenib"
            )
            assert all(fila.fases[f].presente_en_hoja for f in ("explorar", "explicar", "indagar", "reforzar"))
            for celda in fila.celdas:
                assert celda.valor_original == hoja[celda.coordenada].value
                assert celda.valor == hoja[celda.origen_combinado or celda.coordenada].value
                if celda.campo and celda.campo.startswith("fase_"):
                    encabezado = hoja[celda.columna + "2"]
                    for rango in hoja.merged_cells.ranges:
                        if encabezado.coordinate in rango:
                            encabezado = hoja.cell(rango.min_row, rango.min_col)
                            break
                    assert celda.campo.removeprefix("fase_").casefold() in str(encabezado.value).casefold()
            for celda in d["celdas_sin_rol"]:
                assert celda["campo"] is None and celda["valor"] is not None
        unidades = [r for r in registros if r["tipo"] == "paquete_unidad"]
        ids = [f for u in unidades for f in u["documento"]["filas_ids"]]
        assert len(ids) == len(set(ids)) == 1020
        assert set(ids) == set(por_id)
        enriquecida = b.manifest["procedencia"].get("version_diseno") == "preparacion-malla-final-2"
        informe = json.loads(archivos["edicion_base/informe.json" if enriquecida else "informe.json"])
        assert informe["objetivos_identicos"] == 1020
        assert informe["filas_con_cambios_logicos"] >= 884
        assert informe["filas_sin_paquete"] == []
        assert b.manifest["procedencia"]["estado_revision"] == (
            "propuesta_didactica_pendiente_revision" if enriquecida else "pendiente_revision_nueva_edicion"
        )
    finally:
        for libro in libros.values():
            libro.close()


def test_catorce_zip_cubren_todos_los_paquetes_sin_duplicados(entrega_final):
    _, registros, archivos = entrega_final
    ids = []
    indice = json.loads(archivos["paquetes/indice.json"])
    assert len(indice["paquetes"]) == 14
    for e in indice["paquetes"]:
        raw = archivos["paquetes/" + e["archivo"]]
        assert (len(raw), sha(raw)) == (e["bytes"], e["sha256"])
        with ZipFile(BytesIO(raw)) as z:
            manifest = json.loads(z.read("manifest.json"))
            for nombre, esperado in manifest["archivos"].items():
                contenido = z.read(nombre)
                assert (len(contenido), sha(contenido)) == (esperado["bytes"], esperado["sha256"])
                if nombre.startswith(("paquete_objetivo/", "paquete_unidad/")) and nombre.endswith(".json"):
                    d = json.loads(contenido)
                    ids.append(d["id"])
                    assert "pendiente" in d["documento"]["estado_revision"]
    assert len(ids) == len(set(ids)) == len(registros) == 1128


def test_api_final_descarga_todos_los_json_markdown_y_originales(entrega_final):
    b, registros, archivos = entrega_final
    app = FastAPI()
    app.include_router(
        crear_router_materiales(
            BASE, directorio="data/materiales_produccion_v4", prefijo="/api/v4/produccion/materiales"
        )
    )
    c = TestClient(app)
    prefix = "/api/v4/produccion/materiales"
    assert c.get(prefix + "/salud").json()["estado_revision_contenido"] == b.manifest["procedencia"]["estado_revision"]
    assert c.get(prefix + "?tipo=paquete_objetivo").json()["total"] == 1020
    assert c.get(prefix + "?tipo=paquete_unidad").json()["total"] == 108
    for r in registros:
        id_ = r["id"]
        assert c.get(prefix + "/descarga/" + id_).json()["documento"] == r["documento"]
        md = c.get(prefix + "/descarga/" + id_ + "?formato=md")
        assert md.status_code == 200 and "pendiente" in md.text.lower()
    for _, raw in archivos.items():
        respuesta = c.get(prefix + "/archivos/" + sha(raw))
        assert respuesta.status_code == 200 and respuesta.content == raw
    assert b.manifest["materiales"] == 1128
