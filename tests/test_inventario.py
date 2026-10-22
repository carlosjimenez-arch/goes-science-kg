"""La inspección de entregas nuevas es declarativa, trazable y no activa currículos."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from goes_science_kg.ingesta.inventario import inventariar


def _libro(archivo: Path, filas: list[list], nombre: str = "Programa") -> None:
    archivo.parent.mkdir(parents=True, exist_ok=True)
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = nombre
    for fila in filas:
        hoja.append(fila)
    libro.save(archivo)
    libro.close()


def test_inventario_no_muta_fuentes_y_es_determinista(tmp_path):
    """No presupone once archivos, nombres de hojas, columnas ni fila de cabecera."""
    archivo = tmp_path / "ciclo" / "ciencias.xlsx"
    _libro(
        archivo,
        [
            ["Nueva entrega"],
            [],
            ["Indicadores de logro", "Objetivo de aprendizaje", "Tema"],
            ["Explica el fenómeno", "Observar un circuito", "Electricidad"],
        ],
        nombre="Quinto año",
    )
    antes = archivo.read_bytes()
    uno = inventariar(tmp_path, " V4 ")
    dos = inventariar(tmp_path, "V4")
    assert uno == dos and archivo.read_bytes() == antes
    assert uno.version_declarada == "V4" and uno.estado == "listo_inspeccion"
    assert len(uno.archivos) == 1
    libro = uno.archivos[0]
    assert libro.archivo == "ciclo/ciencias.xlsx" and libro.sha256 == hashlib.sha256(antes).hexdigest()
    hoja = libro.hojas[0]
    assert hoja.nombre == "Quinto año" and hoja.filas_con_valores == 3
    assert hoja.primera_fila == 1 and hoja.ultima_fila == 4 and hoja.ultima_columna == 3
    assert hoja.cabeceras[0].fila == 3
    assert [(c.letra, c.campo_sugerido) for c in hoja.cabeceras[0].columnas] == [
        ("A", "indicador"),
        ("B", "objetivo"),
        ("C", "tema"),
    ]
    assert sorted(p.relative_to(tmp_path).as_posix() for p in tmp_path.rglob("*") if p.is_file()) == [
        "ciclo/ciencias.xlsx"
    ]


def test_no_silencia_duplicados_de_cabecera(tmp_path):
    _libro(tmp_path / "ciencias.xlsx", [["Contenido", "Contenidos", "Procedimental"], ["A", "B", "C"]])
    resultado = inventariar(tmp_path, "V4")
    assert resultado.estado == "incompleto"
    assert any("contenido aparece en columnas distintas: A, B" in motivo for motivo in resultado.observaciones)


def test_esquemas_distintos_en_una_hoja_requieren_mapeo(tmp_path):
    _libro(tmp_path / "ciencias.xlsx", [["Tema", "Objetivo"], ["A", "B"], [], ["Objetivo", "Tema"], ["C", "D"]])
    hoja = inventariar(tmp_path, "V4").archivos[0].hojas[0]
    assert hoja.estado == "requiere_mapeo" and len(hoja.cabeceras) == 2
    assert any("campos sugeridos distintos" in mensaje for mensaje in hoja.ambiguedades)


def test_incluye_hojas_sin_esquema_y_vacias_sin_inferir_su_exclusion(tmp_path):
    archivo = tmp_path / "ciencias.xlsx"
    _libro(archivo, [["Tema", "Objetivo"], ["Luz", "Observar"]])
    libro = openpyxl.load_workbook(archivo)
    libro.create_sheet("Vacía")
    libro.create_sheet("Instrucciones").append(["Revisar esta entrega antes de importarla"])
    libro.save(archivo)
    libro.close()
    resultado = inventariar(tmp_path, "V4")
    assert resultado.estado == "incompleto"
    estados = {hoja.nombre: hoja.estado for hoja in resultado.archivos[0].hojas}
    assert estados == {"Programa": "detectada", "Vacía": "vacia", "Instrucciones": "requiere_mapeo"}


def test_libro_vacio_es_incompleto(tmp_path):
    _libro(tmp_path / "vacio.xlsx", [])
    resultado = inventariar(tmp_path, "V4")
    assert resultado.estado == "incompleto"
    assert resultado.archivos[0].hojas[0].filas_con_valores == 0


def test_ignora_archivos_ajenos_y_locks_de_excel(tmp_path):
    _libro(tmp_path / "ciencias.xlsx", [["Tema", "Objetivo"], ["Luz", "Observar"]])
    (tmp_path / "notas.txt").write_text("Descripción de la entrega")
    (tmp_path / "~$ciencias.xlsx").write_bytes(b"lock de Excel")
    resultado = inventariar(tmp_path, "V4")
    assert resultado.estado == "listo_inspeccion"
    assert resultado.archivos_omitidos == ("notas.txt", "~$ciencias.xlsx")


@pytest.mark.parametrize("contenido,nombre", [(b"no es un zip", "corrupto.xlsx"), (b"binario", "legado.xls")])
def test_rechaza_formatos_ilegibles(tmp_path, contenido, nombre):
    (tmp_path / nombre).write_bytes(contenido)
    with pytest.raises(ValueError, match="ilegible|no admitido"):
        inventariar(tmp_path, "V4")


def test_rechaza_directorio_sin_excel(tmp_path):
    with pytest.raises(ValueError, match="no contiene libros"):
        inventariar(tmp_path, "V4")
    with pytest.raises(ValueError, match="No es un directorio"):
        inventariar(tmp_path / "ausente", "V4")


@pytest.mark.parametrize("version", ["", "  ", "x" * 51, None])
def test_rechaza_version_sin_identidad(tmp_path, version):
    with pytest.raises(ValueError, match="versión declarada"):
        inventariar(tmp_path, version)


@pytest.mark.parametrize("anidado", [False, True])
def test_permiso_denegado_no_se_confunde_con_directorio_vacio(tmp_path, monkeypatch, anidado):
    """El sandbox y los permisos pueden impedir enumerar una entrega que sí existe."""
    import os

    from goes_science_kg.ingesta import inventario

    _libro(tmp_path / "ciencias.xlsx", [["Tema", "Objetivo"], ["Luz", "Observar"]])
    protegido = tmp_path / "privado" if anidado else tmp_path
    protegido.mkdir(exist_ok=True)
    original = os.scandir

    def scandir(directorio):
        if Path(directorio) == protegido:
            raise PermissionError(f"Sin acceso a {directorio}")
        return original(directorio)

    monkeypatch.setattr(inventario.os, "scandir", scandir)
    with pytest.raises(PermissionError, match="Sin acceso"):
        inventariar(tmp_path, "V4")


def test_enlaces_fuera_del_directorio_no_se_siguen(tmp_path):
    base = tmp_path / "entrega"
    _libro(base / "ciencias.xlsx", [["Tema", "Objetivo"], ["Luz", "Observar"]])
    externo = tmp_path / "otra_entrega"
    _libro(externo / "fuente.xlsx", [["Tema", "Objetivo"], ["Calor", "Observar"]])
    (base / "enlace_directorio").symlink_to(externo, target_is_directory=True)
    (base / "enlace.xlsx").symlink_to(externo / "fuente.xlsx")
    resultado = inventariar(base, "V4")
    assert len(resultado.archivos) == 1
    assert resultado.estado == "incompleto"
    assert resultado.archivos_omitidos == ("enlace.xlsx", "enlace_directorio")
    assert len(resultado.observaciones) == 2
