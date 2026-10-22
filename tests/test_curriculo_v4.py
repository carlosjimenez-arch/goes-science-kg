"""Regresiones de ingesta exhaustiva: identidad, procedencia y estructura real V4."""

from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from goes_science_kg.curriculo import ingesta
from goes_science_kg.curriculo.contratos import FASES, Curriculo


def _procedencia(base: Path, *, version: str = "V4") -> None:
    registros = []
    for archivo in sorted((base / "insumos").glob("*.xlsx")):
        contenido = archivo.read_bytes()
        registros.append(
            {
                "archivo": archivo.name,
                "sha256": hashlib.sha256(contenido).hexdigest(),
                "bytes": len(contenido),
                "url": "https://example.invalid/original",
            }
        )
    (base / "procedencia.json").write_text(json.dumps({"version": version, "archivos": registros}), encoding="utf-8")


@pytest.fixture
def configuracion(tmp_path: Path) -> Path:
    """Libro pequeño con encabezados dobles, combinaciones y columnas no catalogadas."""
    carpeta = tmp_path / "insumos"
    carpeta.mkdir()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Unidad exacta "
    hoja.append(
        [
            "Eje",
            "Contenido",
            "Subcontenidos",
            "Procedimental",
            "Indicador avanzado",
            "Indicadores de logro",
            "Evidencia de aprendizaje",
            "Microhabilidades locales",
            None,
            None,
            None,
            "Conocimiento PISA\n(tipo y elemento)",
            "Contexto PISA\n(situación y área)",
            "Dato libre",
            "Dato libre",
            None,
        ]
    )
    hoja.append([None] * 7 + ["2. Explorar\nPreguntas", "3. Explicar", "4. Indagar", "5. Reforzar"])
    hoja.merge_cells("B1:B2")
    hoja.append(
        [
            "Eje local",
            "Materia",
            "Estados",
            "Mide el agua",
            "Avanzado distinto",
            "Logro distinto",
            "Tabla de datos",
            "E1, PR2.1 y EP3",
            "D2",
            "I10",
            "EP2",
            "Procedimental",
            "Local: agua",
            "Extra uno",
            "Extra dos",
            "Sin título",
        ]
    )
    hoja.append([None, None, None, "Compara medidas", None, None, "Gráfico", "E2"])
    hoja.merge_cells("A3:A4")
    hoja.merge_cells("B3:B4")
    hoja.append(
        [
            "Nota fuera de tema",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Dato fuera de tema",
        ]
    )
    auxiliar = libro.create_sheet("Notas")
    auxiliar.append(["Declaración local", "Código"])
    auxiliar.append(["No es un estándar oficial", "E1"])
    auxiliar.append(["Fórmula preservada", "=1+1"])
    libro.save(carpeta / "fuente.xlsx")
    libro.close()
    _procedencia(tmp_path)
    alias = {
        "contenido": ["Contenido"],
        "subcontenido": ["Subcontenidos"],
        "procedimental": ["Procedimental"],
        "indicador_avanzado": ["Indicador avanzado"],
        "indicadores_logro": ["Indicadores de logro"],
        "evidencia_aprendizaje": ["Evidencia de aprendizaje"],
        "conocimiento_pisa": ["Conocimiento PISA"],
        "contexto_pisa": ["Contexto PISA"],
        **{f"fase_{fase}": [f"{i}. {fase.title()}"] for i, fase in enumerate(FASES, 1)},
    }
    cfg = {
        "version": "V4",
        "raiz": ".",
        "directorio": "insumos",
        "procedencia": "procedencia.json",
        "alias_campos": alias,
        "fuentes": [
            {
                "archivo": "fuente.xlsx",
                "hojas": [
                    {
                        "nombre": "Unidad exacta ",
                        "tratamiento": "curricular",
                        "motivo": "Matriz original.",
                        "grado": 7,
                        "asignatura_malla": "Ciencias",
                        "unidad": "Unidad configurada",
                        "filas_encabezado": [1, 2],
                        "campos_por_columna": {"A": "eje_integrador"},
                    },
                    {
                        "nombre": "Notas",
                        "tratamiento": "auxiliar",
                        "motivo": "Declaraciones del insumo.",
                        "filas_encabezado": [1],
                    },
                ],
            }
        ],
    }
    archivo = tmp_path / "curriculo.yaml"
    archivo.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    return archivo


def _mutar_libro(configuracion: Path, accion) -> None:
    archivo = configuracion.parent / "insumos/fuente.xlsx"
    libro = load_workbook(archivo)
    accion(libro)
    libro.save(archivo)
    libro.close()
    _procedencia(configuracion.parent)


def _mutar_config(configuracion: Path, accion) -> None:
    cfg = yaml.safe_load(configuracion.read_text(encoding="utf-8"))
    accion(cfg)
    configuracion.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")


def test_conserva_roles_distintos_combinaciones_y_celdas_desconocidas(configuracion: Path):
    curriculo = ingesta.cargar_curriculo(configuracion)
    assert len(curriculo.filas) == 2
    primera, segunda = curriculo.filas
    assert primera.indicador_avanzado == "Avanzado distinto"
    assert primera.indicadores_logro == "Logro distinto"
    assert primera.conocimiento_pisa == "Procedimental"
    assert primera.contexto_pisa == "Local: agua"
    assert primera.fases["aproximar"].presente_en_hoja is False
    assert primera.fases["aproximar"].texto == ""
    assert primera.fases["explorar"].codigos_locales == ("E1", "PR2.1", "EP3")
    assert primera.fases["explorar"].autoridad_codigos == "insumo_local"
    assert segunda.contenido == "Materia"
    assert segunda.subcontenido == ""  # Blanco ordinario: no se hereda "Estados".
    combinada = next(c for c in segunda.celdas if c.coordenada == "B4")
    assert combinada.valor_original is None
    assert combinada.valor == "Materia" and combinada.origen_combinado == "B3"
    extras = {c.coordenada: c.valor for c in primera.celdas if c.campo is None}
    assert extras["N3"] == "Extra uno" and extras["O3"] == "Extra dos" and extras["P3"] == "Sin título"
    assert curriculo.hojas[0].nombre == "Unidad exacta "
    assert len(curriculo.hojas[0].registros) == 5
    assert curriculo.hojas[0].registros[-1].celdas[13].valor_original == "Dato fuera de tema"
    formula = curriculo.hojas[1].registros[-1].celdas[1]
    assert formula.valor_original == "=1+1" and formula.tipo_excel == "f"
    assert len(curriculo.agrupaciones) == 1
    assert curriculo.agrupaciones[0].filas == tuple(f.id for f in curriculo.filas)


def test_identidad_determinista_serializable_y_fuentes_no_mutadas(configuracion: Path):
    archivo = configuracion.parent / "insumos/fuente.xlsx"
    original = archivo.read_bytes()
    primera = ingesta.cargar_curriculo(configuracion)
    segunda = ingesta.cargar_curriculo(configuracion)
    assert primera == segunda
    assert Curriculo.model_validate_json(primera.model_dump_json()) == primera
    assert archivo.read_bytes() == original
    assert primera.fuentes[0].sha256 == hashlib.sha256(original).hexdigest()
    assert primera.fuentes[0].archivo_local == "insumos/fuente.xlsx"
    assert primera.fuentes[0].procedencia["archivo"]["url"].endswith("/original")


def test_cambiar_version_no_reutiliza_identidades(configuracion: Path):
    anterior = ingesta.cargar_curriculo(configuracion)
    _mutar_config(configuracion, lambda cfg: cfg.update(version="V5"))
    _procedencia(configuracion.parent, version="V5")
    nueva = ingesta.cargar_curriculo(configuracion)
    assert nueva.huella != anterior.huella
    assert {f.id for f in anterior.filas}.isdisjoint(f.id for f in nueva.filas)


def test_cambio_de_contenido_original_cambia_identidad_sin_correspondencia_implicita(configuracion: Path):
    anterior = ingesta.cargar_curriculo(configuracion)
    _mutar_libro(configuracion, lambda libro: setattr(libro.worksheets[0]["D3"], "value", "Otro procedimiento"))
    nueva = ingesta.cargar_curriculo(configuracion)
    assert nueva.filas[0].id != anterior.filas[0].id
    assert nueva.filas[0].procedimental == "Otro procedimiento"
    assert "revisado_por" not in nueva.filas[0].model_dump()


def test_ninguna_celda_original_con_valor_se_pierde(configuracion: Path):
    curriculo = ingesta.cargar_curriculo(configuracion)
    libro = load_workbook(configuracion.parent / "insumos/fuente.xlsx", data_only=False)
    try:
        for original, resultado in zip(libro, curriculo.hojas, strict=True):
            esperado = {c.coordinate: c.value for fila in original for c in fila if c.value is not None}
            obtenido = {
                c.coordenada: c.valor_original
                for r in resultado.registros
                for c in r.celdas
                if c.valor_original is not None
            }
            assert obtenido == esperado
    finally:
        libro.close()


def test_alias_compartido_por_roles_se_rechaza(configuracion: Path):
    _mutar_config(configuracion, lambda cfg: cfg["alias_campos"].update(indicadores_logro=["Procedimental"]))
    with pytest.raises(ValueError, match="Alias vacío o ambiguo"):
        ingesta.cargar_curriculo(configuracion)


def test_contrato_no_permite_omitir_fases_ausentes(configuracion: Path):
    datos = ingesta.cargar_curriculo(configuracion).model_dump(mode="json")
    datos["filas"][0]["fases"].pop("aproximar")
    with pytest.raises(ValueError, match="cinco fases"):
        Curriculo.model_validate(datos)


def test_hash_se_verifica_antes_de_abrir_excel(configuracion: Path, monkeypatch):
    archivo = configuracion.parent / "insumos/fuente.xlsx"
    archivo.write_bytes(archivo.read_bytes() + b"alterado")
    llamadas = []
    monkeypatch.setattr(ingesta, "load_workbook", lambda *a, **kw: llamadas.append(a))
    with pytest.raises(ValueError, match="SHA256"):
        ingesta.cargar_curriculo(configuracion)
    assert llamadas == []


def test_parser_recibe_exactamente_snapshot_cotejado(configuracion: Path, monkeypatch):
    archivo = configuracion.parent / "insumos/fuente.xlsx"
    original = archivo.read_bytes()
    abrir = ingesta.load_workbook

    def mutar_despues_del_hash(contenido, **kwargs):
        archivo.write_bytes(b"Cambio concurrente en disco")
        assert contenido.getvalue() == original
        return abrir(contenido, **kwargs)

    monkeypatch.setattr(ingesta, "load_workbook", mutar_despues_del_hash)
    curriculo = ingesta.cargar_curriculo(configuracion)
    assert curriculo.fuentes[0].sha256 == hashlib.sha256(original).hexdigest()
    assert curriculo.filas[0].procedimental == "Mide el agua"


@pytest.mark.parametrize("cambio", ["version", "duplicado", "conjunto", "tamano"])
def test_procedencia_incoherente_rechazada(configuracion: Path, cambio: str):
    ruta = configuracion.parent / "procedencia.json"
    origen = json.loads(ruta.read_text())
    if cambio == "version":
        origen["version"] = "V3"
    elif cambio == "duplicado":
        origen["archivos"].append(origen["archivos"][0])
    elif cambio == "conjunto":
        origen["archivos"][0]["archivo"] = "otra.xlsx"
    else:
        origen["archivos"][0]["bytes"] += 1
    ruta.write_text(json.dumps(origen))
    with pytest.raises(ValueError):
        ingesta.cargar_curriculo(configuracion)


def test_hoja_nueva_no_se_ignora_por_nombre(configuracion: Path):
    def nueva(libro):
        libro.create_sheet("No se parece a una unidad").append(["contenido significativo"])

    _mutar_libro(configuracion, nueva)
    with pytest.raises(ValueError, match="Hojas sin decisión"):
        ingesta.cargar_curriculo(configuracion)


def test_hoja_excluida_requiere_motivo_y_preserva_registros(configuracion: Path):
    def excluir(cfg):
        cfg["fuentes"][0]["hojas"][1].update(tratamiento="excluida", motivo="Fuera de alcance curricular.")

    _mutar_config(configuracion, excluir)
    hoja = ingesta.cargar_curriculo(configuracion).hojas[1]
    assert hoja.tratamiento == "excluida" and len(hoja.registros) == 3
    _mutar_config(configuracion, lambda cfg: cfg["fuentes"][0]["hojas"][1].update(motivo=" "))
    with pytest.raises(ValueError, match="motivo"):
        ingesta.cargar_curriculo(configuracion)


def test_encabezados_semanticos_duplicados_no_sobrescriben(configuracion: Path):
    _mutar_libro(configuracion, lambda libro: setattr(libro.worksheets[0]["N1"], "value", "Procedimental"))
    with pytest.raises(ValueError, match="Campo duplicado"):
        ingesta.cargar_curriculo(configuracion)


def test_encabezado_1_y_2_no_pueden_atribuir_dos_roles_a_una_columna(configuracion: Path):
    _mutar_libro(configuracion, lambda libro: setattr(libro.worksheets[0]["H1"], "value", "Indicadores de logro"))
    with pytest.raises(ValueError, match="Encabezados ambiguos"):
        ingesta.cargar_curriculo(configuracion)


@pytest.mark.parametrize("procedimiento", ["=1+1", 123, True])
def test_procedimental_sin_texto_verificable_se_rechaza(configuracion: Path, procedimiento):
    _mutar_libro(configuracion, lambda libro: setattr(libro.worksheets[0]["D3"], "value", procedimiento))
    with pytest.raises(ValueError, match="Procedimental"):
        ingesta.cargar_curriculo(configuracion)


def test_archivos_inesperados_en_directorio_se_rechazan(configuracion: Path):
    (configuracion.parent / "insumos/nueva.xlsx").write_bytes(b"No inventariada")
    with pytest.raises(ValueError, match="conjunto de libros"):
        ingesta.cargar_curriculo(configuracion)


def test_no_lee_archivo_fuera_de_raiz_por_symlink(configuracion: Path, tmp_path: Path):
    externo = tmp_path.parent / f"fuera-{tmp_path.name}.xlsx"
    externo.write_bytes((tmp_path / "insumos/fuente.xlsx").read_bytes())
    try:
        (tmp_path / "insumos/fuente.xlsx").unlink()
        (tmp_path / "insumos/fuente.xlsx").symlink_to(externo)
        with pytest.raises(ValueError, match="fuera de la raíz"):
            ingesta.cargar_curriculo(configuracion)
    finally:
        externo.unlink()


def test_version_real_v4_inventariada_conserva_todas_las_hojas_y_filas():
    curriculo = ingesta.cargar_curriculo()
    assert len(curriculo.fuentes) == 11 and len(curriculo.hojas) == 126 and len(curriculo.filas) == 1020
    assert Counter(h.tratamiento for h in curriculo.hojas) == {"curricular": 62, "auxiliar": 64}
    assert Counter(f.grado for f in curriculo.filas) == {
        2: 63,
        3: 64,
        4: 63,
        5: 64,
        6: 96,
        7: 94,
        8: 96,
        9: 96,
        10: 192,
        11: 192,
    }
    assert len({f.id for f in curriculo.filas}) == 1020
    assert len(curriculo.agrupaciones) == 108
    assert sum(len(h.registros) for h in curriculo.hojas) == 3313
    assert sum(len(r.celdas) for h in curriculo.hojas for r in h.registros) == 54975
    sin_aproximar = [
        h
        for h in curriculo.hojas
        if h.tratamiento == "curricular" and "fase_aproximar" not in {c.campo for c in h.columnas}
    ]
    assert len(sin_aproximar) == 8
    assert any(f.numero_clase for f in curriculo.filas if f.asignatura_malla == "Química")
    assert any(f.numero_tema for f in curriculo.filas if f.asignatura_malla == "Física")
    quimica_11 = next(h for h in curriculo.hojas if h.nombre == "Química 11.°")
    assert {c.columna for c in quimica_11.columnas if c.campo is None} == {"Z", "AA"}
    assert all(len(r.celdas) == 27 for r in quimica_11.registros)
