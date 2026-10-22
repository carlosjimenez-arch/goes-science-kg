"""Test cell-level context with original bytes, bounded roots, merged anchors, and budgets."""

from __future__ import annotations

import hashlib
import json
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest

from goes_science_kg.curriculo.contexto_fuentes import (
    catalogo_fuentes_v4,
    contexto_celdas,
    documento_de_fila,
    fuentes_del_curriculo,
)
from goes_science_kg.curriculo.contratos import (
    FASES,
    CeldaCurricular,
    Curriculo,
    FaseCurricular,
    FilaCurricular,
    FuenteCurricular,
)
from goes_science_kg.motor.fuentes import FuentesVerificadas
from goes_science_kg.recuperacion import Documento, Resultado


@pytest.fixture
def archivo_v4(tmp_path):
    principal = tmp_path / "historicas"
    v4 = tmp_path / "v4"
    principal.mkdir()
    v4.mkdir()
    archivo = v4 / "malla.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Unidad 1"
    hoja.append(
        [
            "Unidad",
            "Contenido",
            "Procedimental",
            "Indicador",
            "Evidencia",
            "Aproximar",
            "Explorar",
            "Explicar",
            "Indagar",
            "Reforzar",
            "Conocimiento",
            "Contexto",
        ]
    )
    hoja.append(
        [
            "Seres vivos",
            "Fotosíntesis",
            "Identifica las necesidades de agua y luz de las plantas.",
            "Explica las necesidades de una planta con una observación.",
            "Registra una observación y justifica su explicación.",
            "Observa dos dibujos de plantas del entorno.",
            "Formula una pregunta sobre el crecimiento de una planta.",
            "Explica qué necesitan las plantas para crecer.",
            "Compara observaciones de dos plantas.",
            "Aplica la explicación a una planta diferente.",
            "Conocimiento de contenido y procedimiento.",
            "Situación local de la escuela y el hogar.",
        ]
    )
    hoja.append(["Fila ajena", None, None, "Otro objetivo no relacionado", "=2+2"])
    hoja.merge_cells("B2:B3")
    libro.save(archivo)
    libro.close()
    catalogo = {
        "v4_doc": {
            "archivo_local": archivo.name,
            "sha256": hashlib.sha256(archivo.read_bytes()).hexdigest(),
            "raiz_fuente": "v4",
            "version_curriculo": "V4",
        }
    }
    fuentes = FuentesVerificadas(catalogo, raiz_fuentes=principal, raices_adicionales={"v4": v4})
    return archivo, principal, v4, catalogo, fuentes


def _documento(**fuente):
    return Documento(
        id="fila_2",
        texto="Este texto de metadatos NO debe aceptarse como evidencia.",
        fuente={"documento": "v4_doc", "hoja": "Unidad 1", **fuente},
        version="V4",
        grado=5,
    )


def _curriculo(archivo_v4):
    archivo, _, _, catalogo, _ = archivo_v4
    libro = openpyxl.load_workbook(archivo)
    campos = [
        "unidad",
        "contenido",
        "procedimental",
        "indicadores_logro",
        "evidencia_aprendizaje",
        *[f"fase_{fase}" for fase in FASES],
        "conocimiento_pisa",
        "contexto_pisa",
    ]
    celdas = tuple(
        CeldaCurricular(
            coordenada=celda.coordinate,
            fila=celda.row,
            columna=celda.column_letter,
            campo=campo,
            valor_original=celda.value,
            valor=celda.value,
            tipo_excel=celda.data_type,
        )
        for celda, campo in zip(libro.active[2], campos, strict=True)
    )
    libro.close()
    fuente = FuenteCurricular(
        id="v4_doc",
        version="V4",
        archivo=archivo.name,
        archivo_local=f"insumos/v4/{archivo.name}",
        sha256=catalogo["v4_doc"]["sha256"],
        bytes=archivo.stat().st_size,
        procedencia={"origen": "fixture"},
    )
    fila = FilaCurricular(
        id="fila_2",
        fuente_id="v4_doc",
        version="V4",
        hoja="Unidad 1",
        fila=2,
        grado=5,
        asignatura_malla="ciencias",
        unidad="Seres vivos",
        procedimental=str(celdas[2].valor),
        celdas=celdas,
        fases={
            fase: FaseCurricular(presente_en_hoja=True, texto=str(celdas[5 + i].valor)) for i, fase in enumerate(FASES)
        },
    )
    return Curriculo(version="V4", huella="a" * 64, fuentes=(fuente,), filas=(fila,), hojas=(), agrupaciones=())


def test_hidrata_valores_reales_descartando_texto_y_valores_suplantados(archivo_v4):
    _, _, _, catalogo, fuentes = archivo_v4
    documento = _documento(
        celdas=["C2", "D2"],
        fila=999,
        valor_real="Las plantas vuelan",
        raiz_fuente="principal",
        celdas_verificadas=[{"valor_real": "falso"}],
        sha256="0" * 64,
    )
    evidencias = fuentes.hidratar_celdas(documento)
    assert len(evidencias) == 2
    assert evidencias[0].texto == "Identifica las necesidades de agua y luz de las plantas."
    assert evidencias[0].fuente["sha256"] == catalogo["v4_doc"]["sha256"]
    assert evidencias[0].fuente["raiz_fuente"] == "v4"
    assert evidencias[0].fuente["fila"] == 2
    assert "valor_real" not in evidencias[0].fuente
    assert evidencias[0].fuente["celdas_verificadas"][0]["valor_real"] == evidencias[0].texto
    assert "metadatos" not in evidencias[0].texto


def test_una_lectura_y_un_parseo_por_archivo_incluso_concurrencia(archivo_v4):
    archivo, _, _, _, fuentes = archivo_v4
    documentos = [_documento(celdas=[coordenada]) for coordenada in ("C2", "D2", "F2", "L2")]
    with ThreadPoolExecutor(max_workers=4) as ejecutor:
        resultados = list(ejecutor.map(fuentes.hidratar_celdas, documentos * 3))
    assert all(len(resultado) == 1 for resultado in resultados)
    assert fuentes.estadisticas() == {
        "archivos_leidos": 1,
        "bytes_verificados": archivo.stat().st_size,
        "libros_parseados": 1,
        "pdfs_parseados": 0,
    }


def test_combinadas_conservan_solicitada_ancla_y_vacio_original_sin_arrastre(archivo_v4):
    fuentes = archivo_v4[-1]
    resultado = fuentes.hidratar_celdas(_documento(celdas=["B3", "C3"]))
    assert len(resultado) == 1
    celda = resultado[0].fuente["celdas_verificadas"][0]
    assert celda["coordenada_solicitada"] == "B3"
    assert celda["coordenada_origen"] == "B2"
    assert celda["valor_original"] is None
    assert celda["valor_real"] == "Fotosíntesis"
    assert celda["rango_combinado"] == "B2:B3"


def test_rango_acotado_y_cache_no_confunden_localizadores(archivo_v4):
    fuentes = archivo_v4[-1]
    primero = fuentes.hidratar(_documento(celdas=["C2"]))
    segundo = fuentes.hidratar(_documento(rango="D2:E2"))
    assert primero is not None and segundo is not None
    assert primero.texto != segundo.texto
    assert "Registra una observación" in segundo.texto
    assert len(segundo.fuente["celdas_verificadas"]) == 2


def test_localizadores_publicados_reflejan_solo_las_celdas_extraidas(archivo_v4):
    fuentes = archivo_v4[-1]
    una_fila = fuentes.hidratar(_documento(celdas=["C2"], fila=999, pagina=123))
    varias_filas = fuentes.hidratar(_documento(rango="C2:D3", fila=999, pagina=123))
    assert una_fila is not None and varias_filas is not None
    assert una_fila.fuente["fila"] == 2
    assert "fila" not in varias_filas.fuente
    assert "pagina" not in una_fila.fuente and "pagina" not in varias_filas.fuente


@pytest.mark.parametrize(
    "selector",
    [
        {"rango": "A1:XFD1048576"},
        {"rango": "A:A"},
        {"rango": "C3:A1"},
        {"celdas": ["A2"] * 65},
        {"celdas": ["A1048577"]},
        {"celdas": ["Otra!C2"]},
        {"celdas": ["A2"], "rango": "A2:B2"},
    ],
)
def test_selectores_invalidos_no_producen_evidencia(archivo_v4, selector):
    assert archivo_v4[-1].hidratar_celdas(_documento(**selector)) == []


def test_formula_se_preserva_como_formula_sin_inventar_su_resultado(archivo_v4):
    resultado = archivo_v4[-1].hidratar_celdas(_documento(celdas=["E3"]))[0]
    assert resultado.texto == "=2+2"
    celda = resultado.fuente["celdas_verificadas"][0]
    assert celda["valor_es_formula"] is True and celda["tipo_excel"] == "f"
    assert celda["valor_real"] != 4


def test_modificacion_posterior_invalida_tambien_celdas_ya_cacheadas(archivo_v4):
    archivo, _, _, _, fuentes = archivo_v4
    assert fuentes.hidratar_celdas(_documento(celdas=["C2"]))
    archivo.write_bytes("Archivo alterado después del snapshot".encode())
    assert fuentes.hidratar_celdas(_documento(celdas=["C2"])) == []
    assert fuentes.hidratar_celdas(_documento(celdas=["D2"])) == []


def test_raiz_del_payload_no_amplia_permisos_y_symlink_fuera_no_se_lee(archivo_v4, tmp_path):
    archivo, principal, v4, catalogo, _ = archivo_v4
    externo = tmp_path / "fuera.xlsx"
    externo.write_bytes(archivo.read_bytes())
    (v4 / "escape.xlsx").symlink_to(externo)
    catalogo["v4_doc"]["archivo_local"] = "escape.xlsx"
    fuentes = FuentesVerificadas(catalogo, raiz_fuentes=principal, raices_adicionales={"v4": v4})
    assert fuentes.hidratar_celdas(_documento(celdas=["C2"], raiz_fuente=str(tmp_path))) == []
    assert fuentes.estadisticas()["archivos_leidos"] == 0


def test_version_y_sha_proceden_del_catalogo_no_de_metadatos(archivo_v4):
    archivo, principal, v4, catalogo, fuentes = archivo_v4
    documento = _documento(celdas=["C2"]).model_copy(update={"version": "V3"})
    assert fuentes.hidratar_celdas(documento) == []
    catalogo["v4_doc"]["sha256"] = "0" * 64
    incorrectas = FuentesVerificadas(catalogo, raiz_fuentes=principal, raices_adicionales={"v4": v4})
    documento = _documento(celdas=["C2"], sha256=hashlib.sha256(archivo.read_bytes()).hexdigest())
    assert incorrectas.hidratar_celdas(documento) == []


def test_contexto_prioriza_procedimental_evaluacion_y_fases_sin_leer_otra_fila(archivo_v4):
    curriculo = _curriculo(archivo_v4)
    documento = documento_de_fila(curriculo.filas[0])
    fuentes = fuentes_del_curriculo(curriculo, raiz_v4=archivo_v4[2], raiz_historica=archivo_v4[1])
    contexto = contexto_celdas(
        [Resultado(documento=documento, puntaje=1, canales=("bm25",))],
        version="V4",
        fuentes=fuentes,
        max_caracteres=18000,
    )
    assert contexto[0].fuente["celdas"] == ["C2"]
    assert {contexto[1].fuente["celdas"][0], contexto[2].fuente["celdas"][0]} == {"D2", "E2"}
    assert all("Fila ajena" not in e.texto and "Otro objetivo" not in e.texto for e in contexto)
    assert len(json.dumps([e.model_dump(mode="json") for e in contexto], ensure_ascii=False)) <= 18000
    assert catalogo_fuentes_v4(curriculo)["v4_doc"]["archivo_local"] == "malla.xlsx"


def test_contexto_respeta_presupuesto_y_se_abstiene_sin_anclaje_local(archivo_v4):
    curriculo = _curriculo(archivo_v4)
    documento = documento_de_fila(curriculo.filas[0])
    resultado = Resultado(documento=documento, puntaje=1, canales=("bm25",))
    fuentes = archivo_v4[-1]
    contexto = contexto_celdas([resultado], version="V4", fuentes=fuentes, max_caracteres=2000)
    assert contexto and contexto[0].fuente["celdas"] == ["C2"]
    assert len(json.dumps([e.model_dump(mode="json") for e in contexto], ensure_ascii=False)) <= 2000
    assert contexto_celdas([resultado], version="V4", fuentes=fuentes, max_caracteres=100) == []
    referencia = resultado.model_copy(update={"documento": documento.model_copy(update={"version": "referencia"})})
    assert contexto_celdas([referencia], version="V4", fuentes=fuentes) == []
