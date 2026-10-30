"""Change only supported correction cells while preserving the original OOXML package."""

import hashlib
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from goes_science_kg.motor.contratos import huella
from goes_science_kg.producto.enmiendas_excel import exportar_enmiendas


@pytest.fixture
def insumo(tmp_path):
    raiz = tmp_path / "originales"
    raiz.mkdir()
    ruta = raiz / "original.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Biología"
    ws["A1"] = "C = Citocina"
    ws["A1"].font = Font(bold=True, color="123456")
    ws["B1"] = "Valor que debe conservarse."
    ws["C2"] = "=1+1"
    ws.merge_cells("D1:E1")
    ws["D1"] = "Nota combinada"
    wb.save(ruta)
    wb.close()
    digest = hashlib.sha256(ruta.read_bytes()).hexdigest()
    datos = dict(
        fuente_id="fuente",
        sha256=digest,
        hoja="Biología",
        celda_ancla="A1",
        antes="C = Citocina",
        despues="C = Citosina",
        alcance="Sólo copia derivada.",
    )
    e = {"id": "ENMIENDA:" + huella(datos), **datos}
    fuente = SimpleNamespace(id="fuente", archivo=ruta.name, sha256=digest)
    propuesta = SimpleNamespace(enmiendas_aceptadas=(e["id"],))
    c = SimpleNamespace(
        preparacion=SimpleNamespace(correcciones_requeridas=(e,)), resultado=SimpleNamespace(propuesta=propuesta)
    )
    return raiz, ruta, SimpleNamespace(fuentes=(fuente,)), SimpleNamespace(id="archivo", unidades=(c,))


def test_solo_cambia_xml_de_hoja_y_celda_respaldada(insumo, tmp_path):
    raiz, original, curriculo, archivo = insumo
    bytes_antes = original.read_bytes()
    informe = exportar_enmiendas(curriculo, archivo, raiz_v4=raiz, destino=tmp_path / "derivadas")
    copia = tmp_path / "derivadas" / informe["copias"][0]["archivo"]
    assert original.read_bytes() == bytes_antes
    with zipfile.ZipFile(original) as a, zipfile.ZipFile(copia) as b:
        assert a.namelist() == b.namelist()
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == ["xl/worksheets/sheet1.xml"]
    wb = load_workbook(copia, data_only=False)
    assert wb.active["A1"].value == "C = Citosina" and wb.active["A1"].font.bold
    assert wb.active["C2"].value == "=1+1" and wb.active["C2"].data_type == "f"
    wb.close()
    with pytest.raises(ValueError, match="nuevo"):
        exportar_enmiendas(curriculo, archivo, raiz_v4=raiz, destino=tmp_path / "derivadas")


def test_cambio_de_fuente_impide_escribir_derivada(insumo, tmp_path):
    raiz, original, curriculo, archivo = insumo
    original.write_bytes(original.read_bytes() + b" cambiado")
    with pytest.raises(ValueError, match="original"):
        exportar_enmiendas(curriculo, archivo, raiz_v4=raiz, destino=tmp_path / "rechazada")
    assert not (tmp_path / "rechazada").exists()


def test_no_exporta_enmienda_de_unidad_sin_respaldo(insumo, tmp_path):
    raiz, _, curriculo, archivo = insumo
    archivo.unidades[0].resultado.propuesta = None
    informe = exportar_enmiendas(curriculo, archivo, raiz_v4=raiz, destino=tmp_path / "sin_respaldo")
    assert informe["copias"] == [] and not list((tmp_path / "sin_respaldo").glob("*.xlsx"))
