"""Integración de terminal con fuentes locales: controles previos a cualquier proveedor."""

from __future__ import annotations

import hashlib
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from goes_science_kg.cli import app
from goes_science_kg.cli import unidades as cli_unidades
from goes_science_kg.curriculo.catalogacion import CatalogoDisciplinas, ClasificacionFila
from goes_science_kg.curriculo.contexto_fuentes import documento_de_fila, fuentes_del_curriculo
from goes_science_kg.curriculo.contratos import (
    FASES,
    CeldaCurricular,
    Curriculo,
    FaseCurricular,
    FilaCurricular,
    FuenteCurricular,
)
from goes_science_kg.modelos import Nodo, TipoNodo
from goes_science_kg.motor.contratos import huella
from goes_science_kg.producto.conocimiento import InformeConocimiento
from goes_science_kg.producto.contratos import ReferenciasProducto
from goes_science_kg.producto.unidades import (
    ContextoAula,
    PaqueteUnidad,
    ResultadoUnidad,
    SesionEjecutada,
    preparar_unidad,
)
from goes_science_kg.recuperacion import Buscador

RUNNER = CliRunner()
RUTA = ["producto", "unidades"]


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    """Dos filas reales y una publicación en memoria; Vertex queda prohibido por defecto."""
    monkeypatch.setenv("GSKG_RAIZ", str(tmp_path))
    directorio = tmp_path / "insumos/v4"
    directorio.mkdir(parents=True)
    archivo = directorio / "fisica.xlsx"
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Unidad sintética"
    hoja.append(["Contenido", "Procedimental"])
    hoja.append(["Medición", "Compara dos mediciones de longitud sobre papel."])
    hoja.append(["Medición", "Representa dos mediciones mediante dibujos."])
    libro.save(archivo)
    contenido = archivo.read_bytes()
    fuente = FuenteCurricular(
        id="fuente-sintetica",
        version="V4",
        archivo=archivo.name,
        archivo_local=f"insumos/v4/{archivo.name}",
        sha256=hashlib.sha256(contenido).hexdigest(),
        bytes=len(contenido),
        procedencia={"tipo": "prueba sin red"},
    )
    filas = tuple(
        FilaCurricular(
            id=f"curriculo-v4-{numero}",
            fuente_id=fuente.id,
            version="V4",
            grado=10,
            asignatura_malla="Física",
            hoja=hoja.title,
            fila=numero,
            unidad="Unidad sintética",
            contenido=hoja.cell(numero, 1).value,
            procedimental=hoja.cell(numero, 2).value,
            fases={f: FaseCurricular(presente_en_hoja=False) for f in FASES},
            celdas=tuple(
                CeldaCurricular(
                    coordenada=c.coordinate,
                    fila=numero,
                    columna=c.column_letter,
                    campo=campo,
                    valor_original=c.value,
                    valor=c.value,
                    tipo_excel=c.data_type,
                )
                for c, campo in zip(hoja[numero], ("contenido", "procedimental"), strict=True)
            ),
        )
        for numero in (2, 3)
    )
    libro.close()
    curriculo = Curriculo(version="V4", huella="a" * 64, fuentes=(fuente,), filas=filas, hojas=(), agrupaciones=())
    catalogo = CatalogoDisciplinas(
        version_curriculo="V4",
        huella_curriculo=curriculo.huella,
        huella_instrucciones="b" * 64,
        huella_harness="c" * 64,
        modelos={"catalogador_a": "sin-red-a", "catalogador_b": "sin-red-b"},
        modo="completo",
        piloto_superado=True,
        filas=tuple(
            ClasificacionFila(
                fila_id=f.id,
                estado="declarado_malla",
                disciplinas=("fisica",),
                metodo="asignatura_malla",
                justificacion="Asignatura explícita de prueba",
            )
            for f in filas
        ),
    )
    fuentes = fuentes_del_curriculo(curriculo, raiz_v4=directorio)
    from goes_science_kg.producto.consulta_conocimiento import BuscadorProducto

    buscador = BuscadorProducto(
        Buscador(documento_de_fila(f) for f in filas),
        curriculo,
        catalogo,
        None,
        fuentes,
        snapshot_id="fixture-sin-publicar",
    )
    runtime = SimpleNamespace(
        producto=SimpleNamespace(curriculo=curriculo, catalogo=catalogo),
        fuentes=fuentes,
        buscador=buscador,
    )
    llamadas = {"runtime": 0, "vertex": 0}

    def cargar_runtime():
        llamadas["runtime"] += 1
        return runtime

    def vertex_prohibido(*args, **kwargs):
        llamadas["vertex"] += 1
        raise AssertionError("La prueba no permite construir proveedores remotos")

    monkeypatch.setattr("goes_science_kg.curriculo.cargar_curriculo", lambda: curriculo)
    monkeypatch.setattr("goes_science_kg.producto.runtime.cargar_runtime", cargar_runtime)
    monkeypatch.setattr("goes_science_kg.motor.vertex.Vertex", vertex_prohibido)
    alertas = tmp_path / "docs/investigacion/alertas_v4.json"
    alertas.parent.mkdir(parents=True)
    alertas.write_text('{"alertas": []}', encoding="utf-8")
    cfg = {
        "vertex": {"proyecto": "sin-red", "modelos": {}, "tarifas": {}},
        "max_caracteres_contexto": 18000,
        "max_reparaciones": 0,
    }
    monkeypatch.setattr(cli_unidades, "cargar", lambda nombre: cfg)
    return SimpleNamespace(
        base=tmp_path, curriculo=curriculo, catalogo=catalogo, fuentes=fuentes, runtime=runtime, llamadas=llamadas
    )


def _paquete(entorno, *, catalogo=None):
    return preparar_unidad(
        entorno.curriculo,
        catalogo or entorno.catalogo,
        grado=10,
        asignatura_malla="Física",
        unidad="Unidad sintética",
        contexto=ContextoAula(),
        fuentes=entorno.fuentes,
    )


def _guardar(entorno, paquete, *, recalcular=False):
    if recalcular:
        paquete = paquete.model_copy(update={"id": huella(paquete.model_dump(mode="json", exclude={"id"}))})
    archivo = entorno.base / "entrada.json"
    archivo.write_text(json.dumps({"paquete": paquete.model_dump(mode="json")}), encoding="utf-8")
    return archivo


def _ejecutar(entorno, archivo, *opciones):
    return RUNNER.invoke(app, [*RUTA, "ejecutar", str(archivo), "--destino", str(entorno.base / "salida"), *opciones])


def _sin_efectos(entorno):
    assert entorno.llamadas["vertex"] == 0
    assert not (entorno.base / "data/operaciones").exists()
    assert not (entorno.base / "salida").exists()


@pytest.mark.parametrize(
    "comando", [[*RUTA], [*RUTA, "listar"], [*RUTA, "preparar"], [*RUTA, "ejecutar"], ["producto", "conocimiento"]]
)
def test_ayuda_nuevas_rutas_sin_fuentes_ni_proveedores(entorno, monkeypatch, comando):
    monkeypatch.setattr("goes_science_kg.curriculo.cargar_curriculo", lambda: pytest.fail("La ayuda no carga fuentes"))
    resultado = RUNNER.invoke(app, [*comando, "--help"])
    assert resultado.exit_code == 0, resultado.output
    assert "Usage:" in resultado.stdout
    assert entorno.llamadas == {"runtime": 0, "vertex": 0}


def test_listar_filtra_por_literal_de_asignatura_y_grado(entorno):
    resultado = RUNNER.invoke(app, [*RUTA, "listar", "--grado", "10", "--asignatura-malla", "Física"])
    assert resultado.exit_code == 0, resultado.output
    assert json.loads(resultado.stdout) == [
        {"grado": 10, "asignatura_malla": "Física", "unidad": "Unidad sintética", "filas": 2}
    ]
    otro = RUNNER.invoke(app, [*RUTA, "listar", "--asignatura-malla", "fisica"])
    assert otro.exit_code == 0 and json.loads(otro.stdout) == []
    _sin_efectos(entorno)


def test_preparar_piloto_exporta_fuentes_contexto_y_cero_llamadas(entorno):
    piloto = entorno.catalogo.model_copy(update={"modo": "piloto"})
    archivo = entorno.base / "catalogo.json"
    archivo.write_text(json.dumps({"catalogo": piloto.model_dump(mode="json")}), encoding="utf-8")
    resultado = RUNNER.invoke(
        app,
        [
            *RUTA,
            "preparar",
            "--grado",
            "10",
            "--asignatura-malla",
            "Física",
            "--unidad",
            "Unidad sintética",
            "--catalogo",
            str(archivo),
            "--destino",
            str(entorno.base / "salida"),
            "--minutos",
            "25",
            "--recursos",
            "papel",
            "--conectividad",
            "sin_conexion",
            "--apoyo",
            "Representaciones visuales",
            "--seguridad",
            "Supervisión docente",
        ],
    )
    assert resultado.exit_code == 0, resultado.output
    salida = json.loads(resultado.stdout)
    assert salida["llamadas"] == 0 and salida["catalogo_modo"] == "piloto" and salida["filas"] == 2
    datos = json.loads((entorno.base / "salida/paquete.json").read_text(encoding="utf-8"))
    paquete = PaqueteUnidad.model_validate(datos["paquete"])
    assert paquete.filas == entorno.curriculo.filas and datos["ejecucion"] is None
    assert paquete.contexto.minutos_por_sesion == 25 and paquete.contexto.recursos == ("papel",)
    assert paquete.contexto.apoyos == ("Representaciones visuales",)
    assert paquete.contexto.seguridad == ("Supervisión docente",)
    assert all(e.fuente["sha256"] == entorno.curriculo.fuentes[0].sha256 for e in paquete.evidencias)
    assert "Catálogo: piloto" in (entorno.base / "salida/paquete.md").read_text(encoding="utf-8")
    assert entorno.llamadas == {"runtime": 0, "vertex": 0}
    assert not (entorno.base / "data/operaciones").exists()


def test_plan_no_carga_publicacion_configuracion_ni_crea_operaciones(entorno, monkeypatch):
    archivo = _guardar(entorno, _paquete(entorno))
    monkeypatch.setattr(cli_unidades, "cargar", lambda nombre: pytest.fail("Plan sin configuración remota"))
    resultado = _ejecutar(entorno, archivo, "--max-llamadas", "7", "--max-usd", "0.5")
    assert resultado.exit_code == 0, resultado.output
    assert json.loads(resultado.stdout) == {
        "estado": "plan",
        "paquete_id": _paquete(entorno).id,
        "llamadas": 0,
        "catalogo_modo": "completo",
        "sesiones": 2,
        "max_intentos_acumulados": 7,
        "max_reserva_usd": 0.5,
        "max_salida_tokens": 4096,
        "requiere_publicacion_completa": True,
    }
    assert entorno.llamadas["runtime"] == 0
    _sin_efectos(entorno)


def test_recuperacion_muestra_plan_sin_proveedor_y_prohibe_sobrescribir_original(entorno):
    paquete = _paquete(entorno)
    original = ResultadoUnidad(
        paquete_id=paquete.id,
        ejecucion_id="prueba-local",
        huella_entorno="sin-modelos",
        sesiones=(),
        cobertura={},
        tiempo_ms=0,
        consumo_antes={},
        consumo_despues={"llamadas": 2},
        consumo_incremental={},
        limites_presupuesto={"max_llamadas": 7, "max_usd": 0.5},
        checkpoints="sin-modelos",
    )
    archivo = _guardar(entorno, paquete)
    datos = json.loads(archivo.read_bytes())
    datos["ejecucion"] = original.model_dump(mode="json")
    archivo.write_text(json.dumps(datos), encoding="utf-8")
    antes = archivo.read_bytes()
    comando = [*RUTA, "recuperar-transporte", str(archivo), "--destino"]
    plan = RUNNER.invoke(app, [*comando, str(entorno.base / "recuperada")])
    assert plan.exit_code == 0, plan.output
    salida = json.loads(plan.stdout)
    assert salida["llamadas"] == 0 and salida["consumo_original"] == 2
    assert salida["limites_originales"] == {"max_llamadas": 7, "max_usd": 0.5}
    invalida = RUNNER.invoke(app, [*comando, str(archivo.parent), "--ejecutar"])
    assert invalida.exit_code == 2 and "aparte del archivo original" in invalida.output
    assert archivo.read_bytes() == antes
    assert entorno.llamadas == {"runtime": 0, "vertex": 0}
    assert not (entorno.base / "data/operaciones").exists()


@pytest.mark.parametrize("cambio", [{"modo": "piloto"}, {"piloto_superado": False}])
def test_ejecucion_de_piloto_o_sin_superar_se_rechaza_antes_de_runtime(entorno, cambio):
    paquete = _paquete(entorno, catalogo=entorno.catalogo.model_copy(update=cambio))
    resultado = _ejecutar(entorno, _guardar(entorno, paquete), "--ejecutar")
    assert resultado.exit_code == 2 and "catálogo completo" in resultado.output
    assert entorno.llamadas["runtime"] == 0
    _sin_efectos(entorno)


@pytest.mark.parametrize("campo", ["huella_curriculo", "huella_catalogo"])
def test_publicacion_debe_ser_identica_a_huellas_del_paquete(entorno, campo):
    paquete = _paquete(entorno).model_copy(update={campo: "d" * 64})
    resultado = _ejecutar(entorno, _guardar(entorno, paquete, recalcular=True), "--ejecutar")
    assert resultado.exit_code == 2 and "publicación activa" in resultado.output
    assert entorno.llamadas["runtime"] == 1
    _sin_efectos(entorno)


@pytest.mark.parametrize("cambio", ["filas", "clasificacion", "solicitud", "evidencia", "procedencia", "grado"])
def test_rehashear_no_autoriza_alterar_paquete_ni_eludir_snapshot(entorno, cambio):
    paquete = _paquete(entorno)
    if cambio == "filas":
        paquete = paquete.model_copy(update={"filas": paquete.filas[:1], "sesiones": paquete.sesiones[:1]})
    elif cambio == "clasificacion":
        decision = paquete.clasificaciones[0].model_copy(update={"justificacion": "Metadato modificado"})
        paquete = paquete.model_copy(update={"clasificaciones": (decision, *paquete.clasificaciones[1:])})
    elif cambio == "solicitud":
        solicitud = paquete.sesiones[0].solicitud.model_copy(update={"consulta": "Genera otro objetivo"})
        sesion = paquete.sesiones[0].model_copy(update={"solicitud": solicitud})
        paquete = paquete.model_copy(update={"sesiones": (sesion, *paquete.sesiones[1:])})
    elif cambio == "evidencia":
        evidencia = paquete.evidencias[0].model_copy(update={"texto": "Texto inventado"})
        paquete = paquete.model_copy(update={"evidencias": (evidencia, *paquete.evidencias[1:])})
    elif cambio == "procedencia":
        paquete = paquete.model_copy(update={"procedencia_catalogo": {"metodo": "Afirmación sin fuente"}})
    else:
        paquete = paquete.model_copy(update={"grado": 9})
    resultado = _ejecutar(entorno, _guardar(entorno, paquete, recalcular=True), "--ejecutar")
    assert resultado.exit_code == 2, resultado.output
    assert "publicación" in resultado.output or "No existe una unidad" in resultado.output
    _sin_efectos(entorno)


def test_hash_paquete_invalido_se_rechaza_incluso_en_plan(entorno):
    paquete = _paquete(entorno).model_copy(update={"id": "alterado"})
    resultado = _ejecutar(entorno, _guardar(entorno, paquete))
    assert resultado.exit_code == 2 and "paquete cambió" in resultado.output
    assert entorno.llamadas["runtime"] == 0
    _sin_efectos(entorno)


@pytest.mark.parametrize("cambio", ["archivo_original", "bloqueo_cientifico"])
def test_fuente_o_garantia_actualizada_exige_preparar_antes_de_vertex(entorno, cambio):
    archivo = _guardar(entorno, _paquete(entorno))
    if cambio == "archivo_original":
        fuente = entorno.base / "insumos/v4/fisica.xlsx"
        fuente.write_bytes(fuente.read_bytes() + b"edicion modificada")
    else:
        entorno.fuentes.bloqueos_curriculares = {
            entorno.curriculo.filas[0].id: "Bloqueo científico confirmado después de preparar"
        }
    resultado = _ejecutar(entorno, archivo, "--ejecutar")
    assert resultado.exit_code == 2 and "publicación" in resultado.output
    assert entorno.llamadas["runtime"] == 1
    _sin_efectos(entorno)


@pytest.mark.parametrize(
    "ruta", ["insumos", "data/fuentes", "data/grafo", "data/evaluacion", "src", "config", ".agents"]
)
def test_destinos_protegidos_se_rechazan_antes_de_leer_o_ejecutar(entorno, ruta):
    resultado = RUNNER.invoke(
        app,
        [*RUTA, "ejecutar", str(entorno.base / "no-existe.json"), "--destino", str(entorno.base / ruta), "--ejecutar"],
    )
    assert resultado.exit_code == 2 and "fuera de originales" in resultado.output
    assert entorno.llamadas["runtime"] == 0
    _sin_efectos(entorno)


@pytest.mark.parametrize("tipo", ["archivo", "enlace_directorio", "enlace_segundo_artefacto"])
def test_destino_no_exportable_no_gasta_presupuesto_antes_de_fallar(entorno, tipo):
    destino = entorno.base / "salida"
    protegido = entorno.base / "no-modificar.txt"
    protegido.write_text("original", encoding="utf-8")
    if tipo == "archivo":
        destino.write_text("original", encoding="utf-8")
    elif tipo == "enlace_directorio":
        externo = entorno.base / "externo"
        externo.mkdir()
        destino.symlink_to(externo, target_is_directory=True)
    else:
        destino.mkdir()
        (destino / "paquete.md").symlink_to(protegido)
    resultado = _ejecutar(entorno, _guardar(entorno, _paquete(entorno)), "--ejecutar")
    assert resultado.exit_code == 2, resultado.output
    assert entorno.llamadas == {"runtime": 0, "vertex": 0}
    assert protegido.read_text(encoding="utf-8") == "original"
    assert not (entorno.base / "data/operaciones").exists()
    assert not (destino / "paquete.json").exists()


def test_id_ejecucion_invalido_se_rechaza_antes_de_ledger_y_proveedor(entorno):
    resultado = _ejecutar(entorno, _guardar(entorno, _paquete(entorno)), "--ejecutar", "--ejecucion", "../otro")
    assert resultado.exit_code == 2 and "ID de presupuesto" in resultado.output
    assert entorno.llamadas["runtime"] == 0
    _sin_efectos(entorno)


@pytest.mark.parametrize("comando", ["ejecutar", "preparar"])
def test_json_de_entrada_sin_objeto_produce_error_de_uso(entorno, comando):
    archivo = entorno.base / "entrada.json"
    archivo.write_text("[]", encoding="utf-8")
    if comando == "ejecutar":
        resultado = _ejecutar(entorno, archivo, "--ejecutar")
    else:
        resultado = RUNNER.invoke(
            app,
            [
                *RUTA,
                comando,
                "--grado",
                "10",
                "--asignatura-malla",
                "Física",
                "--unidad",
                "Unidad sintética",
                "--catalogo",
                str(archivo),
                "--destino",
                str(entorno.base / "salida"),
            ],
        )
    assert resultado.exit_code == 2, resultado.output
    _sin_efectos(entorno)


def test_paquete_igual_a_publicacion_llega_al_ejecutor_con_un_solo_presupuesto(entorno, monkeypatch):
    paquete = _paquete(entorno)
    invocaciones = []

    def proveedor(configuracion, registro, ejecucion_id):
        entorno.llamadas["vertex"] += 1
        return SimpleNamespace(configuracion=configuracion, registro=registro, ejecucion_id=ejecucion_id)

    def ejecutar_verificado(recibido, *, motor, registro, ejecucion_id, directorio_operaciones):
        assert recibido == paquete and motor.fuentes is entorno.fuentes
        assert motor.modelo.registro is registro and motor.modelo.ejecucion_id == ejecucion_id == "prueba-local"
        assert motor.modelo.configuracion.max_llamadas == 7 and motor.modelo.configuracion.max_usd == 0.5
        assert motor.modelo.configuracion.max_salida_tokens == 8192
        assert directorio_operaciones == entorno.base / "data/operaciones/unidades"
        invocaciones.append(recibido.id)
        return ResultadoUnidad(
            paquete_id=recibido.id,
            ejecucion_id=ejecucion_id,
            huella_entorno="sin-modelos",
            sesiones=tuple(
                SesionEjecutada(sesion_id=s.id, fila_id=s.fila_id, estado="no_ejecutada", motivos=("Prueba de CLI",))
                for s in recibido.sesiones
            ),
            cobertura={"filas_totales": 2, "no_ejecutadas": 2},
            tiempo_ms=0,
            consumo_antes={},
            consumo_despues={"llamadas": 0, "reserva_usd": 0},
            consumo_incremental={"llamadas": 0, "reserva_usd": 0},
            checkpoints=str(directorio_operaciones / "simulado.sqlite"),
        )

    monkeypatch.setattr("goes_science_kg.motor.vertex.Vertex", proveedor)
    monkeypatch.setattr("goes_science_kg.producto.unidades.ejecutar_unidad", ejecutar_verificado)
    resultado = _ejecutar(
        entorno,
        _guardar(entorno, paquete),
        "--ejecutar",
        "--ejecucion",
        "prueba-local",
        "--max-salida-tokens",
        "8192",
        "--max-llamadas",
        "7",
        "--max-usd",
        "0.5",
    )
    assert resultado.exit_code == 0, resultado.output
    assert json.loads(resultado.stdout)["consumo"]["llamadas"] == 0
    assert entorno.llamadas == {"runtime": 1, "vertex": 1} and invocaciones == [paquete.id]
    datos = json.loads((entorno.base / "salida/paquete.json").read_text(encoding="utf-8"))
    assert datos["ejecucion"]["ejecucion_id"] == "prueba-local" and len(datos["paquete"]["filas"]) == 2


@pytest.mark.parametrize("salida", ["config/no.json", "data/fuentes/no.json", "data/evaluacion/no.json", "salida.md"])
def test_conocimiento_valida_destino_antes_de_fuentes_y_generacion(entorno, monkeypatch, salida):
    monkeypatch.setattr("goes_science_kg.curriculo.cargar_curriculo", lambda: pytest.fail("Salida inválida"))
    resultado = RUNNER.invoke(app, ["producto", "conocimiento", "--salida", str(entorno.base / salida)])
    assert resultado.exit_code == 2, resultado.output
    assert not (entorno.base / salida).exists()
    _sin_efectos(entorno)


def test_conocimiento_ejecuta_motor_local_y_exporta_hipotesis_sin_aceptarlas(entorno, monkeypatch):
    referencia = ReferenciasProducto(
        nodos=(Nodo(id="vocabulario-medicion", tipo=TipoNodo.CONCEPTO, etiqueta="Medición"),),
        procedencia={"version_curriculo": "V3", "clase": "historica"},
    )
    recibidas = []

    def cargar_referencias(carpeta):
        recibidas.append(carpeta)
        return referencia

    monkeypatch.setattr("goes_science_kg.producto.referencias.cargar_referencias_historicas", cargar_referencias)
    salida = entorno.base / "reportes-generados/candidatos.json"
    resultado = RUNNER.invoke(
        app, ["producto", "conocimiento", "--grafo-v3", "grafo-referencia", "--salida", str(salida)]
    )
    assert resultado.exit_code == 0, resultado.output
    resumen = json.loads(resultado.stdout)
    informe = InformeConocimiento.model_validate_json(salida.read_text(encoding="utf-8"))
    assert resumen["llamadas"] == 0 and resumen["salida"] == str(salida)
    assert recibidas == [entorno.base / "grafo-referencia"]
    assert set(informe.por_fila) == {f.id for f in entorno.curriculo.filas}
    assert informe.candidatos and not informe.aceptados and not informe.dag_aceptado.aristas
    assert all(e.fuente["sha256"] == entorno.curriculo.fuentes[0].sha256 for e in informe.evidencias.values())
    assert entorno.llamadas == {"runtime": 0, "vertex": 0}
    assert not list(salida.parent.glob("*.tmp"))


def test_conocimiento_no_sustituye_salida_enlazada(entorno, monkeypatch):
    fuente = entorno.base / "conservar.json"
    fuente.write_text('{"original": true}', encoding="utf-8")
    salida = entorno.base / "candidatos.json"
    salida.symlink_to(fuente)
    monkeypatch.setattr("goes_science_kg.curriculo.cargar_curriculo", lambda: pytest.fail("Salida enlazada"))
    resultado = RUNNER.invoke(app, ["producto", "conocimiento", "--salida", str(salida)])
    assert resultado.exit_code == 2 and fuente.read_text(encoding="utf-8") == '{"original": true}'
