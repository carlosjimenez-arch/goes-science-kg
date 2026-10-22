"""Preserve export denominators and prevent curricular text from executing as code."""

from datetime import UTC, datetime
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from test_mejoras_curriculares import caso as caso

from goes_science_kg.curriculo.contratos import FASES, FaseCurricular, FilaCurricular
from goes_science_kg.producto.exportacion_mejoras import csv_curricular, exportar_excel, filas_exportables
from goes_science_kg.producto.mejoras_publicadas import CasoMejora, verificar_preparacion
from goes_science_kg.producto.validacion_mejoras import resolver_mejora
from goes_science_kg.producto.vista_mejoras import vista_mejoras


@pytest.fixture
def entrega(caso):
    p, b, modelos, ronda = caso
    fila = FilaCurricular(
        id="fila1",
        fuente_id="fuente1",
        version="V4",
        hoja="Prueba",
        fila=3,
        grado=3,
        asignatura_malla="Ciencias",
        unidad="Materia",
        procedimental=p.filas_originales[0]["procedimental"],
        fases={fase: FaseCurricular(presente_en_hoja=False) for fase in FASES},
        celdas=(),
    )
    resultado = resolver_mejora(p, (ronda,), ejecucion_id="prueba", modelos=modelos)
    p = p.model_copy(update={"filas_originales": (fila.model_dump(mode="json", exclude={"celdas"}),)})
    c = CasoMejora(preparacion=p, resultado=resultado, entorno="a" * 64, plan_id="b" * 64)
    archivo = SimpleNamespace(
        id="c" * 64,
        unidades=(c,),
        limite="Sólo una prueba técnica.",
        cobertura={"unidades": 1, "filas": 1, "filas_respaldadas": 1},
        progresiones={"catalogo": {"expectativas": []}},
    )
    return SimpleNamespace(filas=(fila,)), archivo


def test_csv_y_excel_incluyen_original_propuesta_y_ausencias(entrega, tmp_path):
    curriculo, archivo = entrega
    filas = list(filas_exportables(curriculo, archivo))
    assert len(filas) == 1 and filas[0]["estado_revision"] == "respaldado_automaticamente"
    assert filas[0]["objetivo_original"] == filas[0]["objetivo_propuesto"]
    assert all(filas[0]["fase_" + f] == "Ausente en el insumo" for f in FASES)
    assert csv_curricular(curriculo, archivo).startswith("\ufefffila_id,")
    destino = tmp_path / "matriz.xlsx"
    exportar_excel(curriculo, archivo, destino)
    with pytest.raises(ValueError, match="nuevo"):
        exportar_excel(curriculo, archivo, destino)
    libro = load_workbook(destino, read_only=True)
    assert libro.worksheets[0].max_row == 2
    libro.close()


def test_hoja_no_ejecuta_formulas_del_texto(entrega, tmp_path):
    curriculo, archivo = entrega
    curriculo = SimpleNamespace(filas=(curriculo.filas[0].model_copy(update={"procedimental": "=1+1"}),))
    assert "'=1+1" in csv_curricular(curriculo, archivo)
    p = tmp_path / "segura.xlsx"
    exportar_excel(curriculo, archivo, p)
    libro = load_workbook(p, read_only=True, data_only=False)
    c = next(c for c in libro.worksheets[0][2] if c.value == "=1+1")
    assert c.data_type == "s"
    libro.close()


def test_exportacion_excel_conserva_bytes_con_reloj_distinto(entrega, tmp_path, monkeypatch):
    from openpyxl.writer import excel

    curriculo, archivo = entrega
    salidas = []
    for anio in (2026, 2027):
        instante = datetime(anio, 5, 10, 12, 30, tzinfo=UTC)
        monkeypatch.setattr(
            excel,
            "datetime",
            SimpleNamespace(
                datetime=SimpleNamespace(now=lambda instante=instante, **_: instante),
                timezone=SimpleNamespace(utc=UTC),
            ),
        )
        destino = tmp_path / f"matriz-{anio}.xlsx"
        exportar_excel(curriculo, archivo, destino)
        salidas.append(destino.read_bytes())
        with ZipFile(destino) as zip_:
            assert {info.date_time for info in zip_.infolist()} == {(2026, 1, 1, 0, 0, 0)}
    assert salidas[0] == salidas[1]


def test_hoja_de_control_no_ejecuta_formulas_de_metadata(entrega, tmp_path):
    curriculo, archivo = entrega
    archivo.limite = "=1+1"
    destino = tmp_path / "control-seguro.xlsx"
    exportar_excel(curriculo, archivo, destino)
    libro = load_workbook(destino, read_only=True, data_only=False)
    try:
        celda = libro["Alcance y control"]["B2"]
        assert celda.value == "=1+1" and celda.data_type == "s"
        assert all(celda.data_type != "f" for hoja in libro for fila in hoja for celda in fila)
    finally:
        libro.close()


def test_vista_escapa_texto_y_mantiene_filtros_y_origen(entrega):
    curriculo, archivo = entrega
    c = archivo.unidades[0]
    b = c.resultado.propuesta
    fila = b.filas[0].model_copy(update={"contexto_salvador": "<script>alert(1)</script>"})
    archivo.unidades = (
        c.model_copy(
            update={"resultado": c.resultado.model_copy(update={"propuesta": b.model_copy(update={"filas": (fila,)})})}
        ),
    )
    vista = vista_mejoras(curriculo, archivo)
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in vista
    assert "<script>alert(1)</script>" not in vista
    assert 'data-grado="3"' in vista and 'data-estado="respaldado_automaticamente"' in vista
    assert "archivo-curricular.json" in vista and "Fragmentos seleccionados" in vista


def test_preparacion_rechaza_texto_cambiado_y_literal_ajeno(caso):
    p = caso[0]
    verificar_preparacion(p)
    for cambios in (
        {"evidencias": (p.evidencias[0].model_copy(update={"texto": "Cambiado"}), *p.evidencias[1:])},
        {"fragmentos": (p.fragmentos[0].model_copy(update={"literal": "Texto ajeno a la fuente"}), *p.fragmentos[1:])},
    ):
        alterada = p.model_copy(update=cambios)
        from goes_science_kg.motor.contratos import huella

        alterada = alterada.model_copy(update={"id": huella(alterada.model_dump(mode="json", exclude={"id"}))})
        with pytest.raises(ValueError):
            verificar_preparacion(alterada)
