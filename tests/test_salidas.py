"""Generadores de salidas (fichas, progresión, alineación, informe y visor internacionales, CLI).

Todas escriben en carpetas temporales: las pruebas no tocan el repo.
"""

from __future__ import annotations

import csv
import json
import re

import pytest

from goes_science_kg.config import cargar, ruta


def _redirigir(monkeypatch, modulo, tmp_path) -> None:
    """Hace que `modulo.ruta(...)` apunte a tmp_path (solo para escribir; lo que lee llega por argumento)."""
    monkeypatch.setattr(modulo, "ruta", lambda rel: tmp_path / rel)


def test_fichas_una_por_asignatura(grafo, tmp_path, monkeypatch):
    from goes_science_kg import fichas

    _redirigir(monkeypatch, fichas, tmp_path)
    escritas = fichas.escribir_fichas(*grafo, "prueba")
    assert len(escritas) == len(cargar("asignaturas")["asignaturas"])
    for p in escritas:
        texto = open(p, encoding="utf-8").read()
        assert texto.startswith("#") and "|" in texto


def test_progresion_autocontenida(grafo, tmp_path, monkeypatch):
    from goes_science_kg import progresion

    _redirigir(monkeypatch, progresion, tmp_path)
    for p in progresion.escribir_todas(*grafo):
        html = open(p, encoding="utf-8").read()
        assert "<script src" not in html                      # sin librerías externas
        datos = json.loads(re.search(r"const D = (.*?);\n", html).group(1).replace("<\\/", "</"))
        assert datos["filas"] and datos["grados"] == list(range(2, 12))


def _catalogo() -> list[dict]:
    return [{"codigo": c, "asignatura": a, "unidad": "U1", "eje": "E", "tema": "T", "objetivo": "O"}
            for c, a in (("BIO1", "biologia"), ("QUI1", "quimica"))]


def test_alineacion_preparar_y_unir(tmp_path, monkeypatch):
    from goes_science_kg import alineacion

    monkeypatch.setattr(alineacion, "DIR", str(tmp_path))
    lotes = alineacion.preparar_bachillerato(_catalogo(), "prueba")
    assert lotes
    for p in sorted((tmp_path / "lotes").glob("lote_*.json")):
        lote = json.loads(p.read_text(encoding="utf-8"))
        codigo = "BIO1" if "biologia" in lote["lote"] else "QUI1"
        salida = [{"id": i["id"], "obj1": codigo, "obj2": "", "confianza": "alta", "justificacion": "prueba"}
                  for i in lote["items"]]
        p.with_name(p.name.replace("lote_", "salida_", 1)).write_text(json.dumps(salida), encoding="utf-8")
    r = alineacion.unir("AUSS_", "union", _catalogo())
    assert r["temas"] > 0 and r["confianza"] == {"alta": r["temas"]}


def test_alineacion_rechaza_codigos_inventados(tmp_path, monkeypatch):
    from goes_science_kg import alineacion

    monkeypatch.setattr(alineacion, "DIR", str(tmp_path))
    alineacion.preparar_bachillerato(_catalogo(), "prueba")
    for p in (tmp_path / "lotes").glob("lote_*.json"):
        items = json.loads(p.read_text(encoding="utf-8"))["items"]
        salida = [{"id": i["id"], "obj1": "INVENTADO", "confianza": "alta", "justificacion": "x"} for i in items]
        p.with_name(p.name.replace("lote_", "salida_", 1)).write_text(json.dumps(salida), encoding="utf-8")
    with pytest.raises(ValueError, match="obj1 inválido"):
        alineacion.unir("AUSS_", "union", _catalogo())


def test_consenso_versionado_al_dia(tmp_path, monkeypatch):
    """El consenso internacional versionado corresponde a los objetivos y etiquetas versionados."""
    import dataclasses

    from goes_science_kg.internacional import consenso, tramo

    monkeypatch.setitem(tramo.TRAMOS, "9_11", dataclasses.replace(tramo.TRAMOS["9_11"], grafo=str(tmp_path)))
    nuevo = consenso.construir()
    versionado = json.loads(ruta("data/grafo/internacional/manifest.json").read_text(encoding="utf-8"))
    assert json.loads(json.dumps(nuevo)) == versionado
    for nombre in ("nodos.jsonl", "aristas.jsonl", "consenso.json"):
        assert (tmp_path / nombre).read_bytes() == ruta(f"data/grafo/internacional/{nombre}").read_bytes(), nombre


def test_informe_internacional(tmp_path, monkeypatch):
    import dataclasses

    from openpyxl import load_workbook

    from goes_science_kg.internacional import contraste, informe, tramo

    monkeypatch.setitem(tramo.TRAMOS, "9_11", dataclasses.replace(tramo.TRAMOS["9_11"], informe=str(tmp_path)))
    r = informe.escribir()
    for asig in contraste.ASIGNATURAS:
        md = (tmp_path / f"{asig}.md").read_text(encoding="utf-8")
        assert "## Resumen" in md and "## Profundidad" in md
        wb = load_workbook(tmp_path / f"Contraste_{asig}.xlsx")
        assert wb.sheetnames == ["Datos", "Contraste", "Resumen"]
        c = wb["Contraste"]
        assert all(str(x.value).startswith("=") for x in c["C"][1:])       # todo número es fórmula
        total = wb["Resumen"].cell(wb["Resumen"].max_row, 2).value
        assert total.startswith("=SUM(")
    with open(tmp_path / "revision_humana.csv", encoding="utf-8") as f:
        filas = list(csv.DictReader(f))
    assert len(filas) == r["revision_humana"]["filas"] and {"decision", "revisado_por"} <= set(filas[0])
    assert (tmp_path / "visor.html").exists()
    # Mismo resultado que el versionado (salidas deterministas).
    assert (tmp_path / "contraste.json").read_bytes() == ruta("internacional/contraste.json").read_bytes()


def test_visor_internacional_escapa_el_cierre_de_script(tmp_path):
    from goes_science_kg.internacional import informe, visor

    fila = {"concepto": "CON:x/y", "nombre": "</script><b>inyección</b>", "asignatura": "fisica", "clase": "alineado",
            "sv_primer_grado": 10, "sv_en_tramo": 10, "n_nucleo_en_tramo": 1, "n_especializacion": 0, "paises": {},
            "sv_temas": []}
    visor.escribir({"filas": [fila], "temas": {}}, informe.clases(), informe.NOMBRE, tmp_path)
    html = (tmp_path / "visor.html").read_text(encoding="utf-8")
    assert html.count("</script>") == 1                       # solo el cierre real del bloque
    datos = json.loads(re.search(r"const D = (.*?);\n", html).group(1).replace("<\\/", "</"))
    assert datos["filas"][0]["c"] == fila["nombre"]


@pytest.mark.parametrize("args", [["--help"], ["internacional", "--help"], ["grafo", "validar"]])
def test_cli(args):
    from typer.testing import CliRunner

    from goes_science_kg.cli import app

    r = CliRunner().invoke(app, args)
    assert r.exit_code == 0, r.output


def test_visor_de_grado_y_comunidades(grafo, tmp_path, monkeypatch):
    from goes_science_kg import comunidades, visor
    from goes_science_kg.grados import subgrafo
    from goes_science_kg.prerrequisitos import evidencia_orden

    nodos, aristas = grafo
    sn, sa, diag = subgrafo(7, nodos, aristas, evidencia_orden(nodos, aristas))
    a, b = comunidades.comunidades_grado(7, sn, sa), comunidades.comunidades_grado(7, sn, sa)
    assert a and a == b                                       # Louvain con semilla fija: mismo resultado
    _redirigir(monkeypatch, visor, tmp_path)
    html = open(visor.escribir_visor(7, sn, sa, diag), encoding="utf-8").read()
    externos = re.findall(r'<script src="([^"]+)"', html)
    assert all(re.search(r"@\d+\.\d+\.\d+/", u) for u in externos), externos   # versión exacta fijada
    assert html.count("</script>") == html.count("<script")


@pytest.mark.parametrize("clave", ["9_11", "2_8"])
def test_formula_lectura_balanceada_en_cada_tramo(tmp_path, clave):
    """La fórmula «Lectura» del Excel cambia según el tramo; un paréntesis de más la vuelve inválida en Numbers."""
    from openpyxl import load_workbook

    from goes_science_kg.internacional import informe, tramo

    tramo.usar(clave)
    try:
        fila = {"nombre": "x", "clase": "alineado", "sv_primer_grado": 9, "sv_en_tramo": None, "paises": {}}
        informe._excel("fisica", [fila], tmp_path)
        formula = load_workbook(tmp_path / "Contraste_fisica.xlsx")["Contraste"]["I2"].value
        assert formula.count("(") == formula.count(")"), formula
        assert ("despues del tramo" in formula) == (clave == "2_8")
    finally:
        tramo.usar("9_11")
